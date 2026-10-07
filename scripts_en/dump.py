import sys,json,re
from pptx import Presentation
P=sys.argv[1]
prs=Presentation(P)
from collections import OrderedDict
U=OrderedDict()
def paras(shape):
    if shape.shape_type==6:
        for s in shape.shapes: yield from paras(s)
    elif shape.has_text_frame:
        for p in shape.text_frame.paragraphs: yield p
    elif getattr(shape,'has_table',False) and shape.has_table:
        for r in shape.table.rows:
            for c in r.cells:
                for p in c.text_frame.paragraphs: yield p
cyr=re.compile('[А-Яа-яІіЇїЄєҐґ]')
n=0
for i,sl in enumerate(prs.slides,1):
    for sh in sl.shapes:
        for p in paras(sh):
            t=''.join(r.text for r in p.runs)
            if t.strip():
                n+=1
                U.setdefault(t,[]).append(i)
print(n,len(U), sum(1 for t in U if cyr.search(t)))
json.dump([[t,v] for t,v in U.items()],open('en/strings.json','w'),ensure_ascii=False,indent=0)
