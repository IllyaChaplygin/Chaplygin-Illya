import sys,re,json
sys.path.insert(0,'/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad/en')
from pptx import Presentation
from rules import *; from d1 import D1; from d2 import D2; from d3 import D3
D={**D1,**D2,**D3}
cyr=re.compile('[А-Яа-яІіЇїЄєҐґ]')
SRC,OUT=sys.argv[1],sys.argv[2]
prs=Presentation(SRC)
def paras(sh):
    if sh.shape_type==6:
        for s in sh.shapes: yield from paras(s)
    elif sh.has_text_frame:
        yield from sh.text_frame.paragraphs
def tr(t):
    if t in D: r=D[t]
    else:
        r=r_weight(t)
        if r is None:
            m=re.match(r'^(\d+(?:,\d+)?) т$',t)
            r=f"{m.group(1)} t" if m else t
    r=re.sub(r'(\d),(\d)',r'\1.\2',r)
    return r
miss=[];n=0
for i,sl in enumerate(prs.slides,1):
    for sh in sl.shapes:
        for p in paras(sh):
            if not p.runs: continue
            t=''.join(r.text for r in p.runs)
            if not t.strip(): continue
            r=tr(t)
            if cyr.search(r): miss.append((i,t))
            if r!=t:
                p.runs[0].text=r
                for x in p.runs[1:]: x.text=''
                for x in p.runs:
                    rp=x._r.get_or_add_rPr(); rp.set('lang','en-US')
                n+=1
print('changed',n,'missing',len(miss))
for m in miss: print(m)
prs.save(OUT)
