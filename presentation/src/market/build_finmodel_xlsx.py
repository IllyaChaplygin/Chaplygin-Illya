# -*- coding: utf-8 -*-
"""Excel finmodel for the snack suppliers, built on the exact layout and
formulas of RICE_RICE_FINMODEL_ALL_MARGINS_Final.xlsx. The rice workbook is
used as the style template (fonts, fills, widths, merges, dropdown), its data
rows are replaced by the 28 retail SKUs from the SelfCost, and every price
column stays a live formula: change the scenario / bonus / margin / markup in
D2:D6 and the whole sheet recalculates."""
import copy, json, os, re, sys
import openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'deck'))
from catalog import SUPPLIERS, DATA  # noqa: E402

TEMPLATE = ('/root/.claude/uploads/c53f34b4-ae43-52db-aeb2-01bdcda48cc3/'
            'e1e1ccb4-RICE_RICE_FINMODEL_ALL_MARGINS_Final..xlsx')
OUT = os.path.join(HERE, '..', '..', 'SNACKS_FINMODEL_ALL_MARGINS.xlsx')
SCEN = [("20'", '20ф'), ("40'", '40ф'), ('LCL 17', 'Збірний 17 м³'), ('LCL 34', 'Збірний 34 м³')]
CHEAPEST = 'Найдешевший'
GROUP = {'singha': 'Singha Kameda — Arare Norimaki', 'thainichi': 'Thai-Nichi — Mizuho Norimaki',
         'tmk': 'TMK — KOKIRI Wow', 'zek': 'HanJin — ZEK'}
BRAND = {'singha': 'Singha Kameda (Таїланд)', 'thainichi': 'Thai-Nichi (Таїланд)',
         'tmk': 'TMK Thailand / KOKIRI', 'zek': 'HanJin / ZEK (Китай)'}


def weight(badge):
    m = re.match(r'([\d,]+)\s*г', badge)
    return float(m.group(1).replace(',', '.'))


def skus():
    rows = []
    for sup in SUPPLIERS:
        for p in sup['products']:
            cost = next(r for r in DATA[sup['sheet']] if r['name'] == p['key'])['cost']
            rows.append(dict(sup=sup['id'], group=GROUP[sup['id']] + ' ' + re.sub(r'.*?(\d[\d,]*\s*г).*', r'\1', p['unit']),
                             title='%s — %s' % (sup['brand'], p['title'].replace('\n', ' ')),
                             grams=weight(p['badge']), cost=cost))
    return rows


def clone(src, dst):
    dst._style = copy.copy(src._style)


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws, base = wb.worksheets
    items = skus()
    n = len(items)

    # ---------------------------------------------------------------- СС_база
    nb = wb.create_sheet('СС_нова')
    nb.sheet_view.showGridLines = base.sheet_view.showGridLines
    nb.freeze_panes = 'E4'
    for col, w in {'A': 26, 'B': 38, 'C': 46, 'D': 9}.items():
        nb.column_dimensions[col].width = w
    for c in 'EFGHIJKLM':
        nb.column_dimensions[c].width = 14
    nb['B1'] = 'Собівартість із розрахункових аркушів Self-Cost_Snacks.xlsx (SelfCost, курс 45 грн/$)'
    clone(base['B1'], nb['B1'])
    nb.merge_cells('E2:H2'); nb.merge_cells('I2:L2')
    nb['E2'] = 'СОБІВАРТІСТЬ, $ / уп.'; clone(base['E2'], nb['E2'])
    nb['I2'] = 'СОБІВАРТІСТЬ SELF_COST, грн / уп.'; clone(base['L2'], nb['I2'])
    for j, h in enumerate(['Производитель', 'Група', 'Позиція', 'Вага, г']):
        c = nb.cell(3, 1 + j, h); clone(base['A3'], c)
    for k, (_, name) in enumerate(SCEN):
        c = nb.cell(3, 5 + k, name); clone(base['E3'], c)
        c = nb.cell(3, 9 + k, name); clone(base['L3'], c)
    for r, it in enumerate(items, start=4):
        for j, v in enumerate([BRAND[it['sup']], it['group'], it['title'], it['grams']]):
            c = nb.cell(r, 1 + j, v); clone(base.cell(4, 1 + j), c)
        for k, (key, _) in enumerate(SCEN):
            c = nb.cell(r, 5 + k, it['cost'][key]['usd']); clone(base['E4'], c)
            c = nb.cell(r, 9 + k, it['cost'][key]['uah']); clone(base['L4'], c)
    nb.row_dimensions[2].height = 22

    # ------------------------------------------------------------------ model
    # data rows 11.. : keep row-11 style, drop the surplus rice rows
    tmpl_style = {col: copy.copy(ws.cell(11, col)._style) for col in range(1, 37)}
    tmpl_h = ws.row_dimensions[11].height
    old_last = ws.max_row
    for r in range(11, old_last + 1):
        for col in range(1, 37):
            ws.cell(r, col).value = None
    if old_last > 10 + n:
        ws.delete_rows(11 + n, old_last - 10 - n)
    ws['B1'] = 'Снеки (норі). Ціна партнеру та полиця за прикладом рису'
    ws['D2'] = CHEAPEST
    ws['E2'] = 'Сценарій з випадаючого списку; «Найдешевший» = мінімум із 4 сценаріїв для кожного SKU (як у презентації).'
    ws['E3'] = 'Курс 45 — з Self-Cost_Snacks.'
    ws['E4'] = 'Бонус 25%; полиця = ціна партнера × 1,40.'
    ws['E5'] = '35% = (ціна партнеру − собівартість − бонус) / повна ціна партнеру.'
    ws['E6'] = 'Прибуток після бонусу = ціна партнеру − собівартість − бонус мережі.'
    ws['E7'] = 'ПДВ як у прикладі рису: СС збережено повністю; ціна вже з ПДВ. Множення на 1,20 немає.'
    ws.data_validations.dataValidation = []
    dv = DataValidation(type='list', formula1='"%s"' % ','.join([s[1] for s in SCEN] + [CHEAPEST]), allow_blank=False)
    dv.add('D2'); ws.add_data_validation(dv)
    ws.conditional_formatting = type(ws.conditional_formatting)()
    for i, it in enumerate(items):
        r = 11 + i; s = 4 + i
        vals = {1: i + 1, 2: BRAND[it['sup']], 3: it['group'], 4: it['title'], 5: it['grams'],
                6: '=IF($D$2="%s",MIN(СС_нова!$E%d:$H%d),INDEX(СС_нова!$E%d:$H%d,1,MATCH($D$2,СС_нова!$E$3:$H$3,0)))' % (CHEAPEST, s, s, s, s),
                7: '=IF(AND(ISNUMBER(F{r}),ISNUMBER($D$8)),F{r}*$D$3,NA())',
                8: '=G{r}/$D$8', 9: '=H{r}*$D$4', 10: '=H{r}-G{r}-I{r}', 11: '=J{r}/H{r}', 12: '=H{r}*$D$6',
                13: '=G{r}/(1-$D$4)', 14: '=M{r}*$D$6', 15: '=H{r}-G{r}', 16: '=O{r}-I{r}'}
        for blk, col0 in enumerate([17, 22, 27, 32]):
            m = '$%s$9' % L(col0)
            a, b, c_, d, e = [L(col0 + k) for k in range(5)]
            vals[col0] = '=$G{r}/(1-$D$4-%s)' % m
            vals[col0 + 1] = '=%s{r}*$D$4' % a
            vals[col0 + 2] = '=%s{r}-$G{r}-%s{r}' % (a, b)
            vals[col0 + 3] = '=%s{r}/%s{r}' % (c_, a)
            vals[col0 + 4] = '=%s{r}*$D$6' % a
        for col in range(1, 37):
            c = ws.cell(r, col); c._style = copy.copy(tmpl_style[col])
            v = vals.get(col)
            c.value = v.replace('{r}', str(r)) if isinstance(v, str) else v
        ws.row_dimensions[r].height = tmpl_h
    ws.auto_filter.ref = 'A10:AJ%d' % (10 + n)

    wb.remove(base)
    nb.title = 'СС_база'
    # formulas were written against the temporary name
    for row in ws.iter_rows(min_row=11, max_row=10 + n, min_col=6, max_col=6):
        for c in row:
            c.value = c.value.replace('СС_нова', 'СС_база')
    wb.save(OUT)
    print('saved', OUT, n, 'SKUs')


if __name__ == '__main__':
    main()
