# -*- coding: utf-8 -*-
"""Morskyi Dim slide kit — the layout language of the SIAS and rice market
studies (navy header band, gold rule, white logo, light-blue panels, segment
colours), re-expressed for the 10 x 5.625 in supplier deck. The samples are
13.333 x 7.5 in — the same 16:9 ratio — so every coordinate below is written in
sample units and scaled by K, which keeps the geometry identical to the
studies while the slide size of the user's own deck stays untouched."""
import os, sys
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
K = 10.0 / 13.3333333
FONT = 'Segoe UI'


def rgb(h):
    return RGBColor.from_string(h)


NAVY, GOLD, GOLD_TX, PAGE_NO = rgb('303A5D'), rgb('F9A50B'), rgb('FFC95C'), rgb('465382')
PANEL, MUTED, INK, WHITE = rgb('F1F4FA'), rgb('717890'), rgb('1C2235'), rgb('FFFFFF')
AMBER, PURPLE, TEAL, RED, GREEN, GREY = rgb('C4820A'), rgb('5D4B96'), rgb('1F9AA6'), rgb('B23A3A'), rgb('3E8E5A'), rgb('A5ABBD')
RULE = rgb('D9DEEA')
SEG_COLOR = {'mini': NAVY, 'chips': AMBER, 'tempura': PURPLE, 'topping': TEAL, 'ricecr': RED, 'ricechip': GREY}
SUP_COLOR = {'singha': rgb('B36214'), 'thainichi': rgb('1D5B63'), 'tmk': rgb('2F7A3E'), 'zek': rgb('B23A22')}


def mixc(a, b, t):
    return RGBColor(*(round(x * t + y * (1 - t)) for x, y in zip(a, b)))


def tint(c, t=0.15):
    return mixc(c, WHITE, t)


def fs(pt):
    return max(6.0, pt * 0.80)


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(kind, Inches(x * K), Inches(y * K), Inches(w * K), Inches(h * K))
    if radius:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.shadow.inherit = False
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    return sh


def text(s, x, y, w, h, t, size=9, color=INK, bold=False, align='l', anchor='t', italic=False, wrap=True, spc=0, ls=1.0):
    """t: str or list of (str, {size,color,bold}) runs on ONE paragraph; '\\n' splits paragraphs."""
    box = s.shapes.add_textbox(Inches(x * K), Inches(y * K), Inches(w * K), Inches(h * K))
    tf = box.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    paras = t if isinstance(t, list) and t and isinstance(t[0], list) else [t]
    first = True
    for para in paras:
        runs = para if isinstance(para, list) else [(para, {})]
        # split plain strings on newlines into paragraphs
        chunks = []
        for r in runs:
            txt, o = r if isinstance(r, tuple) else (r, {})
            parts = txt.split('\n')
            for i, part in enumerate(parts):
                chunks.append((i > 0, part, o))
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        p.line_spacing = ls
        for newp, part, o in chunks:
            if newp:
                p = tf.add_paragraph(); p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]; p.line_spacing = ls
            r = p.add_run(); r.text = part
            f = r.font; f.name = FONT; f.size = Pt(fs(o.get('size', size)))
            f.bold = o.get('bold', bold); f.italic = o.get('italic', italic); f.color.rgb = o.get('color', color)
            if spc or o.get('spc'):
                r.font._rPr.set('spc', str(int(o.get('spc', spc) * 100)))
    return box


def caps(s, x, y, w, t, size=7, color=MUTED, align='l'):
    return text(s, x, y, w, 0.2, t.upper(), size=size, color=color, bold=True, align=align, spc=0.4)


def picture(s, path, x, y, w, h):
    """Fit (contain) inside the box, centred."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    return s.shapes.add_picture(path, Inches((x + (w - pw) / 2) * K), Inches((y + (h - ph) / 2) * K), Inches(pw * K), Inches(ph * K))


class Deck:
    """Page counter + slide factory for the appended section."""
    def __init__(self, prs, first_page):
        self.prs, self.page = prs, first_page

    def slide(self, eyebrow, title):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.page += 1
        rect(s, 0, 0, 13.3333, 7.5, fill=WHITE)
        rect(s, 0, 0, 13.3333, 1.16, fill=NAVY)
        rect(s, 0, 1.16, 13.3333, 0.06, fill=GOLD)
        s.shapes.add_picture(os.path.join(HERE, 'assets', 'logo.png'), Inches(0.66 * K), Inches(0.26 * K), Inches(1.52 * K), Inches(0.66 * K))
        text(s, 2.52, 0.30, 9.0, 0.2, eyebrow.upper(), size=9.5, color=GOLD_TX, bold=True, spc=0.3)
        text(s, 2.52, 0.55, 9.0, 0.5, title, size=19, color=WHITE, bold=True, anchor='m')
        text(s, 11.55, 0.44, 1.12, 0.33, '%02d' % self.page, size=17, color=PAGE_NO, bold=True, align='r')
        return s


def footnote(s, t):
    text(s, 0.62, 7.08, 12.1, 0.3, t, size=7, color=MUTED, italic=True)


def kpi(s, x, y, w, h, label, value, sub=None, color=NAVY, vsize=22):
    rect(s, x, y, w, h, fill=PANEL)
    rect(s, x, y, 0.04, h, fill=color)
    caps(s, x + 0.26, y + 0.17, w - 0.4, label, size=7.6)
    text(s, x + 0.26, y + 0.4, w - 0.4, 0.5, value, size=vsize, color=color, bold=True)
    if sub:
        text(s, x + 0.26, y + h - 0.3, w - 0.4, 0.26, sub, size=7.4, color=MUTED)


def section_tag(s, x, y, w, label, color=NAVY):
    rect(s, x, y, 0.04, 0.21, fill=color)
    text(s, x + 0.16, y + 0.02, w, 0.15, label.upper(), size=6.8, color=color, bold=True, spc=0.3)


def hbars(s, x, y, w, rows, color=NAVY, row_h=0.24, name_w=1.45, val_w=0.5, maxv=None, size=7.8, fmt=lambda v: '%d' % v):
    """rows: [(name, value)] or [(name, value, color)]."""
    maxv = maxv or max(r[1] for r in rows)
    bw = w - name_w - val_w
    for i, r in enumerate(rows):
        c = r[2] if len(r) > 2 else color
        yy = y + i * row_h
        text(s, x, yy, name_w - 0.08, row_h, r[0], size=size, color=NAVY, anchor='m')
        rect(s, x + name_w, yy + row_h * 0.27, max(0.03, bw * r[1] / maxv), row_h * 0.46, fill=c)
        text(s, x + name_w + max(0.03, bw * r[1] / maxv) + 0.07, yy, val_w + 0.3, row_h, fmt(r[1]), size=size, color=MUTED, bold=True, anchor='m')


def columns(s, x, y, w, h, bins, color=NAVY, labels=None, size=7.2):
    """Vertical histogram; bins: list of (label, count)."""
    n = len(bins); gap = 0.1; bw = (w - gap * (n - 1)) / n; mx = max(b[1] for b in bins) or 1
    for i, (lab, v) in enumerate(bins):
        bh = (h - 0.55) * v / mx
        bx = x + i * (bw + gap)
        if v:
            rect(s, bx, y + h - 0.3 - bh, bw, bh, fill=color)
        text(s, bx - 0.1, y + h - 0.3 - bh - 0.19, bw + 0.2, 0.17, str(v) if v else '–', size=size + 0.4, color=NAVY, bold=True, align='c')
        text(s, bx - 0.15, y + h - 0.26, bw + 0.3, 0.2, lab, size=size, color=MUTED, align='c')
