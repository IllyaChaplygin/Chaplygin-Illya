# -*- coding: utf-8 -*-
"""Cost build-up of one SKU, read straight from the user's SelfCost workbook, so the deck can show
FOB -> duty -> import VAT -> transport -> credit -> SS and the reader can check it cell by cell."""
import os, sys
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..'))
from refresh_prices import LAYOUT  # noqa: E402

SRC = '/root/.claude/uploads/c53f34b4-ae43-52db-aeb2-01bdcda48cc3/7076a2bd-Self-Cost_Snacks..xlsx'
_wb = None


def wb():
    global _wb
    if _wb is None:
        _wb = openpyxl.load_workbook(SRC, data_only=True)
    return _wb


def breakdown(sheet, idx, scen="40'"):
    blocks, _ = LAYOUT[sheet]
    start = next(st for st, sc in blocks.items() if sc == scen)
    r = start + idx
    ws = wb()[sheet]
    units = ws['M%d' % r].value or ws['H%d' % r].value
    rate = ws['AJ%d' % r].value
    ah, f, x, y, ad, af = (ws[c + str(r)].value for c in ('AH', 'F', 'X', 'Y', 'AD', 'AF'))
    per = lambda v: v / ah * ws['AK%d' % r].value          # $ per unit
    parts = dict(fob=per(f), duty=per(x), vat=per(y), transport=per(ad), other=per(af), credit=ws['AL%d' % r].value)
    ss_usd, ss_uah = ws['AN%d' % r].value, ws['AO%d' % r].value
    return dict(row=r, name=ws['C%d' % r].value, rate=rate, fob_unit=ws['E%d' % r].value, parts=parts, ss_usd=ss_usd, ss_uah=ss_uah,
                check=sum(parts.values()))
