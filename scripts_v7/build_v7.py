# -*- coding: utf-8 -*-
"""V7: правки поверх файлу користувача (його видалення збережено). Нові слайди додаються в кінець, потім переставляються."""
import sys,os,re,math,json,statistics as st,collections as co
sys.path.insert(0,"/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad/v5")
import kit2
from kit2 import *
from pptx import Presentation
from pptx.util import Inches,Emu
from deckkit import _rect
UP="/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/b2d4af9e-SIAS_MARKET_RESEARCH_UA_V6.pptx"
prs=Presentation(UP)
N_USER=len(prs.slides._sldIdLst)
CREAM=RGBColor(0xFF,0xF6,0xE3); SHADOW=RGBColor(0xDF,0x91,0x00)
GREEN=RGBColor(0x4E,0x8F,0x6B); BLUE2=RGBColor(0x3E,0x59,0xA8); GREYB=RGBColor(0x9A,0xA1,0xB8)
BC={"Nongshim":NAVY,"Samyang":RED,"Paldo":PURPLE,"Ottogi":TEAL,"Bibigo":BLUE2,"O'Food":DARKAMBER,"Otoki":GREYB,"fire bull":GREYB,"Choi's":AMBER}
def rr(s,l,t,w,h,fill,rad=0.08):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    sh.line.fill.background(); sh.shadow.inherit=False; sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.adjustments[0]=rad; return sh
def oval(s,l,t,d,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(l),Inches(t),Inches(d),Inches(d))
    sh.line.fill.background(); sh.shadow.inherit=False; sh.fill.solid(); sh.fill.fore_color.rgb=fill; return sh
def photo(s,path,l,t,w,h,pad=0.06):
    from PIL import Image
    im=Image.open(path); ar=im.width/im.height; bw,bh=w-2*pad,h-2*pad
    if bw/bh>ar: ph=bh; pw=ph*ar
    else: pw=bw; ph=pw/ar
    s.shapes.add_picture(path,Inches(l+(w-pw)/2),Inches(t+(h-ph)/2),Inches(pw),Inches(ph))
def newslide(kicker,title,sub=None): return slide(prs,kicker,title,"00",sub)
def legend(s,x,y,items,gap=0.25,fs=8.5):
    for lab,col in items:
        _rect(s,x,y+0.02,0.16,0.14,col); w=0.075*len(lab)+0.1
        txt(s,x+0.22,y,w,0.18,lab,fs,False,NAVY); x+=0.22+w+gap
    return x
def dig(x,n=1): return f"{x:.{n}f}".replace('.',',')

# ============ ОБКЛАДИНКА ============
s=prs.slides.add_slide(prs.slide_layouts[6])
_rect(s,0,0,13.333,7.5,CREAM)
_rect(s,7.25,0,6.083,7.5,AMBER)
oval(s,7.5,0.7,5.6,GOLD); oval(s,6.6,4.7,2.5,AMBER); oval(s,10.4,4.6,2.8,GOLD)
_rect(s,0,7.22,13.333,0.28,NAVY)
rr(s,0.80,0.62,2.55,1.12,NAVY,0.12)
s.shapes.add_picture(LOGO,Inches(0.98),Inches(0.72),Inches(2.19),Inches(0.95))
txt(s,0.85,2.55,6.2,0.25,"ДОСЛІДЖЕННЯ РИНКУ ТА ПОЗИЦІОНУВАННЯ · ЖОВТЕНЬ 2026",11,True,DARKAMBER)
txt(s,0.82,2.95,6.4,2.0,"Корейський\nрамен\nв Україні",58,True,NAVY,lh=1.0)
_rect(s,0.88,6.05,1.6,0.09,AMBER)
cover_ph=[("08801073110502",7.55,0.55,2.55,2.45),("08801043150620",7.55,3.2,2.55,2.45),
          ("00648436310685",10.5,1.45,2.55,2.45),("03760344350847",10.5,4.15,2.55,2.45)]
for e,l,t,w,h in cover_ph:
    rr(s,l+0.06,t+0.07,w,h,SHADOW,0.06); rr(s,l,t,w,h,WHITE,0.06)
    photo(s,f"{SC}/img/{e}.jpg",l,t,w,h,0.10)
NEW={'cover':s}

# ============ ПОХОДЖЕННЯ: УСІ НАПРЯМКИ ============
TOT=838.480003
CT=[("В’єтнам",166.687505),("Корея",118.497356),("Румунія",86.919064),("Велика Британія",58.564844),
    ("Латвія",57.30452),("Чехія",50.800755),("Польща",48.761462)]
OTH=TOT-sum(v for _,v in CT)
CT=CT+[("Інші країни разом",OTH)]
CPAL=[NAVY,BLUE2,TEAL,PURPLE,DARKAMBER,RED,GREEN,GREYB]
s=newslide("МИТНА БАЗА · ПОХОДЖЕННЯ","Звідки їде локшина: усі напрямки ввозу","Січень–травень 2026 · вага нетто, тонн · частка в загальному ввозі 838,5 т")
txt(s,0.62,1.78,6,0.16,"СТРУКТУРА ВВОЗУ, %",8.4,True,NAVY)
x=0.62; W=12.10
for (nm,v),col in zip(CT,CPAL):
    w=W*v/TOT; _rect(s,x,2.02,w-0.02,0.50,col)
    if w>0.6: txt(s,x,2.17,w-0.02,0.2,f"{v/TOT*100:.1f}".replace('.',',')+" %",9.5,True,WHITE,C)
    x+=w
txt(s,0.62,2.72,6,0.16,"ПО КРАЇНАХ ПОХОДЖЕННЯ",8.4,True,NAVY)
y=3.02; RH=0.44
for (nm,v),col in zip(CT,CPAL):
    _rect(s,0.62,y+0.10,0.14,0.14,col)
    txt(s,0.88,y+0.06,2.5,0.24,nm,10.5,True,NAVY)
    bw=6.4*v/300
    _rect(s,3.5,y+0.04,max(bw,0.03),0.27,col)
    txt(s,3.5+bw+0.12,y+0.06,1.3,0.24,dig(v)+" т",10.5,True,NAVY)
    txt(s,11.4,y+0.06,1.3,0.24,dig(v/TOT*100)+" %",10.5,True,GREY,R)
    y+=RH
_rect(s,0.62,y+0.04,12.10,0.012,NAVY)
txt(s,0.88,y+0.12,2.5,0.24,"Усього",10.5,True,DARKAMBER); txt(s,3.5,y+0.12,2.0,0.24,dig(TOT)+" т",10.5,True,DARKAMBER); txt(s,11.4,y+0.12,1.3,0.24,"100 %",10.5,True,DARKAMBER,R)
txt(s,0.62,7.04,12.1,0.2,"Митна база користувача, код УКТ ЗЕД 1902 30 10 00. «Інші країни» — одним рядком: розбивки за країнами в базі немає.",7.4,False,PALE)
NEW['origin']=s

# ============ ЗВЕДЕННЯ БРЕНДИ × МЕРЕЖІ ============
SV_CH=[("auchan","Ашан"),("grono","Grono"),("onde","Onde"),("metro","METRO"),("zaraz","«Зараз»"),("kharkiv","«Клас»"),
       ("vostorg","«Восторг»"),("cosmos","«Космос»"),("tavriav","«Таврія»"),("novus","NOVUS"),("chudomarket","«ЧудоМ.»"),
       ("torba","«Торба»"),("silpo","«Сільпо»")]
SV_BR=["Nongshim","Samyang","Paldo","Ottogi","O'Food","Bibigo","Otoki","fire bull","Choi's"]
SILPO_P={"Samyang":[89.99,89.99,199.0],"Nongshim":[139.0]}
SILPO_BR={b:(len(v),st.median(v)) for b,v in SILPO_P.items()}
_EAN={r['ean']:r for r in SD_ROWS}
def _cell(b,ch):
    if ch=="silpo": return SILPO_BR.get(b,(0,None))
    v=[z for z in ZREC if z['chain']==ch and _EAN[z['ean']]['tm']==b]
    return (len({z['ean'] for z in v}),st.median([z['price'] for z in v]) if v else None)
SV=[[_cell(b,ch) for ch,_ in SV_CH] for b in SV_BR]
def _tot(b):
    rb=[r for r in SD_ROWS if r['tm']==b]
    chs={c for r in rb for c in r['chains']}|({"silpo"} if b in SILPO_BR else set())
    n=len(rb)+(SILPO_BR[b][0] if b in SILPO_BR else 0)
    ws=sorted(r['w'] for r in rb); pm=[r['pmed'] for r in rb if r.get('pmed')]
    wt=(f"{ws[0]:g} г" if ws[0]==ws[-1] else f"{ws[0]:g}–{ws[-1]:g} г").replace('.',',')
    return n,len(chs),(st.median(pm) if pm else None),wt
s=newslide("ЗВЕДЕННЯ · БРЕНДИ × МЕРЕЖІ","Усі бренди в усіх мережах")
# легенда: що означає клітинка
_rect(s,0.62,1.38,0.62,0.44,SEQT[4]); txt(s,0.62,1.42,0.62,0.2,"17",11,True,WHITE,C); txt(s,0.62,1.62,0.62,0.16,"89",8,False,WHITE,C)
txt(s,1.34,1.38,5.6,0.2,"верхнє число — кількість позицій (SKU) бренду в мережі",9,True,NAVY)
txt(s,1.34,1.60,5.6,0.2,"нижнє число — медіанна ціна пачки, грн  ·  чим темніше, тим більше позицій",9,False,GREY)
L0=0.62; RL=1.30; CW=0.60; T0=1.98; HH=0.42; RH=0.405
x_tail=L0+RL+CW*13+0.10; TWS=[0.62,0.70,0.82,0.98]
for j,(ch,nm) in enumerate(SV_CH):
    _rect(s,L0+RL+j*CW+0.02,T0,CW-0.04,HH-0.04,NAVY); txt(s,L0+RL+j*CW+0.02,T0+0.12,CW-0.04,0.18,nm,6.8,True,WHITE,C)
xt=x_tail
for hd,tw in zip(("SKU","МЕРЕЖ","МЕДІАНА, ГРН","ВАГА ПАЧКИ"),TWS):
    _rect(s,xt,T0,tw-0.03,HH-0.04,DARKAMBER); txt(s,xt,T0+0.12,tw-0.03,0.18,hd,6.6,True,WHITE,C); xt+=tw
mx=max(n for row in SV for n,_ in row); y=T0+HH
for i,b in enumerate(SV_BR):
    ours=(b=="Choi's")
    txt(s,L0,y+RH/2-0.09,RL-0.10,0.18,("Choi’s (SIAS)" if ours else b),8.8,True,(DARKAMBER if ours else NAVY),R)
    for j,(n,p) in enumerate(SV[i]):
        xx=L0+RL+j*CW; step=0 if n==0 else 1+int(round((len(SEQT)-2)*(n/mx)))
        _rect(s,xx+0.02,y+0.02,CW-0.04,RH-0.04,SEQT[step])
        if n:
            col=WHITE if step>=4 else NAVY
            txt(s,xx+0.02,y+0.035,CW-0.04,0.18,str(n),10.5,True,col,C); txt(s,xx+0.02,y+0.23,CW-0.04,0.14,f"{p:.0f}",7.8,False,col,C)
        else: txt(s,xx+0.02,y+0.12,CW-0.04,0.18,"—",9,False,PALE,C)
    n,nc,pm,wt=_tot(b); xt=x_tail
    for val,tw in zip((str(n),str(nc),f"{pm:.0f}" if pm else "—",wt),TWS):
        _rect(s,xt,y+0.02,tw-0.03,RH-0.04,LIGHT); txt(s,xt,y+0.12,tw-0.03,0.18,val,8.8,True,NAVY,C); xt+=tw
    y+=RH
_rect(s,L0+RL,y+0.04,CW*13+sum(TWS)+0.10,0.012,NAVY)
txt(s,L0,y+0.14,RL-0.10,0.18,"УСЬОГО",8.8,True,DARKAMBER,R)
for j,(ch,_) in enumerate(SV_CH):
    n=sum(SV[i][j][0] for i in range(len(SV_BR)))
    pr=[z['price'] for z in ZREC if z['chain']==ch] if ch!="silpo" else SILPO_P["Samyang"]+SILPO_P["Nongshim"]
    txt(s,L0+RL+j*CW,y+0.10,CW,0.18,str(n),10.5,True,DARKAMBER,C); txt(s,L0+RL+j*CW,y+0.30,CW,0.14,f"{st.median(pr):.0f}",7.8,False,GREY,C)
tn=len(SD_ROWS)+4; xt=x_tail
for val,tw in zip((str(tn),"13",f"{st.median([r['pmed'] for r in SD_ROWS]):.0f}",f"{min(r['w'] for r in SD_ROWS):g}–{max(r['w'] for r in SD_ROWS):g} г".replace('.',',')),TWS):
    txt(s,xt,y+0.20,tw-0.03,0.18,val,8.8,True,DARKAMBER,C); xt+=tw
slide_tail_sources(s,"Каталоги мереж (zakaz.ua) і «Сільпо», 29.09.2026. Клітинка = унікальні EAN × медіана цін по магазинах мережі. «Сільпо» — 4 позиції без EAN (3 Samyang, 1 Nongshim); 93 = 89 позицій з EAN + 4 «Сільпо». Вага — мін.–макс. пачки бренду.",7.04)
NEW['svod']=s

# ============ ФОРМАТ × ВАГА: ГРАФІК ============
B=F['bands']; MX=F['matrix']
s=newslide("ФОРМАТ × ВАГА","Де насправді щільна полиця","Кількість позицій і медіанна ціна пачки за ваговими смугами")
legend(s,0.62,1.74,[("Пакет",NAVY),("Стакан",PURPLE)])
txt(s,0.62,2.05,6.8,0.16,"ПОЗИЦІЙ У ВАГОВІЙ СМУЗІ",8.4,True,NAVY)
txt(s,8.35,2.05,3.4,0.16,"МЕДІАННА ЦІНА ПАЧКИ, ГРН",8.4,True,NAVY)
T=2.40; RH=0.62
_rect(s,0.62,T+RH*2-0.02,12.10,RH*2,AMBLIGHT)
for i,b in enumerate(B):
    y=T+i*RH; pk,cp=MX[0][i],MX[1][i]; unit=4.3/38
    txt(s,0.72,y+0.19,1.3,0.24,b['lab'],11,True,NAVY)
    x=2.05
    for v,col in ((pk,NAVY),(cp,PURPLE)):
        if v:
            w=unit*v; _rect(s,x,y+0.12,w,0.38,col)
            if w>0.28: txt(s,x,y+0.20,w,0.2,str(v),10,True,WHITE,C)
            x+=w
    txt(s,x+0.10,y+0.20,1.4,0.22,f"{pk+cp}  ·  {b['share']} %",10,True,NAVY)
    pw=3.05*b['pmed']/150; _rect(s,8.35,y+0.12,pw,0.38,DARKAMBER)
    txt(s,8.35,y+0.20,pw-0.12,0.22,f"{b['pmed']}",10.5,True,WHITE,R)
xl=8.35+3.05*120/150
_rect(s,xl-0.006,T-0.02,0.014,RH*5+0.04,RED); txt(s,xl-0.45,T+RH*5+0.04,0.9,0.16,"120 грн",8,True,RED,C)
txt(s,11.55,T+RH*2+0.45,1.15,0.4,"наші SKU:\n112,5–131 г",8.5,True,DARKAMBER,C,lh=1.1)
ecall(s,0.62,6.14,5.92,0.86,"ВАГА НАШИХ SKU — РИНКОВА",
 f"112,5–131 г потрапляє у дві найщільніші смуги: {B[2]['share']+B[3]['share']} % усіх позицій ринку.",TEAL,9.6,8.6)
ecall(s,6.80,6.14,5.92,0.86,"ЦІНА НАШИХ SKU ПРОТИ ЦИХ СМУГ",
 f"При полиці 120 грн наші SKU на {(120/B[2]['pmed']-1)*100:.0f} % дорожчі за медіану смуги 110–125 г ({B[2]['pmed']} грн), при поточних 140,55 грн — на {(140.55/B[2]['pmed']-1)*100:.0f} %.",RED,9.6,8.6)
NEW['fmt']=s

# ============ КАНАЛИ: ЦІНИ ============
CHP=co.defaultdict(dict)                       # мережа → {EAN: медіанна ціна по магазинах}
_tmp=co.defaultdict(list)
for z in ZREC: _tmp[(z['chain'],z['ean'])].append(z['price'])
for (c,e),v in _tmp.items(): CHP[c][e]=st.median(v)
CHP['silpo']={f"sp{i}":p for i,p in enumerate([89.99,89.99,199.0,139.0])}
_sh=json.load(open(f"{SC}/shops_raw.json",encoding='utf-8'))
SHOP={k:[r['price'] for r in _sh if r['shop']==k and r['price'] and r['price']<=250] for k in ("SushiPovar","AsiaFoods","TaiyakiMart")}
_h=open(f"{SC}/prom.html",encoding='utf-8',errors='ignore').read()
PROM=[p for p in (float(x) for x in re.findall(r'"price":\s*"?([\d\.]+)"?',_h)) if p<=250]
NAT_={'auchan','novus','metro','zaraz','silpo'}
CHANNELS=[("Національні мережі","5 мереж · Ашан, METRO, NOVUS…",[p for c in NAT_ for p in CHP[c].values()],NAVY),
          ("Регіональні мережі","8 мереж · Grono, Onde, «Космос»…",[p for c in CHP if c not in NAT_ for p in CHP[c].values()],BLUE2),
          ("Суші Повар","фудшоп · категорія «лапша рамен»",SHOP["SushiPovar"],TEAL),
          ("Asia Foods","фудшоп · рамен швидкого приготування",SHOP["AsiaFoods"],TEAL),
          ("Тайякі Март","фудшоп · корейський рамьон",SHOP["TaiyakiMart"],TEAL),
          ("Prom.ua","маркетплейс · пошук «рамен»",PROM,DARKAMBER)]
def q4(a): 
    q=st.quantiles(sorted(a),n=4); return q[0],q[1],q[2]
s=newslide("КАНАЛИ · ЦІНИ","Ціна пачки по каналах: де стоїть 120 грн","Мін.–макс. (лінія), 25–75 % пропозицій (блок), медіана (риска); ціна пачки, грн")
X0,X1=40,220; PX=3.35; PW=6.7
X=lambda v: PX+PW*(min(max(v,X0),X1)-X0)/(X1-X0)
T=2.35; RH=0.64
for i in range(0,len(CHANNELS),2): _rect(s,0.62,T+i*RH,12.10,RH-0.02,LIGHT)
for tv in range(40,221,20):
    _rect(s,X(tv),T,0.006,RH*len(CHANNELS),HAIR); txt(s,X(tv)-0.3,T+RH*len(CHANNELS)+0.05,0.6,0.14,str(tv),7.5,False,GREY,C)
txt(s,10.55,1.95,1.0,0.3,"ПРОПОЗИЦІЙ",7.4,True,GREY,R); txt(s,11.65,1.95,1.07,0.3,"ДО 120 ГРН",7.4,True,GREY,R)
for i,(nm,sub,pr,col) in enumerate(CHANNELS):
    y=T+i*RH; lo,hi=min(pr),max(pr); a,m,b_=q4(pr)
    txt(s,0.74,y+0.09,2.6,0.2,nm,10,True,NAVY); txt(s,0.74,y+0.31,2.6,0.2,sub,7.2,False,GREY)
    _rect(s,X(lo),y+RH/2-0.025,max(X(hi)-X(lo),0.03),0.05,tint(col,0.4))
    _rect(s,X(a),y+RH/2-0.12,max(X(b_)-X(a),0.03),0.24,tint(col,0.15))
    _rect(s,X(m)-0.02,y+RH/2-0.17,0.04,0.34,col)
    txt(s,X(m)-0.4,y+0.01,0.8,0.16,f"{m:.0f}",8.8,True,NAVY,C)
    txt(s,X(hi)+0.08,y+RH/2-0.09,0.6,0.18,f"{hi:.0f}",7.6,False,GREY)
    txt(s,10.55,y+0.20,1.0,0.22,str(len(pr)),9.6,False,GREY,R)
    txt(s,11.65,y+0.20,1.07,0.22,f"{sum(1 for p in pr if p<=120)/len(pr)*100:.0f} %",10.5,True,NAVY,R)
H=RH*len(CHANNELS)
_rect(s,X(120)-0.008,T-0.06,0.018,H+0.10,RED); txt(s,X(120)-0.5,T-0.24,1.0,0.16,"ціль 120",7.8,True,RED,C)
_rect(s,X(140.55)-0.008,T-0.06,0.018,H+0.10,DARKAMBER); txt(s,X(140.55)-0.2,T-0.24,1.0,0.16,"зараз 140,55",7.8,True,DARKAMBER,L)
slide_tail_sources(s,"Мережі: каталоги zakaz.ua і «Сільпо» 29.09.2026, медіана ціни позиції по магазинах мережі. Фудшопи й Prom.ua: сторінки категорій, жовтень 2026; без пропозицій дорожче 250 грн (мультипаки). Prom.ua — одна сторінка видачі, Rozetka, MAUDAU і Bigl цінових даних не дали.",7.04)
NEW['chA']=s

# ============ КАНАЛИ: ЦІНОВІ РІВНІ ============
ROWS=[]
for c,d in CHP.items():
    ROWS.append((NM.get(c,c),list(d.values())))
ROWS+= [("Суші Повар",SHOP["SushiPovar"]),("Asia Foods",SHOP["AsiaFoods"]),("Тайякі Март",SHOP["TaiyakiMart"]),("Prom.ua",PROM)]
TIERS=[("до 100 грн",lambda p:p<=100,TEAL),("100–120",lambda p:100<p<=120,tint(TEAL,0.55)),("120–140",lambda p:120<p<=140,tint(AMBER,0.45)),("понад 140 грн",lambda p:p>140,RED)]
def shares(pr): return [sum(1 for p in pr if f(p))/len(pr) for _,f,_ in TIERS]
ROWS=[(n,p,shares(p)) for n,p in ROWS if p]
ROWS.sort(key=lambda r:-(r[2][0]+r[2][1]))
s=newslide("КАНАЛИ · ЦІНОВІ РІВНІ","Частка пропозицій за ціною пачки в кожному каналі")
legend(s,0.62,1.46,[(l,c) for l,_,c in TIERS],gap=0.35,fs=9)
txt(s,10.4,1.46,2.32,0.18,"ДО 120 ГРН · ПОЗИЦІЙ",8,True,GREY,R)
T=1.92; RH=0.285
for i,(nm,pr,sh) in enumerate(ROWS):
    y=T+i*RH
    txt(s,0.62,y+0.04,2.1,0.2,nm,9.4,True,NAVY,R)
    x=2.9
    for v,(lab,_,col) in zip(sh,TIERS):
        w=7.2*v
        if w>0.005:
            _rect(s,x,y+0.03,w,RH-0.06,col)
            if w>0.42: txt(s,x,y+0.055,w,0.18,f"{v*100:.0f}",8.5,True,(WHITE if col in (TEAL,RED) else NAVY),C)
        x+=w
    txt(s,10.4,y+0.04,1.5,0.2,f"{(sh[0]+sh[1])*100:.0f} %",10,True,NAVY,R)
    txt(s,11.95,y+0.05,0.77,0.2,str(len(pr)),8.6,False,GREY,R)
slide_tail_sources(s,"Мережі — частка унікальних позицій (EAN) за медіанною ціною в мережі, 29.09.2026; «Сільпо» — 4 позиції. Фудшопи й Prom.ua — пропозиції зі сторінок категорій (до 250 грн), жовтень 2026. Числа в сегментах — % пропозицій.",7.04)
NEW['chB']=s

# ============ ЦІНА НА ПОЛИЦІ ЗАРАЗ ============
F33=M.forward_current("33"); F6=M.forward_current("6")
C_CC=RGBColor(0x46,0x53,0x82); C_BON=DARKAMBER; C_MAR=AMBER; C_MK=RGBColor(0xA9,0xB0,0xC8)
s=newslide("ФІНМОДЕЛЬ · ПОТОЧНА ПОЛИЦЯ","Ціна на полиці зараз проти цілі 120 грн")
legend(s,0.62,1.42,[("Собівартість (СС)",C_CC),("Бонус мережі 25 %",C_BON),("Наша маржа 30 %",C_MAR),("Націнка магазину 40 %",C_MK)],gap=0.4,fs=9)
BX=3.95; PWd=7.0; XM=225.0
bx=lambda v: BX+PWd*v/XM
def panel(y0,head,col,fc,labels):
    txt(s,0.62,y0,5,0.2,head,9.4,True,col)
    for k,(idx,(names,sub)) in enumerate(zip((2,0),labels)):
        y=y0+0.32+k*0.74; f=fc[idx]; p=f['partner']
        txt(s,0.62,y+0.02,3.25,0.2,names,9.6,True,NAVY); txt(s,0.62,y+0.24,3.25,0.2,sub,8,False,GREY)
        x=BX
        for v,c_,tc in ((f['cc'],C_CC,WHITE),(p*BONUS,C_BON,WHITE),(p*MARGIN,C_MAR,NAVY),(p*(MARKUP-1),C_MK,NAVY)):
            w=PWd*v/XM; _rect(s,x,y,w-0.015,0.46,c_)
            txt(s,x,y+0.12,w-0.015,0.22,f"{v:.0f}",10.5,True,tc,C); x+=w
        txt(s,bx(f['shelf'])+0.12,y+0.02,1.7,0.28,dig(f['shelf'],2)+" грн",16,True,RED)
        txt(s,bx(f['shelf'])+0.12,y+0.31,1.9,0.18,f"+{f['shelf']-120:.2f}".replace('.',',')+f" грн · +{(f['shelf']/120-1)*100:.0f} %",8.4,False,GREY)
BONUS,MARGIN,MARKUP=M.BONUS,M.MARGIN,M.MARKUP
LAB=[("Vegetable · Spicy · Beef · Chicken","112,5–113 г · закупівля SIAS €0,65"),("Gouda Cheese · Carbonara Spicy","123 і 131 г · закупівля SIAS €0,75")]
panel(1.92,"33 ПАЛЕТИ · ПОВНА МАШИНА",TEAL,F33,LAB)
panel(3.92,"6 ПАЛЕТ",RED,F6,LAB)
_rect(s,bx(120)-0.008,1.86,0.018,3.82,RED); txt(s,bx(120)-0.6,5.72,1.2,0.16,"ціль 120 грн",8.4,True,RED,C)
for tv in (0,50,100,150,200): txt(s,bx(tv)-0.3,5.52,0.6,0.14,str(tv),7.5,False,GREY,C)
_rect(s,0.62,6.12,12.10,0.62,LIGHT); _rect(s,0.62,6.12,0.045,0.62,NAVY)
txt(s,0.9,6.2,11.6,0.2,"Як рахується:  ціна партнеру = СС ÷ 0,45  (1 − бонус 25 % − маржа 30 %);   полиця = ціна партнеру × 1,40",9,True,NAVY)
txt(s,0.9,6.46,11.6,0.2,f"СС 33 палети: {dig(F33[2]['cc'],2)} і {dig(F33[0]['cc'],2)} грн;  6 палет: {dig(F6[2]['cc'],2)} і {dig(F6[0]['cc'],2)} грн (Self-Cost, ПДВ вже в ціні).  Курс 51,20 грн/€. Числа в смугах — грн за пачку.",8.4,False,GREY)
slide_tail_sources(s,"Фінмодель SIAS (вкладки «Финмодель 6 паллет» / «33 паллет», маржа 30 %); СС — Self-Cost_Sias.xlsx.",7.04)
NEW['shelf']=s

# ============ ПОПУЛЯРНІ УПАКОВКИ: ПОЗИЦІЇ ============
def _clean_name(r):
    t=r['title']; tm=r['tm']
    pat=re.compile(r"nong\s?shim" if tm=="Nongshim" else re.escape(tm).replace("'","['’]?"),re.I)
    m=pat.search(t); rest=t[m.end():] if m else t
    rest=re.sub(r"\s*\d+[.,]?\d*\s*(г|гр|g)\b\s*$","",rest.strip(),flags=re.I)
    rest=re.sub(r"швидкого приготування|локшина|рамьон|рамен","",rest,flags=re.I)
    rest=re.sub(r"\s+"," ",rest).strip(" ,.")
    return rest or "рамен"
def _short(x,n): return x if len(x)<=n else x[:n-1].rstrip()+"…"
PK=[r for r in SD_ROWS if r['seg'] in (1,2) and r.get('pmed')]
OURS={120:[("SIAS · Vegetable/Spicy/Beef/Chicken 113 г",140.55),("SIAS · Gouda Cheese 123 г",162.18)],
      130:[("SIAS · Carbonara Spicy 131 г",162.18)],140:[]}
def ladder_panel(l,t,w,wt,rows,ours,xmax=210,rowh=0.185,title=None):
    items=[dict(kind='m',p=r['pmed'],brand=r['tm'],name=_short(_clean_name(r),27),nch=r['nch']) for r in rows]
    items+=[dict(kind='o',p=120.0,name=n,now=nw) for n,nw in ours]
    items.sort(key=lambda d:(d['p'],d['kind']=='m'))
    pm=[r['pmed'] for r in rows]
    txt(s,l,t,w,0.2,f"ПАКЕТ {wt} Г · {len(rows)} ПОЗИЦІЙ",9.2,True,NAVY)
    txt(s,l,t+0.2,w,0.16,f"{min(pm):.0f}–{max(pm):.0f} грн · до 120 грн включно: {sum(1 for p in pm if p<=120)}",7.8,False,GREY)
    lab_w=2.8; bl=l+lab_w+0.05; bw=w-lab_w-0.05-0.62
    X=lambda v: bl+bw*v/xmax
    y=t+0.46; H=rowh*len(items)
    _rect(s,bl,y-0.02,0.006,H+0.04,HAIR); _rect(s,X(120)-0.007,y-0.05,0.015,H+0.08,RED)
    for d in items:
        if d['kind']=='m':
            txt(s,l,y+0.015,lab_w,0.16,f"{d['brand']} · {d['name']}",7.9,False,NAVY)
            _rect(s,bl,y+0.03,X(d['p'])-bl,rowh-0.07,BC.get(d['brand'],GREYB))
            txt(s,X(d['p'])+0.06,y+0.01,0.6,0.16,f"{d['p']:.0f}",8,True,NAVY)
        else:
            _rect(s,l,y-0.005,w,rowh,AMBLIGHT)
            txt(s,l+0.04,y+0.015,lab_w,0.16,d['name'],7.9,True,DARKAMBER)
            _rect(s,bl,y+0.03,X(120)-bl,rowh-0.07,AMBER); _rect(s,X(120),y+0.03,X(d['now'])-X(120),rowh-0.07,tint(AMBER,0.55))
            txt(s,X(d['now'])+0.06,y+0.01,1.3,0.16,f"120 → {dig(d['now'],2)}",8,True,DARKAMBER)
        y+=rowh
    for tv in (0,60,120,180): txt(s,X(tv)-0.25,y+0.04,0.5,0.14,str(tv),7,False,GREY,C)
    return y+0.2
s=newslide("КОНКУРЕНТИ · ПОПУЛЯРНІ ПАКЕТИ","Популярні пакети: кожна позиція та її ціна")
brands=[b for b in ("Nongshim","Samyang","Paldo","Ottogi","Bibigo","O'Food","Otoki") ]
x=0.62
for b in brands:
    _rect(s,x,1.44,0.16,0.14,BC[b]); txt(s,x+0.22,1.42,1.0,0.18,b,8.8,False,NAVY); x+=0.32+0.085*len(b)+0.25
_rect(s,x,1.44,0.16,0.14,AMBER); txt(s,x+0.22,1.42,2.3,0.18,"SIAS: ціль 120 → зараз",8.8,True,DARKAMBER)
byw=lambda w: [r for r in PK if round(r['w'])==w]
ladder_panel(0.62,1.80,5.95,120,byw(120),OURS[120])
yy=ladder_panel(6.82,1.80,5.90,130,byw(130),OURS[130])
ladder_panel(6.82,yy+0.02,5.90,140,byw(140),OURS[140])
slide_tail_sources(s,"Каталоги мереж 29.09.2026: медіанна ціна пачки по мережах, де позиція є. Червона лінія — 120 грн. SIAS: світла частина смуги — поточна полиця за фінмоделлю (33 палети).",7.06)
NEW['pop']=s

# ============ КОНКУРЕНТИ НИЖЧЕ 120 ============
KEY=[('08801073113428','Samyang Buldak 2× spicy','140 г'),('08801073110502','Samyang Buldak Hot Chicken','140 г'),
     ('08801073113381','Samyang Buldak тушкована курка','145 г'),('08801073116474','Samyang Buldak з сиром','130 г'),
     ('08801043150620','Nongshim Рамен зі спеціями','120 г'),('08801043157742','Nongshim Kimchi Ramyun','120 г'),
     ('08801043157728','Nongshim Chapaghetti','140 г'),('08801043018470','Nongshim Toomba Spicy & Creamy','137 г'),
     ('00648436310685','Paldo Volcano Carbonara','130 г')]
_pk_e={r['ean'] for r in SD_ROWS if r['seg'] in (1,2)}
krow=[]
for e,lab,wt in KEY:
    d={c:CHP[c][e] for c in CHP if e in CHP[c]}
    lo=min(d,key=d.get); hi=max(d,key=d.get)
    krow.append(dict(label=lab,wt=wt,lo=d[lo],lochain=NM.get(lo,lo),hi=d[hi],hichain=NM.get(hi,hi),n=len(d)))
krow.sort(key=lambda r:r['lo'])
# акції: найбільша знижка по позиції (є «стара ціна»)
pr={}
for z in ZREC:
    if z['ean'] in _pk_e and z.get('old_price') and z['price']<z['old_price']-0.005:
        dsc=1-z['price']/z['old_price']
        if z['ean'] not in pr or dsc>pr[z['ean']][0]: pr[z['ean']]=(dsc,z['price'],z['old_price'],z['chain'])
promo=sorted(pr.items(),key=lambda kv:-kv[1][0])[:5]
promo_rows=[dict(label="Samyang Buldak курка та сир · 140 г",chain="«Сільпо»",new=89.99,old=174.0),
            dict(label="Samyang Buldak карбонара · 130 г",chain="«Сільпо»",new=89.99,old=174.0)]
for e,(dsc,new,old,ch) in promo[:4]:
    r=_EAN[e]; promo_rows.append(dict(label=f"{r['tm']} · {_short(_clean_name(r),24)} · {r['w']:g} г".replace('.',','),chain=NM.get(ch,ch),new=new,old=old))
s=newslide("КОНКУРЕНТИ · НИЖЧЕ 120 ГРН","Ключові конкуренти продаються нижче 120 грн")
# ліва панель: діапазон цін по мережах
txt(s,0.62,1.52,7,0.18,"ДІАПАЗОН ЦІНИ ПАЧКИ ПО МЕРЕЖАХ, ГРН",8.8,True,NAVY)
legend(s,0.62,1.78,[("дешевше за ціль 120",TEAL),("дорожче за ціль",GREYB)],gap=0.3,fs=8.5)
LX=3.05; LW=3.55; X0,X1=60,200
X=lambda v: LX+LW*(v-X0)/(X1-X0)
T=2.28; RH=0.47
for tv in range(60,201,20):
    _rect(s,X(tv),T,0.006,RH*len(krow),HAIR); txt(s,X(tv)-0.3,T+RH*len(krow)+0.05,0.6,0.14,str(tv),7.5,False,GREY,C)
for i,r in enumerate(krow):
    y=T+i*RH
    txt(s,0.62,y+0.05,2.4,0.2,r['label'],8.8,True,NAVY); txt(s,0.62,y+0.25,2.4,0.16,f"{r['wt']} · {r['n']} мереж",7.4,False,GREY)
    if r['lo']<120: _rect(s,X(r['lo']),y+0.13,X(min(120,r['hi']))-X(r['lo']),0.2,TEAL)
    if r['hi']>120: _rect(s,X(max(120,r['lo'])),y+0.13,X(r['hi'])-X(max(120,r['lo'])),0.2,GREYB)
    txt(s,X(r['lo'])-0.95,y+0.02,0.9,0.16,f"{r['lo']:.0f}",8.6,True,(TEAL if r['lo']<120 else NAVY),R)
    txt(s,X(r['lo'])-1.25,y+0.2,1.2,0.15,r['lochain'],7,False,GREY,R)
    txt(s,X(r['hi'])+0.06,y+0.02,0.8,0.16,f"{r['hi']:.0f}",8.6,True,NAVY)
    txt(s,X(r['hi'])+0.06,y+0.2,1.1,0.15,r['hichain'],7,False,GREY)
_rect(s,X(120)-0.008,T-0.05,0.018,RH*len(krow)+0.08,RED); txt(s,X(120)-0.5,T-0.24,1.0,0.16,"ціль 120",8,True,RED,C)
_rect(s,X(140.55)-0.008,T-0.05,0.018,RH*len(krow)+0.08,DARKAMBER); txt(s,X(140.55)-0.1,T-0.24,1.2,0.16,"зараз 140,55",8,True,DARKAMBER,L)
# права панель: акції
RX=8.55
txt(s,RX,1.52,4.2,0.18,"АКЦІЇ: ЗВИЧАЙНА ЦІНА → АКЦІЙНА, ГРН",8.8,True,NAVY)
legend(s,RX,1.78,[("звичайна",GREYB),("акційна",AMBER)],gap=0.3,fs=8.5)
PH_=4.0; PXM=190
y=2.28
for r in promo_rows:
    txt(s,RX,y,4.2,0.18,r['label'],8.4,True,NAVY)
    _rect(s,RX,y+0.22,PH_*r['old']/PXM,0.15,GREYB); txt(s,RX+PH_*r['old']/PXM+0.06,y+0.2,0.6,0.16,f"{r['old']:.0f}",8,False,GREY)
    _rect(s,RX,y+0.40,PH_*r['new']/PXM,0.15,AMBER); txt(s,RX+PH_*r['new']/PXM+0.06,y+0.38,1.6,0.16,f"{r['new']:.0f}  (−{(1-r['new']/r['old'])*100:.0f} %) · {r['chain']}",8.2,True,DARKAMBER)
    y+=0.70
slide_tail_sources(s,"Каталоги мереж 29.09.2026: ціна позиції в мережі (медіана по магазинах); «стара ціна» віддається не всіма мережами (здебільшого «Ашан», «Сільпо»), тому акцій насправді більше. Червона лінія — 120 грн.",7.04)
NEW['comp']=s

# ============ ПОТРІБНА ЗАКУПІВЛЯ: ВОДОСПАД ============
def wf(sc):
    n=M.units(sc); mcc=M.max_cc_uah()/M.RATE; v1=mcc/(1+M.FIN); b=M.batch_costs(sc)/n; v2=v1-b; v3=v2/((1+M.DUTY)*(1+M.VAT))
    return [("Макс. СС\nпачки",mcc,'start'),("Фінан-\nсування\n2,5 %",-(mcc-v1),'d'),("Витрати\nпартії на\nпачку",-b,'d'),("ПДВ 20 %",-(v2-v3),'d'),("Макс.\nзакупівля",v3,'end')]
s=newslide("ФІНМОДЕЛЬ · ПОТРІБНА ЗАКУПІВЛЯ","Яка закупівля потрібна для полиці 120 грн")
def waterfall(l,t,w,h,sc,head,col):
    steps=wf(sc); ask=M.ask_price(sc)
    txt(s,l,t,w-2.0,0.2,head,9.6,True,col)
    rr(s,l+w-1.85,t-0.06,1.85,0.5,col,0.12)
    txt(s,l+w-1.85,t-0.02,1.1,0.4,"запит\nSIAS",8,True,WHITE,C,lh=1.0); txt(s,l+w-0.8,t+0.02,0.8,0.36,"€"+f"{ask:.2f}".replace('.',','),17,True,WHITE,C)
    top=t+0.85; ph=h-1.65; ymax=0.80
    Y=lambda v: top+ph*(1-v/ymax)
    bw=0.78; gap=(w-0.3-5*bw)/4; x=l+0.1; cur=0
    for v,lab in ((0.75,"зараз €0,75 (Gouda, Carbonara)"),(0.65,"зараз €0,65 (Vegetable, Spicy, Beef, Chicken)")):
        _rect(s,l,Y(v),w,0.012,DARKAMBER); txt(s,l,Y(v)-0.19,w,0.16,lab,7.8,True,DARKAMBER,R)
    _rect(s,l,Y(0),w,0.01,NAVY)
    for lab,v,kind in steps:
        if kind=='start': y0,y1=Y(v),Y(0); cur=v; c=C_CC
        elif kind=='d': y0,y1=Y(cur),Y(cur+v); cur+=v; c=GREYB
        else: y0,y1=Y(v),Y(0); c=col
        _rect(s,x,min(y0,y1),bw,abs(y1-y0),c)
        sv=f"{abs(v):.3f}".replace('.',',')
        txt(s,x-0.2,max(min(y0,y1),top-0.1)+(-0.2 if kind!='d' else abs(y1-y0)+0.03),bw+0.4,0.18,(f"−€{sv}" if kind=='d' else f"€{sv}"),9.4,True,NAVY,C)
        txt(s,x-0.15,Y(0)+0.06,bw+0.3,0.5,lab,8,False,GREY,C,lh=1.05)
        x+=bw+gap
waterfall(0.62,1.45,5.95,4.45,"33","33 ПАЛЕТИ · ПОВНА МАШИНА · 52 800 ПАЧОК",TEAL)
waterfall(6.77,1.45,5.95,4.45,"6","6 ПАЛЕТ · 9 600 ПАЧОК",RED)
# знижка по позиціях
A33=M.ask_price("33"); A6=M.ask_price("6")
txt(s,0.62,5.95,6,0.18,"ЗНИЖЕННЯ ВІД ПОТОЧНОЇ ЗАКУПІВЛІ",8.4,True,NAVY)
for k,(names,cur) in enumerate((("Vegetable · Spicy · Beef · Chicken",0.65),("Gouda Cheese · Carbonara Spicy",0.75))):
    y=6.22+k*0.40
    txt(s,0.62,y+0.05,3.2,0.22,names,9.4,True,NAVY)
    txt(s,3.9,y+0.05,1.5,0.22,f"зараз €{cur:.2f}".replace('.',','),9.4,False,GREY)
    rr(s,5.5,y,3.4,0.32,TEAL,0.2); txt(s,5.5,y+0.06,3.4,0.2,f"33 палети: €{A33:.2f} (−{(1-A33/cur)*100:.0f} %)".replace('.',','),9.4,True,WHITE,C)
    rr(s,9.1,y,3.4,0.32,RED,0.2); txt(s,9.1,y+0.06,3.4,0.2,f"6 палет: €{A6:.2f} (−{(1-A6/cur)*100:.0f} %)".replace('.',','),9.4,True,WHITE,C)
slide_tail_sources(s,"120 грн ÷ 1,40 = 85,71 грн ціна партнеру → × 45 % = 38,57 грн макс. СС = €0,753 (курс 51,20). Бонус 25 %, маржа 30 %, ПДВ 20 %, мито 0 %. Запит округлено вниз до євроцента; згоди SIAS немає.",7.1)
NEW['need']=s

# ============ ЗБІРКА: порядок, видалення замінених, нумерація ============
lst=prs.slides._sldIdLst; ids=list(lst)
U=lambda k: ids[k-1]
newids={k:ids[N_USER+i] for i,k in enumerate(['cover','origin','svod','fmt','chA','chB','shelf','pop','comp','need'])}
assert len(ids)==N_USER+10,len(ids)
order=[newids['cover'],U(2),newids['origin']]+[U(k) for k in range(4,18)]+[newids['svod'],newids['fmt'],U(20),U(21),newids['chA'],newids['chB'],newids['shelf'],U(23),newids['pop'],newids['comp'],newids['need'],U(27)]
drop=[U(1),U(3),U(18),U(19),U(22),U(24),U(25),U(26)]
for e in drop: prs.part.drop_rel(e.rId)
for e in ids: lst.remove(e)
for e in order: lst.append(e)
# нумерація
for i,sl in enumerate(prs.slides,1):
    for sh in sl.shapes:
        if sh.has_text_frame and re.fullmatch(r"\d\d",sh.text_frame.text.strip()) and abs(Emu(sh.left).inches-11.55)<0.05:
            r=sh.text_frame.paragraphs[0].runs[0]; r.text=f"{i:02d}"
OUT=os.environ.get("OUT","/home/user/Chaplygin-Illya/SIAS_MARKET_RESEARCH_UA_V7.pptx")
prs.save(OUT); print("saved",OUT,len(prs.slides._sldIdLst))
