import openpyxl,math,sys,zipfile
from lxml import etree
F=sys.argv[1]
nv=openpyxl.load_workbook(F,data_only=True); nf=openpyxl.load_workbook(F)
z=zipfile.ZipFile(F); assert z.testzip() is None
for n in z.namelist():
    if n.endswith('.xml') or n.endswith('.rels'): etree.fromstring(z.read(n))
errs=[(ws.title,c.coordinate,nv[ws.title][c.coordinate].value) for ws in nf.worksheets for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('=') and (nv[ws.title][c.coordinate].value is None or (isinstance(nv[ws.title][c.coordinate].value,str) and nv[ws.title][c.coordinate].value.startswith('#')))]
print('formula errors/empty:',errs)
fail=0
def chk(n,a,b,tol=1e-9):
    global fail
    ok=abs(a-b)<=tol*max(1,abs(b)); fail+=0 if ok else 1
    if not ok: print('FAIL',n,a,b)
RATE=51.2;MARK=1.4;FIN=.025;VAT=.2;SHELF=120
def ask(freight,pallets,b,m,x):
    P=SHELF/MARK; n=pallets*1600; batch=freight+288+42+freight*0.7*VAT
    return ((P*(1-b-m)-x)/RATE/(1+FIN)-batch/n)/(1+VAT)
def cc(freight,pallets,p,x):
    n=pallets*1600; prod=p*RATE; logi=(freight+330)/n*RATE; vat=(p+freight*0.7/n)*VAT*RATE
    return (prod+logi+vat)*(1+FIN)+x
S=nv['Сценарии']
for r in range(8,40):
    nm=S.cell(r,1).value; b=S.cell(r,2).value
    if not isinstance(b,(int,float)): continue
    m=S.cell(r,3).value; x=S.cell(r,4).value
    for (cm,ca,fr,pal) in ((9,10,3700,33),(12,13,2600,6)):
        mp=ask(fr,pal,b,m,x); chk(f'sc r{r} max{pal}',S.cell(r,cm).value,mp); chk(f'sc r{r} ask{pal}',S.cell(r,ca).value,math.floor(mp*100+1e-9)/100)
    print(f'{nm[:62]:62s} b{b:.2f} m{m:.2f} X{x:5.2f} | 33p {S.cell(r,10).value:.2f} ({S.cell(r,11).value:+.0%})  6p {S.cell(r,13).value:.2f} ({S.cell(r,14).value:+.0%})')
for sheet,pal,fr,ar in (('Цена 120 грн — 33 паллет',33,3700,19),('Цена 120 грн — 6 паллет',6,2600,20)):
    w=nv[sheet]; b=w['D14'].value; m=w['C14'].value; x=w['G14'].value
    mp=ask(fr,pal,b,m,x); a=math.floor(mp*100+1e-9)/100
    chk(sheet+' ask',w[f'B{ar}'].value,a); chk(sheet+' max',w[f'B{ar+1}'].value,mp); chk(sheet+' CCafter',w[f'B{ar+2}'].value,cc(fr,pal,a,x))
    print(sheet,'b,m,x',b,m,x,'ask',w[f'B{ar}'].value,'| E3',repr(w['E3'].value),'| F3',repr(w['F3'].value),'| C3',repr(w['C3'].value))
    for r,pur in ((4,0.65),(6,0.75),(8,a)):
        C=w[f'C{r}'].value; D=w[f'D{r}'].value; E=w[f'E{r}'].value; Fp=w[f'F{r}'].value; G=w[f'G{r}'].value; H=w[f'H{r}'].value
        chk('partner',D,C/(1-b-m)); chk('bonus',E,D*b); chk('profit',Fp,D-C-E); chk('profit%',Fp/D,m); chk('shelf',H,D*MARK)
        print(f'  r{r}: buy {w[f"B{r}"].value:.2f}₴  CC {C:.2f}  partner {D:.2f}  bonus {E:.2f}  profit {Fp:.2f}  markup {G:.2f}  shelf {H:.2f}')
for n in ('Финмодель 6 паллет','Финмодель 33 паллет'):
    w=nv[n]; print(n,w['D4'].value,w['D5'].value,w['D8'].value,w['E4'].value,'|',w['E8'].value,'| K11',round(w['K11'].value,2),'K13',round(w['K13'].value,2),'T13',round(w['T13'].value,2))
print('FAILS',fail)
