import openpyxl,sys
F,SRC,LANG=sys.argv[1],sys.argv[2],sys.argv[3]
N=openpyxl.load_workbook(F,data_only=True); NF=openpyxl.load_workbook(F); O=openpyxl.load_workbook(SRC,data_only=True)
if LANG=='ua': PAIRS=(('Финмодель 6 паллет','Цена 120 грн — 6 паллет',20),('Финмодель 33 паллет','Цена 120 грн — 33 паллет',19))
else: PAIRS=(('Model 6 pallets','Price 120 UAH - 6 pallets',20),('Model 33 pallets','Price 120 UAH - 33 pallets',19))
fails=0
def chk(n,a,b,tol=1e-9):
    global fails
    if a is None or abs(a-b)>tol*max(1,abs(b)): fails+=1; print('FAIL',n,a,b)
bad=[(ws.title,c.coordinate) for ws in NF for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('=') and (N[ws.title][c.coordinate].value is None or str(N[ws.title][c.coordinate].value).startswith('#'))]
print('formula errors',bad)
for mn,pn,ar in PAIRS:
    m=N[mn]; p=N[pn]; b,mg,mk=m['D4'].value,m['D5'].value,m['D6'].value
    chk(mn+' margin link',mg,p['C14'].value); chk(mn+' bonus',b,p['D14'].value)
    for k,r in enumerate(range(11,17)):
        o=46+k; two=r in (11,12)
        chk(f'{mn} before buy r{o}',m[f'E{o}'].value,p['B7'].value if two else p['B5'].value)
        chk(f'{mn} before shelf r{o}',m[f'J{o}'].value,p['H6'].value if two else p['H4'].value)
        chk(f'{mn} before cost r{o}',m[f'F{o}'].value,p['C6'].value if two else p['C4'].value)
        chk(f'{mn} before bonus r{o}',m[f'H{o}'].value,p['E6'].value if two else p['E4'].value)
        chk(f'{mn} before left r{o}',m[f'I{o}'].value,p['F6'].value if two else p['F4'].value)
        chk(f'{mn} after buy r{o}',m[f'K{o}'].value,p[f'B{ar}'].value)
        for col,pc in (('L','C8'),('M','D8'),('N','E8'),('O','F8'),('P','H8')): chk(f'{mn} after {col} r{o}',m[f'{col}{o}'].value,p[pc].value)
        chk(f'{mn} d shelf r{o}',m[f'Q{o}'].value,m[f'P{o}'].value-m[f'J{o}'].value); chk(f'{mn} d pct r{o}',m[f'R{o}'].value,m[f'Q{o}'].value/m[f'J{o}'].value)
        chk(f'{mn} main table shelf = before block r{o}',m[f'K{r}'].value,m[f'J{o}'].value)
    print(mn,'| before shelf 2SKU/4SKU',round(m['J46'].value,2),round(m['J48'].value,2),'| after',round(m['P46'].value,2),'| d%',round(m['R46'].value,3),round(m['R48'].value,3),'| main-table margin',round(mg,4))
# прочие листы не изменились (кроме D5/E4/E5/E8 на листах модели)
d=0
for ws in O:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,(int,float)):
                v=N[ws.title][c.coordinate].value
                if not isinstance(v,(int,float)) or abs(v-c.value)>1e-9*max(1,abs(c.value)):
                    d+=1
                    if d<8: print('changed',ws.title,c.coordinate,round(c.value,4),'->',round(v,4) if isinstance(v,(int,float)) else v)
print('changed numeric cells vs source:',d,'| FAILS',fails)
