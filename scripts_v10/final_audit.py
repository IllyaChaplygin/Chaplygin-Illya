import openpyxl,math,zipfile,re,sys
from lxml import etree
F='/home/user/Chaplygin-Illya/RAMEN_FINMODEL_Final_NEW_STRUCTURE.xlsx'
SCF='/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/87c2996d-Self-Cost_Sias.xlsx'
nv=openpyxl.load_workbook(F,data_only=True); nf=openpyxl.load_workbook(F)
SC=openpyxl.load_workbook(SCF,data_only=True).worksheets[0]
fail=[0]
def chk(n,a,b,tol=1e-9):
    if a is None or b is None or abs(a-b)>tol*max(1,abs(b)): fail[0]+=1; print('FAIL',n,a,b)
# 0. package integrity
z=zipfile.ZipFile(F); assert z.testzip() is None
for n in z.namelist():
    if n.endswith('.xml') or n.endswith('.rels'): etree.fromstring(z.read(n))
# number formats with non-ascii literals
for ws in nf:
    for row in ws.iter_rows():
        for c in row:
            if c.number_format and re.search(r'[^\x00-\x7f]',c.number_format.replace('₴','').replace('€','').replace('×','')): print('NONASCII FMT',ws.title,c.coordinate,c.number_format)
# 1. cached value exists for every formula, no errors
bad=[(ws.title,c.coordinate) for ws in nf for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('=') and (nv[ws.title][c.coordinate].value is None or str(nv[ws.title][c.coordinate].value).startswith('#'))]
print('formulas w/o value or error:',bad)
# 2. tabs 1–2 independent recompute from Self-Cost
RATE=51.2
for name,rows,pal in (('Финмодель 6 паллет',(6,11,12,13,14,15,16),6),('Финмодель 33 паллет',(30,31,32,33,34,35),33)):
    w=nv[name]; bon,mar,mk=w['D4'].value,w['D5'].value,w['D6'].value
    chk(name+' rate',w['D3'].value,RATE); chk(name+' D8',w['D8'].value,1-bon-mar)
    # SC rows: 6p: 6..11 ; 33p: 30..35 (order Gouda,Carbonara,Veg,Spicy,Beef,Chicken)
    scrows=range(6,12) if pal==6 else range(30,36)
    for k,r in enumerate(scrows):
        row=11+k
        ccE=SC[f'AJ{r}'].value
        chk(f'{name} r{row} E',w[f'E{row}'].value,ccE,1e-9)
        F_=ccE*RATE; G=F_/(1-bon-mar); H=G*bon; I=G-F_-H
        for col,exp in (('F',F_),('G',G),('H',H),('I',I),('J',I/G),('K',G*mk),('L',F_/(1-bon)),('M',F_/(1-bon)*mk),('N',G-F_),('O',G-F_-H)):
            chk(f'{name} {col}{row}',w[f'{col}{row}'].value,exp)
        # scenario blocks
        for c0,mg in (('P',.30),('U',.25),('Z',.20),('AE',.15)):
            ci=openpyxl.utils.column_index_from_string(c0)
            P_=F_/(1-bon-mg); 
            exp=[P_,P_*bon,P_-F_-P_*bon,(P_-F_-P_*bon)/P_,P_*mk]
            for d,e in enumerate(exp): chk(f'{name} {c0}+{d} r{row}',w.cell(row,ci+d).value,e)
    # Self-Cost vs inputs: purchase price, freight
print('tabs1-2 recompute done, fails so far',fail[0])
# 3. tabs 3–4 + scenarios independent
SHELF=120;MK=1.4;FIN=.025;VAT=.2
def batch(fr): return fr+288+42+fr*0.7*VAT
def ask_mp(fr,pal,b,m,x): n=pal*1600; return (((SHELF/MK)*(1-b-m)-x)/RATE/(1+FIN)-batch(fr)/n)/(1+VAT)
def ccf(fr,pal,p,x): n=pal*1600; return ((p*RATE+(fr+330)/n*RATE+(p+fr*.7/n)*VAT*RATE))*(1+FIN)+x
for sheet,pal,fr,ar,scr in (('Цена 120 грн — 33 паллет',33,3700,19,(32,30)),('Цена 120 грн — 6 паллет',6,2600,20,(8,6))):
    w=nv[sheet]; b,m,x=w['D14'].value,w['C14'].value,w['G14'].value
    chk(sheet+' B14',w['B14'].value,120); chk(sheet+' F14',w['F14'].value,RATE); chk(sheet+' E14',w['E14'].value,MK)
    assert (w['B16'].value,w['C16'].value,w['D16'].value,w['E16'].value,w['F16'].value,w['G16'].value,w['B18'].value,w['C18'].value,w['D18'].value)==(.2,.025,0,fr,288,42,.7,pal,1600),sheet
    mp=ask_mp(fr,pal,b,m,x); a=math.floor(mp*100+1e-9)/100
    chk(sheet+' ask',w[f'B{ar}'].value,a); chk(sheet+' max',w[f'B{ar+1}'].value,mp); chk(sheet+' ccafter',w[f'B{ar+2}'].value,ccf(fr,pal,a,x))
    # СС до — из Self-Cost напрямую
    cc4=SC[f'AJ{scr[0]}'].value*RATE+x; cc2=SC[f'AJ{scr[1]}'].value*RATE+x
    for r,buy,cc in ((4,0.65,cc4),(6,0.75,cc2),(8,a,ccf(fr,pal,a,x))):
        P_=cc/(1-b-m)
        for col,e in (('B',buy*RATE),('C',cc),('D',P_),('E',P_*b),('F',P_-cc-P_*b),('G',P_*(MK-1)),('H',P_*MK)):
            chk(f'{sheet} {col}{r}',w[f'{col}{r}'].value,e); chk(f'{sheet} {col}{r+1} eur',w[f'{col}{r+1}'].value,e/RATE)
    assert abs(SC[f'J{scr[0]}'].value-0.65)<1e-12 and abs(SC[f'J{scr[1]}'].value-0.75)<1e-12
    for col in 'BCDEFGH': chk(f'{sheet} Δ{col}',w[f'{col}10'].value,w[f'{col}8'].value-w[f'{col}4'].value)
    print(sheet,'b,m,x=',b,m,x,'| ask',w[f'B{ar}'].value,'| shelf до',round(w['H4'].value,2),round(w['H6'].value,2),'после',round(w['H8'].value,2),'| ПОСЛЕ прибыль/партнёру',round(w['F8'].value/w['D8'].value,4),'бонус/партнёру',round(w['E8'].value/w['D8'].value,4))
S=nv['Сценарии']; cnt=0
for r in range(8,40):
    b=S.cell(r,2).value
    if not isinstance(b,(int,float)): continue
    m,x=S.cell(r,3).value,S.cell(r,4).value; cnt+=1
    P_=120/MK
    chk(f'sc{r} bonus₴',S.cell(r,5).value,P_*b); chk(f'sc{r} profit₴',S.cell(r,6).value,P_*m); chk(f'sc{r} budget',S.cell(r,7).value,P_*(1-b-m)); chk(f'sc{r} cc goods',S.cell(r,8).value,P_*(1-b-m)-x)
    for cm,ca,fr,pal in ((9,10,3700,33),(12,13,2600,6)):
        mp=ask_mp(fr,pal,b,m,x); chk(f'sc{r} max{pal}',S.cell(r,cm).value,mp); a=math.floor(mp*100+1e-9)/100; chk(f'sc{r} ask{pal}',S.cell(r,ca).value,a)
        chk(f'sc{r} chg{pal}',S.cell(r,ca+1).value,a/0.65-1)
        # обратная проверка: при запросе a полка не выше 120
        sh=(ccf(fr,pal,a,x)+0)/(1-b-m)*MK  # СС уже включает x
        assert sh<=120+1e-6,(r,pal,sh)
print('scenario rows checked',cnt,'| TOTAL FAILS',fail[0])
