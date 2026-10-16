# -*- coding: utf-8 -*-
import openpyxl,copy,sys
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.workbook.properties import CalcProperties
SRC,OUT,LANG=sys.argv[1],sys.argv[2],sys.argv[3]
wb=openpyxl.load_workbook(SRC)
EN=(LANG=='en')
if EN: PAIRS=(('Model 6 pallets','Price 120 UAH - 6 pallets',6),('Model 33 pallets','Price 120 UAH - 33 pallets',33))
else:  PAIRS=(('Финмодель 6 паллет','Цена 120 грн — 6 паллет',6),('Финмодель 33 паллет','Цена 120 грн — 33 паллет',33))
CONST={6:dict(fr=2600,n=9600),33:dict(fr=3700,n=52800)}
MARGINS=[0.35,0.30,0.25,0.20,0.15,0.10]
T=dict(
 en=dict(title='MARGIN SCENARIOS · BEFORE vs AFTER the SIAS request',
   sub='For each margin: shelf price at the current purchase price (BEFORE) and the purchase price SIAS must accept so that the shelf price is 120 ₴ (AFTER). Chain bonus and markup as on the "Price 120 UAH" sheet.',
   g1='BEFORE · current purchase price',g2='AFTER · request for a 120 ₴ shelf',g3='CHANGE vs BEFORE',
   h=['Margin','','Share of price for cost','Shelf 4 SKUs, ₴','Shelf 2 SKUs, ₴','Request, €/pack','Landed cost, ₴','Shelf, ₴','Δ 4 SKUs, ₴','Δ 4 SKUs, %','Δ 2 SKUs, ₴','Δ 2 SKUs, %'],cur='◄ ≈ current'),
 ua=dict(title='СЦЕНАРІЇ МАРЖІ · ДО та ПІСЛЯ запиту SIAS',
   sub='Для кожної маржі: полиця при поточній закупівлі (ДО) і закупівля, яку має прийняти SIAS, щоб полиця була 120 ₴ (ПІСЛЯ). Бонус мережі та націнка — як на аркуші «Цена 120 грн».',
   g1='ДО · поточна закупівля',g2='ПІСЛЯ · запит для полиці 120 ₴',g3='ЗМІНА відносно ДО',
   h=['Маржа','','Частка ціни на СС','Полиця 4 SKU, ₴','Полиця 2 SKU, ₴','Запит, €/пач.','СС, ₴','Полиця, ₴','Δ 4 SKU, ₴','Δ 4 SKU, %','Δ 2 SKU, ₴','Δ 2 SKU, %'],cur='◄ поточна'))['en' if EN else 'ua']
thin=Side(style='thin',color='C9D3E0'); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
def st(c,sz=10,b=False,col='FF222222',fill=None,fmt=None,al='center'):
    c.font=Font(name='Arial',sz=sz,b=b,color=col)
    if fill: c.fill=PatternFill('solid',fgColor=fill)
    if fmt: c.number_format=fmt
    c.alignment=Alignment(horizontal=al,vertical='center',wrap_text=True); c.border=BR
UAH='#,##0.00" ₴"'; DUAH='+#,##0.00" ₴";-#,##0.00" ₴"'; DPCT='+0.0%;-0.0%'
def grid(ws,r0,cols,mn,pn,pal,price_tab):
    """cols: dict name->column letter. Формулы ссылаются на лист модели (mn) и лист цен (pn)."""
    M=f"'{mn}'!"; P=f"'{pn}'!"; c=CONST[pal]; fr=c['fr']; n=c['n']
    u=f"{P}$G$14" if EN else "0"
    vat,fin,dut,br,ot,sh=0.2,0.025,0,288,42,0.7
    if EN:
        batch=f"({P}$E$16+{P}$F$16+{P}$G$16+{P}$E$16*{P}$B$18*({P}$D$16+{P}$B$16*(1+{P}$D$16)))/({P}$C$18*{P}$D$18)"
        logi=f"({P}$E$16+{P}$F$16+{P}$G$16)/({P}$C$18*{P}$D$18)"; frs=f"{P}$E$16*{P}$B$18/({P}$C$18*{P}$D$18)"
        VAT,FIN,DUT=f"{P}$B$16",f"{P}$C$16",f"{P}$D$16"
    else:
        batch=f"({fr}+{br}+{ot}+{fr}*{sh}*({dut}+{vat}*(1+{dut})))/{n}"
        logi=f"({fr}+{br}+{ot})/{n}"; frs=f"{fr}*{sh}/{n}"; VAT,FIN,DUT=str(vat),str(fin),str(dut)
    ws.merge_cells(f"{cols['m']}{r0}:{cols['dp2']}{r0}"); c0=ws[f"{cols['m']}{r0}"]; c0.value=T['title']; c0.font=Font(name='Arial',sz=14,b=True,color='FF1F3864'); c0.alignment=Alignment(horizontal='left',vertical='center'); ws.row_dimensions[r0].height=26
    ws.merge_cells(f"{cols['m']}{r0+1}:{cols['dp2']}{r0+1}"); c1=ws[f"{cols['m']}{r0+1}"]; c1.value=T['sub']; c1.font=Font(name='Arial',sz=9,color='FF5E6D7E'); c1.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); ws.row_dimensions[r0+1].height=30
    g=r0+2
    for a,b_,txt,fill in ((cols['b4'],cols['b2'],T['g1'],'FF203864'),(cols['ask'],cols['sa'],T['g2'],'FFB07A12'),(cols['d4'],cols['dp2'],T['g3'],'FFA33A3A')):
        ws.merge_cells(f"{a}{g}:{b_}{g}")
        for cc in range(openpyxl.utils.column_index_from_string(a),openpyxl.utils.column_index_from_string(b_)+1): st(ws.cell(g,cc),10,True,'FFFFFFFF',fill)
        ws[f"{a}{g}"].value=txt
    ws.row_dimensions[g].height=24
    h=g+1; keys=['m','cur','share','b4','b2','ask','lc','sa','d4','dp4','d2','dp2']
    for k,t in zip(keys,T['h']):
        i=openpyxl.utils.column_index_from_string(cols[k]); fill='FF44546A' if k in('m','cur','share') else ('FF2F75B5' if k in('b4','b2') else ('FFD9A441' if k in('ask','lc','sa') else 'FFC26A6A'))
        st(ws.cell(h,i,t),9,True,'FFFFFFFF',fill)
    ws.row_dimensions[h].height=38
    for k_,m in enumerate(MARGINS):
        r=h+1+k_; cc=lambda k: f"{cols[k]}{r}"
        mc=cc('m'); b=f"{M}$D$4"; mk=f"{M}$D$6"; rate=f"{M}$D$3"; tgt=f"{P}$B$14/{P}$E$14"
        f={}
        f['m']=m
        f['cur']=f'=IF(ABS({mc}-{P}$C$14)<{0.01 if EN else 0.0000001},"{T["cur"]}","")'
        f['share']=f"=1-{b}-{mc}"
        f['b4']=f"={M}$F$13/{cc('share')}*{mk}"; f['b2']=f"={M}$F$11/{cc('share')}*{mk}"
        f['ask']=f"=ROUNDDOWN((({tgt}*{cc('share')}/(1+{u}))/{rate}/(1+{FIN})-{batch})/((1+{DUT})*(1+{VAT})),2)"
        f['lc']=f"=({cc('ask')}*{rate}+{logi}*{rate}+({cc('ask')}+{frs})*{DUT}*{rate}+({cc('ask')}+{frs})*(1+{DUT})*{VAT}*{rate})*(1+{FIN})*(1+{u})"
        f['sa']=f"={cc('lc')}/{cc('share')}*{mk}"
        f['d4']=f"={cc('sa')}-{cc('b4')}"; f['dp4']=f"={cc('d4')}/{cc('b4')}"; f['d2']=f"={cc('sa')}-{cc('b2')}"; f['dp2']=f"={cc('d2')}/{cc('b2')}"
        for k,v in f.items():
            c_=ws[cc(k)]; c_.value=v
            if k=='m': st(c_,11,True,'FF203864','FFF2F5F9','0%')
            elif k=='cur': st(c_,9,True,'FFA33A3A','FFF2F5F9')
            elif k=='share': st(c_,10,False,'FF222222','FFF2F5F9','0%')
            elif k in('b4','b2'): st(c_,11,k=='b4','FF203864','FFE7EEF6',UAH)
            elif k=='ask': st(c_,12,True,'FF203864','FFFCE3A6','"€"0.00')
            elif k in('lc','sa'): st(c_,11,k=='sa','FF203864','FFFEF4DC',UAH)
            elif k in('d4','d2'): st(c_,11,True,'FFA33A3A','FFFBE9E9',DUAH)
            else: st(c_,11,True,'FFA33A3A','FFFBE9E9',DPCT)
        ws.row_dimensions[r].height=28
    return h+len(MARGINS)+1
for mn,pn,pal in PAIRS:
    ws=wb[mn]
    grid(ws,56,dict(m='B',cur='C',share='D',b4='E',b2='F',ask='G',lc='H',sa='I',d4='J',dp4='K',d2='L',dp2='M'),mn,pn,pal,False)
    wp=wb[pn]
    grid(wp,28,dict(m='A',cur='B',share='C',b4='D',b2='E',ask='F',lc='G',sa='H',d4='I',dp4='J',d2='K',dp2='L'),mn,pn,pal,True)
    for col in 'IJKL': wp.column_dimensions[col].width=14
# Scenarios (EN): до/после полки по сценариям
if EN and 'Scenarios' in wb.sheetnames:
    sc=wb['Scenarios']
    heads=[('O','33 pallets: shelf BEFORE 4 SKUs, ₴'),('P','33 pallets: shelf BEFORE 2 SKUs, ₴'),('Q','33 pallets: shelf AFTER request, ₴'),('R','6 pallets: shelf BEFORE 4 SKUs, ₴'),('S','6 pallets: shelf BEFORE 2 SKUs, ₴'),('T','6 pallets: shelf AFTER request, ₴')]
    for col,t in heads:
        sc[f'{col}7']=t; sc[f'{col}7']._style=copy.copy(sc['J7']._style); sc.column_dimensions[col].width=16
    for r in range(8,40):
        b=sc.cell(r,2).value
        if not isinstance(b,(int,float)): 
            continue
        for (cb4,cb2,caf,mn,pn,ask_col) in (('O','P','Q','Model 33 pallets','Price 120 UAH - 33 pallets','J'),('R','S','T','Model 6 pallets','Price 120 UAH - 6 pallets','M')):
            M=f"'{mn}'!"; P=f"'{pn}'!"; share=f"(1-$B{r}-$C{r})"; mk=f"{M}$D$6"; rate=f"{M}$D$3"; u=f"$D{r}"
            n=f"({P}$C$18*{P}$D$18)"
            sc[f'{cb4}{r}']=f"={M}$AL$13*{rate}*(1+{u})/{share}*{mk}"
            sc[f'{cb2}{r}']=f"={M}$AL$11*{rate}*(1+{u})/{share}*{mk}"
            ask=f"${ask_col}{r}"
            lc=f"({ask}*{rate}+({P}$E$16+{P}$F$16+{P}$G$16)/{n}*{rate}+({ask}+{P}$E$16*{P}$B$18/{n})*{P}$D$16*{rate}+({ask}+{P}$E$16*{P}$B$18/{n})*(1+{P}$D$16)*{P}$B$16*{rate})*(1+{P}$C$16)*(1+{u})"
            sc[f'{caf}{r}']=f"={lc}/{share}*{mk}"
            for col in (cb4,cb2,caf): sc[f'{col}{r}']._style=copy.copy(sc[f'H{r}']._style); sc[f'{col}{r}'].number_format=UAH
            sc[f'{caf}{r}'].font=Font(name='Arial',sz=11,b=True,color='FF203864')
    for r in range(8,40):
        if isinstance(sc.cell(r,1).value,str) and sc.cell(r,2).value is None:
            pass
wb.calculation=CalcProperties(fullCalcOnLoad=True); wb.save(OUT); print('saved',OUT)
