import openpyxl,copy,sys
from openpyxl.workbook.properties import CalcProperties
SRC="/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/4ba8d4c5-RAMEN_FINMODEL_Final..xlsx"; OUT=sys.argv[1]
wb=openpyxl.load_workbook(SRC); wv=openpyxl.load_workbook(SRC,data_only=True)
for n,ar in (('Цена 120 грн — 33 паллет',19),('Цена 120 грн — 6 паллет',20)):
    ws=wb[n]; v=wv[n]; cc=ar+2
    vat,fin,dut,fr,br,ot,sh,pl,pk=[v[a].value for a in ('B16','C16','D16','E16','F16','G16','B18','C18','D18')]
    n_=pl*pk
    batch=f"({fr}+{br}+{ot}+{fr}*{sh}*({dut}+{vat}*(1+{dut})))/{n_}"
    core=f"((B14/E14*(1-D14-C14))/F14/(1+{fin})-{batch})/((1+{dut})*(1+{vat}))"
    ws[f'B{ar}']=f"=ROUNDDOWN({core},2)"; ws[f'B{ar+1}']=f"={core}"
    ws[f'B{cc}']=f"=(B{ar}*F14+({fr}+{br}+{ot})/{n_}*F14+(B{ar}+{fr}*{sh}/{n_})*{dut}*F14+(B{ar}+{fr}*{sh}/{n_})*(1+{dut})*{vat}*F14)*(1+{fin})"
    for r in (15,16,17,18):
        for c in 'ABCDEFGH':
            ws[f'{c}{r}'].value=None; ws[f'{c}{r}']._style=copy.copy(ws['I1']._style)
        ws.row_dimensions[r].height=3
wb.calculation=CalcProperties(fullCalcOnLoad=True); wb.save(OUT); print('saved')
