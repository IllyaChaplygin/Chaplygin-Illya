# -*- coding: utf-8 -*-
import openpyxl,copy,sys
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties
SRC="/home/user/Chaplygin-Illya/RAMEN_FINMODEL_INTERNAL_UA.xlsx"; OUT=sys.argv[1]
wb=openpyxl.load_workbook(SRC); wv=openpyxl.load_workbook(SRC,data_only=True)
for n in ('Финмодель 6 паллет','Финмодель 33 паллет','Сценарии'): del wb[n]
CFG=(('Цена 120 грн — 33 паллет','33 pallets',33,19,'33 pallets (Full Truck)'),('Цена 120 грн — 6 паллет','6 pallets',6,20,'6 pallets'))
for old,new,pal,ar,ttl in CFG:
    ws=wb[old]; v=wv[old]; cc=ar+2; r1=cc+2; r2=cc+3
    ws.title=new
    # константы вместо ссылок на внутренние вкладки
    ws['D14']=v['D14'].value; ws['E14']=v['E14'].value; ws['F14']=v['F14'].value
    for a in ('D14','E14','F14'): ws[a]._style=copy.copy(ws['C14']._style)
    ws['D14'].number_format='0.0%'; ws['E14'].number_format='0.00"×"'; ws['F14'].number_format='0.00'
    ws['B%d'%r1]=v['B%d'%r1].value; ws['B%d'%r2]=v['B%d'%r2].value
    for a in ('B%d'%r1,'B%d'%r2):
        f=ws[a].font; ws[a].font=Font(name=f.name,sz=f.sz,b=f.b,color='FF0000FF'); ws[a].fill=copy.copy(ws['C14'].fill)
    # убираем доп. расходы (внутренний параметр)
    for a in ('G13','G14'): ws[a].value=None; ws[a]._style=copy.copy(ws['H13']._style)
    ws['C4']=f"=$B${r1}*F14"; ws['C6']=f"=$B${r2}*F14"
    batch="(E16+F16+G16+E16*B18*(D16+B16*(1+D16)))/(C18*D18)"
    core=f"((B14/E14*(1-D14-C14))/F14/(1+C16)-{batch})/((1+D16)*(1+B16))"
    ws[f'B{ar}']=f"=ROUNDDOWN({core},2)"; ws[f'B{ar+1}']=f"={core}"
    ws[f'B{cc}']=f"=(B{ar}*F14+(E16+F16+G16)/(C18*D18)*F14+(B{ar}+E16*B18/(C18*D18))*D16*F14+(B{ar}+E16*B18/(C18*D18))*(1+D16)*B16*F14)*(1+C16)"
    # английские подписи
    T={'A1':f'SIAS · from purchase price to shelf price, one line · {ttl}','A2':'REQUEST TO SIAS, €/pack','H2':'one price for all 6 SKUs','A3':'Scenario',
       'B3':'SIAS purchase\nper pack','C3':'Landed cost in Ukraine\nproduct + costs in Ukraine','D3':'Partner price\ncost ÷ (1 − bonus − our share)',
       'G3':'Store markup\n+40 % on partner price','H3':'SHELF PRICE incl. VAT\npartner price × 1.40',
       'A4':'BEFORE · Vegetable, Spicy, Beef, Chicken\n(4 SKUs)','A6':'BEFORE · Gouda Cheese, Carbonara Spicy\n(2 SKUs)','A8':'AFTER · SIAS request\n(all 6 SKUs)','A10':'CHANGE: AFTER − BEFORE (4 SKUs)',
       'A12':'INPUTS — yellow cells can be changed','B13':'Shelf price, ₴','C13':'Our share of partner price','D13':'Chain bonus','E13':'Shelf multiplier','F13':'Rate ₴/€','A14':'Target and model terms',
       'B15':'Import VAT','C15':'Financing','D15':'Customs duty','E15':'Batch freight, €','F15':'Broker, €','G15':'Other, €','B17':'Freight share in customs value','C17':'Pallets in batch','D17':'Packs per pallet',
       f'A{ar}':'Request to SIAS, €/pack (rounded down to the euro cent)',f'A{ar+1}':'Maximum purchase price before rounding, €/pack',f'A{cc}':'Landed cost in Ukraine at the request, ₴/pack',
       f'A{r1}':'Landed cost today, 4 SKUs (Vegetable, Spicy, Beef, Chicken), €/pack',f'A{r2}':'Landed cost today, 2 SKUs (Gouda, Carbonara), €/pack'}
    for k,t in T.items(): ws[k]=t
    ws.row_dimensions[13].height=26
    ws['D2']='="now €"&B5&" (4 SKUs) / €"&B7&" (2 SKUs)"'
    ws['F2']=f'="reduction −"&ROUND((1-B{ar}/B5)*100,0)&" % / −"&ROUND((1-B{ar}/B7)*100,0)&" %"'
    ws['E3']='="Chain bonus"&CHAR(10)&ROUND(D14*100,0)&" % of partner price"'
    ws['F3']='="Left for us"&CHAR(10)&"costs & margin"'
    for a in ('B4','B6','B8'): pass
# проверка: нет ссылок на удалённые вкладки
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('=') and ('Финмодель' in c.value or 'Сценарии' in c.value): print('BAD REF',ws.title,c.coordinate,c.value)
            if isinstance(c.value,str) and not c.value.startswith('=') and any('а'<=ch<='я' or 'А'<=ch<='Я' for ch in c.value): print('CYR',ws.title,c.coordinate,c.value[:60])
wb.properties.title='SIAS financial model'; wb.properties.creator=''; wb.properties.lastModifiedBy=''; wb.properties.description=None; wb.properties.keywords=None
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print('saved',OUT)
