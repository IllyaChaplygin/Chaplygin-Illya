# -*- coding: utf-8 -*-
import openpyxl,re,sys,copy
from openpyxl.workbook.properties import CalcProperties
import os
SRC=os.environ.get("SRC","/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad/xl/scn_raw.xlsx"); OUT=sys.argv[1]; REAL=os.environ.get("MODE")=="real"; EXTRA=os.environ.get("MODE")=="extra"; PROP=os.environ.get("MODE")=="prop"; U0=float(os.environ.get("UPLIFT","0.22"))
wb=openpyxl.load_workbook(SRC)
SHEETS={'Финмодель 6 паллет':'Model 6 pallets','Финмодель 33 паллет':'Model 33 pallets','Цена 120 грн — 6 паллет':'Price 120 UAH - 6 pallets','Цена 120 грн — 33 паллет':'Price 120 UAH - 33 pallets','Сценарии':'Scenarios'}
# 1) имена листов в формулах
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='):
                v=c.value
                for a,b in SHEETS.items(): v=v.replace(a,b)
                c.value=v
for a,b in SHEETS.items(): wb[a].title=b
D={
'Сценарій поставки':'Delivery scenario','6 палет':'6 pallets','33 палети':'33 pallets',
'Курс, грн/EUR':'Rate, UAH/EUR','Бонуси мережі, %':'Chain bonus, %','Цільова маржа, %':'Target margin, %','Націнка мережі (x)':'Chain markup (x)',
'Прибуток після бонусу = ціна партнеру − собівартість − бонус мережі.':'Profit after bonus = partner price − landed cost − chain bonus.',
'Ціна партнера':'Partner price','вже з ПДВ':'incl. VAT','ПДВ як у прикладі: СС збережено повністю; ціна вже з ПДВ. Множення на 1,20 немає.':'VAT as in the example: landed cost kept in full; the price already includes VAT. No multiplication by 1.20.',
'Частка ціни для собівартості':'Share of price for landed cost',
'ТОВАР':'PRODUCT','СОБІВАРТІСТЬ':'LANDED COST','ЦІНА ПАРТНЕРУ ТА ПРИБУТОК':'PARTNER PRICE AND PROFIT','ПОЛИЦЯ':'SHELF','МОДЕЛЬ: НУЛЬОВА МАРЖА':'MODEL: ZERO MARGIN','ПРИБУТОК З ОДНІЄЇ УПАКОВКИ':'PROFIT PER PACK',
'Група':'Group','Позиція':'Item','Вага, г':'Weight, g','СС, EUR/уп.':'Landed cost, EUR/pack','1. СС, грн/уп.':'1. Landed cost, UAH/pack','2. Ціна партнеру з ПДВ':'2. Partner price incl. VAT','3. Бонус мережі, грн':'3. Chain bonus, UAH',
'4. Прибуток після бонусу, грн':'4. Profit after bonus, UAH','5. Маржа за прикладом':'5. Margin per example','6. Полиця з ПДВ, грн':'6. Shelf price incl. VAT, UAH','Мін. ціна партнеру':'Min. partner price','Мін. полиця':'Min. shelf price',
'Валовий прибуток за прикладом, до бонусу, грн/уп.':'Gross profit per example, before bonus, UAH/pack','Прибуток після бонусу за прикладом, грн/уп.':'Profit after bonus per example, UAH/pack',
'Ціна партнеру з ПДВ, грн':'Partner price incl. VAT, UAH','Бонус мережі, грн':'Chain bonus, UAH','Прибуток після бонусу, грн':'Profit after bonus, UAH','Перевірка маржі':'Margin check','Полиця з ПДВ, грн':'Shelf price incl. VAT, UAH',
'SIAS — рамен у пакеті':'SIAS — pack ramen','Кожен крок формування ціни — за вашим прикладом':'Each step of price formation — per your example','Що рахуємо':'What we calculate','Дія':'Action',
'Собівартість в Україні, EUR/уп.':'Landed cost in Ukraine, EUR/pack','Із Self-Cost':'From Self-Cost','Умова в D3':'Input in D3','Собівартість в Україні, грн/уп.':'Landed cost in Ukraine, UAH/pack','Крок 1 × крок 2':'Step 1 × step 2',
'Повна ціна партнеру, %':'Full partner price, %','База для всіх відсотків':'Base for all percentages','Бонус мережі, %':'Chain bonus, %','Умова в D4':'Input in D4','Частка ціни після бонусу, %':'Share of price after bonus, %','Крок 4 − крок 5':'Step 4 − step 5',
'Наша цільова маржа, %':'Our target margin, %','Умова в D5':'Input in D5','Частка ціни для собівартості, %':'Share of price for landed cost, %','Крок 6 − крок 7':'Step 6 − step 7','Ціна партнеру з ПДВ, грн/уп.':'Partner price incl. VAT, UAH/pack','Крок 3 ÷ крок 8':'Step 3 ÷ step 8',
'Валовий прибуток до бонусу, грн/уп.':'Gross profit before bonus, UAH/pack','Крок 9 − крок 3':'Step 9 − step 3','Бонус мережі, грн/уп.':'Chain bonus, UAH/pack','Крок 9 × крок 5':'Step 9 × step 5',
'Нам залишається після собівартості та бонусу, грн/уп.':'Left for us after landed cost and bonus, UAH/pack','Крок 10 − крок 11':'Step 10 − step 11','Перевірка нашої маржі після бонусу, %':'Check of our margin after bonus, %','Крок 12 ÷ крок 9':'Step 12 ÷ step 9',
'Націнка ритейлера для полиці, %':'Retailer markup for shelf, %','Множник D6 − 1':'Multiplier D6 − 1','Множник ритейлера':'Retailer multiplier','1 + крок 14':'1 + step 14','Ціна на полиці з ПДВ, грн/уп.':'Shelf price incl. VAT, UAH/pack','Крок 9 × крок 15':'Step 9 × step 15',
'Ціна партнеру при нульовому нашому прибутку, грн/уп.':'Partner price at zero profit for us, UAH/pack','Крок 3 ÷ крок 6':'Step 3 ÷ step 6','Полиця при нульовому нашому прибутку, грн/уп.':'Shelf price at zero profit for us, UAH/pack','Крок 17 × крок 15':'Step 17 × step 15',
'Рамен. Ціна партнеру та полиця — 6 палет':'Ramen. Partner price and shelf price — 6 pallets','Рамен. Ціна партнеру та полиця — 33 палет / Full Truck':'Ramen. Partner price and shelf price — 33 pallets / Full Truck',
'Готова СС із Self-Cost_Sias: AJ6:AJ11. Склад СС збережено.':'Ready-made landed cost from Self-Cost_Sias: AJ6:AJ11. Cost composition preserved.','Готова СС із Self-Cost_Sias: AJ30:AJ35. Склад СС збережено.':'Ready-made landed cost from Self-Cost_Sias: AJ30:AJ35. Cost composition preserved.',
'Курс 51,20 грн/EUR — із Self-Cost_Sias.':'Rate 51.20 UAH/EUR — from Self-Cost_Sias.',
# tabs 3-4
'ЗАПРОС ДЛЯ SIAS, €/пачку':'REQUEST TO SIAS, €/pack','одна цена на 6 SKU':'one price for all 6 SKUs','Вариант':'Scenario','Закупка SIAS\nза пачку':'SIAS purchase\nper pack','Цена партнёру\nСС ÷ (1 − бонус − маржа)':'Partner price\ncost ÷ (1 − bonus − margin)',
'Наценка магазина\n+40 % к цене партнёру':'Store markup\n+40 % on partner price','ПОЛКА с НДС\nцена партнёру × 1,40':'SHELF PRICE incl. VAT\npartner price × 1.40',
'ДО · Vegetable, Spicy, Beef, Chicken\n(4 SKU)':'BEFORE · Vegetable, Spicy, Beef, Chicken\n(4 SKUs)','ДО · Gouda Cheese, Carbonara Spicy\n(2 SKU)':'BEFORE · Gouda Cheese, Carbonara Spicy\n(2 SKUs)','ПОСЛЕ · запрос SIAS\n(все 6 SKU)':'AFTER · SIAS request\n(all 6 SKUs)','ИЗМЕНЕНИЕ: ПОСЛЕ − ДО (4 SKU)':'CHANGE: AFTER − BEFORE (4 SKUs)',
'Полка, ₴':'Shelf price, ₴','Наша маржа':'Our margin','Бонус сети':'Chain bonus (shown)','Множитель полки':'Shelf multiplier','Курс ₴/€':'Rate ₴/€','Цель расчёта и условия модели':'Target and model terms',
'Импортный ПДВ':'Import VAT','Финансирование':'Financing','Пошлина':'Customs duty','Перевозка партии, €':'Batch freight, €','Брокер, €':'Broker, €','Прочие, €':'Other, €','Доля фрахта в таможне':'Freight share in customs value','Паллет в партии':'Pallets in batch','Пачек на паллете':'Packs per pallet',
'Запрос SIAS, €/уп. (вниз до евроцента)':'Request to SIAS, €/pack (rounded down to the euro cent)','Предельная закупка до округления, €/уп.':'Maximum purchase price before rounding, €/pack','Себестоимость в Украине при запросе (с доп. расходами), ₴/уп.':'Landed cost in Ukraine at the request (incl. extra cost), ₴/pack',
'СС из Self-Cost, 4 SKU (Vegetable…), €/уп. — без доп. расходов':'Landed cost per Self-Cost, 4 SKUs (Vegetable…), €/pack — without extra cost','СС из Self-Cost, 2 SKU (Gouda, Carbonara), €/уп. — без доп. расходов':'Landed cost per Self-Cost, 2 SKUs (Gouda, Carbonara), €/pack — without extra cost',
# scenarios
'Сценарии структуры цены: как меняется запрос к SIAS при другой марже, бонусе сети и доп. расходах':'INTERNAL FILE — for us only · Price-structure scenarios: the request to SIAS at a different margin, chain bonus and extra cost',
'ВНУТРЕННИЙ ФАЙЛ — только для нас · Сценарии структуры цены: запрос к SIAS при другой марже, бонусе сети и доп. расходах':'INTERNAL FILE — for us only · Price-structure scenarios: the request to SIAS at a different margin, chain bonus and extra cost',
'Цена партнёру при полке, ₴':'Partner price at shelf price, ₴','Сценарий':'Scenario','Маржа, которую показываем':'Margin shown','Доп. расходы в СС, ₴/уп.':'Extra cost in landed cost, ₴/pack','Бонус сети, ₴':'Chain bonus, ₴','Наша прибыль, ₴':'Our profit, ₴',
'Бюджет СС (товар + расходы + доп.), ₴':'Landed-cost budget (product + costs + extra), ₴','СС товара в Украине, ₴':'Product landed cost in Ukraine, ₴','33 палеты: макс. закупка, €':'33 pallets: max purchase, €','33 палеты: ЗАПРОС SIAS, €':'33 pallets: REQUEST TO SIAS, €','к текущим €0,65, %':'vs current €0.65, %',
'6 палет: макс. закупка, €':'6 pallets: max purchase, €','6 палет: ЗАПРОС SIAS, €':'6 pallets: REQUEST TO SIAS, €',
'БАЗА · как было в прошлой версии модели':'BASE · previous version of the model','Было: бонус 25 % + маржа 30 %':'Previous: bonus 25 % + margin 30 %',
'ТА ЖЕ ЦЕНА ЗАПРОСА, другая структура (сумма «бонус + маржа» и расходов не меняется)':'SAME REQUEST PRICE, different structure (the sum of bonus + margin + extra cost does not change)',
'★ Бонус сети 35 % + маржа 20 % (без доп. расходов)':'★ Chain bonus 35 % + margin 20 % (no extra cost)','Бонус 30 % + маржа 25 %':'Bonus 30 % + margin 25 %',
'Бонус 25 % + маржа 20 % + доп. расходы, подобранные под тот же запрос':'Bonus 25 % + margin 20 % + extra cost matched to the same request','Бонус 25 % + маржа 25 % + доп. расходы, подобранные под тот же запрос':'Bonus 25 % + margin 25 % + extra cost matched to the same request',
'ПРИМЕР: +10 ₴ в себестоимость (запрос становится ниже)':'EXAMPLE: +10 ₴ in landed cost (the request gets lower)','Бонус 25 % + маржа 20 % + 10 ₴':'Bonus 25 % + margin 20 % + 10 ₴','Бонус 25 % + маржа 25 % + 10 ₴':'Bonus 25 % + margin 25 % + 10 ₴','Бонус 35 % + маржа 20 % + 10 ₴':'Bonus 35 % + margin 20 % + 10 ₴','Бонус 35 % + маржа 25 % + 10 ₴':'Bonus 35 % + margin 25 % + 10 ₴',
'БОНУС 35 %: если маржа должна остаться выше 20 %, запрос падает':'CHAIN BONUS 35 %: if the margin must stay above 20 %, the request falls','Бонус 35 % + маржа 25 %':'Bonus 35 % + margin 25 %','Бонус 35 % + маржа 30 %':'Bonus 35 % + margin 30 %',
'ТОЛЬКО СНИЗИЛИ МАРЖУ (бонус 25 %, без доп. расходов): допустимая закупка растёт — аргумент для SIAS ослабляется':'ONLY THE MARGIN LOWERED (bonus 25 %, no extra cost): the allowed purchase price rises — the argument toward SIAS weakens','Бонус 25 % + маржа 25 %':'Bonus 25 % + margin 25 %','Бонус 25 % + маржа 20 %':'Bonus 25 % + margin 20 %',
'Как читать: жёлтые ячейки (бонус, маржа, доп. расходы) можно менять — запросы пересчитаются. Полка и параметры партии берутся из вкладок «Цена 120 грн».':'How to read: yellow cells (bonus, margin, extra cost) can be changed — requests recalculate. Shelf price and batch parameters come from the "Price 120 UAH" sheets.',
'Допущения: бонус и маржа считаются от цены партнёру (с НДС) — как в основной модели; доп. расходы добавляются к СС уже после финансирования. Текущая закупка SIAS: €0,65 (4 SKU) / €0,75 (2 SKU).':'Assumptions: bonus and margin are taken from the partner price (incl. VAT), as in the main model; extra cost is added to landed cost after financing. Current SIAS purchase price: €0.65 (4 SKUs) / €0.75 (2 SKUs).',
'Что показывать SIAS — решение за вами. Показанные маржа, бонус и доп. расходы должны соответствовать реальной структуре цены: SIAS может запросить подтверждение.':'What to show SIAS is your decision. Any margin, bonus and extra cost shown externally should match the real price structure: SIAS may ask for confirmation.',
}
cyr=re.compile('[А-Яа-яІіЇїЄєҐґ]')
def tr(v):
    if v in D: return D[v]
    v2=re.sub(r'(\d) г\b',r'\1 g',v)
    return v2
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and not c.value.startswith('=') and cyr.search(c.value) or (isinstance(c.value,str) and re.search(r'\d г\b',c.value) and not c.value.startswith('=')):
                c.value=tr(c.value)
# формулы с кириллицей
M6,M33='Model 6 pallets','Model 33 pallets'
for n in (M6,M33):
    ws=wb[n]
    ws['E4']='="Chain bonus "&ROUND(D4*100,0)&"%; shelf = partner price × "&D6&"."'
    ws['E5']='=ROUND(D5*100,0)&"% = (partner price − landed cost − bonus) / full partner price."'
    ws['E8']='="100% − bonus − margin. Current terms: 100% − "&ROUND(D4*100,0)&"% − "&ROUND(D5*100,0)&"% = "&ROUND(D8*100,0)&"%."'
    for a in ('P9','U9','Z9','AE9'): ws[a].number_format='"MARGIN "0%'
for n,ar,t in (('Price 120 UAH - 33 pallets',19,'33 pallets (Full Truck)'),('Price 120 UAH - 6 pallets',20,'6 pallets')):
    ws=wb[n]
    ws['A1']=(f'INTERNAL — not for SIAS · real chain bonus shown, landed cost not adjusted · {t}' if REAL else f'INTERNAL — not for SIAS · extra cost compensates the chain-bonus gap · {t}')
    ws['D2']='="now €"&B5&" (4 SKUs) / €"&B7&" (2 SKUs)"'
    ws['F2']=f'="reduction −"&ROUND((1-B{ar}/B5)*100,0)&" % / −"&ROUND((1-B{ar}/B7)*100,0)&" %"'
    ws['C3']='="Landed cost in Ukraine"&CHAR(10)&"product + costs in Ukraine"&IF(G14>0," + "&ROUND(G14,2)&" ₴ extra","")'
    ws['E3']='="Chain bonus (shown)"&CHAR(10)&ROUND(D14*100,0)&" % of partner price"'
    ws['F3']='="Our profit"&CHAR(10)&ROUND(C14*100,0)&" % of partner price"'
    mname='Model 33 pallets' if '33' in n else 'Model 6 pallets'
    ws['A12']=f'INPUTS — yellow cells can be changed, green ones are linked from the sheet "{mname}"'
    # компенсация бонуса: реальный бонус сети (внутренний ввод) и доп. расходы = разрыв × цена партнёру при целевой полке
    ws['H13']='Real chain bonus (internal)'; ws['H13']._style=copy.copy(ws['F13']._style)
    ws['H14']=0.35; ws['H14']._style=copy.copy(ws['C14']._style); ws['H14'].number_format='0.0%'
    ws['G13']='Extra cost, ₴/pack (bonus-gap compensation)'
    ws['G14']='=ROUND(B14/E14*(H14-D14),2)'
    f=ws['G14'].font; ws['G14'].font=openpyxl.styles.Font(name=f.name,sz=f.sz,b=f.b,color='FF172B46'); ws['G14'].fill=copy.copy(ws['E14'].fill)
    ws.column_dimensions['H'].width=17; ws.row_dimensions[13].height=36; ws.row_dimensions[3].height=78
    if EXTRA:
        ws['A1']=f'INTERNAL — not for SIAS · real chain bonus 35 %, extra cost in landed cost = assumption (edit G14) · {t}'
        ws['G14']=10; ws['G14']._style=copy.copy(ws['C14']._style); ws['G14'].number_format='#,##0.00" ₴"'
        ws['G13']='Extra cost in landed cost, ₴/pack (assumption)'
        ws['H13']=None; ws['H14']=None; ws['H13']._style=copy.copy(ws['I13']._style); ws['H14']._style=copy.copy(ws['I14']._style)
    ws.title=n
# сценарии: A2
sc=wb['Scenarios']
sc['A2']=('="The request depends on the sum of bonus + margin and on extra cost: allowed landed cost = partner price × (1 − bonus − margin) − extra cost. '
          '1 p.p. of margin or bonus = "&ROUND(E5/100,2)&" ₴ per pack = €"&ROUND(E5/100/D5/(1+C5)/((1+G5)*(1+F5)),4)&" on the request."')
for a,b in (('A4','Shelf price, ₴'),('B4','Shelf multiplier'),('C4','Financing'),('D4','Rate ₴/€'),('E4','Partner price at shelf price, ₴'),('F4','Import VAT'),('G4','Customs duty'),('B7','Chain bonus'),('C7','Margin shown'),('D7','Extra cost in landed cost, ₴/pack')): sc[a]=b

if PROP:
    for n,ar,t in (('Price 120 UAH - 33 pallets',19,'33 pallets (Full Truck)'),('Price 120 UAH - 6 pallets',20,'6 pallets')):
        ws=wb[n]; cc=ar+2; r1=cc+2; r2=cc+3
        ws['A1']=f'INTERNAL — not for SIAS · real bonus 35 %, proportional cost uplift (G14) · {t}'
        ws['G13']='Cost uplift, % of landed cost (assumption)'
        ws['G14']=U0; ws['G14']._style=copy.copy(ws['C14']._style); ws['G14'].number_format='0.0%'
        ws['H13']='Base share of price for cost (real model: 1 − 25 % − 30 %)'; ws['H13']._style=copy.copy(ws['F13']._style)
        ws['H14']=0.45; ws['H14']._style=copy.copy(ws['G14']._style); ws['H14'].number_format='0.0%'
        ws['C13']='Margin shown (auto: offsets the cost uplift)'
        ws['C14']='=1-D14-H14*(1+G14)'; f_=ws['C14'].font; ws['C14'].font=openpyxl.styles.Font(name=f_.name,sz=f_.sz,b=f_.b,color='FF172B46'); ws['C14'].fill=copy.copy(ws['E14'].fill); ws['C14'].number_format='0.0%'
        ws.column_dimensions['H'].width=22
        ws['C4']=f"=$B${r1}*F14*(1+$G$14)"; ws['C6']=f"=$B${r2}*F14*(1+$G$14)"
        batch="(E16+F16+G16+E16*B18*(D16+B16*(1+D16)))/(C18*D18)"
        core=f"((B14/E14*(1-D14-C14)/(1+G14))/F14/(1+C16)-{batch})/((1+D16)*(1+B16))"
        ws[f'B{ar}']=f"=ROUNDDOWN({core},2)"; ws[f'B{ar+1}']=f"={core}"
        ws[f'B{cc}']=f"=(B{ar}*F14+(E16+F16+G16)/(C18*D18)*F14+(B{ar}+E16*B18/(C18*D18))*D16*F14+(B{ar}+E16*B18/(C18*D18))*(1+D16)*B16*F14)*(1+C16)*(1+G14)"
        ws['C3']='="Landed cost in Ukraine"&CHAR(10)&"product + costs in Ukraine"&IF(G14>0," + "&ROUND(G14*100,1)&" % uplift","")'
        ws[f'A{ar}']='Request to SIAS, €/pack (rounded down to the euro cent)'; ws[f'A{cc}']='Landed cost in Ukraine at the request (incl. uplift), ₴/pack'
        for rr_ in (r1,r2):
            ws[f'C{rr_}']=f"=B{rr_}*(1+$G$14)"; ws[f'C{rr_}']._style=copy.copy(ws[f'B{rr_}']._style); ws[f'C{rr_}'].number_format='"€"0.000'
            f_=ws[f'C{rr_}'].font; ws[f'C{rr_}'].font=openpyxl.styles.Font(name=f_.name,sz=f_.sz,b=True,color='FF172B46'); ws[f'C{rr_}'].fill=copy.copy(ws['E14'].fill)
        ws[f'A{r1}']='Landed cost, 4 SKUs (Vegetable…), €/pack — Self-Cost (col. B) → after uplift (col. C)'
        ws[f'A{r2}']='Landed cost, 2 SKUs (Gouda, Carbonara), €/pack — Self-Cost (col. B) → after uplift (col. C)'
    sc=wb['Scenarios']
    sc['A1']='INTERNAL FILE — for us only · Price-structure scenarios with a proportional cost uplift'
    sc['A2']=('="Allowed (shown) landed cost = partner price × (1 − bonus − margin); product landed cost = that budget ÷ (1 + uplift). '
              '1 p.p. of margin or bonus = "&ROUND(E5/100,2)&" ₴ per pack = €"&ROUND(E5/100/D5/(1+C5)/((1+G5)*(1+F5)),4)&" on the request."')
    sc['D7']='Cost uplift, % of landed cost'; sc['G7']='Landed-cost budget (shown cost), ₴'; sc['H7']='Product landed cost allowed (budget ÷ (1 + uplift)), ₴'
    ROWS={9:('Previous: bonus 25 % + margin 30 %',.25,.30,0),11:('★ Chain bonus 35 % + margin 20 % (no uplift)',.35,.20,0),12:('Bonus 30 % + margin 25 % (no uplift)',.30,.25,0),
          13:('Bonus 25 % + margin 20 % + uplift matched to the same request',.25,.20,'=(1-B13-C13)/(1-$B$9-$C$9)-1'),14:('Bonus 25 % + margin 25 % + uplift matched to the same request',.25,.25,'=(1-B14-C14)/(1-$B$9-$C$9)-1'),
          16:('Bonus 25 % + margin 20 % + 10 % uplift',.25,.20,.10),17:('Bonus 25 % + margin 25 % + 10 % uplift',.25,.25,.10),18:('Bonus 35 % + margin 20 % + 10 % uplift',.35,.20,.10),19:('Bonus 35 % + margin 25 % + 10 % uplift',.35,.25,.10),
          21:('Bonus 35 % + margin 25 % (no uplift)',.35,.25,0),22:('Bonus 35 % + margin 30 % (no uplift)',.35,.30,0),24:('Bonus 25 % + margin 25 % (no uplift)',.25,.25,0),25:('Bonus 25 % + margin 20 % (no uplift)',.25,.20,0)}
    for r,(nm,b,m,u) in ROWS.items():
        sc.cell(r,1,nm); sc.cell(r,2,b); sc.cell(r,3,m); sc.cell(r,4,u); sc.cell(r,4).number_format='0.0%'
        sc.cell(r,8,f"=G{r}/(1+D{r})")
    sc['A15']='EXAMPLE: +10 % proportional uplift of landed cost (the request gets lower)'
    sc['A27']='How to read: yellow cells (bonus, margin, cost uplift) can be changed — requests recalculate. Shelf price and batch parameters come from the "Price 120 UAH" sheets.'
    sc['A28']='Assumptions: bonus and margin are taken from the partner price (incl. VAT), as in the main model; the uplift is applied to the whole landed cost (after financing). Current SIAS purchase price: €0.65 (4 SKUs) / €0.75 (2 SKUs).'
    sc['A10']='SAME REQUEST PRICE, different structure (uplift matched to the sum of bonus + margin)'


if PROP:
    for mn,pn,r1,r2 in (('Model 6 pallets','Price 120 UAH - 6 pallets',24,25),('Model 33 pallets','Price 120 UAH - 33 pallets',23,24)):
        m=wb[mn]; p=wb[pn]
        m['AL10']='Self-Cost source, EUR/pack (before uplift)'; m['AL10']._style=copy.copy(m['O10']._style); m.column_dimensions['AL'].width=18
        for r in range(11,17):
            m[f'AL{r}']=m[f'E{r}'].value; m[f'AL{r}']._style=copy.copy(m[f'E{r}']._style)
            m[f'E{r}']=f"=AL{r}*(1+'{pn}'!$G$14)"
        m['D5']=f"='{pn}'!C14"
        m['E10']='Landed cost incl. uplift, EUR/pack'
        m['E2']='Landed cost = Self-Cost_Sias (col. AL) × (1 + uplift from G14 on the "Price 120 UAH" sheet). Cost composition preserved.'
        p[f'B{r1}']=f"='{mn}'!AL13"; p[f'B{r2}']=f"='{mn}'!AL11"
if PROP:
    import json
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    BL=json.load(open("/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad/v10/blocks.json"))
    bs=wb.create_sheet('Landed cost blocks',0)
    thin=Side(style='thin',color='C9D3E0'); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
    def st(c,sz=10,b=False,col='FF172B46',fill=None,fmt=None,al='center'):
        c.font=Font(name='Arial',sz=sz,b=b,color=col)
        if fill: c.fill=PatternFill('solid',fgColor=fill)
        if fmt: c.number_format=fmt
        c.alignment=Alignment(horizontal=al,vertical='center',wrap_text=True); c.border=BR
    bs.sheet_view.showGridLines=False
    bs.merge_cells('A1:I1'); bs['A1']='INTERNAL — not for SIAS · Landed-cost blocks with the proportional uplift · ₴ per pack'; bs['A1'].font=Font(name='Arial',sz=16,b=True,color='FF172B46'); bs.row_dimensions[1].height=26
    bs.merge_cells('A2:I2'); bs['A2']='Purchase + delivery + import VAT + financing = real landed cost; the uplift (assumption, cell G14 on the "Price 120 UAH" sheets) is added on top as a separate block.'
    bs['A2'].font=Font(name='Arial',sz=10,color='FF5E6D7E'); bs['A2'].alignment=Alignment(wrap_text=True,vertical='center'); bs.row_dimensions[2].height=30
    HEAD=['','Purchase\npack price × rate','Delivery\nfreight, broker, terminal','Import VAT\n20 % on goods and freight','Financing\n2.5 % of the sum','Real landed cost, ₴','Uplift (assumption), ₴','Landed cost with uplift, ₴','Check vs model sheet']
    r=4
    for sh,key,pal,fr_ask in (('Price 120 UAH - 33 pallets','33',33,'B19'),('Price 120 UAH - 6 pallets','6',6,'B20')):
        bs.cell(r,1,f'{pal} PALLETS'+(' · FULL TRUCK' if pal==33 else '')); bs.cell(r,1).font=Font(name='Arial',sz=11,b=True,color='FF223A5E'); r+=1
        for i,h in enumerate(HEAD,1): st(bs.cell(r,i,h),9,True,'FFFFFFFF','3E5A85')
        bs.row_dimensions[r].height=44; r+=1
        T=f"'{sh}'!"
        for nm,vals,ref in (('BEFORE · 4 SKUs\n(Vegetable, Spicy, Beef, Chicken)',BL[key]['before4'],f"{T}C4"),('BEFORE · 2 SKUs\n(Gouda, Carbonara)',BL[key]['before2'],f"{T}C6"),('AFTER · SIAS request\n(all 6 SKUs)',None,f"{T}C8")):
            bs.cell(r,1,nm); st(bs.cell(r,1),10,True,'FF172B46','FCE3A6' if vals is None else 'F2F5F9',None,'left')
            if vals is not None:
                for j in range(4): bs.cell(r,2+j,vals[j]); st(bs.cell(r,2+j),11,False,'FF0000FF','FFF4D2','#,##0.00" ₴"')
            else:
                ask=f"{T}{fr_ask}"; n=f"({T}C18*{T}D18)"
                bs.cell(r,2,f"={ask}*{T}F14"); bs.cell(r,3,f"=({T}E16+{T}F16+{T}G16)/{n}*{T}F14")
                bs.cell(r,4,f"=({ask}+{T}E16*{T}B18/{n})*(1+{T}D16)*{T}B16*{T}F14")
                bs.cell(r,5,f"=(B{r}+C{r}+D{r}+({ask}+{T}E16*{T}B18/{n})*{T}D16*{T}F14)*{T}C16")
                for j in range(4): st(bs.cell(r,2+j),11,False,'FF172B46','FCE3A6','#,##0.00" ₴"')
            bs.cell(r,6,f"=SUM(B{r}:E{r})"); st(bs.cell(r,6),11,True,'FF172B46','E7EDF5','#,##0.00" ₴"')
            bs.cell(r,7,f"=F{r}*{T}G14"); st(bs.cell(r,7),11,True,'FFB03A3A','FDE8E8','#,##0.00" ₴"')
            bs.cell(r,8,f"=F{r}+G{r}"); st(bs.cell(r,8),13,True,'FF172B46','FCE3A6','#,##0.00" ₴"')
            bs.cell(r,9,f"=ROUND(H{r}-{ref},4)"); st(bs.cell(r,9),9,False,'FF5E6D7E','F2F5F9','0.00;-0.00;"OK"')
            bs.row_dimensions[r].height=34; r+=1
        r+=1
    bs.column_dimensions['A'].width=38
    for col in 'BCDEFGHI': bs.column_dimensions[col].width=17
wb.properties.title='Ramen financial model — internal (EN)'; wb.properties.creator=''; wb.properties.lastModifiedBy=''
wb.calculation=CalcProperties(fullCalcOnLoad=True); wb.save(OUT)
left=[(ws.title,c.coordinate,c.value[:50]) for ws in wb for r in ws.iter_rows() for c in r if isinstance(c.value,str) and cyr.search(c.value)]
nf=[(ws.title,c.coordinate,c.number_format) for ws in wb for r in ws.iter_rows() for c in r if cyr.search(c.number_format or '')]
print('saved; cyrillic left:',left[:20],'| in formats:',nf[:5])
