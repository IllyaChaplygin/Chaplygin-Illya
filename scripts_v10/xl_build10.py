# -*- coding: utf-8 -*-
import openpyxl,copy,sys
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.workbook.properties import CalcProperties
SRC="/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/4ba8d4c5-RAMEN_FINMODEL_Final..xlsx"
OUT=sys.argv[1]
BONUS,MARGIN,XEXTRA=0.35,0.20,0.0
wb=openpyxl.load_workbook(SRC)
# ---------- вкладки 1–2: бонус 35 %, маржа 20 % ----------
for n in ('Финмодель 6 паллет','Финмодель 33 паллет'):
    ws=wb[n]
    ws['D4']=BONUS; ws['D5']=MARGIN
    ws['E4']='="Бонус мережі "&ROUND(D4*100,0)&"%; полиця = ціна партнера × "&D6&"."'
    ws['E5']='=ROUND(D5*100,0)&"% = (ціна партнеру − собівартість − бонус) / повна ціна партнеру."'
    ws['E8']='="100% − бонус − маржа. За поточних умов: 100% − "&ROUND(D4*100,0)&"% − "&ROUND(D5*100,0)&"% = "&ROUND(D8*100,0)&"%."'
# ---------- вкладки 3–4 ----------
YEL=PatternFill('solid',fgColor='FFF4D2'); GRY=PatternFill('solid',fgColor='F2F5F9')
for n,askrow,ccrow,sc_ in (('Цена 120 грн — 33 паллет',19,21,"'Финмодель 33 паллет'"),('Цена 120 грн — 6 паллет',20,22,"'Финмодель 6 паллет'")):
    ws=wb[n]
    ws['C14']=MARGIN
    # доп. расходы в СС
    for a,b in (('F13','G13'),('F14','G14')):
        ws[b]._style=copy.copy(ws[a]._style)
    ws['G13']='Доп. расходы в СС, ₴/уп.'
    ws['G14']=XEXTRA; ws['G14']._style=copy.copy(ws['C14']._style); ws['G14'].number_format='#,##0.00" ₴"'
    ws['A12']=ws['A12'].value  # без изменений
    # источники СС из Self-Cost (вкладка «Финмодель») — ссылки вынесены вниз
    r1=ccrow+2; r2=ccrow+3
    lab=ws['A%d'%askrow]._style
    ws['A%d'%r1]='СС из Self-Cost, 4 SKU (Vegetable…), €/уп. — без доп. расходов'; ws['A%d'%r2]='СС из Self-Cost, 2 SKU (Gouda, Carbonara), €/уп. — без доп. расходов'
    for r,src in ((r1,'E13'),(r2,'E11')):
        ws['A%d'%r]._style=copy.copy(ws['A%d'%askrow]._style)
        ws['B%d'%r]=f"={sc_}!{src}"; ws['B%d'%r]._style=copy.copy(ws['C5']._style); ws['B%d'%r].number_format='"€"0.000'
    ws['C4']=f"=$B${r1}*F14+$G$14"; ws['C5']="=C4/F14"
    ws['C6']=f"=$B${r2}*F14+$G$14"; ws['C7']="=C6/F14"
    for a in ('C5','C7'):
        f=copy.copy(ws[a].font); ws[a].font=Font(name=f.name,sz=f.sz,b=f.b,color='FF172B46')
    # запрос SIAS с учётом доп. расходов
    batch="(E16+F16+G16+E16*B18*(D16+B16*(1+D16)))/(C18*D18)"
    core=f"((B14/E14*(1-D14-C14)-G14)/F14/(1+C16)-{batch})/((1+D16)*(1+B16))"
    ws[f'B{askrow}']=f"=ROUNDDOWN({core},2)"
    ws[f'B{askrow+1}']=f"={core}"
    ws[f'B{ccrow}']=f"=(B{askrow}*F14+(E16+F16+G16)/(C18*D18)*F14+(B{askrow}+E16*B18/(C18*D18))*D16*F14+(B{askrow}+E16*B18/(C18*D18))*(1+D16)*B16*F14)*(1+C16)+G14"
    ws[f'A{ccrow}']='Себестоимость в Украине при запросе (с доп. расходами), ₴/уп.'
    # заголовки — динамические
    ws['C3']='="Себестоимость в Украине"&CHAR(10)&"товар + расходы в Украине"&IF(G14>0," + "&ROUND(G14,2)&" ₴ доп.","")'
    ws['E3']='="Бонус сети"&CHAR(10)&ROUND(D14*100,0)&" % цены партнёру"'
    ws['F3']='="Наша прибыль"&CHAR(10)&ROUND(C14*100,0)&" % цены партнёру"'
    for rr_ in (ccrow,r1,r2): ws.row_dimensions[rr_].height=30
    ws['A12']='ИСХОДНЫЕ ДАННЫЕ — жёлтое можно менять, зелёное из вкладки «'+sc_.strip("'")+'»'
# ---------- вкладка «Сценарии» ----------
ws=wb.create_sheet('Сценарии')
T33="'Цена 120 грн — 33 паллет'!"; T6="'Цена 120 грн — 6 паллет'!"
thin=Side(style='thin',color='C9D3E0'); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
def st(c,sz=10,b=False,col='FF172B46',fill=None,fmt=None,al='center'):
    c.font=Font(name='Arial',sz=sz,b=b,color=col)
    if fill: c.fill=PatternFill('solid',fgColor=fill)
    if fmt: c.number_format=fmt
    c.alignment=Alignment(horizontal=al,vertical='center',wrap_text=True); c.border=BR
ws.sheet_view.showGridLines=False
ws.merge_cells('A1:N1'); ws['A1']='Сценарии структуры цены: как меняется запрос к SIAS при другой марже, бонусе сети и доп. расходах'
ws['A1'].font=Font(name='Arial',sz=16,b=True,color='FF172B46'); ws['A1'].alignment=Alignment(wrap_text=True,vertical='center'); ws.row_dimensions[1].height=26
ws.merge_cells('A2:N2')
ws['A2']=('="Запрос зависит от суммы «бонус + маржа» и доп. расходов: допустимая СС = цена партнёру × (1 − бонус − маржа) − доп. расходы. '
          '1 п.п. маржи или бонуса = "&ROUND(E5/100,2)&" ₴ на пачку = €"&ROUND(E5/100/D5/(1+C5)/((1+E6)*(1+F6)),4)&" к запросу."')
ws['A2'].font=Font(name='Arial',sz=10,color='FF5E6D7E'); ws['A2'].alignment=Alignment(wrap_text=True,vertical='center'); ws.row_dimensions[2].height=30
# общие параметры
hdr=['Полка, ₴','Множитель полки','Курс ₴/€','Финансирование','Цена партнёру при полке, ₴','Пошлина','Импортный ПДВ']
# row 4 headers, row 5 values (links)
labels=[('A','Полка, ₴',f"={T33}B14",'#,##0.00" ₴"'),('B','Множитель полки',f"={T33}E14",'0.00"×"'),('C','Финансирование',f"={T33}C16",'0.0%'),('D','Курс ₴/€',f"={T33}F14",'0.00'),
        ('E','Цена партнёру при полке, ₴','=A5/B5','#,##0.00" ₴"'),('F','Импортный ПДВ',f"={T33}B16",'0.0%'),('G','Пошлина',f"={T33}D16",'0.0%')]
# переупорядочим, чтобы формулы выше совпали: C5=fin, D5=rate, E5=partner, F5=VAT, G5=duty ; E6/F6 в A2 -> VAT/duty
# (A2 использует E6,F6 — заменим на корректные ссылки ниже)
ws['A2']=('="Запрос зависит от суммы «бонус + маржа» и доп. расходов: допустимая СС = цена партнёру × (1 − бонус − маржа) − доп. расходы. '
          '1 п.п. маржи или бонуса = "&ROUND(E5/100,2)&" ₴ на пачку = €"&ROUND(E5/100/D5/(1+C5)/((1+G5)*(1+F5)),4)&" к запросу."')
for col,lab,f,fmt in labels:
    c=ws[f'{col}4']; c.value=lab; st(c,8,True,'FF5E6D7E',None,None); c.border=Border()
    v=ws[f'{col}5']; v.value=f; st(v,11,True,'FF008000' if col!='E' else 'FF172B46','F2F5F9',fmt)
ws.row_dimensions[4].height=24; ws.row_dimensions[5].height=22
# заголовок таблицы
H=['Сценарий','Бонус сети','Маржа, которую показываем','Доп. расходы в СС, ₴/уп.','Бонус сети, ₴','Наша прибыль, ₴','Бюджет СС (товар + расходы + доп.), ₴',
   'СС товара в Украине, ₴','33 палеты: макс. закупка, €','33 палеты: ЗАПРОС SIAS, €','к текущим €0,65, %','6 палет: макс. закупка, €','6 палет: ЗАПРОС SIAS, €','к текущим €0,65, %']
for i,h in enumerate(H,1):
    c=ws.cell(7,i,h); st(c,9,True,'FFFFFFFF','3E5A85')
ws.row_dimensions[7].height=52
def batchref(T): return f"({T}E16+{T}F16+{T}G16+{T}E16*{T}B18*($G$5+$F$5*(1+$G$5)))/({T}C18*{T}D18)"
P='$E$5'
rows=[
 ('sec','БАЗА · как было в прошлой версии модели'),
 ('Было: бонус 25 % + маржа 30 %',0.25,0.30,0,False),
 ('sec','ТА ЖЕ ЦЕНА ЗАПРОСА, другая структура (сумма «бонус + маржа» и расходов не меняется)'),
 ('★ Бонус сети 35 % + маржа 20 % (без доп. расходов)',0.35,0.20,0,True),
 ('Бонус 30 % + маржа 25 %',0.30,0.25,0,False),
 ('Бонус 25 % + маржа 20 % + доп. расходы, подобранные под тот же запрос',0.25,0.20,'=ROUND($E$5*((1-B{r}-C{r})-(1-$B$9-$C$9)),2)',False),
 ('Бонус 25 % + маржа 25 % + доп. расходы, подобранные под тот же запрос',0.25,0.25,'=ROUND($E$5*((1-B{r}-C{r})-(1-$B$9-$C$9)),2)',False),
 ('sec','ПРИМЕР: +10 ₴ в себестоимость (запрос становится ниже)'),
 ('Бонус 25 % + маржа 20 % + 10 ₴',0.25,0.20,10,False),
 ('Бонус 25 % + маржа 25 % + 10 ₴',0.25,0.25,10,False),
 ('Бонус 35 % + маржа 20 % + 10 ₴',0.35,0.20,10,False),
 ('Бонус 35 % + маржа 25 % + 10 ₴',0.35,0.25,10,False),
 ('sec','БОНУС 35 %: если маржа должна остаться выше 20 %, запрос падает'),
 ('Бонус 35 % + маржа 25 %',0.35,0.25,0,False),
 ('Бонус 35 % + маржа 30 %',0.35,0.30,0,False),
 ('sec','ТОЛЬКО СНИЗИЛИ МАРЖУ (бонус 25 %, без доп. расходов): допустимая закупка растёт — аргумент для SIAS ослабляется'),
 ('Бонус 25 % + маржа 25 %',0.25,0.25,0,False),
 ('Бонус 25 % + маржа 20 %',0.25,0.20,0,False),
]
r=8
for row in rows:
    if row[0]=='sec':
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=14)
        c=ws.cell(r,1,row[1]); c.font=Font(name='Arial',sz=9,b=True,color='FF223A5E'); c.fill=PatternFill('solid',fgColor='E7EDF5'); c.alignment=Alignment(vertical='center')
        ws.row_dimensions[r].height=20; r+=1; continue
    name,b,m,x,star=row
    ws.cell(r,1,name); st(ws.cell(r,1),10,True,'FF172B46','FCE3A6' if star else 'F2F5F9',None,'left')
    ws.cell(r,2,b); st(ws.cell(r,2),11,True,'FF0000FF','FFF4D2','0.0%')
    ws.cell(r,3,m); st(ws.cell(r,3),11,True,'FF0000FF','FFF4D2','0.0%')
    ws.cell(r,4,x.replace('{r}',str(r)) if isinstance(x,str) else x); st(ws.cell(r,4),11,True,'FF0000FF','FFF4D2','#,##0.00" ₴"')
    ws.cell(r,5,f"=$E$5*B{r}"); ws.cell(r,6,f"=$E$5*C{r}"); ws.cell(r,7,f"=$E$5-E{r}-F{r}"); ws.cell(r,8,f"=G{r}-D{r}")
    for c_ in (5,6,7,8): st(ws.cell(r,c_),10,False,'FF172B46','F2F5F9','#,##0.00" ₴"')
    for (cm,ca,cd,T) in ((9,10,11,T33),(12,13,14,T6)):
        core=f"(($H{r}/$D$5/(1+$C$5)-{batchref(T)})/((1+$G$5)*(1+$F$5)))"
        ws.cell(r,cm,"="+core); st(ws.cell(r,cm),10,False,'FF172B46','F2F5F9','"€"0.0000')
        ws.cell(r,ca,f"=ROUNDDOWN({core},2)"); st(ws.cell(r,ca),12,True,'FF172B46','FCE3A6','"€"0.00')
        ws.cell(r,cd,f"={openpyxl.utils.get_column_letter(ca)}{r}/0.65-1"); st(ws.cell(r,cd),10,False,'FF172B46','F2F5F9','+0%;−0%;0%')
    ws.row_dimensions[r].height=34; r+=1
last=r-1
r+=1
notes=['Как читать: жёлтые ячейки (бонус, маржа, доп. расходы) можно менять — запросы пересчитаются. Полка и параметры партии берутся из вкладок «Цена 120 грн».',
 'Допущения: бонус и маржа считаются от цены партнёру (с НДС) — как в основной модели; доп. расходы добавляются к СС уже после финансирования. Текущая закупка SIAS: €0,65 (4 SKU) / €0,75 (2 SKU).',
 'Что показывать SIAS — решение за вами. Показанные маржа, бонус и доп. расходы должны соответствовать реальной структуре цены: SIAS может запросить подтверждение.']
for t in notes:
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=14); c=ws.cell(r,1,t)
    c.font=Font(name='Arial',sz=9,color='FF5E6D7E'); c.alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[r].height=28; r+=1
widths={'A':46,'B':12,'C':14,'D':15,'E':14,'F':14,'G':18,'H':16,'I':15,'J':15,'K':12,'L':15,'M':15,'N':12}
for k,v in widths.items(): ws.column_dimensions[k].width=v
ws.freeze_panes='B8'
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print('saved',OUT)
