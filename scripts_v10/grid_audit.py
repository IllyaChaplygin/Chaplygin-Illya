import sys,re,math
from pptx import Presentation
from pptx.util import Emu
P=Presentation(sys.argv[1]); sl=list(P.slides)[28]
RATE=51.2;MK=1.4;FIN=.025;VAT=.2
def ask(fr,pal,sh,m,b=.35):
    n=pal*1600; batch=fr+330+fr*.7*VAT
    mp=(sh/MK*(1-b-m)/RATE/(1+FIN)-batch/n)/(1+VAT); return math.floor(mp*100+1e-9)/100
cells=[]
for sh in sl.shapes:
    if sh.has_text_frame and re.fullmatch(r'€\d,\d\d|€\d\.\d\d',sh.text_frame.text.strip()):
        cells.append((Emu(sh.left).inches,Emu(sh.top).inches,float(sh.text_frame.text.strip()[1:].replace(',','.'))))
print(len(cells))
fails=0
for xs,(fr,pal) in (((0,6.5),(3700,33)),((6.5,14),(2600,6))):
    sub=[c for c in cells if xs[0]<=c[0]<xs[1]]
    xsort=sorted({round(c[0],2) for c in sub}); ysort=sorted({round(c[1],2) for c in sub})
    assert len(xsort)==5 and len(ysort)==5,(xsort,ysort)
    for yi,m in enumerate((.30,.25,.20,.15,.10)):
        row=[]
        for xi,shf in enumerate((105,110,120,130,140)):
            v=[c[2] for c in sub if round(c[0],2)==xsort[xi] and round(c[1],2)==ysort[yi]][0]
            e=ask(fr,pal,shf,m); row.append(v)
            if abs(v-e)>1e-9: fails+=1; print('FAIL',pal,m,shf,v,e)
        print(pal,f'{m:.0%}',row)
print('grid fails',fails)
