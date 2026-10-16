import openpyxl,math,sys
F=sys.argv[1]
nv=openpyxl.load_workbook(F,data_only=True); nf=openpyxl.load_workbook(F)
SC=openpyxl.load_workbook('/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/87c2996d-Self-Cost_Sias.xlsx',data_only=True).worksheets[0]
fail=0
def chk(n,a,b,tol=1e-9):
    global fail
    if a is None or abs(a-b)>tol*max(1,abs(b)): fail+=1; print('FAIL',n,a,b)
bad=[(ws.title,c.coordinate) for ws in nf for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('=') and (nv[ws.title][c.coordinate].value is None or str(nv[ws.title][c.coordinate].value).startswith('#'))]
print('formula errors:',bad)
RATE=51.2;MK=1.4;FIN=.025;VAT=.2;SHELF=120
def ask_mp(fr,pal,b,m,u): n=pal*1600; batch=fr+330+fr*.7*VAT; return ((SHELF/MK*(1-b-m)/(1+u))/RATE/(1+FIN)-batch/n)/(1+VAT)
def ccf(fr,pal,p,u): n=pal*1600; return (p*RATE+(fr+330)/n*RATE+(p+fr*.7/n)*VAT*RATE)*(1+FIN)*(1+u)
for sheet,pal,fr,ar,scr in (('Price 120 UAH - 33 pallets',33,3700,19,(32,30)),('Price 120 UAH - 6 pallets',6,2600,20,(8,6))):
    w=nv[sheet]; b,m,u=w['D14'].value,w['C14'].value,w['G14'].value
    mp=ask_mp(fr,pal,b,m,u); a=math.floor(mp*100+1e-9)/100
    chk(sheet+' ask',w[f'B{ar}'].value,a); chk(sheet+' max',w[f'B{ar+1}'].value,mp); chk(sheet+' cc after',w[f'B{ar+2}'].value,ccf(fr,pal,a,u))
    for r,buy,cc in ((4,.65,SC[f'AJ{scr[0]}'].value*RATE*(1+u)),(6,.75,SC[f'AJ{scr[1]}'].value*RATE*(1+u)),(8,a,ccf(fr,pal,a,u))):
        P=cc/(1-b-m)
        for col,e in (('B',buy*RATE),('C',cc),('D',P),('E',P*b),('F',P-cc-P*b),('G',P*(MK-1)),('H',P*MK)): chk(f'{sheet} {col}{r}',w[f'{col}{r}'].value,e)
    print(sheet,'b/m/u',b,m,u,'| ask',w[f'B{ar}'].value,'| СС до',round(w['C4'].value,2),round(w['C6'].value,2),'после',round(w['C8'].value,2),'| полка',round(w['H4'].value,2),round(w['H6'].value,2),round(w['H8'].value,2),'| бонус₴ / остаток₴ после',round(w['E8'].value,2),round(w['F8'].value,2))
S=nv['Scenarios']
for r in range(8,40):
    b=S.cell(r,2).value
    if not isinstance(b,(int,float)): continue
    m,u=S.cell(r,3).value,S.cell(r,4).value
    chk(f'sc{r} H',S.cell(r,8).value,120/MK*(1-b-m)/(1+u))
    for cm,ca,fr,pal in ((9,10,3700,33),(12,13,2600,6)):
        mp=ask_mp(fr,pal,b,m,u); chk(f'sc{r} max{pal}',S.cell(r,cm).value,mp); chk(f'sc{r} ask{pal}',S.cell(r,ca).value,math.floor(mp*100+1e-9)/100)
    print(f'{S.cell(r,1).value[:66]:66s} u={u:6.1%} | 33p {S.cell(r,10).value:.2f}  6p {S.cell(r,13).value:.2f}')
print('FAILS',fail)
