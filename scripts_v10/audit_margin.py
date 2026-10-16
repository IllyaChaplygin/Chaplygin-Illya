import openpyxl,math,sys
F,LANG=sys.argv[1],sys.argv[2]; EN=LANG=='en'
N=openpyxl.load_workbook(F,data_only=True); NF=openpyxl.load_workbook(F)
SC=openpyxl.load_workbook('/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/87c2996d-Self-Cost_Sias.xlsx',data_only=True).worksheets[0]
R=51.2;MK=1.4;FIN=.025;VAT=.2
PAIRS=((('Model 6 pallets','Price 120 UAH - 6 pallets',6),('Model 33 pallets','Price 120 UAH - 33 pallets',33)) if EN else (('Финмодель 6 паллет','Цена 120 грн — 6 паллет',6),('Финмодель 33 паллет','Цена 120 грн — 33 паллет',33)))
SCR={6:(8,6),33:(32,30)}; FR={6:2600,33:3700}
fails=0
def chk(n,a,b,tol=1e-9):
    global fails
    if a is None or abs(a-b)>tol*max(1,abs(b)): fails+=1; print('FAIL',n,a,b)
bad=[(ws.title,c.coordinate) for ws in NF for r in ws.iter_rows() for c in r if isinstance(c.value,str) and c.value.startswith('=') and 'cur' not in c.value and 'IF(ABS' not in c.value and (N[ws.title][c.coordinate].value is None or str(N[ws.title][c.coordinate].value).startswith('#'))]
print('formula errors',bad)
MARG=[.35,.30,.25,.20,.15,.10]
for mn,pn,pal in PAIRS:
    m=N[mn]; p=N[pn]; b=m['D4'].value; u=p['G14'].value if EN else 0; fr=FR[pal]; n=pal*1600
    cost4=SC[f'AJ{SCR[pal][0]}'].value*R*(1+u); cost2=SC[f'AJ{SCR[pal][1]}'].value*R*(1+u)
    for tab,ws,h0,cols in (('model',m,60,dict(m='B',share='D',b4='E',b2='F',ask='G',lc='H',sa='I',d4='J',dp4='K',d2='L',dp2='M')),('price',p,32,dict(m='A',share='C',b4='D',b2='E',ask='F',lc='G',sa='H',d4='I',dp4='J',d2='K',dp2='L'))):
        for k,mg in enumerate(MARG):
            r=h0+k; g=lambda c: ws[f"{cols[c]}{r}"].value
            chk(f'{tab} {mn} m r{r}',g('m'),mg); share=1-b-mg; chk('share',g('share'),share)
            b4=cost4/share*MK; b2=cost2/share*MK; chk('b4',g('b4'),b4); chk('b2',g('b2'),b2)
            mp=((120/MK*share/(1+u))/R/(1+FIN)-(fr+330+fr*.7*VAT)/n)/(1+VAT); ask=math.floor(mp*100+1e-9)/100; chk(f'ask {tab} {mn} r{r}',g('ask'),ask)
            lc=(ask*R+(fr+330)/n*R+(ask+fr*.7/n)*VAT*R)*(1+FIN)*(1+u); chk('lc',g('lc'),lc); sa=lc/share*MK; chk('sa',g('sa'),sa)
            chk('d4',g('d4'),sa-b4); chk('dp4',g('dp4'),(sa-b4)/b4); chk('d2',g('d2'),sa-b2); chk('dp2',g('dp2'),(sa-b2)/b2)
        if tab=='price': print(pn,'| 30%: shelf before',round(ws[f'D{h0+1}'].value,2),round(ws[f'E{h0+1}'].value,2),'ask',ws[f'F{h0+1}'].value,'after',round(ws[f'H{h0+1}'].value,2),'| 20%: ask',ws[f'F{h0+3}'].value,'| 15%: ask',ws[f'F{h0+4}'].value,'| marker:',[ws[f'B{h0+k}'].value for k in range(6)])
if EN:
    s=N['Scenarios']; cnt=0
    for r in range(8,40):
        b=s.cell(r,2).value
        if not isinstance(b,(int,float)): continue
        mg,u=s.cell(r,3).value,s.cell(r,4).value; share=1-b-mg; cnt+=1
        for col4,col2,cola,pal,askc in (('O','P','Q',33,10),('R','S','T',6,13)):
            c4=SC[f'AJ{SCR[pal][0]}'].value*R*(1+u); c2=SC[f'AJ{SCR[pal][1]}'].value*R*(1+u); n=pal*1600; fr=FR[pal]; ask=s.cell(r,askc).value
            chk(f'scn b4 r{r}',s[f'{col4}{r}'].value,c4/share*MK); chk(f'scn b2 r{r}',s[f'{col2}{r}'].value,c2/share*MK)
            lc=(ask*R+(fr+330)/n*R+(ask+fr*.7/n)*VAT*R)*(1+FIN)*(1+u); chk(f'scn after r{r}',s[f'{cola}{r}'].value,lc/share*MK)
            assert lc/share*MK<=120+1e-6
    print('scenario rows',cnt)
print('FAILS',fails)
