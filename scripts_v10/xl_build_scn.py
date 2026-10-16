import openpyxl,sys
from openpyxl.workbook.properties import CalcProperties
SRC="/home/user/Chaplygin-Illya/RAMEN_FINMODEL_INTERNAL_UA.xlsx"; OUT=sys.argv[1]
wb=openpyxl.load_workbook(SRC)
for n in ('Финмодель 6 паллет','Финмодель 33 паллет'):
    wb[n]['D4']=0.25; wb[n]['D5']=0.20
for n,t in (('Цена 120 грн — 33 паллет','33 паллеты (Full Truck)'),('Цена 120 грн — 6 паллет','6 паллет')):
    ws=wb[n]; ws['C14']=0.20; ws['G14']=10
    ws['A1']=f'ВНУТРЕННИЙ СЦЕНАРИЙ — только для нас, не для SIAS · бонус 25 %, маржа 20 %, доп. расходы 10 ₴ (допущение) · {t}'
    ws['G13']='Доп. расходы в СС, ₴/уп. (допущение)'
wb['Сценарии']['A1']='ВНУТРЕННИЙ ФАЙЛ — только для нас · Сценарии структуры цены: запрос к SIAS при другой марже, бонусе сети и доп. расходах'
wb.calculation=CalcProperties(fullCalcOnLoad=True); wb.save(OUT); print('saved')
