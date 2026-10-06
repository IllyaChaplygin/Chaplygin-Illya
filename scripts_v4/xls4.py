# -*- coding: utf-8 -*-
import json,os,collections,statistics as st
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as CL
import calc120 as M
SC=os.path.dirname(os.path.abspath(__file__))
SJ=json.load(open(f"{SC}/segmented.json",encoding='utf-8')); ROWS=SJ['rows']
SS=json.load(open(f"{SC}/segstats.json",encoding='utf-8'))
NAVY="FF303A5D"; LIGHT="FFF1F4FA"
HF=Font(name="Segoe UI",size=9,bold=True,color="FFFFFFFF"); BF=Font(name="Segoe UI",size=9,bold=True,color=NAVY)
RF=Font(name="Segoe UI",size=9); TF=Font(name="Segoe UI",size=12,bold=True,color=NAVY); GF=Font(name="Segoe UI",size=8,color="FF717890")
fH=PatternFill("solid",fgColor=NAVY); fL=PatternFill("solid",fgColor=LIGHT); IN=PatternFill("solid",fgColor="FFFFF3D6")
OUT=PatternFill("solid",fgColor="FFE3F3F1")
thin=Side(style="thin",color="FFD8DEEC"); BD=Border(bottom=thin)
wb=Workbook()
def title(ws,t,sub=""):
    ws["A1"]=t; ws["A1"].font=TF
    if sub: ws["A2"]=sub; ws["A2"].font=GF
def head(ws,row,cols,widths=None):
    for j,c in enumerate(cols,1):
        x=ws.cell(row=row,column=j,value=c); x.font=HF; x.fill=fH; x.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[row].height=28
    if widths:
        for j,w in enumerate(widths,1): ws.column_dimensions[CL(j)].width=w
def put(ws,ref,v,fmt=None,font=RF,fill=None):
    c=ws[ref]; c.value=v; c.font=font
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    return c

# ============ 1. ОБЕРНЕНИЙ РОЗРАХУНОК ============
ws=wb.active; ws.title="1. Запит 120 грн"
title(ws,"Запит до SIAS: ціна закупівлі для полиці 120 грн","Повторює вкладки «Цена 120 грн — 6 паллет» і «— 33 паллет» файлу RAMEN_FINMODEL_4_TABS_PRESENTATION.xlsx. Жовті клітинки — вхідні дані; решта — формули.")
for c,w in zip("ABCD",(46,16,16,60)): ws.column_dimensions[c].width=w
ws["A4"]="ПАРАМЕТРИ"; ws["A4"].font=BF
P=[("Цільова полиця, грн",120,"0.00","B5"),("Наша маржа після бонусу",0.30,"0 %","B6"),("Бонус мережі",0.25,"0 %","B7"),
   ("Націнка мережі (×)",1.40,"0.00","B8"),("Курс, грн/EUR",51.2,"0.00","B9"),("Фінансування СС",0.025,"0.0 %","B10"),
   ("Імпортний ПДВ",0.20,"0 %","B11"),("Мито",0.0,"0 %","B12"),("Частка фрахту в митній базі",0.70,"0 %","B13"),
   ("Брокер партії, EUR",288,"0","B14"),("Інші витрати партії, EUR",42,"0","B15"),("Пачок на палеті",1600,"0","B16"),("Пачок у коробі",20,"0","B17")]
for i,(lab,v,fmt,ref) in enumerate(P):
    ws.cell(row=5+i,column=1,value=lab).font=RF; put(ws,ref,v,fmt,BF,IN)
ws["A19"]="СЦЕНАРІЇ ПОСТАВКИ"; ws["A19"].font=BF
for j,h in enumerate(("","6 палет","33 палети · повна машина","Пояснення")):
    c=ws.cell(row=20,column=1+j,value=h); c.font=HF; c.fill=fH
put(ws,"A21","Фрахт партії, EUR"); put(ws,"B21",2600,"0",BF,IN); put(ws,"C21",3700,"0",BF,IN); put(ws,"D21","Self-Cost: збірне 2 600, ціле авто 3 700",font=GF)
put(ws,"A22","Палет у партії"); put(ws,"B22","='2. Запит до SIAS'!H11","0"); put(ws,"C22","='2. Запит до SIAS'!I11","0")
put(ws,"A23","Пачок у партії"); 
for col in "BC": put(ws,f"{col}23",f"={col}22*$B$16","#,##0")
put(ws,"A24","Витрати партії, EUR"); put(ws,"D24","фрахт + брокер + інші + ПДВ (і мито) на фрахт у митній базі",font=GF)
for col in "BC": put(ws,f"{col}24",f"={col}21+$B$14+$B$15+{col}21*$B$13*($B$12+$B$11*(1+$B$12))","#,##0")
put(ws,"A25","Ціна партнеру для цільової полиці, грн"); put(ws,"D25","полиця ÷ націнка мережі",font=GF)
put(ws,"A26","Максимальна СС, грн/уп."); put(ws,"D26","ціна партнеру × (1 − бонус − маржа)",font=GF)
for col in "BC":
    put(ws,f"{col}25","=$B$5/$B$8","0.00"); put(ws,f"{col}26",f"={col}25*(1-$B$7-$B$6)","0.00")
put(ws,"A27","ПРЕДЕЛЬНА ЗАКУПІВЛЯ, EUR/уп.",font=BF); put(ws,"D27","(СС ÷ курс ÷ (1+фін.) − витрати/пачок) ÷ ((1+мито)(1+ПДВ))",font=GF)
for col in "BC": put(ws,f"{col}27",f"=({col}26/$B$9/(1+$B$10)-{col}24/{col}23)/((1+$B$12)*(1+$B$11))","0.0000",BF)
put(ws,"A28","ЗАПИТ ДО SIAS, EUR/уп. (вниз до євроцента)",font=BF)
for col in "BC": put(ws,f"{col}28",f"=ROUNDDOWN({col}27,2)","0.00",Font(name="Segoe UI",size=11,bold=True,color="FF1E9EA6"),OUT)
ws["A30"]="РОЗКЛАД ЦІНИ ПРИ ЗАПИТІ (грн за пачку)"; ws["A30"].font=BF
rows=[("Товар від постачальника","={c}28*$B$9"),
      ("Перевезення, брокер і прочі","=({c}21+$B$14+$B$15)/{c}23*$B$9"),
      ("Мито","=({c}28+{c}21*$B$13/{c}23)*$B$12*$B$9"),
      ("Імпортний ПДВ","=({c}28+{c}21*$B$13/{c}23)*(1+$B$12)*$B$11*$B$9"),
      ("Фінансування закупівлі","=SUM({c}31:{c}34)*$B$10"),
      ("Повна собівартість в Україні","=SUM({c}31:{c}35)"),
      ("Ціна партнеру з маржею","={c}36/(1-$B$7-$B$6)"),
      ("Бонус мережі","={c}37*$B$7"),
      ("Наша прибуток після бонусу","={c}37-{c}36-{c}38"),
      ("Націнка магазину","={c}37*($B$8-1)"),
      ("ПОЛИЦЯ З ПДВ","={c}37*$B$8"),
      ("Перевірка маржі","={c}39/{c}37"),
      ("Запас до цільової полиці, грн","=$B$5-{c}41")]
for i,(lab,fm) in enumerate(rows):
    r=31+i; put(ws,f"A{r}",lab,font=(BF if lab.isupper() or lab.startswith("Повна") else RF))
    for col in "BC": put(ws,f"{col}{r}",fm.replace("{c}",col),"0.0 %" if "маржі" in lab else "0.00",(BF if lab=="ПОЛИЦЯ З ПДВ" else RF))

# ============ 2. ЗАПИТ ДО SIAS ============
w2=wb.create_sheet("2. Запит до SIAS")
title(w2,"Запит до SIAS по позиціях","Одна запитувана ціна для всіх шести SKU. Ціна «зараз» — з Self-Cost_Sias.xlsx. Згода постачальника не отримана.")
head(w2,4,["Позиція SIAS","Вага, г","Ціна зараз, €/уп.","Запит 6 палет, €","Зниження","Запит 33 палети, €","Зниження","Палет у 6","Палет у 33","Пачок у 33","Товар у 33, €"],[24,9,16,16,12,18,12,10,10,12,14])
for i,(nm,wt,w_,cur) in enumerate(M.SKUS):
    r=5+i
    put(w2,f"A{r}",nm,font=BF); put(w2,f"B{r}",w_,"0.0"); put(w2,f"C{r}",cur,"0.00",BF,IN)
    put(w2,f"D{r}","='1. Запит 120 грн'!$B$28","0.00"); put(w2,f"E{r}",f"=1-D{r}/C{r}","0 %")
    put(w2,f"F{r}","='1. Запит 120 грн'!$C$28","0.00"); put(w2,f"G{r}",f"=1-F{r}/C{r}","0 %")
    put(w2,f"H{r}",1,"0",BF,IN); put(w2,f"I{r}",M.SC["33"]["pallets"][i],"0",BF,IN)
    put(w2,f"J{r}",f"=I{r}*'1. Запит 120 грн'!$B$16","#,##0"); put(w2,f"K{r}",f"=J{r}*F{r}","#,##0")
put(w2,"A11","Усього",font=BF); put(w2,"H11","=SUM(H5:H10)","0",BF); put(w2,"I11","=SUM(I5:I10)","0",BF)
put(w2,"J11","=SUM(J5:J10)","#,##0",BF); put(w2,"K11","=SUM(K5:K10)","#,##0",BF)

# ============ 3. ПОЛИЦЯ ЗАРАЗ ============
w3=wb.create_sheet("3. Полиця зараз")
title(w3,"Полиця за поточною закупівлею","СС — із Self-Cost_Sias.xlsx (колонка AJ, EUR/уп.). Маржа й бонус — з аркуша 1; ціна партнеру вже з ПДВ.")
head(w3,4,["Позиція","Вага, г","СС 6 палет, €","СС 33 палети, €","Ціна партнеру 6п, грн","Полиця 6п, грн","Ціна партнеру 33п, грн","Полиця 33п, грн","До 120, 33п","Маржа при полиці 120, 6п","Маржа при полиці 120, 33п"],[24,9,15,15,18,14,19,15,12,20,20])
for i,(nm,wt,w_,cur) in enumerate(M.SKUS):
    r=5+i
    put(w3,f"A{r}",nm,font=BF); put(w3,f"B{r}",w_,"0.0")
    put(w3,f"C{r}",M.SC["6"]["cc"][i],"0.0000",RF,IN); put(w3,f"D{r}",M.SC["33"]["cc"][i],"0.0000",RF,IN)
    R_="'1. Запит 120 грн'!"
    put(w3,f"E{r}",f"=C{r}*{R_}$B$9/(1-{R_}$B$7-{R_}$B$6)","0.00"); put(w3,f"F{r}",f"=E{r}*{R_}$B$8","0.00")
    put(w3,f"G{r}",f"=D{r}*{R_}$B$9/(1-{R_}$B$7-{R_}$B$6)","0.00"); put(w3,f"H{r}",f"=G{r}*{R_}$B$8","0.00")
    put(w3,f"I{r}",f"=H{r}-{R_}$B$5","0.00")
    put(w3,f"J{r}",f"=({R_}$B$5/{R_}$B$8*(1-{R_}$B$7)-C{r}*{R_}$B$9)/({R_}$B$5/{R_}$B$8)","0.0 %")
    put(w3,f"K{r}",f"=({R_}$B$5/{R_}$B$8*(1-{R_}$B$7)-D{r}*{R_}$B$9)/({R_}$B$5/{R_}$B$8)","0.0 %")

# ============ 4. ЧУТЛИВІСТЬ ============
w4=wb.create_sheet("4. Чутливість")
title(w4,"Яка закупівля потрібна при іншій полиці й маржі","Формула H66 вкладок «Цена 120 грн». EUR/уп., вниз до євроцента. Параметри — з аркуша 1.")
w4.column_dimensions['A'].width=22
R_="'1. Запит 120 грн'!"
def block(top,col,label,sc_col):
    w4.cell(row=top,column=1,value=label).font=BF
    w4.cell(row=top+1,column=1,value="Маржа \\ Полиця, грн").font=GF
    for j,sh in enumerate([105,110,120,130,140]):
        c=w4.cell(row=top+1,column=2+j,value=sh); c.font=HF; c.fill=fH
        w4.column_dimensions[CL(2+j)].width=11
    for i,m in enumerate([0.35,0.30,0.25,0.20,0.15]):
        r=top+2+i
        c=w4.cell(row=r,column=1,value=m); c.number_format="0 %"; c.font=BF
        for j in range(5):
            cc=CL(2+j)
            f=(f"=ROUNDDOWN((({cc}${top+1}/{R_}$B$8*(1-{R_}$B$7-$A{r}))/{R_}$B$9/(1+{R_}$B$10)-{R_}{sc_col}$24/{R_}{sc_col}$23)"
               f"/((1+{R_}$B$12)*(1+{R_}$B$11)),2)")
            x=w4.cell(row=r,column=2+j,value=f); x.number_format="0.00"; x.font=RF
block(4,2,"33 ПАЛЕТИ · ПОВНА МАШИНА","$C")
block(13,2,"6 ПАЛЕТ","$B")

# ============ 5. РЕЄСТР SKU ============
w5=wb.create_sheet("5. Реєстр SKU")
SN={1:"Пакет · з бульйоном",2:"Пакет · із соусом",3:"Стакан",4:"Рисові галушки"}
NM={'auchan':'Ашан','novus':'NOVUS','metro':'METRO','megamarket':'МегаМаркет','ultramarket':'Ultramarket','ekomarket':'ЕКО','cosmos':'Космос','vostorg':'Восторг','tavriav':'Таврія В','torba':'Торба','epicentr':'Епіцентр','zaraz':'Зараз','grono':'Grono','ideal':'Ідеал','onde':'Onde','chudomarket':'ЧудоМаркет','kharkiv':'Клас'}
title(w5,f"Реєстр корейського рамену · {len(ROWS)} SKU","Каталоги 18 мереж, 29.09.2026. Ціна — роздріб за пачку, з ПДВ. Сегмент — після ручного аудиту (Buldak і Toomba — із соусом).")
head(w5,4,["Сегмент","Бренд","Назва","Вага, г","EAN","Мереж","Мережі","Мін, грн","Медіана, грн","Макс, грн","грн/100 г"],[20,12,64,9,16,8,36,10,12,10,11])
for i,r in enumerate(sorted(ROWS,key=lambda x:(x['seg'],x['pmed']))):
    rr=5+i
    vals=[SN[r['seg']],r['tm'],r['title'],r['w'],r['ean'],r['nch'],", ".join(NM.get(c,c) for c in r['chains']),r['pmin'],r['pmed'],r['pmax'],r['per100']]
    for j,v in enumerate(vals,1):
        c=w5.cell(row=rr,column=j,value=v); c.font=RF
        if j in (8,9,10,11): c.number_format="0.0"
w5.freeze_panes="A5"

# ============ 6. СЕГМЕНТИ ============
w6=wb.create_sheet("6. Сегменти")
title(w6,"Статистика сегментів","Рахується за аркушем «5. Реєстр SKU»; квартилі — метод exclusive.")
head(w6,4,["Сегмент","SKU","Брендів","Мереж","Мін, грн","P25","Медіана","P75","Макс, грн","грн/100 г","Вага, г"],[24,8,10,8,10,8,10,8,10,12,14])
for i,k in enumerate("1234"):
    v=SS[k]; r=5+i
    vals=[SN[int(k)],v['n'],v['brands'],v['chains'],v['lo'],v['p25'],v['med'],v['p75'],v['hi'],v['g'],f"{v['wlo']}–{v['whi']}"]
    for j,x in enumerate(vals,1): w6.cell(row=r,column=j,value=x).font=RF
pk=sorted(r['pmed'] for r in ROWS if r['seg'] in (1,2) and r.get('pmed')); q=st.quantiles(pk,n=4)
w6["A10"]="Пакет разом (сегменти 1+2)"; w6["A10"].font=BF
for j,x in enumerate([len(pk),"","",min(pk),round(q[0],1),round(st.median(pk),1),round(q[2],1),max(pk)],2): w6.cell(row=10,column=j,value=x).font=RF
w6["A12"]="Позиція ціни 120 грн у пакеті"; w6["A12"].font=BF
w6["A13"]="Дорожчих за 120 грн"; w6["B13"]=sum(1 for p in pk if p>120)
w6["A14"]="Дорожчих за 140,55 грн (4 SKU, 33 палети)"; w6["B14"]=sum(1 for p in pk if p>140.5534)
w6["A15"]="Дорожчих за 162,18 грн (2 SKU, 33 палети)"; w6["B15"]=sum(1 for p in pk if p>162.177)

# ============ 7. ІМПОРТ ============
w7=wb.create_sheet("7. Імпорт")
title(w7,"Імпорт за митною базою, УКТ ЗЕД 1902 30 10 00","Джерело: митна база користувача (дані з SIAS_MARKET_RESEARCH_UA_V3.pptx). Січень–травень.")
head(w7,4,["Показник","2025, 01–05","2026, 01–05","Зміна"],[34,16,16,14])
data=[("Обсяг, т",239.67149,838.480003),("Інвойс, млн $",1.02458998,3.20062004),("Корея: обсяг, т",54.8,118.497356),("Корея: інвойс, тис. $",348.6,809.2)]
for i,(l,a,b) in enumerate(data):
    r=5+i; w7.cell(row=r,column=1,value=l).font=BF
    for j,x in ((2,a),(3,b)): w7.cell(row=r,column=j,value=x).number_format="0.00"
    c=w7.cell(row=r,column=4,value=f"=C{r}/B{r}-1"); c.number_format="+0.0 %"
w7["A10"]="Весь 2025 рік, т"; w7["B10"]=711.6
w7["A11"]="Митна вартість, $/кг"; w7["B11"]="=B6*1000/B5"; w7["C11"]="=C6*1000/C5"; w7["D11"]="=C11/B11-1"
w7["B11"].number_format="0.00"; w7["C11"].number_format="0.00"; w7["D11"].number_format="+0.0 %"
w7["A12"]="Корея: $/кг"; w7["B12"]="=B8/B7"; w7["C12"]="=C8/C7"; w7["D12"]="=C12/B12-1"
w7["B12"].number_format="0.00"; w7["C12"].number_format="0.00"; w7["D12"].number_format="+0.0 %"
w7["A13"]="Частка Кореї в обсязі"; w7["B13"]="=B7/B5"; w7["C13"]="=C7/C5"; w7["B13"].number_format="0.0 %"; w7["C13"].number_format="0.0 %"
w7["A15"]="Країни походження, 2026 01–05, т"; w7["A15"].font=BF
for i,(k,v) in enumerate([("В’єтнам",166.687505),("Корея",118.497356),("Румунія",86.919064),("Велика Британія",58.564844),("Латвія",57.30452),("Чехія",50.800755),("Польща",48.761462)]):
    w7.cell(row=16+i,column=1,value=k); w7.cell(row=16+i,column=2,value=v).number_format="0.0"

# ============ 8. CHOI'S ТА АКЦІЇ ============
w8=wb.create_sheet("8. Choi's і акції")
title(w8,"Choi’s Carbonara 131 г у «Таврії В» і «Космосі»; акції конкурентів","Каталоги 29.09.2026; ціна «Таврії В» 05–06.10 — за V3; ЄС — dotasia.eu.")
head(w8,4,["Джерело","Дата","Ціна, грн","Примітка"],[40,14,12,60])
ch=[("«Таврія В» (zakaz.ua, Івано-Франківськ)","29.09.2026",125.20,"стара ціна = ціна, акції немає"),
    ("«Космос» (zakaz.ua, Київ)","29.09.2026",125.20,"стара ціна = ціна"),
    ("«Таврія В» (за V3)","05–06.10.2026",131.50,"+5 % до 29.09"),
    ("ЄС · dotasia.eu · акція €1,80","жовтень 2026",None,"перерахунок за курсом 51,20; ПДВ не вказано; товару немає в наявності"),
    ("ЄС · dotasia.eu · звичайна €2,20","жовтень 2026",None,"те саме")]
for i,(a,b,c_,d) in enumerate(ch):
    r=5+i
    for j,x in enumerate((a,b,c_,d),1): w8.cell(row=r,column=j,value=x).font=RF
    if c_ is not None: w8.cell(row=r,column=3).number_format="0.00"
w8["C8"]="=1.8*'1. Запит 120 грн'!B9"; w8["C9"]="=2.2*'1. Запит 120 грн'!B9"
w8["C8"].number_format="0.00"; w8["C9"].number_format="0.00"
w8["A11"]="Ключові конкуренти: найнижча і найвища ціна по мережах"; w8["A11"].font=BF
Z=json.load(open(f"{SC}/zakaz_raw.json",encoding='utf-8')); E={r['ean']:r for r in ROWS}
KEY=[('08801073113428','Samyang Buldak 2× spicy 140 г'),('08801073110502','Samyang Buldak Hot Chicken 140 г'),('08801073113381','Samyang Buldak тушкована курка 145 г'),
     ('08801073116474','Samyang Buldak з сиром 130 г'),('08801043150620','Nongshim Shin Ramyun 120 г'),('08801043157742','Nongshim Kimchi Ramyun 120 г'),
     ('08801043157728','Nongshim Chapaghetti 140 г'),('08801043018470','Nongshim Toomba Spicy & Creamy 137 г'),('00648436310685','Paldo Volcano Carbonara 130 г')]
head(w8,12,["Позиція","Мереж","Найнижча, грн","Найвища, грн","Різниця"])
for i,(e,l) in enumerate(KEY):
    v=[z for z in Z if z['ean']==e]; r=13+i
    w8.cell(row=r,column=1,value=l).font=RF; w8.cell(row=r,column=2,value=len({z['chain'] for z in v}))
    w8.cell(row=r,column=3,value=min(z['price'] for z in v)).number_format="0.00"; w8.cell(row=r,column=4,value=max(z['price'] for z in v)).number_format="0.00"
    c=w8.cell(row=r,column=5,value=f"=1-C{r}/D{r}"); c.number_format="0 %"
out="/home/user/Chaplygin-Illya/Korean_Ramen_Model_and_Data.xlsx"
wb.save(out); print("saved",out)
