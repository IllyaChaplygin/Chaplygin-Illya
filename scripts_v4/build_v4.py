# -*- coding: utf-8 -*-
"""Спільний модуль колоди «Корейський рамен в Україні» за моделлю рисової колоди."""
import sys,json,os,re,math,statistics as st,collections as co
SC="/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad"
sys.path.insert(0,SC)
from deckkit import *
from deckkit import _rect,_S
from pptx.enum.text import PP_ALIGN as A
import calc120 as M
L=A.LEFT; R=A.RIGHT; C=A.CENTER

# ---------------- дані ----------------
F=json.load(open(f"{SC}/facts.json",encoding='utf-8'))
SEG=json.load(open(f"{SC}/cards_seg.json",encoding='utf-8'))
CK=json.load(open(f"{SC}/cards_kor.json",encoding='utf-8'))
CM=json.load(open(f"{SC}/cards_mass.json",encoding='utf-8'))
SS=json.load(open(f"{SC}/segstats.json",encoding='utf-8'))
_SJ=json.load(open(f"{SC}/segmented.json",encoding='utf-8')); SD_ROWS=_SJ['rows']
BT={r['tm']:r for r in json.load(open(f"{SC}/brandtiers.json",encoding='utf-8'))}
if "без тм" in CK: CK["Choi's"]=CK.pop("без тм")
def g(b): return CK.get(b,[])
def tl(tm,label=None):
    r=BT[tm]; p=f"{SC}/img/{r['best']['ean']}.jpg"
    return {"img":p if os.path.exists(p) else None,"name":label or tm,
            "price":f"{r['pmed']} грн","sub":f"{r['g']} грн/100 г · {r['nch']} мереж · {r['sku']} SKU"}
def img_for(ean):
    p=f"{SC}/img/{ean}.jpg"; return p if os.path.exists(p) else None

# ---------------- форматування ----------------
def d1(x): return f"{x:.1f}".replace('.',',')
def d2(x): return f"{x:.2f}".replace('.',',')
def d0(x): return f"{x:.0f}"
def th(x): return f"{x:,.0f}".replace(',',' ')
def th1(x): return f"{x:,.1f}".replace(',',' ').replace('.',',')
def eur(x,n=2): return "€"+f"{x:.{n}f}".replace('.',',')
def pct(x,n=1): return (f"{x:.{n}f}".replace('.',',').replace('-','−'))+" %"
def sgn(x,n=1): return (f"{x:+.{n}f}".replace('.',',').replace('-','−'))+" %"
NM={'auchan':'Ашан','novus':'NOVUS','metro':'METRO','cosmos':'«Космос»','vostorg':'«Восторг»',
    'tavriav':'«Таврія В»','torba':'«Торба»','zaraz':'«Зараз»','grono':'Grono','onde':'Onde',
    'chudomarket':'«ЧудоМаркет»','kharkiv':'«Клас»','silpo':'«Сільпо»'}
NATIONAL={'auchan','novus','metro','zaraz','silpo'}

# сегменти
SNAME={'1':'РАМЕН У ПАКЕТІ · З БУЛЬЙОНОМ','2':'РАМЕН У ПАКЕТІ · ІЗ СОУСОМ','3':'РАМЕН У СТАКАНІ','4':'РИСОВІ ГАЛУШКИ ТА ЛОКШИНА'}
SSHORT={'1':'Пакет · з бульйоном','2':'Пакет · із соусом','3':'Стакан','4':'Рисові'}
SCOL={'1':NAVY,'2':DARKAMBER,'3':PURPLE,'4':TEAL}
STITLE={'1':'Пакет з бульйоном','2':'Пакет із соусом','3':'Стакан','4':'Рисові галушки'}
SPREP={'1':'Варіння 4–5 хв · подається з бульйоном','2':'Варіння 4–5 хв · воду злити, змішати із соусом',
       '3':'Залити окропом 3–4 хв · без плити','4':'Рисова основа · окремий привід споживання'}
PB=[(0,80,"до 80"),(80,95,"80–95"),(95,110,"95–110"),(110,130,"110–130"),(130,150,"130–150"),(150,180,"150–180"),(180,9e9,"180 +")]
def band_of(p):
    for i,(lo,hi,_) in enumerate(PB):
        if lo<=p<hi: return i
    return len(PB)-1

PACK=[r for r in SD_ROWS if r['seg'] in (1,2) and r.get('pmed')]
PACK_P=sorted(r['pmed'] for r in PACK)
PACK_MED=st.median(PACK_P)
def quart(arr):
    """Квартили методом statistics.quantiles (exclusive) — як у статистиці сегментів."""
    return st.quantiles(sorted(arr),n=4)
PACK_Q=quart(PACK_P)

# ---------------- базові примітиви колоди ----------------
def estat(s,l,t,label,big,note,acc=DARKAMBER,w=3.88,h=1.46,bigsz=26):
    _rect(s,l,t,w,h,LIGHT); _rect(s,l,t,0.045,h,acc)
    txt(s,l+0.28,t+0.17,w-0.52,0.14,label,7.6,True,GREY)
    txt(s,l+0.28,t+0.40,w-0.52,0.46,big,bigsz,True,acc)
    ny=(t+h-0.52) if h>=1.40 else (t+0.40+bigsz*1.25/72+0.06)
    txt(s,l+0.28,ny,w-0.52,0.40,note,8.4,False,GREY,lh=1.3)

def ecall(s,l,t,w,h,head,body,acc=AMBER,hs=10,bs=8.8):
    _rect(s,l,t,w,h,LIGHT); _rect(s,l,t,0.045,h,acc)
    txt(s,l+0.30,t+0.18,w-0.60,0.18,head,hs,True,NAVY)
    txt(s,l+0.30,t+0.48,w-0.60,h-0.66,body,bs,False,GREY,lh=1.42)

def estrip(s,t,head,body,acc=NAVY,h=0.60):
    _rect(s,0.62,t,12.10,h,LIGHT); _rect(s,0.62,t,0.045,h,acc)
    txt(s,0.92,t+0.10,11.50,0.16,head,9.2,True,NAVY)
    txt(s,0.92,t+0.32,11.50,0.20,body,8.6,False,GREY,lh=1.3)

def verdict(s,t,head,body,acc=RED,h=0.92):
    _rect(s,0.62,t,12.10,h,NAVY); _rect(s,0.62,t,0.06,h,acc)
    txt(s,0.96,t+0.14,11.50,0.22,head,11,True,GOLD)
    txt(s,0.96,t+0.42,11.50,0.40,body,13.5,True,WHITE,lh=1.22)

def cardrow(s,t,items,acc=BLUE,n=7):
    for i,c in enumerate(items[:n]):
        card(s,CARD_X[i],t,c.get("img"),c["name"],c["w"],c["price"],c["sellers"],acc)

def minibar(s,l,t,w,rows,xmax,acc=BLUE,lab_w=1.05,val_w=0.62,rowh=0.205):
    bw=w-lab_w-val_w
    for lab,v,extra in rows:
        txt(s,l,t+0.02,lab_w-0.08,0.14,lab,7.8,False,NAVY)
        _rect(s,l+lab_w,t+0.045,max(bw*v/xmax,0.02),0.105,acc)
        txt(s,l+w-val_w,t+0.02,val_w,0.14,extra,7.8,True,GREY,PP_ALIGN.RIGHT)
        t+=rowh
    return t

def tint(c,k=0.55):
    r,g_,b=c[0],c[1],c[2]
    return RGBColor(int(r+(255-r)*k),int(g_+(255-g_)*k),int(b+(255-b)*k))
AMBLIGHT=RGBColor(0xFD,0xF3,0xDF)

def rangebars(s,l,t,w,rows,xmax,ticks,rowh=0.285,lab_w=2.60,w_col=1.05,n_col=0.45,val_w=1.15,band=None,bandlab=None,
              head=None,headcol=NAVY):
    """Діапазонні смуги за брендами: тонка лінія min–max, блок — медіана. rows=[dict(label,wt,n,lo,med,hi,col,bold)]."""
    px=l+lab_w+w_col+n_col+0.12; pw=w-(lab_w+w_col+n_col+0.12)-val_w
    def X(v): return px+pw*min(max(v,0),xmax)/xmax
    H=len(rows)*rowh
    if head:
        _rect(s,l,t,w,0.30,headcol)
        txt(s,l+0.14,t+0.07,lab_w,0.18,head,8.2,True,WHITE)
        txt(s,l+lab_w,t+0.08,w_col,0.14,"ВАГА",6.6,True,WHITE)
        txt(s,l+lab_w+w_col,t+0.08,n_col,0.14,"SKU",6.6,True,WHITE,PP_ALIGN.RIGHT)
        t+=0.36
    if band:
        _rect(s,X(band[0]),t-0.02,X(band[1])-X(band[0]),H+0.04,AMBLIGHT)
    for tv in ticks:
        _rect(s,X(tv),t-0.02,0.006,H+0.04,HAIR)
    y=t
    for rw in rows:
        b=rw.get('bold',False); col=rw['col']
        if b: _rect(s,l,y,w-0.02,rowh-0.02,AMBLIGHT)
        txt(s,l+0.10,y+rowh/2-0.095,lab_w-0.12,0.19,rw['label'],8.4,True,DARKAMBER if b else NAVY)
        txt(s,l+lab_w,y+rowh/2-0.085,w_col,0.17,rw.get('wt',''),7.6,False,GREY)
        txt(s,l+lab_w+w_col,y+rowh/2-0.085,n_col,0.17,str(rw.get('n','')),8,True,GREY,PP_ALIGN.RIGHT)
        lo,med,hi=rw['lo'],rw['med'],rw['hi']
        if hi>lo: _rect(s,X(lo),y+rowh/2-0.03,max(X(hi)-X(lo),0.02),0.06,tint(col,0.45) if not b else tint(col,0.25))
        _rect(s,X(med)-0.07,y+rowh/2-0.095,0.14,0.19,col)
        txt(s,X(hi)+0.10,y+rowh/2-0.09,val_w+0.3,0.18,rw.get('txt') or (f"{lo:.0f}–{hi:.0f}" if hi>lo+0.5 else f"{med:.0f}"),8.2,True,DARKAMBER if b else NAVY)
        y+=rowh
    for tv in ticks:
        txt(s,X(tv)-0.35,y+0.02,0.70,0.14,f"{tv}",7.2,False,GREY,PP_ALIGN.CENTER)
    if bandlab:
        txt(s,X(band[0])-1.6,y+0.18,X(band[1])-X(band[0])+3.2,0.14,bandlab,7.2,True,DARKAMBER,PP_ALIGN.CENTER)
    return y+0.34

def brand_rows(rows_,col,cap=None):
    """Групування SKU за брендом → рядки для rangebars (за зростанням медіани)."""
    d=co.defaultdict(list)
    for r in rows_:
        if r.get('pmed'): d[r['tm']].append(r)
    out=[]
    for tm,v in d.items():
        pm=[x['pmed'] for x in v]; ws=sorted(set(round(x['w']) for x in v))
        wt=(f"{ws[0]} г" if len(ws)==1 else f"{ws[0]}–{ws[-1]} г")
        out.append(dict(label=tm,wt=wt,n=len(v),lo=min(pm),med=st.median(pm),hi=max(pm),col=col))
    out.sort(key=lambda x:x['med'])
    return out[:cap] if cap else out

def slide_tail_sources(s,text,y=7.12):
    txt(s,0.62,y,12.10,0.30,text,7.2,False,PALE,lh=1.25)

prs=new_deck(); N=[0]
def num(): N[0]+=1; return f"{N[0]:02d}"

# ============ 01 ОБКЛАДИНКА ============
s=prs.slides.add_slide(prs.slide_layouts[6]); N[0]=1
_rect(s,0,0,13.333,7.5,NAVY); _rect(s,0,7.20,13.333,0.30,AMBER)
s.shapes.add_picture(LOGO,Inches(0.90),Inches(0.80),Inches(2.10),Inches(0.91))
txt(s,0.92,2.52,11.0,0.25,"ДОСЛІДЖЕННЯ РИНКУ ТА ПОЗИЦІОНУВАННЯ · ЖОВТЕНЬ 2026",10.5,True,GOLD)
txt(s,0.90,2.88,11.5,1.60,"Корейський рамен\nв Україні",46,True,WHITE,lh=1.05)
txt(s,0.92,4.98,11.0,0.70,
 "Оцінка шести SKU SIAS — рамену в пакеті — проти ринкової полиці та ціна закупівлі для полиці 120 грн:\n"
 f"{F['n_sku']} SKU корейського рамену · {F['n_brands']} брендів · {F['n_chains']} мереж із {F['chains_checked']} · зріз каталогів 29.09.2026 · митна база 2025–2026",
 11,False,PALE,lh=1.5)
_rect(s,0.92,5.96,0.05,0.92,AMBER)
txt(s,1.18,5.99,11.0,0.88,
 "Усі ціни — роздріб за одну пачку, з ПДВ, на дату зрізу. Опт, ящики та мультипаки до медіан не входять.\n"
 "Фото товарів — з каталогів мереж, як ідентифікація SKU. Фінансова модель — RAMEN_FINMODEL_4_TABS_PRESENTATION.xlsx.",
 9,False,PALE,lh=1.45)

# ============ 02 ШІСТЬ ЦИФР ============
cust_x=838.480003/239.67149
s=slide(prs,"ГОЛОВНЕ","Шість цифр, які описують ринок",num(),
 "Усі ціни в колоді — за одну упаковку, з ПДВ. Опт і ящики до розрахунку не входять.")
estat(s,0.62,1.78,"ПОЗИЦІЙ У ПРОДАЖУ","96","11 брендів · 4 сегменти · 13 мереж із 18. Кожна позиція має ціну в каталозі.",DARKAMBER,3.88,1.42,26)
estat(s,4.72,1.78,"ІМПОРТ ЗА МИТНИЦЕЮ","838,5 т","Січень–травень 2026, код УКТ ЗЕД 1902 30 10 00. За весь 2025 рік — 711,6 т.",TEAL,3.88,1.42,26)
estat(s,8.82,1.78,"2026 ПРОТИ 2025",f"×{d1(cust_x)}","Ті самі п’ять місяців: 239,7 → 838,5 т (+249,8 %); інвойс $1,02 → $3,20 млн.",PURPLE,3.90,1.42,26)
estat(s,0.62,3.34,"ІМПОРТ З КОРЕЇ","118,5 т","Січень–травень 2026: +116,3 % до 54,8 т торік. Це 14 % коду; лідер — В’єтнам, 166,7 т.",NAVY,3.88,1.42,26)
estat(s,4.72,3.34,"РОЗРИВ ЦІН ФОРМАТІВ","×1,8","За 100 г: стакан 154 грн проти 86 грн у пакеті з бульйоном.",DARKAMBER,3.88,1.42,26)
estat(s,8.82,3.34,"ЦІЛЬОВА ПОЛИЦЯ SIAS","120 грн",f"P75 серед {len(PACK)} пакетів: +{(120/PACK_MED-1)*100:.0f} % до медіани {d1(PACK_MED)} грн; дорожчих — 18 із 71.",RED,3.90,1.42,26)
ecall(s,0.62,4.94,5.92,1.80,"ЩО ЦЕ ОЗНАЧАЄ ДЛЯ ПОЛИЦІ",
 "1. Категорія повністю імпортна і швидко росте: ввіз за п’ять місяців 2026 року вже більший за весь 2025.\n"
 "2. Корея росте повільніше за код в цілому: її частка впала з 23 % до 14 %.\n"
 "3. Полиця концентрована: Samyang і Nongshim дають 54 позиції з 96 і стоять у 10–11 мережах.",AMBER)
ecall(s,6.80,4.94,5.92,1.80,"ЩО ПОТРІБНО ВІД SIAS",
 "1. Полиця 120 грн за моделлю з маржею 30 % вимагає закупівлі €0,54 (повна машина) або €0,32 (6 палет).\n"
 "2. Зараз SIAS пропонує €0,65–0,75: потрібне зниження — на 17–28 % для повної машини і на 51–57 % для 6 палет.\n"
 "3. Тому торг іде про ціну закупівлі, а не про ціну на полиці.",RED)
slide_tail_sources(s,"Джерела: каталоги мереж 29.09.2026 (96 SKU); митна база користувача, код 1902 30 10 00, січень–травень 2025 і 2026; фінмодель RAMEN_FINMODEL_4_TABS (вкладки «Цена 120 грн»).",6.92)

# ============ 03 ОБСЯГ РИНКУ · МИТНА БАЗА ============
s=slide(prs,"ОБСЯГ РИНКУ · МИТНА БАЗА","Скільки локшини ввозять в Україну",num(),
 "Код УКТ ЗЕД 1902 30 10 00: обсяг товарного коду, а не роздрібні продажі рамену. Однакові п’ять місяців для порівняння.")
txt(s,0.62,1.78,6.0,0.16,"ОБСЯГ ІМПОРТУ, ТОНН",8.6,True,NAVY)
hbars(s,0.62,2.10,6.20,[("2025 · січень–травень",239.67,GREYBAR,"239,7 т"),
                       ("2026 · січень–травень",838.48,AMBBAR,"838,5 т"),
                       ("2025 · увесь рік",711.6,NAVYBAR,"711,6 т")],
      900,[0,300,600,900],lab_w=1.95,val_w=1.20,rowh=0.50,bar_h=0.26)
txt(s,0.62,4.02,6.0,0.16,"ІНВОЙС, МЛН ДОЛАРІВ",8.6,True,NAVY)
hbars(s,0.62,4.34,6.20,[("2025 · січень–травень",1.02459,GREYBAR,"$1,02 млн"),
                       ("2026 · січень–травень",3.20062,AMBBAR,"$3,20 млн")],
      3.6,[0,1,2,3],lab_w=1.95,val_w=1.20,rowh=0.50,bar_h=0.26)
estat(s,7.10,1.78,"ОБСЯГ","+249,8 %","239,7 → 838,5 т за січень–травень: ×3,5.",AMBER,2.72,1.30,21)
estat(s,10.00,1.78,"ІНВОЙС","+212,4 %","$1,02 → $3,20 млн: ×3,1.",TEAL,2.72,1.30,21)
ecall(s,7.10,3.22,5.62,1.78,"ЩО КАЖЕ СЕРЕДНЯ ЦІНА ІМПОРТУ",
 "Митна вартість за кілограм: $4,27 у 2025 → $3,82 у 2026, тобто −10,7 %. Ввіз росте швидше за гроші: ймовірно, в обсязі додається дешевший товар.\n"
 "Для SIAS це означає: категорія розширюється за рахунок цінового низу, а не премії.",RED)
ecall(s,7.10,5.12,5.62,1.62,"ЧОГО ЦЕ НЕ ПОКАЗУЄ",
 "Код охоплює товар усіх країн походження, не лише корейський. Частки брендів і форматів у митній базі не видно — вони взяті з каталогів мереж.",NAVY)
ecall(s,0.62,5.80,6.20,0.96,"ЯК ЧИТАТИ ЦИФРИ",
 "Порівняння — за однакові п’ять місяців (01–05). За увесь 2025 рік — 711,6 т: за п’ять місяців 2026-го ввезено вже на 18 % більше.",AMBER)
slide_tail_sources(s,"Митна база користувача, код УКТ ЗЕД 1902 30 10 00, січень–травень 2025 і 2026; повний 2025 рік — 711,6 т. Митна вартість на кілограм — власний розрахунок: інвойс ÷ вага.",6.92)

# ============ 04 МИТНА БАЗА · ПОХОДЖЕННЯ ============
s=slide(prs,"МИТНА БАЗА · ПОХОДЖЕННЯ","Звідки їде: Корея — друга за обсягом країна",num(),
 "Січень–травень 2026, код УКТ ЗЕД 1902 30 10 00. Країни за вагою нетто, тонн.")
ctry=[("В’єтнам",166.687505),("Корея",118.497356),("Румунія",86.919064),("Велика Британія",58.564844),
      ("Латвія",57.30452),("Чехія",50.800755),("Польща",48.761462)]
oth=838.480003-sum(v for _,v in ctry)
txt(s,0.62,1.78,6.0,0.16,"ПО КРАЇНАХ ПОХОДЖЕННЯ, ТОНН",8.6,True,NAVY)
hbars(s,0.62,2.10,6.30,[(k,v,(AMBBAR if k=="Корея" else GREYBAR),f"{th1(v)} т") for k,v in ctry]+
      [("Інші країни разом",oth,NAVYBAR,f"{th1(oth)} т")],300,[0,100,200,300],lab_w=1.95,val_w=1.20,rowh=0.42,bar_h=0.22)
txt(s,7.10,1.78,5.6,0.16,"КОРЕЯ → УКРАЇНА · ПОРІВНЯННЯ ОДНАКОВИХ ПЕРІОДІВ",8.6,True,NAVY)
table(s,2.10,[("СІЧЕНЬ–ТРАВЕНЬ",7.10,1.90,L),("2025",9.00,1.20,R),("2026",10.30,1.20,R),("ЗМІНА",11.60,1.12,R)],
 [["Обсяг, т","54,8","118,5","+116,3 %"],
  ["Інвойс, тис. $","348,6","809,2","+132,1 %"],
  ["$ за кг","6,36","6,83","+7,4 %"],
  ["Частка в імпорті","22,9 %","14,1 %","−8,7 п.п."]],
 head_size=7.4,row_size=10,rh=0.38,left=7.10,bandw=5.62)
ecall(s,7.10,3.98,5.62,1.62,"КОРЕЯ ДОРОЖЧА ЗА СЕРЕДНЄ В 1,8 РАЗА",
 "Митна вартість корейського ввозу — $6,83/кг проти $3,82 в цілому по коду. Покупець у категорії вже платить премію за корейський товар.",TEAL)
ecall(s,7.10,5.78,5.62,1.00,"АЛЕ ЧАСТКА КОРЕЇ ПАДАЄ",
 "Код в цілому виріс на 250 %, Корея — на 116 %: різницю забрав дешевший ввіз.",AMBER)
ecall(s,0.62,5.78,6.30,1.00,"СКІЛЬКИ ЦЕ В ПАЧКАХ",
 "118,5 т при середній вазі пачки 120 г — близько 0,99 млн пачок за п’ять місяців. Це оцінка порядку величини, не продажі.",NAVY)
slide_tail_sources(s,"Митна база користувача, код УКТ ЗЕД 1902 30 10 00, січень–травень. Частка Кореї, $/кг і кількість пачок — власний розрахунок. Країна походження за митною декларацією.",6.92)

# ============ 05 СЕГМЕНТАЦІЯ ============
def pick_photo(k,must=None):
    arr=[c for c in SEG[k] if c.get('img') and os.path.exists(c['img']) and (not must or re.search(must,c['name'],re.I))]
    arr.sort(key=lambda c:-(c.get('nch') or 0)); return arr[0] if arr else None
PH={'1':pick_photo('1','Shin Ramyun'),'2':pick_photo('2','Carbonara'),'3':pick_photo('3'),'4':pick_photo('4','Yopokki')}
s=slide(prs,"СЕГМЕНТАЦІЯ","Сегментація за способом приготування",num(),
 "Чотири сегменти, 96 позицій: спершу спосіб приготування, далі вага, ціна й бренди. Належність наших шести SKU до сегментів 1 чи 2 потребує інструкцій приготування.")
_x=[0.62,3.68,6.74,9.80]
for i,k in enumerate(("1","2","3","4")):
    v=SS[k]; l=_x[i]; col=SCOL[k]
    _rect(s,l,1.78,2.90,0.40,col)
    txt(s,l+0.14,1.86,2.66,0.26,SNAME[k],7.8,True,WHITE,lh=1.1)
    _rect(s,l+0.07,2.26,2.76,1.14,LIGHT)
    ph=PH[k]
    if ph:
        try:
            from PIL import Image
            im=Image.open(ph['img']); ar=im.width/im.height; h=1.0; w=min(h*ar,2.6)
            if w/h>2.6: w=2.6; h=w/ar
            s.shapes.add_picture(ph['img'],Inches(l+0.07+(2.76-w)/2),Inches(2.26+(1.14-h)/2),Inches(w),Inches(h))
        except Exception: pass
    txt(s,l+0.14,3.46,1.30,0.40,f"{v['n']}",26,True,col)
    txt(s,l+1.10,3.62,1.76,0.20,f"позицій · {round(v['n']*100/F['n_sku'])} %",8.6,False,GREY)
    for j,(lab,val) in enumerate((("ВАГА",f"{v['wlo']}–{v['whi']} г"),("МЕДІАНА ЦІНИ",f"{v['med']} грн"),("ЦІНА ЗА 100 Г",f"{v['g']} грн"),("БРЕНДІВ",f"{v['brands']}"))):
        yy=4.00+j*0.27
        txt(s,l+0.14,yy,1.40,0.16,lab,7,True,GREY)
        txt(s,l+1.20,yy-0.02,1.58,0.20,val,9.6,True,NAVY,PP_ALIGN.RIGHT)
txt(s,0.62,5.14,8.0,0.16,"СЕГМЕНТ × ЦІНОВА СМУГА · ПОЗИЦІЙ І МЕДІАНА, грн за пачку",7.8,True,NAVY)
_grid=[[sum(1 for r in SD_ROWS if r['seg']==int(k) and r.get('pmed') and lo<=r['pmed']<hi) for lo,hi,_ in PB] for k in ("1","2","3","4")]
heat(s,0.62,5.30,[p[2] for p in PB],[SSHORT[k] for k in ("1","2","3","4")],_grid,rl=1.60,cw=1.10,ch=0.31,
     tail=[f"{SS[k]['n']} SKU · медіана {SS[k]['med']} грн" for k in ("1","2","3","4")],tail_w=2.30,tail_head="УСЬОГО")
slide_tail_sources(s,"Сегмент визначено за способом приготування й форматом упаковки на фото картки; медіана — за медіанними цінами позицій у каталогах 18 мереж, 29.09.2026. Сільпо (4 позиції без ваги) до сегментів не віднесено.",6.98)

# ============ 06 ЦІНОВІ ПОВЕРХИ ============
s=slide(prs,"ЦІНОВІ ПОВЕРХИ","Три поверхи локшини швидкого приготування",num(),
 "Поверх визначено за ціною за 100 г — це єдиний коректний вимір, коли пачки важать від 50 до 200 г. Межі взяті з фактичного розподілу брендів, а не задані наперед.")
tierpanel(s,0.62,1.74,3.90,3.40,BLUE,"ПОВЕРХ 1 · ЕКОНОМ","9–25 грн  ·  14–37 грн/100 г",
 "Українське виробництво і приватні марки мереж · 15 брендів",
 [tl("Мівіна"),tl("Reeva"),tl("Роллтон"),tl("Щедро")])
tierpanel(s,4.72,1.74,3.90,3.40,TEAL,"ПОВЕРХ 2 · МІДЛ","38–124 грн  ·  58–86 грн/100 г",
 "Азія та імпорт середньої ціни — сюди потрапляє й дешева Корея · 9 брендів",
 [tl("O'Food"),tl("Ottogi"),tl("Paldo"),tl("Yum Yum")])
tierpanel(s,8.82,1.74,3.90,3.40,DARKAMBER,"ПОВЕРХ 3 · ІМПОРТНИЙ ПРЕМІУМ","34–204 грн  ·  93–146 грн/100 г",
 "Samyang, Nongshim і в'єтнамська рисова локшина · 8 брендів",
 [tl("Samyang"),tl("Nongshim"),tl("Мама"),tl("Kool Cung Dinh","Cung Dinh Kool")])
_rect(s,0.62,5.24,12.10,0.76,LIGHT); _rect(s,0.62,5.24,0.045,0.76,RED)
per_t=[120/w*100 for w in (112.5,113,123,131)]
txt(s,0.90,5.32,3.50,0.14,"ДЕ ОПИНЯЮТЬСЯ НАШІ ШІСТЬ SKU",8.8,True,RED)
txt(s,0.90,5.52,3.50,0.40,f"{min(per_t):.0f}–{max(per_t):.0f} грн/100 г при полиці 120 грн;\n124–132 за поточною ціною",9.0,True,NAVY,lh=1.25)
txt(s,4.60,5.32,7.90,0.66,
 "При полиці 120 грн наші SKU стають між Samyang (93) і Nongshim (107 грн/100 г) — двома брендами, що тримають 54 позиції з 96. "
 "Поточна розрахункова полиця 140–162 грн виводить їх вище за обох.",8.6,False,NAVY,lh=1.38)
ecall(s,0.62,6.14,12.10,1.28,"ЩО ЦЕ ЗМІНЮЄ В КАРТИНІ КОНКУРЕНЦІЇ",
 "Корейський рамен — не однорідний преміум. Ottogi (75), Paldo (77), O’Food (67) і Bibigo (70 грн/100 г) стоять у МІДЛІ поруч із Yum Yum, а не поруч із Samyang і Nongshim; саме ця дешева половина (67–89 грн за пачку) — прямий ціновий конкурент новому бренду.\n"
 "На пропозицію впливає й фабрика Nestlé в Смолигові (Волинь): запущена в липні 2025, близько $50 млн; ≈ 75 % випуску — на експорт до ЄС під Maggi, решта — на внутрішній ринок (Latifundist, Just Food).",NAVY,9.4,8.2)

# ============ БЛОКИ СЕГМЕНТІВ ============
SEGROWS={k:[r for r in SD_ROWS if r['seg']==int(k) and r.get('pmed')] for k in "1234"}
NOW_LOW=M.forward_current("33")[2]['shelf']      # 140,55  (4 SKU, 33 палети, маржа 30 %)
NOW_HIGH=M.forward_current("33")[0]['shelf']     # 162,18  (2 SKU)
def n_above(arr,x): return sum(1 for p in arr if p>x)
def sample(k,n=7):
    arr=sorted([c for c in SEG.get(k,[]) if c.get('pmed')],key=lambda c:c['pmed']); out=[]
    for i in range(n):
        j=round(i*(len(arr)-1)/(n-1)) if len(arr)>1 else 0
        if arr[j] not in out: out.append(arr[j])
    return out
def band_text(k):
    v=SS[k]; return f"СЕГМЕНТ {k} · {SNAME[k]}   ·   {SPREP[k]}   ·   вага {v['wlo']}–{v['whi']} г"

def segoverview(k,title,sub,note,keys=None):
    v=SS[k]; col=SCOL[k]
    s=slide(prs,f"СЕГМЕНТ {k} · {SNAME[k]}",title,num(),sub)
    estat(s,0.62,1.78,"ПОЗИЦІЙ У СЕГМЕНТІ",f"{v['n']} SKU",f"{v['brands']} брендів у {v['chains']} мережах.\n{round(v['n']*100/F['n_sku'])} % усієї категорії.",col,3.10,1.22,22)
    estat(s,3.88,1.78,"МЕДІАНА ПОЛИЦІ",f"{v['med']} грн",f"P25–P75: {v['p25']}–{v['p75']} грн.\nДіапазон: {v['lo']}–{v['hi']} грн.",col,3.10,1.22,22)
    estat(s,7.14,1.78,"ЦІНА ЗА ГРАМ",f"{v['g']} грн","за 100 г — медіана сегмента.\nНаші SKU при полиці 120 грн: 92–107.",col,3.10,1.22,22)
    estat(s,10.40,1.78,"ВАГОВИЙ ДІАПАЗОН",f"{v['wlo']}–{v['whi']} г","фактичні межі сегмента\nза зрізом 29.09.2026.",col,2.32,1.22,19)
    txt(s,0.62,3.18,5.60,0.14,"ХТО ТРИМАЄ СЕГМЕНТ · SKU НА БРЕНД",8,True,NAVY)
    mx=max(n for _,n in v['blist'])
    minibar(s,0.62,3.42,5.30,[(b,n,f"{n}") for b,n in v['blist'][:6]],mx,col,lab_w=1.30,val_w=0.50,rowh=0.205)
    txt(s,6.60,3.18,6.12,0.14,"ЩО ЦЕ ОЗНАЧАЄ ДЛЯ НАС",8,True,NAVY)
    txt(s,6.60,3.42,6.12,1.30,note,8.8,False,GREY,lh=1.45)
    y=section_head(s,4.82,"ПРЕДСТАВНИКИ СЕГМЕНТА · ВІД НАЙДЕШЕВШОЇ ДО НАЙДОРОЖЧОЇ ПОЗИЦІЇ","фото — з каталогів мереж")
    cardrow(s,y+0.06,sample(k),col)
    return s

def catalog(k,items,i,m,extra_band=None):
    v=SS[k]; col=SCOL[k]
    pm=[c['pmed'] for c in items]; g_=[c['per100'] for c in items]
    lo,hi=min(pm),max(pm)
    ws=sorted(set(c['w'] for c in items))
    title=f"{STITLE[k]}: каталог {i} з {m} — {lo:.0f}–{hi:.0f} грн"
    s=slide(prs,f"СЕГМЕНТ {k} · {SNAME[k]}   ·   {v['n']} ПОЗИЦІЙ",title,num(),
            f"{len(items)} позицій за медіанною ціною {lo:.0f}–{hi:.0f} грн ({min(g_):.0f}–{max(g_):.0f} грн/100 г). Відсортовано від найдешевшої позиції.")
    band(s,1.68,band_text(k),col)
    cards(s,items,col)
    brands=co.Counter(c['tm'] for c in items).most_common(3)
    bt=", ".join(f"{b} ({n})" for b,n in brands)
    a=sum(1 for p in pm if p<=120); b_=len(pm)-a
    if 120<lo: pos=f"Уся ця частина дорожча за нашу ціль 120 грн (мінімум {lo:.0f} грн)."
    elif 120>hi: pos=f"Уся ця частина дешевша за нашу ціль 120 грн (максимум {hi:.0f} грн)."
    else: pos=f"Наша ціль 120 грн потрапляє в цю частину: дешевших за неї {a} із {len(pm)}, дорожчих — {b_}."
    mg=st.median(g_); ratio=99/mg
    gtxt=(f"Ціна за 100 г у цій частині — медіана {mg:.0f} грн; наші SKU при полиці 120 грн мають 92–107 — "+
          (f"у {d1(ratio)} раза більше." if ratio>=1.15 else ("на рівні." if ratio>0.87 else f"на {(1-ratio)*100:.0f} % менше.")))
    ecall(s,0.62,6.05,12.10,0.90,"ЩО ЦЕ ОЗНАЧАЄ ДЛЯ ПОЛИЦІ",
          f"Бренди частини: {bt}. {pos} {gtxt}",col,9.4,8.4)
    return s
def catalog_series(k,items=None):
    items=items if items is not None else SEG[k]
    m=math.ceil(len(items)/14); size=math.ceil(len(items)/m)
    for i in range(m): catalog(k,items[i*size:(i+1)*size],i+1,m)

ALLCH=[("auchan","Ашан",1),("novus","NOVUS",1),("metro","METRO",1),("zaraz","«Зараз»",1),("silpo","«Сільпо»",1),
       ("megamarket","МегаМаркет",1),("ultramarket","Ultramarket",1),("ekomarket","ЕКО",1),("epicentr","Епіцентр",1),
       ("grono","Grono",0),("onde","Onde",0),("kharkiv","«Клас»",0),("vostorg","«Восторг»",0),("cosmos","«Космос»",0),
       ("tavriav","«Таврія В»",0),("chudomarket","«ЧудоМаркет»",0),("torba","«Торба»",0),("ideal","«Ідеал»",0)]
SILPO={'1':1,'2':3,'3':0,'4':0}
def chipflow(s,x,y,w,items,col,chip_h=0.30,gap=0.10):
    cx=x; cy=y
    for name,cnt,star in items:
        cw=0.24+0.078*len(name)+(0.46 if cnt is not None else 0.0)
        if cx+cw>x+w: cx=x; cy+=chip_h+0.08
        _rect(s,cx,cy,cw,chip_h,LIGHT)
        txt(s,cx+0.10,cy+0.065,cw-(0.5 if cnt is not None else 0.14),0.18,name,8.4,False,INK)
        if cnt is not None:
            _rect(s,cx+cw-0.40,cy+0.04,0.34,0.22,col)
            txt(s,cx+cw-0.40,cy+0.065,0.34,0.17,(f"{cnt}*" if star else f"{cnt}"),8,True,WHITE,PP_ALIGN.CENTER)
        cx+=cw+gap
    return cy+chip_h
def wheresold(keys,title,kick,note):
    col=SCOL[keys[0]]
    cnt=co.Counter()
    for r in SD_ROWS:
        if str(r['seg']) in keys:
            for ch in r['chains']: cnt[ch]+=1
    star=set()
    for k in keys:
        if SILPO[k]: cnt['silpo']+=SILPO[k]; star.add('silpo')
    tot=sum(SS[k]['n'] for k in keys)
    present=[(c,n) for c,n in cnt.most_common() if n>0]
    s=slide(prs,kick,title,num(),
            f"{len(present)} мереж із 18 мають позиції цього сегмента; число поруч — скільки позицій сегмента є в мережі.")
    txt(s,0.62,1.80,5.0,0.16,"ПОЗИЦІЙ СЕГМЕНТА В МЕРЕЖІ",8.2,True,NAVY)
    nm={c:n for c,n,_ in ALLCH}
    rows=[(nm[c],n,col if c not in star else tint(col,0.35),(f"{n}*" if c in star else f"{n}")) for c,n in present]
    hbars(s,0.62,2.12,5.70,rows,max(n for _,n in present)*1.1,[0],lab_w=1.55,val_w=0.70,rowh=min(0.34,4.45/len(rows)),bar_h=0.19)
    # праворуч — групи
    nat=[(n_,cnt.get(c,0),c in star) for c,n_,t in ALLCH if t==1 and cnt.get(c,0)>0]
    reg=[(n_,cnt.get(c,0),c in star) for c,n_,t in ALLCH if t==0 and cnt.get(c,0)>0]
    zero=[(n_,None,False) for c,n_,t in ALLCH if cnt.get(c,0)==0]
    y=1.80
    def grp(y,head,tag,tagcol,items,h):
        _rect(s,6.60,y,0.045,h,tagcol)
        txt(s,6.82,y+0.06,3.80,0.20,head,10.5,True,NAVY)
        _rect(s,11.10,y+0.07,1.62,0.22,tagcol); txt(s,11.10,y+0.10,1.62,0.16,tag,6.4,True,WHITE,PP_ALIGN.CENTER)
        chipflow(s,6.82,y+0.42,5.88,items,tagcol)
    ny=0.42+0.38*math.ceil(max(len(nat),1)*1.9/5.9)+0.15
    grp(y,f"Національні мережі · {len(nat)}","ЦІЛЬОВИЙ КАНАЛ",RED,nat,1.52); y+=1.52+0.10
    grp(y,f"Регіональні мережі · {len(reg)}","РЕГІОНАЛЬНИЙ ВХІД",TEAL,reg,1.52); y+=1.52+0.10
    grp(y,f"Без позицій сегмента · {len(zero)}","ВІЛЬНА ПОЛИЦЯ",PALE,zero,1.62)
    slide_tail_sources(s,note,6.98)
    return s

# ============ СЕГМЕНТ 1 ============
_p1=[r['pmed'] for r in SEGROWS['1']]; _p2=[r['pmed'] for r in SEGROWS['2']]
S1,S2,S3,S4=SS['1'],SS['2'],SS['3'],SS['4']
segoverview("1","Пакет з бульйоном: найбільший сегмент",
 "Основна частина корейської полиці: локшина й порошковий бульйон, вариться в каструлі. Показано сім позицій, рівновіддалених по ціні.",
 f"Сюди найімовірніше потраплять чотири наші SKU по 112,5–113 г — якщо їх спосіб приготування варіння з бульйоном (не підтверджено).\n\n"
 f"Ціль 120 грн — на {(120/S1['med']-1)*100:.0f} % вище за медіану сегмента ({S1['med']} грн) і вище за P75 ({S1['p75']} грн): дорожчих за 120 лише {n_above(_p1,120)} із {S1['n']}. "
 f"Поточна розрахункова полиця {d2(NOW_LOW)} грн дорожча за всіх, крім {n_above(_p1,NOW_LOW)}.")
catalog_series("1")
wheresold(["1"],"Пакет з бульйоном: де продається","СЕГМЕНТ 1 · ДЕ ПРОДАЄТЬСЯ",
 "Суцільна перевірка каталогів 18 мереж, 29.09.2026. * «Сільпо» — позиція без ваги в каталозі, віднесена за назвою (Nongshim Shin Ramyun).")

# ============ СЕГМЕНТ 2 ============
segoverview("2","Пакет із соусом: 30 позицій, медіана 106 грн",
 "Локшина, з якої воду зливають, а смак дає соус: Buldak, Toomba, Chapaghetti, Volcano, Carbonara. Показано сім позицій, рівновіддалених по ціні.",
 f"Сюди найімовірніше потраплять Gouda Cheese 123 г і Carbonara Spicy 131 г — якщо їх спосіб приготування це підтвердить.\n\n"
 f"Ціль 120 грн — на {(120/S2['med']-1)*100:.0f} % вище за медіану сегмента ({S2['med']} грн), але нижче P75 ({S2['p75']} грн): дорожчих за 120 — {n_above(_p2,120)} із {S2['n']}. "
 f"Прямі аналоги за вагою: Paldo Volcano Carbonara 130 г — 114 грн, Choi’s Carbonara 131 г — 125,20 грн. Поточна полиця {d2(NOW_HIGH)} грн дорожча за {S2['n']-n_above(_p2,NOW_HIGH)} із {S2['n']}.")
catalog_series("2")
wheresold(["2"],"Пакет із соусом: де продається","СЕГМЕНТ 2 · ДЕ ПРОДАЄТЬСЯ",
 "Суцільна перевірка каталогів 18 мереж, 29.09.2026. * «Сільпо» — три позиції Samyang Buldak без ваги в каталозі, віднесені за назвою.")

# ============ СЕГМЕНТ 3 І 4 ============
segoverview("3","Стакан: найдорожчий формат за грам",
 "Заливається окропом, готується без плити: формат для споживання поза домом. Наші шість SKU — пакет, тому це ринковий контекст, а не пряма конкуренція.",
 f"Ринок вже платить премію за зручність: {S3['g']} грн/100 г проти {S1['g']} у пакеті з бульйоном — у {S3['g']/S1['g']:.1f} раза більше.\n\n"
 f"Сегмент вузький: {S3['n']} позицій у {S3['chains']} мережах і лише {S3['brands']} бренди. Якщо в лінійці SIAS з’явиться стакан, це єдиний формат, де премія за грам уже прийнята ринком; у поточній пропозиції стакана немає.")
_cups=SEG['3']; _rice=SEG['4']
catalog("3",_cups[:14],1,2)
# каталог 2 з 2: решта стаканів + увесь сегмент 4
_rest=_cups[14:]
s=slide(prs,"СЕГМЕНТИ 3 І 4 · КАТАЛОГ 2 З 2","Стакан (решта) і рисові галушки",num(),
        f"Решта стакана — {len(_rest)} позиції; сегмент 4 — {len(_rice)} позицій: рисові галушки й рисова локшина. Відсортовано від найдешевшої позиції.")
band(s,1.68,f"СЕГМЕНТ 3 · {SNAME['3']}   ·   решта сегмента   ·   {SPREP['3']}",SCOL['3'])
for i,c in enumerate(_rest[:7]): card(s,CARD_X[i],2.06,c.get("img"),c["name"],c["w"],c["price"],c["sellers"],SCOL['3'])
band(s,3.98,f"СЕГМЕНТ 4 · {SNAME['4']}   ·   {S4['n']} позицій   ·   медіана {S4['med']} грн   ·   {S4['g']} грн/100 г   ·   вага {S4['wlo']}–{S4['whi']} г",SCOL['4'])
for i,c in enumerate(_rice[:7]): card(s,CARD_X[i],4.28,c.get("img"),c["name"],c["w"],c["price"],c["sellers"],SCOL['4'])
ecall(s,0.62,6.14,12.10,0.84,"ЧОМУ СЕГМЕНТ 4 НЕ Є ОРІЄНТИРОМ",
 f"Yopokki за 198–204 грн — рисовий снек (топокі), а не локшина. Використовувати цю ціну як доказ, що «корейське може коштувати 200 грн», не можна: у сегменті {S4['n']} позицій і {S4['chains']} мережі.",SCOL['4'],9.4,8.4)
wheresold(["3","4"],"Стакан і рисові: де продається","СЕГМЕНТИ 3 І 4 · ДЕ ПРОДАЄТЬСЯ",
 "Суцільна перевірка каталогів 18 мереж, 29.09.2026. Сегменти 3 і 4 разом; «Сільпо» позицій цих сегментів не має.")

# ============ ФОРМАТ × ВАГА ============
s=slide(prs,"ФОРМАТ × ВАГА","Де насправді щільна полиця",num(),
 "Кожна клітинка — скільки позицій корейського рамену має цей формат у цій ваговій смузі. Насиченість кольору пропорційна кількості.")
txt(s,0.62,1.78,12.10,0.15,"МАТРИЦЯ ФОРМАТ × ВАГА · КІЛЬКІСТЬ SKU",8,True,NAVY)
y=matrix(s,0.62,2.02,[b['lab'] for b in F['bands']],["Пакет (пачка)","Стакан"],F['matrix'],cw=1.70,ch=0.54,rl=2.30)
WB={'40–90 г':'стакани й міні-пачки: дорого за грам, дешево за чек','90–110 г':'перехідна смуга, мало позицій',
 '110–125 г':'найщільніша смуга ринку — Shin Ramyun, Ottogi, Paldo','125–145 г':'смуга преміальних лінійок і соусних SKU',
 '145 г +':'рисові галушки та нестандарт, інший привід споживання'}
y2=section_head(s,y+0.30,"ВАГОВІ СМУГИ РИНКУ","частка SKU, медіана полиці та ціна за грам")
cols=[("ВАГОВА СМУГА",0.62,2.10,L),("SKU",2.76,0.70,R),("ЧАСТКА",3.56,1.00,R),("МЕДІАНА",4.66,1.20,R),
      ("ГРН/100 Г",5.96,1.20,R),("ЩО ЦЕ ЗА СМУГА",7.32,5.40,L)]
table(s,y2,cols,[[b['lab'],str(b['n']),f"{b['share']} %",f"{b['pmed']} грн",str(b['g']),WB[b['lab']]] for b in F['bands']],
      head_size=7.4,row_size=9.6,rh=0.345)
ecall(s,0.62,6.46,5.92,0.92,"ВАГА НАШИХ SKU — РИНКОВА",
 f"112,5–131 г потрапляє у дві найщільніші смуги: {F['bands'][2]['share']+F['bands'][3]['share']} % усіх позицій ринку. За форматом і грамажем заперечень немає.",TEAL)
ecall(s,6.80,6.46,5.92,0.92,"ЦІНА ЗА ГРАМ — ВИЩА ЗА ВІДПОВІДНІ СМУГИ",
 f"При полиці 120 грн наші SKU — 92–107 грн/100 г проти {F['bands'][2]['g']}–{F['bands'][3]['g']} у їхніх смугах: на 7–37 % вище. За поточною розрахунковою ціною 124–132 грн/100 г.",RED)

# ============ НАЙПОПУЛЯРНІШІ ФОРМАТИ ============
_cc=co.defaultdict(list)
for _r in SD_ROWS:
    _f="стакан" if _r['seg']==3 else ("рисові галушки" if _r['seg']==4 else "пакет")
    _cc[(_f,round(_r['w']))].append(_r)
_combos=sorted([(k,v) for k,v in _cc.items() if len(v)>=2],key=lambda x:-len(x[1]))
_cov=sum(len(v) for _,v in _combos)
s=slide(prs,"НАЙПОПУЛЯРНІШІ ФОРМАТИ","Топ поєднань формат + вага",num(),
 f"Поєднання з двома і більше позиціями. У {len(_combos)} таких клітинках лежать {_cov} позицій із {len(SD_ROWS)} — решта розсіяна по унікальних грамажах.")
_top=_combos[0][1]; _top_med=st.median([x['pmed'] for x in _top])
estat(s,0.62,1.76,"УПАКОВКА №1 НА РИНКУ","пакет 120 г","22 позиції, медіана 93 грн, 77 грн/100 г.\nNongshim, Ottogi, Paldo — 11 мереж.",DARKAMBER,3.88,1.28,24)
estat(s,4.72,1.76,"УПАКОВКА №2","пакет 130 г","9 позицій, медіана 114 грн, 87 грн/100 г.\nСюди лягає наша Carbonara 131 г.",TEAL,3.88,1.28,24)
estat(s,8.82,1.76,"НАЙДЕШЕВШИЙ ФОРМАТ","стакан 80 г","4 позиції, медіана 70 грн.\nSamyang Buldak у стакані.",PURPLE,3.88,1.28,24)
estat(s,0.62,3.18,"ВАГА ЧОТИРЬОХ НАШИХ SKU","113 г","Жодної позиції ринку в цій вазі.\nНайближча щільна клітинка — 120 г.",RED,3.88,1.28,24)
estat(s,4.72,3.18,"РІЗНИЦЯ ДО СТАНДАРТУ","113 проти 120 г","На 7 г легше за стандарт категорії.\nЗа однакової ціни це −6 % продукту.",RED,3.88,1.28,24)
estat(s,8.82,3.18,"ПОКРИТТЯ СІТКИ",f"{_cov} з {len(SD_ROWS)}","позицій лежать у повторюваних\nпоєднаннях формату й ваги.",NAVY,3.88,1.28,24)
y=section_head(s,4.64,"ТОП ПОЄДНАНЬ ФОРМАТ × ВАГА","хто їх тримає і чи є там наші SKU")
_CH={"пакет:120":"ні — наші 112,5–113 г","пакет:130":"так — Carbonara 131 г","пакет:140":"ні",
     "пакет:105":"ні","рисові галушки:145":"ні","стакан:67":"ні","стакан:80":"ні",
     "пакет:110":"ні","стакан:68":"ні","стакан:62":"ні","пакет:122":"поруч — Gouda 123 г","пакет:137":"ні"}
cols=[("ФОРМАТ",0.62,1.70,L),("ВАГА",2.40,0.80,R),("SKU",3.30,0.60,R),("МЕРЕЖ",4.00,0.80,R),
      ("МЕДІАНА",4.90,1.10,R),("ГРН/100 Г",6.10,1.10,R),("ХТО ТРИМАЄ",7.34,3.30,L),("ЧИ Є ТУТ НАШІ SKU",10.70,2.02,L)]
rows=[]
for (f,w),v in _combos[:5]:
    br=co.Counter(x['tm'] for x in v).most_common(3)
    rows.append([f,f"{w} г",str(len(v)),str(len(set(c for x in v for c in x['chains']))),
        f"{st.median([x['pmed'] for x in v]):.0f} грн",f"{st.median([x['per100'] for x in v]):.0f}",
        " · ".join(f"{b} ({n})" for b,n in br),_CH.get(f"{f}:{w}","ні")])
table(s,y,cols,rows,head_size=7.4,row_size=9.6,rh=0.31)
estrip(s,6.74,"ГОЛОВНЕ ПРО УПАКОВКУ",
 f"Стандарт категорії — пакет 120 г. Чотири наші SKU на 7 г легші; при полиці 120 грн вони на {(120/_top_med-1)*100:.0f} % дорожчі за медіану цієї клітинки ({_top_med:.0f} грн), за поточної розрахункової 140,55 грн — на {(NOW_LOW/_top_med-1)*100:.0f} %.",RED,0.62)

# ============ ЦІНА ЗА ОДНУ УПАКОВКУ ============
NOWS={"33":M.forward_current("33"),"6":M.forward_current("6")}
def ours_rows(shelfnow33,shelfnow6,wt,nsku):
    return [dict(label="НАШ · SIAS · ціль 120 грн",wt=wt,n=nsku,lo=120,med=120,hi=120,col=AMBBAR,bold=True,txt="120 грн"),
            dict(label="НАШ · SIAS · зараз, 33 палети",wt=wt,n=nsku,lo=shelfnow33,med=shelfnow33,hi=shelfnow33,col=tint(AMBBAR,0.35),bold=True,txt=f"{d2(shelfnow33)} грн"),
            dict(label="НАШ · SIAS · зараз, 6 палет",wt=wt,n=nsku,lo=shelfnow6,med=shelfnow6,hi=shelfnow6,col=tint(AMBBAR,0.35),bold=True,txt=f"{d2(shelfnow6)} грн")]
def priceslide(part,title,sub,k,ours=None,xmax=240,note=None):
    keys=[k] if isinstance(k,str) else list(k)
    rows_=[r for r in SD_ROWS if str(r['seg']) in keys and r.get('pmed')]
    rr=brand_rows(rows_,SCOL[keys[0]])
    sl=slide(prs,f"ЦІНА ЗА ОДНУ УПАКОВКУ · {part}",title,num(),sub)
    nrows=len(rr)+(3 if ours else 0)
    rowh=min(0.31,4.7/(nrows+1))
    v=SS[keys[0]]
    y=rangebars(sl,0.62,1.76,12.10,rr,xmax,[0,40,80,120,160,200,240][:int(xmax//40)+1],rowh=rowh,
                band=(v['p25'],v['p75']),bandlab=f"вікно P25–P75 сегмента: {v['p25']}–{v['p75']} грн",
                head=f"СЕГМЕНТ {keys[0]} · {STITLE[keys[0]].upper()}",headcol=SCOL[keys[0]])
    if ours:
        y=rangebars(sl,0.62,y+0.04,12.10,ours,xmax,[],rowh=rowh,band=(v['p25'],v['p75']),
                    head="НАШІ SKU · ПОЛИЦЯ",headcol=NAVY)
    if ours:
        cover=sum(1 for r in rr if r['lo']<=120<=r['hi'])
        a33=ours[1]['med']; v=SS[keys[0]]
        ecall(sl,0.62,5.98,12.10,0.96,"ДЕ ЦЕ СТАВИТЬ НАС",
              f"Ціль 120 грн — на {(120/v['med']-1)*100:.0f} % вище за медіану сегмента ({v['med']} грн) і {'вище' if 120>v['p75'] else 'нижче'} за P75 ({v['p75']} грн); діапазон перетинає 120 грн у {cover} із {len(rr)} брендів. "
              f"Поточна розрахункова ціна {d2(a33)} грн (33 палети) — на {(a33/v['med']-1)*100:.0f} % вище за медіану.",AMBER,9.6,8.6)
    if note: slide_tail_sources(sl,note,7.04)
    return sl
priceslide("1 З 3","Ціна: пакет з бульйоном і наші SKU",
 f"Сегмент 1 — {S1['n']} позицій: кожен бренд від найдешевшої до найдорожчої. Блок — медіана бренду, лінія — діапазон, тло — вікно P25–P75.",
 "1",ours_rows(NOWS["33"][2]['shelf'],NOWS["6"][2]['shelf'],"112,5–113 г",4),240,
 "Роздрібна ціна за пачку з ПДВ, медіана по мережах (29.09.2026). Наші SKU — полиця за фінмоделлю: ціль 120 грн; «зараз» — поточна собівартість Self-Cost, маржа 30 %, бонус 25 %, націнка 1,40.")
priceslide("2 З 3","Ціна: пакет із соусом і наші SKU",
 f"Сегмент 2 — {S2['n']} позицій. Choi’s — окремий конкурент, не наш товар. Блок — медіана бренду, лінія — діапазон, тло — вікно P25–P75.",
 "2",ours_rows(NOWS["33"][0]['shelf'],NOWS["6"][0]['shelf'],"123–131 г",2),240,
 "Роздрібна ціна за пачку з ПДВ, медіана по мережах (29.09.2026). Наші SKU — Gouda 123 г і Carbonara 131 г: ціль 120 грн; «зараз» — Self-Cost, маржа 30 %, бонус 25 %, націнка 1,40.")
s=slide(prs,"ЦІНА ЗА ОДНУ УПАКОВКУ · 3 З 3","Ціна: стакан і рисові галушки",num(),
 "Суміжні сегменти: наших SKU тут немає. Показано, як ринок ціноутворює формат стакана і рисовий снек.")
rr3=brand_rows([r for r in SD_ROWS if r['seg']==3 and r.get('pmed')],SCOL['3'])
y=rangebars(s,0.62,1.76,12.10,rr3,240,[0,40,80,120,160,200,240],rowh=0.30,band=(S3['p25'],S3['p75']),
            bandlab=f"вікно P25–P75 сегмента 3: {S3['p25']}–{S3['p75']} грн",head="СЕГМЕНТ 3 · СТАКАН",headcol=SCOL['3'])
rr4=brand_rows([r for r in SD_ROWS if r['seg']==4 and r.get('pmed')],SCOL['4'])
y=rangebars(s,0.62,y+0.10,12.10,rr4,240,[],rowh=0.30,head="СЕГМЕНТ 4 · РИСОВІ ГАЛУШКИ",headcol=SCOL['4'])
ecall(s,0.62,5.70,5.92,1.20,"СТАКАН: ДЕШЕВИЙ ЗА ЧЕК, ДОРОГИЙ ЗА ГРАМ",
 f"Ціна за пачку 65–195 грн, медіана {S3['med']} грн — як у пакеті, але вага вдвічі менша: {S3['g']} грн/100 г проти {S1['g']} у пакеті з бульйоном.",SCOL['3'],9.6,8.6)
ecall(s,6.80,5.70,5.92,1.20,"РИСОВІ ГАЛУШКИ — ІНШИЙ ПРИВІД",
 f"Yopokki 198–204 грн — снек, не локшина; без нього сегмент коштує 84–112 грн. Орієнтиром для рамену в пакеті він бути не може.",SCOL['4'],9.6,8.6)
slide_tail_sources(s,"Роздрібна ціна за пачку з ПДВ, медіана по мережах (29.09.2026). Вага стакана — 62–120 г, рисових — 80–145 г.",7.02)

# ============ РИТЕЙЛ-АУДИТ · МЕРЕЖІ ============
SCALE={'auchan':'національна','novus':'національна','metro':'національна · C&C','zaraz':'національна'}
BRANDS_IN={'auchan':'Samyang · Nongshim · Paldo · Ottogi · O’Food · Bibigo','grono':'Samyang · Nongshim · Paldo · Yopokki',
 'onde':'Samyang · Nongshim · Ottogi','metro':'Samyang · Nongshim','zaraz':'Samyang · Nongshim · Ottogi',
 'vostorg':'Samyang · Nongshim','kharkiv':'Samyang · Nongshim','cosmos':'Samyang · Nongshim · Ottogi · Choi’s',
 'tavriav':'Samyang · Nongshim · Ottogi · Choi’s','novus':'Samyang · Nongshim','chudomarket':'Samyang · Nongshim','torba':'Samyang'}
s=slide(prs,"РИТЕЙЛ-АУДИТ · МЕРЕЖІ","Що з корейського рамену стоїть у мережах",num(),
 "Суцільна перевірка онлайн-каталогів 18 мереж, зріз 29.09.2026. SKU — унікальні позиції корейського рамену в каталозі мережі.")
cols=[("МЕРЕЖА",0.62,2.60,L),("ПОКРИТТЯ",3.22,1.80,L),("SKU",5.02,0.70,R),("МЕДІАНА",5.82,1.10,R),
      ("ДІАПАЗОН",7.02,1.50,R),("БРЕНДИ НА ПОЛИЦІ",8.70,4.02,L)]
rows=[]
for r in F['chain_rows']:
    rows.append([NM.get(r['chain'],r['chain']),SCALE.get(r['chain'],'регіональна'),str(r['n']),
                 f"{r['med']} грн",f"{r['lo']}–{r['hi']}",BRANDS_IN.get(r['chain'],'')])
rows.append(["«Сільпо»","національна · №1","4","114 грн","90–199","Samyang (3) · Nongshim (1)"])
rows.append(["МегаМаркет · Ultramarket · ЕКО","національна","0","—","—","категорії немає"])
rows.append(["Епіцентр · «Ідеал»","національна / регіон.","0","—","—","категорії немає"])
table(s,1.86,cols,rows,head_size=7.4,row_size=9.2,rh=0.282)
ecall(s,0.62,6.52,3.88,0.90,"ДІРКА №1 · «СІЛЬПО»",
 "Найбільша мережа країни тримає 4 позиції. Полиця відкрита, але вхід там 90–199 грн.",AMBER,9.2,8.4)
ecall(s,4.72,6.52,3.88,0.90,"ДІРКА №2 · П’ЯТЬ МЕРЕЖ З НУЛЕМ",
 "МегаМаркет, Ultramarket, ЕКО, Епіцентр, «Ідеал» — жодної позиції у зрізі.",AMBER,9.2,8.4)
ecall(s,8.82,6.52,3.90,0.90,"ОДИН ПОСТАЧАЛЬНИК НА ДВІ МЕРЕЖІ",
 "«Космос» і «Таврія В» мають ідентичні 12 позицій і той самий діапазон цін.",NAVY,9.2,8.4)

# ============ РИТЕЙЛ-АУДИТ · ПОЗИЦІЇ ============
CHCOLS=[("auchan","Ашан"),("grono","Grono"),("onde","Onde"),("metro","METRO"),("zaraz","«Зараз»"),("kharkiv","«Клас»"),
        ("vostorg","«Восторг»"),("cosmos","Космос"),("tavriav","Таврія"),("novus","NOVUS"),("chudomarket","«ЧудоМ.»"),
        ("torba","«Торба»"),("silpo","«Сільпо»")]
BR=["Samyang","Nongshim","Paldo","Ottogi","O'Food","Bibigo","Choi's"]
def brand_chain_grid():
    grid=[]; tails=[]
    for b in BR+["Інші"]:
        row=[]
        rows_b=[r for r in SD_ROWS if (r['tm']==b if b!="Інші" else r['tm'] not in BR)]
        for ch,_ in CHCOLS:
            n=sum(1 for r in rows_b if ch in r['chains'])
            if ch=="silpo": n={"Samyang":3,"Nongshim":1}.get(b,0)
            row.append(n)
        grid.append(row)
        nch=len({c for r in rows_b for c in r['chains']}|({"silpo"} if b in("Samyang","Nongshim") else set()))
        tails.append(f"{len(rows_b)} SKU · {nch} мер.")
    return grid,tails
G,TAIL=brand_chain_grid()
s=slide(prs,"РИТЕЙЛ-АУДИТ · ПОЗИЦІЇ","Полиця не порожня — вона зайнята і поділена",num(),
 "Скільки позицій кожного бренду стоїть у кожній мережі. Насиченість кольору пропорційна кількості. Зріз 29.09.2026.")
heat(s,0.62,1.86,[c[1] for c in CHCOLS],BR[:6]+["Choi’s","Інші бренди"],G,rl=1.36,cw=0.68,ch=0.44,tail=TAIL,tail_w=1.66,tail_head="ВСЬОГО У ЗРІЗІ")
estat(s,0.62,5.78,"МЕРЕЖ БЕЗ КАТЕГОРІЇ","5 із 18","МегаМаркет, Ultramarket, ЕКО,\nЕпіцентр, «Ідеал» — нуль позицій.",TEAL,3.88,1.14,22)
estat(s,4.72,5.78,"У «СІЛЬПО»","4 SKU","Найбільша мережа країни.\nДіапазон входу — 90–199 грн.",DARKAMBER,3.88,1.14,22)
estat(s,8.82,5.78,"ДВА ЛІДЕРИ ДАЮТЬ","54 з 96 SKU","Samyang і Nongshim — у 10–11 мережах.\nНова марка заходить проти них.",RED,3.90,1.14,22)
slide_tail_sources(s,"Вхід у категорію вже пробитий: імпортер і місце на полиці під Samyang і Nongshim існують; вільне місце є лише там, де категорії немає зовсім. «Сільпо» — за назвою позицій (без ваги).",7.04)

# ============ КАНАЛИ ПРОДАЖУ ============
s=slide(prs,"КАНАЛИ ПРОДАЖУ","Де продається категорія і кому",num(),
 "Ліворуч — скільки точок у кожному каналі перевірено у зрізі. Праворуч — хто туди приходить і чи це наш покупець.")
txt(s,0.62,1.80,5.60,0.14,"ТОЧОК ПРОДАЖУ В КАНАЛІ · ПЕРЕВІРЕНО У ЗРІЗІ",8,True,NAVY)
hbars(s,0.62,2.12,5.90,[("Продуктові мережі",18,NAVYBAR,"18  ·  13 мають рамен"),("Азійські фудшопи",5,TEALBAR,"5"),("Маркетплейси",4,AMBBAR,"4")],
      20,[0,5,10,15,20],lab_w=1.95,val_w=1.70,rowh=0.46,bar_h=0.22)
ecall(s,0.62,4.10,5.90,2.70,"КАНАЛ І ПОКУПЕЦЬ — РІЗНІ ОСІ",
 "Продуктові мережі — єдиний канал з обсягом, але саме там уже стоять Samyang і Nongshim.\n\n"
 "Маркетплейси й азійські фудшопи дають орієнтир ціни, але не обсяг: медіана на Prom — 135 грн, у фудшопах 61–158 грн. Умови бонусу в цих каналах не перевірялися: у моделі закладено 25 % для мережевого роздробу.\n\n"
 "На відміну від готового рису, outdoor- і мілітарі-рітейл для рамену не канал: рамен потребує окропу, тому в туристичних магазинах його немає.\n\n"
 "Практичний вхід — «Сільпо» і п’ять нацмереж із нулем.",NAVY,10,8.8)
def chcard(t2,tag,tagcol,head,sub,who,body):
    _rect(s,6.80,t2,5.92,1.26,LIGHT); _rect(s,6.80,t2,0.045,1.26,tagcol)
    txt(s,7.06,t2+0.10,3.60,0.20,head,11,True,NAVY)
    _rect(s,10.98,t2+0.10,1.64,0.22,tagcol); txt(s,10.98,t2+0.13,1.64,0.16,tag,6.2,True,WHITE,PP_ALIGN.CENTER)
    txt(s,7.06,t2+0.36,5.42,0.13,sub,6.8,True,tagcol)
    txt(s,7.06,t2+0.54,5.42,0.44,body,7.6,False,NAVY,lh=1.25)
    txt(s,7.06,t2+1.02,5.42,0.13,"ПОКУПЕЦЬ:  "+who,6.8,True,GREY)
chcard(1.80,"ЦІЛЬОВИЙ КАНАЛ",RED,"Продуктовий роздріб","МЕРЕЖІ НАЦІОНАЛЬНОГО ПОКРИТТЯ · 18 ПЕРЕВІРЕНО","Масовий покупець · щоденне харчування",
 "13 мереж мають корейський рамен, 5 — жодної позиції. «Сільпо» тримає 4 позиції на всю мережу.")
chcard(3.18,"ЦІЛЬОВИЙ КАНАЛ",DARKAMBER,"Онлайн-маркетплейси","УНІВЕРСАЛЬНІ МАЙДАНЧИКИ · 4 МАЙДАНЧИКИ","Масовий покупець · планована закупівля",
 "Prom · Rozetka · MAUDAU · Bigl. Медіана в категорії «локшина рамен» на Prom — 135 грн, мінімум 80 грн.")
chcard(4.56,"ВХІД",TEAL,"Азійські фудшопи","СПЕЦІАЛІЗОВАНИЙ ФУДРІТЕЙЛ · 5 МАГАЗИНІВ","Покупець азійської кухні · знає категорію",
 "OMG! Asia · Тайякі Март · Суші Повар · Asia Foods · Panda Asian. Корейський рамен 61–158 грн, ядро 92–135 грн.")
chcard(5.94,"НЕ НАШ КАНАЛ",PALE,"Outdoor- і мілітарі-рітейл","ТУРИСТИЧНІ ТА ВІЙСЬКОВІ МАГАЗИНИ · 0 ПОЗИЦІЙ","Турист і військовий · автономне харчування",
 "Рамену там немає: він потребує окропу й каструлі.")

# ============ СУМІЖНА ПОЛИЦЯ ============
s=slide(prs,"СУМІЖНА ПОЛИЦЯ","Не-корейська локшина: підлога, від якої рахує покупець",num(),
 "Чотирнадцять найпоширеніших не-корейських позицій локшини швидкого приготування — по одній на бренд. Вони формують уявлення про «нормальну» ціну в категорії.")
band(s,1.74,"УКРАЇНА ТА АЗІЯ БЕЗ КОРЕЇ   ·   Варіння або окріп   ·   50–200 г   ·   9–144 грн   ·   18–146 грн/100 г",BLUE)
cards(s,CM,BLUE)
ecall(s,0.62,6.08,5.92,1.26,"ЩО ЦЕ ОЗНАЧАЄ ДЛЯ ЦІНИ",
 "Український мас-маркет — 9–22 грн за 50–60 г, тобто 18–33 грн/100 г. Корейський рамен коштує у 2,5–3 рази дорожче за грам, і ринок це приймає. Але прийнята премія має межу.",BLUE)
ecall(s,6.80,6.08,5.92,1.26,"ДЕ ЦЯ МЕЖА ПРОХОДИТЬ",
 "У корейському рамені в пакеті ціна за грам — від 49 до 158 грн/100 г, медіана 80. При полиці 120 грн наші SKU — 92–107 грн/100 г; на рівні 119 грн/100 г і вище стоять лише 6 із 71 пакетних позицій.",RED)

# ============ CHOI'S ============
CH_NEED=M.max_cc_uah(125.2)
CH_P33=M.max_price_eur("33",shelf=125.2); CH_P6=M.max_price_eur("6",shelf=125.2)
s=slide(prs,"ОКРЕМИЙ КОНКУРЕНТ","Choi’s: ринковий орієнтир, окремий продукт",num(),
 "Choi’s Carbonara 131 г — позиція вихідного зрізу 29.09.2026. Це окремий конкурент, а не наш товар SIAS.")
ch=g("Choi's")
if ch: card(s,0.62,1.86,ch[0].get("img"),ch[0]["name"],ch[0]["w"],ch[0]["price"],ch[0]["sellers"],RED)
estat(s,2.50,1.86,"РОЗДРІБНА ЦІНА НА ПОЛИЦІ","125,20 грн","«Таврія В» і «Космос» на 29.09.\nЗа V3 (05–06.10) у «Таврії В» — 131,50 грн.",NAVY,3.26,1.86,25)
estat(s,5.98,1.86,"НАШІ SIAS · GOUDA / CARBONARA","120 грн","Ціль за моделлю: на 4 % нижче за Choi’s 125,20.\nЗараз за моделлю — 162,18 грн (+30 %).",RED,3.40,1.86,25)
estat(s,9.60,1.86,"СС ДЛЯ ПАРИТЕТУ 125,20",f"{d2(CH_NEED)} грн",f"Закупівля SIAS: {eur(CH_P33)} (машина), {eur(CH_P6)} (6 палет) — при маржі 30 %.",DARKAMBER,3.12,1.86,25)
y=section_head(s,4.06,"ЩО З ЦЬОГО МОЖНА І ЧОГО НЕ МОЖНА ВИВЕСТИ","")
ecall(s,0.62,y,3.88,2.30,"1 · ЦЕ ПОРІВНЯННЯ РІЗНИХ ПРОДУКТІВ",
 "Choi’s Carbonara коштує 125,20–131,50 грн. Наші SIAS Gouda / Carbonara за поточною собівартістю — 162,18 грн за моделлю, тобто на 30 % дорожче.\n\n"
 f"Щоб стати на 125,20 грн з маржею 30 %, потрібна СС {d2(CH_NEED)} грн. Це ціновий сценарій, а не ціна закупівлі конкурента.",AMBER)
ecall(s,4.72,y,3.88,2.30,"2 · ЩО ПЕРЕВІРИТИ ДЛЯ SIAS",
 "Для наших шести SKU потрібно підтвердити:\n\n· склад та інструкції приготування;\n· що саме входить у собівартість — товар і всі витрати;\n· умови поставки й канал продажу;\n· підтримку запуску і попит за нашою ціною.",NAVY)
ecall(s,8.82,y,3.90,2.30,"3 · МЕЖА ВИСНОВКУ",
 "Ціна Choi’s не розкриває його собівартості, бонусу мережі чи маржі.\n\nЗ неї не можна вивести закупівельну ціну SIAS, спільного імпортера або доступність ексклюзиву для нашого продукту.",RED)
slide_tail_sources(s,"Джерело: каталоги tavriav.zakaz.ua і cosmos.zakaz.ua, «Локшина Choi’s рамен карбон 131 г»; 05–06.10 — за V3. Паритет — власний розрахунок за формулою вкладок 3–4.",6.94)

# ============ CHOI'S У «ТАВРІЇ В» · ЦІНА ПРОДАЖУ ============
EU_PROMO=1.80*M.RATE; EU_REG=2.20*M.RATE
_dearer=lambda x: sum(1 for p in PACK_P if p>x)
s=slide(prs,"ОКРЕМИЙ КОНКУРЕНТ · ЦІНА ПРОДАЖУ","Choi’s у «Таврії В»: по чому продавали",num(),
 "Той самий бренд SIAS уже стоїть на українській полиці. Три дати й три рівні цін — щоб порівняти з нашою ціллю та поточною моделлю.")
txt(s,0.62,1.80,6.2,0.16,"CHOI’S CARBONARA 131 Г · ЦІНА ЗА ПАЧКУ, ГРН",8.4,True,NAVY)
hbars(s,0.62,2.10,6.60,[
  ("ЄС · dotasia.eu · акція €1,80",EU_PROMO,GREYBAR,f"{d1(EU_PROMO)} грн"),
  ("ЄС · dotasia.eu · звичайна €2,20",EU_REG,GREYBAR,f"{d1(EU_REG)} грн"),
  ("«Таврія В» · 29.09.2026",125.2,NAVYBAR,"125,20 грн"),
  ("«Космос» · 29.09.2026",125.2,NAVYBAR,"125,20 грн"),
  ("«Таврія В» · 05–06.10 (за V3)",131.5,NAVYBAR,"131,50 грн"),
  ("НАШ SIAS · ціль за моделлю",120.0,AMBBAR,"120,00 грн"),
  ("НАШ SIAS · зараз, 33 палети",NOWS["33"][0]['shelf'],tint(AMBBAR,0.35),f"{d2(NOWS['33'][0]['shelf'])} грн"),
  ("НАШ SIAS · зараз, 6 палет",NOWS["6"][0]['shelf'],tint(AMBBAR,0.35),f"{d2(NOWS['6'][0]['shelf'])} грн"),
 ],220,[0,40,80,120,160,200],lab_w=2.80,val_w=1.10,rowh=0.40,bar_h=0.23)
txt(s,7.55,1.80,5.2,0.16,"ПОРІВНЯННЯ · «ЗАРАЗ» = 33 ПАЛЕТИ, МАРЖА 30 %",8.0,True,NAVY)
cols=[("",7.55,1.70,L),("CHOI’S",9.25,1.15,R),("ЦІЛЬ 120",10.40,1.15,R),("ЗАРАЗ",11.55,1.17,R)]
_c=NOWS["33"][0]['shelf']
tab=[["Ціна, грн","125,20 / 131,50","120,00",d2(_c)],
     ["За 100 г, грн","95,6 / 100,4","91,6",d1(_c/131*100)],
     ["До Choi’s 125,20",  "—", sgn((120/125.2-1)*100), sgn((_c/125.2-1)*100)],
     ["До медіани пакета 102,8",sgn((125.2/PACK_MED-1)*100),sgn((120/PACK_MED-1)*100),sgn((_c/PACK_MED-1)*100)],
     ["Дорожчих у категорії*",f"{_dearer(125.2)} з 71",f"{_dearer(120)} з 71",f"{_dearer(_c)} з 71"]]
table(s,2.10,cols,tab,head_size=7.0,row_size=9.0,rh=0.44,left=7.55,bandw=5.17)
ecall(s,7.55,4.80,5.17,2.06,"ЩО ЦЕ ЗНАЧИТЬ ДЛЯ ПЕРЕГОВОРІВ",
 "1. Наша ціль 120 грн — на 4–9 % нижче за те, по чому цей бренд уже продають у «Таврії В». Якщо це той самий EAN 3760344350847, діючий імпортер матиме ціну вищу за нашу.\n"
 "2. За 29.09 → 05–06.10 ціна в «Таврії В» зросла на 5 %, акційної ціни немає (стара ціна дорівнює новій).\n"
 "3. Історії цін немає — лише дві дати. Її варто запросити в «Таврії В» і SIAS.",RED,9.6,8.2)
ecall(s,0.62,5.74,6.60,1.14,"ЄВРОПЕЙСЬКИЙ ОРІЄНТИР — ТІЛЬКИ ДЛЯ ПОРІВНЯННЯ",
 "dotasia.eu продає ту саму позицію за €1,80 замість €2,20 (−18 %); Gouda Cheese 123 г — так само. Це інтернет-магазин в ЄС, ПДВ на сторінці не вказано, товар «немає в наявності» — прямо порівнювати з українською полицею не можна.",AMBER,9.4,8.2)
slide_tail_sources(s,"* З 71 пакетної позиції зрізу 29.09.2026. Курс для перерахунку EUR→грн — 51,20. Джерела: tavriav.zakaz.ua, cosmos.zakaz.ua; ціна «Таврії В» 05–06.10 — за V3; dotasia.eu (Choi’s Carbonara 131 г, Gouda 123 г).",7.00)

# ============ АКЦІЇ І МІНІМАЛЬНІ ЦІНИ ============
Z=json.load(open(f"{SC}/zakaz_raw.json",encoding='utf-8'))
_eans={r['ean']:r for r in SD_ROWS}
_rec=[z for z in Z if z['ean'] in _eans and _eans[z['ean']]['seg'] in (1,2)]
KEY=[('08801073113428','Samyang Buldak 2× spicy','140 г'),('08801073110502','Samyang Buldak Hot Chicken','140 г'),
     ('08801073113381','Samyang Buldak тушкована курка','145 г'),('08801073116474','Samyang Buldak з сиром','130 г'),
     ('08801043150620','Nongshim Shin Ramyun','120 г'),('08801043157742','Nongshim Kimchi Ramyun','120 г'),
     ('08801043157728','Nongshim Chapaghetti','140 г'),('08801043018470','Nongshim Toomba Spicy & Creamy','137 г'),
     ('00648436310685','Paldo Volcano Carbonara','130 г')]
krow=[]
for e,lab,wt in KEY:
    v=[z for z in _rec if z['ean']==e]
    krow.append((lab,wt,len({z['chain'] for z in v}),min(z['price'] for z in v),max(z['price'] for z in v)))
_ap=[z for z in _rec if z['chain']=='auchan']
_app=[z for z in _ap if z.get('old_price') and z['price']<z['old_price']-0.005]
s=slide(prs,"АКЦІЇ ТА МІНІМАЛЬНІ ЦІНИ","Ключові конкуренти продаються нижче 120 грн",num(),
 "Одна й та сама позиція конкурента коштує по-різному в різних мережах і під час акцій. Для SIAS це діапазон конкурентного тиску на полиці.")
cols=[("ПОЗИЦІЯ КОНКУРЕНТА",0.62,3.10,L),("ВАГА",3.72,0.80,R),("МЕРЕЖ",4.52,0.70,R),("НАЙНИЖЧА, грн",5.22,1.30,R),("НАЙВИЩА, грн",6.52,1.30,R),("РІЗНИЦЯ",7.82,0.90,R)]
table(s,1.86,cols,[[l,w,str(n),d2(lo),d2(hi),pct((1-lo/hi)*100,0)] for l,w,n,lo,hi in krow],head_size=7.4,row_size=9.6,rh=0.33,bandw=8.30)
estat(s,9.10,1.78,"НИЖЧІ ЗА 120 ГРН",f"{sum(1 for k in krow if k[3]<120)} з {len(krow)}","ключових позицій конкурентів\nмають мінімальну ціну нижче 120 грн.",RED,3.62,1.22,24)
estat(s,9.10,3.14,"РОЗКИД МІЖ МЕРЕЖАМИ","34 %","медіанний розкид найдешевшої й найдорожчої\nмережі для 18 позицій у трьох і більше мережах.",DARKAMBER,3.62,1.22,24)
estat(s,9.10,4.50,"АКЦІЯ В «АШАНІ»",f"{len(_app)} з {len(_ap)}","записів пакетного рамену на акції;\nзнижка найчастіше 20–23 %, максимум 39 %.",TEAL,3.62,1.22,24)
ecall(s,0.62,5.46,8.30,1.38,"СКІЛЬКИ КОШТУЄ АКЦІЯ",
 "«Сільпо»: Samyang Buldak зі смаком курки та сиру і Buldak Carbonara — 89,99 грн замість 174 грн (−48 %). Samyang Tangle: від 77,50 до 157,00 грн залежно від мережі (−51 %).\n"
 "Отже, конкурент може виставити ту саму пачку за 86–105 грн на акції; наша ціна 120 грн — вище за цей рівень, але нижче за повну ціну більшості позицій.",NAVY,9.6,8.6)
ecall(s,9.10,5.86,3.62,0.98,"ОБМЕЖЕННЯ ДАНИХ",
 "Стару ціну віддають лише деякі мережі (в основному «Ашан»): частка акцій тут — нижня межа.",AMBER,8.6,7.8)
slide_tail_sources(s,"Каталоги мереж 29.09.2026: ціна за пачку і стара ціна. «Різниця» = 1 − найнижча ÷ найвища ціна позиції по мережах. «Сільпо» — ціна й стара ціна в каталозі.",7.00)

# ============ АРГУМЕНТИ ДЛЯ SIAS ============
NEED120=M.max_cc_uah(120)
s=slide(prs,"ЩО ПОКАЗАТИ SIAS","Чому полиця 120 грн — і що це означає для ціни закупівлі",num(),
 "Чотири кроки аргументації для переговорів: хто на полиці, по чому вони продають, де стоїмо ми і яка закупівля для цього потрібна.")
def arg(l,n,head,big,body,acc):
    _rect(s,l,1.80,2.78,4.30,LIGHT); _rect(s,l,1.80,0.045,4.30,acc)
    txt(s,l+0.26,1.96,0.6,0.40,str(n),26,True,acc)
    txt(s,l+0.82,2.04,1.86,0.40,head,10,True,NAVY,lh=1.15)
    txt(s,l+0.26,2.66,2.40,0.52,big,19,True,acc,lh=1.05)
    txt(s,l+0.26,3.40,2.32,2.66,body,9.4,False,GREY,lh=1.42)
arg(0.62,1,"Хто на полиці","54 з 96",
 "Samyang і Nongshim — 54 позиції з 96 і 10–11 мереж. Медіана пакетного рамену — 102,8 грн; стандарт категорії — пакет 120 г.\n\nНаші 113 г на 6 % легші за нього.",NAVY)
arg(3.72,2,"По чому вони продають","86–105 грн",
 "Так 8 із 9 ключових позицій Samyang, Nongshim і Paldo продаються за мінімальною ціною: в акції та в найдешевшій мережі.\n\n«Ашан»: 30 із 52 записів — акція, знижка 20–23 %. «Сільпо»: Buldak 89,99 замість 174 грн.",DARKAMBER)
arg(6.82,3,"Де стоїмо ми","P75 · 18 з 71",
 "Полиця 120 грн — верхній квартиль пакетного рамену: дорожчих лише 18 із 71, на 17 % вище за медіану. Поточна розрахункова полиця (140,55 і 162,18 грн) — це вже верхні 10 % і 4 %.\n\n120 грн — стеля, а не нижня межа.",RED)
arg(9.92,4,"Яка закупівля потрібна","€0,54 / €0,32",
 f"Модель SIAS: бонус мережі 25 %, наша маржа 30 %, націнка 1,40. Полиця 120 грн = ціна партнеру 85,71 грн → СС не більше {d2(NEED120)} грн.\n\nЦе закупівля €0,54 за повну машину (33 палети) і €0,32 за 6 палет — проти €0,65–0,75 зараз.",TEAL)
verdict(s,6.26,"ЗАПИТ ДО SIAS","Ціна закупівлі, за якої ми ставимо 120 грн на полицю і зберігаємо маржу 30 %: €0,54 — повна машина, €0,32 — 6 палет.",RED,0.92)
slide_tail_sources(s,"Джерела: каталоги мереж 29.09.2026; фінмодель RAMEN_FINMODEL_4_TABS (вкладки «Цена 120 грн»). Згода SIAS на ці ціни не отримана — це ціна для обговорення.",7.16)

# ============ ЦІНОУТВОРЕННЯ ============
C_CC=RGBColor(0x46,0x53,0x82); C_BON=DARKAMBER; C_MAR=RGBColor(0xF9,0xA5,0x0B); C_MK=RGBColor(0xA9,0xB0,0xC8)
def fwd_parts(sc,idx,margin=0.30):
    f=M.forward_current(sc,margin)[idx]; part=f['partner']
    return dict(cc=f['cc'],bonus=part*M.BONUS,profit=part*margin,markup=part*(M.MARKUP-1),partner=part,shelf=f['shelf'])
A33=M.breakdown("33",0.54); A6=M.breakdown("6",0.32)
rowsP=[("Ціль · 33 палети","закупівля €0,54 · усі 6 SKU",A33,"€0,54","запит"),
       ("Ціль · 6 палет","закупівля €0,32 · усі 6 SKU",A6,"€0,32","запит"),
       ("Зараз · 33 палети","4 SKU · закупівля €0,65",fwd_parts("33",2),"€0,65","зараз"),
       ("Зараз · 33 палети","2 SKU · закупівля €0,75",fwd_parts("33",0),"€0,75","зараз"),
       ("Зараз · 6 палет","4 SKU · закупівля €0,65",fwd_parts("6",2),"€0,65","зараз"),
       ("Зараз · 6 палет","2 SKU · закупівля €0,75",fwd_parts("6",0),"€0,75","зараз")]
s=slide(prs,"ЦІНОУТВОРЕННЯ","З чого складається ціна на полиці 120 грн",num(),
 "Як у фінмоделі: ціна партнеру = 100 %, а мережа додає до неї 40 % — так виходить полиця. Закупівля у SIAS — окремо.")
_rect(s,0.62,1.74,12.10,1.10,LIGHT)
txt(s,0.84,1.84,8.0,0.14,"ЦІНА ПАРТНЕРУ — ЗА НЕЮ МИ ПРОДАЄМО МЕРЕЖІ (100 %)",7.4,True,GREY)
bx=0.84; bw=9.0; unit=bw/140.0
for lab,v,col,tc in (("Собівартість 45 %",45,C_CC,WHITE),("Бонус мережі 25 %",25,C_BON,WHITE),("Наша маржа 30 %",30,C_MAR,NAVY),("Націнка магазину +40 %",40,C_MK,NAVY)):
    _rect(s,bx,2.04,v*unit,0.34,col); txt(s,bx,2.11,v*unit,0.2,lab,8.8,True,tc,PP_ALIGN.CENTER); bx+=v*unit
_rect(s,0.84,2.42,100*unit,0.012,NAVY); txt(s,0.84,2.44,100*unit,0.16,"= ціна партнеру 85,71 грн",8,True,NAVY,PP_ALIGN.CENTER)
txt(s,0.84,2.62,11.6,0.18,"ПОЛИЦЯ = ціна партнеру × 1,40 = 120 грн.  Бонус 25 % і маржа 30 % — частки ціни партнеру; націнка 40 % — понад ціну партнеру.",8.6,True,NAVY)
# заголовки колонок
txt(s,0.62,3.00,2.6,0.14,"ПОЗИЦІЯ",6.8,True,GREY)
txt(s,3.10,3.00,6.4,0.14,"ГРН ЗА ПАЧКУ: ЦІНА ПАРТНЕРУ + НАЦІНКА МАГАЗИНУ",6.8,True,GREY)
txt(s,10.00,3.00,0.9,0.14,"ПАРТНЕРУ",6.8,True,GREY,PP_ALIGN.RIGHT)
txt(s,11.05,3.00,0.9,0.14,"ЗАКУПІВЛЯ",6.8,True,GREY,PP_ALIGN.RIGHT)
txt(s,11.95,3.00,0.8,0.14,"СС, ГРН",6.8,True,GREY,PP_ALIGN.RIGHT)
px0=3.10; XM=225.0; PW=6.70
def X(v): return px0+PW*v/XM
y=3.22; RH=0.50
for i,(a,b,p,buy,kind) in enumerate(rowsP):
    if i%2==0: _rect(s,0.62,y,12.10,RH-0.02,LIGHT)
    txt(s,0.76,y+0.05,2.3,0.18,a,9.4,True,NAVY); txt(s,0.76,y+0.26,2.3,0.14,b,7.2,False,GREY)
    pr=p['profit']; cc=p['cc']; bo=p['bonus']; mk=p['markup']
    x=px0
    for val,col,tc in ((cc,C_CC,WHITE),(bo,C_BON,WHITE),(pr,C_MAR,NAVY),(mk,C_MK,NAVY)):
        w=PW*val/XM; _rect(s,x,y+0.09,w,0.30,col)
        if w>0.34: txt(s,x,y+0.14,w,0.18,f"{val:.0f}",8.4,True,tc,PP_ALIGN.CENTER)
        x+=w
    sh=p['shelf']; shs=(f"{sh:.1f}".replace(".",",") if kind=="запит" else f"{sh:.0f}"); txt(s,X(sh)+0.08,y+0.11,0.9,0.26,shs,15,True,(RED if sh>125 else DARKAMBER),L)
    _rect(s,X(p['partner'])-0.01,y+0.05,0.025,0.38,NAVY)
    txt(s,10.00,y+0.14,0.9,0.22,f"{p['partner']:.0f}",11,True,NAVY,PP_ALIGN.RIGHT)
    txt(s,11.05,y+0.14,0.9,0.22,buy,10.5,True,(TEAL if kind=="запит" else NAVY),PP_ALIGN.RIGHT)
    txt(s,11.95,y+0.14,0.8,0.22,f"{cc:.1f}".replace('.',','),10,False,GREY,PP_ALIGN.RIGHT)
    y+=RH
# лінія цілі 120
_rect(s,X(120)-0.006,3.16,0.012,RH*6+0.10,RED); txt(s,X(120)-0.45,y+0.02,0.9,0.14,"ціль 120",7.2,True,RED,PP_ALIGN.CENTER)
# легенда
lx=0.76
for lab,col in (("Собівартість (СС)",C_CC),("Бонус мережі",C_BON),("Наша маржа",C_MAR),("Націнка магазину",C_MK)):
    _rect(s,lx,y+0.26,0.18,0.14,col); txt(s,lx+0.26,y+0.24,1.9,0.16,lab,8,False,NAVY); lx+=2.15
txt(s,9.4,y+0.24,3.3,0.16,"Чорна риска — ціна партнеру.",8,False,GREY)
slide_tail_sources(s,"Фінмодель SIAS (курс 51,20 грн/€, бонус 25 %, маржа 30 %, націнка 1,40, ціна вже з ПДВ): ціна партнеру = СС ÷ 0,45; полиця = ціна партнеру × 1,40. «Ціль» — вкладки «Цена 120 грн»; «Зараз» — СС за Self-Cost_Sias (вкладки 1–2).",6.90)

# ============ ЗАПИТ ДО SIAS ============
s=slide(prs,"ЗАПИТ ДО SIAS","Ціна закупівлі за одну пачку для полиці 120 грн",num(),
 "Одна запитувана ціна для всіх шести SKU. Умови відвантаження — як у Self-Cost (FCA Roye); згода постачальника ще не отримана.")
estat(s,0.62,1.76,"33 ПАЛЕТИ · ПОВНА МАШИНА · 52 800 ПАЧОК","€0,54 / пачка","Товар у партії — €28 512. Ціна короба (20 пачок) — €10,80.\nПредельна ціна за моделлю — €0,5407.",TEAL,5.92,1.56,28)
estat(s,6.80,1.76,"6 ПАЛЕТ · 9 600 ПАЧОК","€0,32 / пачка","Товар у партії — €3 072. Ціна короба (20 пачок) — €6,40.\nПредельна ціна за моделлю — €0,3265.",RED,5.92,1.56,28)
cols=[("ПОЗИЦІЯ SIAS",0.62,2.50,L),("ВАГА",3.12,0.80,R),("ЗАРАЗ, €/пач.",3.92,1.15,R),("ЗАПИТ 6 ПАЛЕТ",5.07,1.35,R),("ЗНИЖЕННЯ",6.42,1.05,R),
      ("ЗАПИТ 33 ПАЛЕТИ",7.47,1.55,R),("ЗНИЖЕННЯ",9.02,1.05,R),("ПАЛЕТ 6 / 33",10.20,1.30,R),("ПАЧОК У 33",11.55,1.17,R)]
rows=[]
for (nm,wt,w_,cur),p6,p33 in zip(M.SKUS,[1]*6,M.SC["33"]["pallets"]):
    rows.append([nm,wt,eur(cur),eur(0.32),pct((1-0.32/cur)*100,0),eur(0.54),pct((1-0.54/cur)*100,0),f"1 / {p33}",th(p33*M.PACK_PALLET)])
rows.append(["ВСЬОГО","","","","","","","6 / 33","9 600 / 52 800"])
table(s,3.48,cols,rows,head_size=7.0,row_size=9.6,rh=0.325)
ecall(s,0.62,6.20,6.00,0.84,"ЯК ЧИТАТИ ЗНИЖЕННЯ",
 "Це різниця між поточною ціною SIAS (€0,65 для чотирьох SKU і €0,75 для Gouda та Carbonara) і запитом. Для машини — 17–28 %, для 6 палет — 51–57 %.",NAVY,9.2,8.2)
ecall(s,6.80,6.20,5.92,0.84,"ЩО ЦЕ ОЗНАЧАЄ",
 "Повна машина — реалістичніший торг: потрібне зниження вдвічі менше. 6 палет за ціною €0,32 вимагає майже половини вартості товару.",RED,9.2,8.2)
slide_tail_sources(s,"Джерело: RAMEN_FINMODEL_4_TABS_PRESENTATION.xlsx, вкладки «Цена 120 грн — 6 паллет» і «— 33 паллет»; ціна «зараз» — з Self-Cost_Sias. Ціна для обговорення, згоди SIAS немає.",7.14)

# ============ ЗВІДКИ €0,54 І €0,32 ============
x120=120/M.MARKUP; ccmax=x120*(1-M.BONUS-M.MARGIN); e1=ccmax/M.RATE; e2=e1/(1+M.FIN)
def stp(sc):
    n=M.units(sc); bc=M.batch_costs(sc); pp=bc/n; mp=M.max_price_eur(sc); a=M.ask_price(sc); b=M.breakdown(sc,a)
    return dict(pp=pp,after=e2-pp,mp=mp,ask=a,shelf=b['shelf'],margin=b['margin'],batch=bc,n=n)
T6,T33=stp("6"),stp("33")
s=slide(prs,"ЗВІДКИ €0,54 І €0,32","Розрахунок закупівлі від полиці 120 грн, крок за кроком",num(),
 "Обернена задача вкладок «Цена 120 грн»: від цільової полиці вгору по ланцюгу — до максимальної ціни, яку ми можемо заплатити SIAS.")
cols=[("КРОК",0.62,0.50,L),("ЩО РАХУЄМО",1.12,4.60,L),("ДІЯ",5.72,3.00,L),("6 ПАЛЕТ",8.72,1.90,R),("33 ПАЛЕТИ",10.62,2.10,R)]
rows=[["1","Цільова полиця з ПДВ, грн","умова","120,00","120,00"],
      ["2","Ціна партнеру з ПДВ, грн/уп.","полиця ÷ 1,40",d2(x120),d2(x120)],
      ["3","Частка ціни для собівартості","100 % − бонус 25 % − маржа 30 %","45 %","45 %"],
      ["4","Максимальна СС, грн/уп.","крок 2 × крок 3",d2(ccmax),d2(ccmax)],
      ["5","Те саме в євро, €/уп.","крок 4 ÷ курс 51,20",f"{e1:.4f}".replace('.',','),f"{e1:.4f}".replace('.',',')],
      ["6","Без фінансування закупівлі, €/уп.","крок 5 ÷ (1 + 2,5 %)",f"{e2:.4f}".replace('.',','),f"{e2:.4f}".replace('.',',')],
      ["7","Мінус витрати партії на пачку, €","партія ÷ кількість пачок",f"− {T6['pp']:.4f}".replace('.',','),f"− {T33['pp']:.4f}".replace('.',',')],
      ["8","Предельна закупівля, €/уп.","(крок 6 − крок 7) ÷ (1 + ПДВ 20 %)",f"{T6['mp']:.4f}".replace('.',','),f"{T33['mp']:.4f}".replace('.',',')],
      ["9","ЗАПИТ ДО SIAS, €/уп.","крок 8, вниз до євроцента",f"{T6['ask']:.2f}".replace('.',','),f"{T33['ask']:.2f}".replace('.',',')],
      ["10","Перевірка: полиця, грн","розклад ціни при запиті",d2(T6['shelf']),d2(T33['shelf'])],
      ["11","Перевірка: наша маржа","прибуток після бонусу ÷ ціна партнеру",pct(T6['margin']*100),pct(T33['margin']*100)]]
table(s,1.84,cols,rows,head_size=7.4,row_size=9.4,rh=0.318)
_rect(s,0.62,5.78,12.10,0.012,NAVY)
ecall(s,0.62,5.90,5.92,1.10,"ЩО В ВИТРАТАХ ПАРТІЇ (КРОК 7)",
 f"6 палет: фрахт €2 600 + брокер €288 + інші €42 + ПДВ на фрахт у митній базі €364 = €{th(T6['batch'])}, на 9 600 пачок.\n"
 f"33 палети: фрахт €3 700 + брокер €288 + інші €42 + ПДВ на фрахт €518 = €{th(T33['batch'])}, на 52 800 пачок.",NAVY,9.4,8.2)
ecall(s,6.80,5.90,5.92,1.10,"ЧОМУ ЗАПИТ РІЗНИЙ",
 "Фрахт і брокер ділять на пачки: при 6 палетах це €0,34 на пачку, при 33 палетах — €0,09. Тому для однакової полиці з малої партії закупівля має бути нижчою на €0,22.",RED,9.4,8.2)
slide_tail_sources(s,"Джерело: «Цена 120 грн — 6 паллет» і «— 33 паллет» (клітинки H62–H82). Імпортний ПДВ і фінансування входять у собівартість, як у Self-Cost; мито 0 % (EUR.1). Лістинг окремою статтею не виділено.",7.10)

# ============ ПОЛИЦЯ ЗАРАЗ ============
s=slide(prs,"ПОЛИЦЯ ЗАРАЗ","Наша ціна на полиці за поточною закупівлею: усі смаки",num(),
 "Шість SKU SIAS при поточних цінах закупівлі €0,65–0,75 і маржі 30 %. Ціна партнеру й полиця — з ПДВ, за одну пачку.")
def nowtable(l,sc,head,acc):
    txt(s,l,1.78,5.9,0.18,head,9.6,True,acc)
    cols=[("ПОЗИЦІЯ",l,1.90,L),("СС, грн",l+1.90,0.70,R),("ПАРТНЕРУ",l+2.60,0.85,R),("ПОЛИЦЯ",l+3.45,0.78,R),("ДО 120",l+4.23,0.62,R),("МАРЖА ПРИ 120",l+4.85,1.00,R)]
    rows=[]
    for (nm,wt,w_,cur),f in zip(M.SKUS,M.forward_current(sc)):
        rows.append([f"{nm} {wt}",d2(f['cc']),d2(f['partner']),d2(f['shelf']),f"+{f['shelf']-120:.0f}",pct(M.margin_at_shelf(f['cc'])*100)])
    table(s,2.04,cols,rows,head_size=6.8,row_size=9.0,rh=0.36,left=l,bandw=5.95)
nowtable(0.62,"33","33 ПАЛЕТИ · ПОВНА МАШИНА · ПОТОЧНА СС",TEAL)
nowtable(6.80,"6","6 ПАЛЕТ · ПОТОЧНА СС",RED)
f33=M.forward_current("33"); f6=M.forward_current("6")
ecall(s,0.62,4.50,5.90,1.30,"ПОВНА МАШИНА: ВИЩЕ ЗА 120 НА 21–42 ГРН",
 f"Чотири SKU — {d2(f33[2]['shelf'])} грн, Gouda і Carbonara — {d2(f33[0]['shelf'])} грн. Якщо поставити на полицю 120 грн за поточної закупівлі, маржа буде {pct(M.margin_at_shelf(f33[2]['cc'])*100)} і {pct(M.margin_at_shelf(f33[0]['cc'])*100)} замість 30 %.",TEAL,9.8,8.6)
ecall(s,6.82,4.50,5.90,1.30,"6 ПАЛЕТ: ВИЩЕ ЗА 120 НА 61–88 ГРН",
 f"Чотири SKU — {d2(f6[2]['shelf'])} грн, Gouda і Carbonara — {d2(f6[0]['shelf'])} грн. Полиця 120 грн за поточної закупівлі дає маржу {pct(M.margin_at_shelf(f6[2]['cc'])*100)} і {pct(M.margin_at_shelf(f6[0]['cc'])*100)} — для пари це збиток.",RED,9.8,8.6)
estrip(s,6.00,"ВИСНОВОК",
 "З малої партії полиця 120 грн недосяжна за поточної закупівлі. З повної машини — досяжна, але з маржею 14–22 % замість 30 %, або з пониженням закупівлі до €0,54.",RED,0.70)
slide_tail_sources(s,"Фінмодель: вкладки «Финмодель 6 паллет» і «Финмодель 33 паллет», маржа 30 % (колонки T). СС — із Self-Cost_Sias.xlsx. «Маржа при 120» = (ціна партнеру при полиці 120 − СС − бонус) ÷ ціна партнеру.",7.00)

# ============ ЧУТЛИВІСТЬ ============
SHELVES=[105,110,120,130,140]; MARGINS=[0.35,0.30,0.25,0.20,0.15]
def gridsc(sc): return [[math.floor(M.max_price_eur(sc,shelf=sh,margin=m)*100+1e-9)/100 for sh in SHELVES] for m in MARGINS]
def cellcol(v):
    if v>=0.75: return (TEAL,WHITE)
    if v>=0.65: return (tint(TEAL,0.55),NAVY)
    if v>=0.50: return (tint(AMBER,0.55),NAVY)
    return (tint(RED,0.55),NAVY)
s=slide(prs,"ЧУТЛИВІСТЬ","Яка закупівля потрібна при іншій полиці й іншій маржі",num(),
 "Максимальна ціна закупівлі, €/пачку, за формулою вкладок «Цена 120 грн». Колір — порівняння з поточною ціною SIAS: €0,75 (Gouda, Carbonara) і €0,65 (решта).")
def sgrid(l,sc,head,acc):
    G_=gridsc(sc)
    txt(s,l,1.78,5.9,0.18,head,9.6,True,acc)
    cw=0.96; rl=1.00; t=2.12
    txt(s,l,t,rl,0.16,"МАРЖА \\ ПОЛИЦЯ",6.4,True,GREY)
    for j,sh in enumerate(SHELVES):
        txt(s,l+rl+j*cw,t-0.02,cw,0.20,f"{sh} грн",8.6,True,(RED if sh==120 else NAVY),PP_ALIGN.CENTER)
    y=t+0.26
    for i,m in enumerate(MARGINS):
        txt(s,l,y+0.12,rl-0.1,0.22,f"{int(m*100)} %",10,True,(RED if abs(m-0.30)<1e-9 else NAVY),PP_ALIGN.RIGHT)
        for j,sh in enumerate(SHELVES):
            v=G_[i][j]; bg,tc=cellcol(v)
            _rect(s,l+rl+j*cw+0.03,y+0.03,cw-0.06,0.50,bg)
            txt(s,l+rl+j*cw+0.03,y+0.16,cw-0.06,0.26,f"€{v:.2f}".replace('.',','),12,True,tc,PP_ALIGN.CENTER)
            if sh==120 and abs(m-0.30)<1e-9:
                fx=l+rl+j*cw+0.03; fy=y+0.03; fw=cw-0.06; fh=0.50
                for (a_,b_,c_,d_) in ((fx,fy,fw,0.04),(fx,fy+fh-0.04,fw,0.04),(fx,fy,0.04,fh),(fx+fw-0.04,fy,0.04,fh)): _rect(s,a_,b_,c_,d_,RED)
        y+=0.56
    return y
sgrid(0.62,"33","33 ПАЛЕТИ · ПОВНА МАШИНА",TEAL)
sgrid(6.80,"6","6 ПАЛЕТ",RED)
y0=5.16
lx=0.76
for lab,col in (("≥ €0,75 — досяжно для всіх SKU",TEAL),("€0,65–0,75 — для чотирьох SKU",tint(TEAL,0.55)),("€0,50–0,65 — потрібне зниження",tint(AMBER,0.55)),("< €0,50 — суттєве зниження",tint(RED,0.55))):
    _rect(s,lx,y0,0.18,0.14,col); txt(s,lx+0.24,y0-0.02,2.6,0.16,lab,7.8,False,NAVY); lx+=2.98
txt(s,0.76,y0+0.20,11.8,0.16,"Червона рамка — ціль: полиця 120 грн, маржа 30 %.",7.8,True,RED)
ecall(s,0.62,5.62,5.92,1.34,"ПОВНА МАШИНА: ДВА ШЛЯХИ ДО 120 ГРН",
 "При маржі 30 % потрібне €0,54 — зниження на 17–28 %. При маржі 20 % достатньо €0,67, що вище за поточні €0,65 для чотирьох SKU; при 15 % — €0,74, майже поточні €0,75.",TEAL,9.8,8.6)
ecall(s,6.80,5.62,5.92,1.34,"6 ПАЛЕТ: ЖОДНА КЛІТИНКА НЕ ДОСЯГАЄ ПОТОЧНОЇ ЦІНИ",
 "Для полиці 120 грн навіть при маржі 15 % потрібне €0,53, тобто мінус 18–29 % до поточної ціни. Щоб тримати 30 %, потрібне €0,32 — мінус 51–57 %.",RED,9.8,8.6)
slide_tail_sources(s,"Розрахунок за формулою H66 вкладок «Цена 120 грн» (курс 51,20; бонус 25 %; націнка 1,40; фінансування 2,5 %; ПДВ 20 %). Ціна вниз до євроцента.",7.06)

# ============ ХТО З КИМ СТОЇТЬ · КАРТА ============
PM_BANDS=[(0,80,"до 80 грн"),(80,95,"80–95"),(95,110,"95–110"),(110,130,"110–130"),(130,150,"130–150"),(150,180,"150–180"),(180,9e9,"180 +")]
def pm_band(p):
    for i,(lo,hi,_) in enumerate(PM_BANDS):
        if lo<=p<hi: return i
    return len(PM_BANDS)-1
s=slide(prs,"ХТО З КИМ СТОЇТЬ · КАРТА","Карта цін: сегмент × цінова смуга",num(),
 "Усі бренди в клітинках за медіанною ціною позиції; помаранчеве — наші SKU: ціль 120 грн і поточна розрахункова полиця (33 палети). ×n — кількість смаків.")
colx=[1.72,4.50,7.28,10.06]; cwid=2.66; y0=1.78
for i,k in enumerate(("1","2","3","4")):
    _rect(s,colx[i],y0,cwid,0.34,SCOL[k]); txt(s,colx[i],y0+0.08,cwid,0.20,f"{STITLE[k]} · {SS[k]['n']}",9,True,WHITE,PP_ALIGN.CENTER)
RHm=0.665; ytop=y0+0.40
OURS={('1',3):"НАШІ 4 SKU · ціль 120",('1',4):f"НАШІ 4 SKU · зараз {d2(NOW_LOW)}",('2',3):"НАШІ 2 SKU · ціль 120",('2',5):f"НАШІ 2 SKU · зараз {d2(NOW_HIGH)}"}
for bi,(lo,hi,lab) in enumerate(PM_BANDS):
    yy=ytop+bi*RHm
    txt(s,0.62,yy+RHm/2-0.12,1.05,0.24,lab,9.4,True,NAVY)
    for ci,k in enumerate(("1","2","3","4")):
        _rect(s,colx[ci],yy,cwid,RHm-0.05,LIGHT)
        cnt=co.Counter(r['tm'] for r in SD_ROWS if str(r['seg'])==k and r.get('pmed') and pm_band(r['pmed'])==bi)
        items=cnt.most_common()
        shown=[(f"{b} ×{n}" if n>1 else b) for b,n in items[:5]]
        if len(items)>5: shown.append(f"+{len(items)-5}")
        txt(s,colx[ci]+0.08,yy+0.05,cwid-0.16,0.34,(" · ".join(shown) if shown else "—"),7.8,False,(NAVY if shown else PALE),lh=1.18)
        if (k,bi) in OURS:
            _rect(s,colx[ci]+0.06,yy+RHm-0.30,cwid-0.12,0.22,AMBBAR)
            txt(s,colx[ci]+0.06,yy+RHm-0.27,cwid-0.12,0.16,OURS[(k,bi)],7.4,True,NAVY,PP_ALIGN.CENTER)
slide_tail_sources(s,"Чорний текст — бренди конкурентів і кількість їхніх позицій у смузі (медіанна ціна по мережах, 29.09.2026); помаранчеві плашки — наші SKU за фінмоделлю. Наших 4 SKU показано в сегменті 1, двох — у сегменті 2 за припущенням, що спосіб приготування це підтвердить.",6.92)

# ============ ХТО З КИМ СТОЇТЬ · СУСІДИ ============
PK=sorted([c for k in ("1","2") for c in SEG[k] if c.get('pmed')],key=lambda c:c['pmed'])
def near(p,n=3):
    lo=[c for c in PK if c['pmed']<p][-n:][::-1]; hi=[c for c in PK if c['pmed']>=p][:n]
    return lo,hi
def cl(c): return f"{c['name']} · {c['w'].replace(' Г',' г')} · {c['pmed']:.0f}"
s=slide(prs,"ХТО З КИМ СТОЇТЬ · СУСІДИ","Найближчі за ціною до кожної нашої групи",num(),
 "Конкуренти сегментів 1–2 (пакет): три найближчі нижче й вище, плюс коридор ±20 %. Формат запису: бренд, назва, вага, медіанна ціна, грн.")
cols=[("НАША ГРУПА",0.62,2.70,L),("ДЕШЕВШЕ ЗА НАС",3.40,3.55,L),("ДОРОЖЧЕ ЗА НАС",7.05,3.55,L),("КОРИДОР ±20 %",10.70,2.02,L)]
for lab,x,w,al in cols: txt(s,x,1.80,w,0.14,lab,7.2,True,GREY,al)
groups=[("4 SKU · 112,5–113 г","ціль 120 грн",120.0),("4 SKU · 112,5–113 г","зараз 140,55 грн",NOW_LOW),
        ("Gouda 123 г · Carbonara 131 г","ціль 120 грн",120.0),("Gouda 123 г · Carbonara 131 г","зараз 162,18 грн",NOW_HIGH)]
yy=2.04; RHn=1.18
for gi,(g1,g2,p) in enumerate(groups):
    if gi%2==0: _rect(s,0.62,yy,12.10,RHn-0.06,LIGHT)
    _rect(s,0.62,yy,0.05,RHn-0.06,AMBBAR)
    txt(s,0.82,yy+0.10,2.5,0.34,g1,9.6,True,NAVY,lh=1.1)
    txt(s,0.82,yy+0.52,2.5,0.24,g2,12,True,DARKAMBER)
    lo,hi=near(p)
    txt(s,3.40,yy+0.10,3.60,1.00,"\n".join(cl(c) for c in lo),8.6,False,NAVY,lh=1.30)
    txt(s,7.05,yy+0.10,3.60,1.00,"\n".join(cl(c) for c in hi),8.6,False,NAVY,lh=1.30)
    cor=[c for c in PK if 0.8*p<=c['pmed']<=1.2*p]
    br=co.Counter(c['tm'] for c in cor).most_common(3)
    txt(s,10.70,yy+0.10,2.0,0.24,f"{len(cor)} позицій",11,True,NAVY)
    txt(s,10.70,yy+0.42,2.0,0.60,", ".join(b for b,_ in br)+f"\n({int(0.8*p)}–{int(1.2*p)} грн)",7.8,False,GREY,lh=1.25)
    yy+=RHn
slide_tail_sources(s,"Сусіди — з 71 пакетної позиції зрізу 29.09.2026 за медіанною ціною. Коридор ±20 % від ціни нашої групи. Стакани й рисові — інший привід, до сусідів не входять.",6.84)

# ============ РИЗИКИ ============
def _p(sc,**kw): return M.max_price_eur(sc,**kw)
_old=M.RATE; M.RATE=55.0; r55=(_p("33"),_p("6")); M.RATE=_old
mk13=(_p("33",markup=1.30),_p("6",markup=1.30)); mk155=(_p("33",markup=1.55),_p("6",markup=1.55))
bn20=(_p("33",bonus=0.20),_p("6",bonus=0.20)); bn30=(_p("33",bonus=0.30),_p("6",bonus=0.30))
s=slide(prs,"РИЗИКИ","Що може змінити розрахунок і яких даних бракує",num(),
 "Запит €0,54 / €0,32 тримається на умовах моделі. Нижче — припущення, які варто перевірити до підписання, і їхній вплив на допустиму ціну закупівлі.")
def risk(l,t,tag,tagcol,head,body,impact):
    _rect(s,l,t,3.88,2.14,LIGHT); _rect(s,l,t,0.045,2.14,tagcol)
    _rect(s,l+3.02,t+0.14,0.74,0.22,tagcol)
    txt(s,l+3.02,t+0.175,0.74,0.15,tag,6.4,True,WHITE,PP_ALIGN.CENTER)
    txt(s,l+0.28,t+0.14,2.70,0.34,head,10.5,True,NAVY,lh=1.12)
    txt(s,l+0.28,t+0.58,3.38,0.92,body,8.4,False,GREY,lh=1.38)
    txt(s,l+0.28,t+1.56,3.38,0.13,"ВПЛИВ",6.9,True,GREY)
    txt(s,l+0.28,t+1.72,3.38,0.36,impact,8.8,True,tagcol,lh=1.2)
risk(0.62,1.80,"ВИСОКИЙ",RED,"Згода SIAS",
 "€0,54 і €0,32 — це запит, а не угода. Для повної машини потрібне зниження 17–28 %, для 6 палет — 51–57 %.",
 "Без згоди: полиця 140,55–162,18 грн (машина)")
risk(4.72,1.80,"ВИСОКИЙ",RED,"ПДВ у моделі",
 f"Імпортний ПДВ ({d1(A33['vat'])} грн/пачку) входить у СС, а ціна партнеру — з ПДВ (вихідний ПДВ {d1(A33['partner']*(1-1/1.2))} грн). Чистий ПДВ до сплати модель не показує.",
 f"Якщо ми платник ПДВ: −{d1((A33['partner']*(1-1/1.2)-A33['vat'])/A33['partner']*100)} п.п. маржі (потребує підтвердження бухгалтерії)")
risk(8.82,1.80,"ВИСОКИЙ",RED,"Маржа 30 % — валова",
 "Внутрішня логістика, лістинг, промо й курсові різниці у моделі окремою статтею не виділені. Лістинг мережі лише зменшує допустиму закупівлю.",
 "Реальний прибуток менший за 25,7 грн/пачку")
risk(0.62,4.08,"СЕРЕДНІЙ",DARKAMBER,"Діючий імпортер Choi’s",
 "Choi’s Carbonara вже стоїть у «Таврії В» і «Космосі» за 125,20–131,50 грн. Наша ціль 120 грн нижча — потрібні ексклюзив або чіткий розподіл каналів.",
 "Ризик конфлікту каналів і цін")
risk(4.72,4.08,"СЕРЕДНІЙ",DARKAMBER,"Націнка 1,40 і бонус 25 %",
 f"Обидва — умови моделі. Націнка 1,30 дає допустиму закупівлю {eur(mk13[0])}, 1,55 — {eur(mk155[0])}. Бонус 20 % — {eur(bn20[0])}, 30 % — {eur(bn30[0])} (машина).",
 f"Допустима закупівля {eur(min(mk155[0],bn30[0]))}–{eur(max(mk13[0],bn20[0]))}")
risk(8.82,4.08,"СЕРЕДНІЙ",DARKAMBER,"Курс 51,20 грн/€",
 f"Закупівля в євро, полиця в гривні — курсовий ризик лягає на нас. При курсі 55 допустима закупівля — {eur(r55[0])} (машина) і {eur(r55[1])} (6 палет).",
 f"Кожні +4 грн до курсу = −{(1-r55[0]/M.max_price_eur('33'))*100:.0f} % до ціни")
estrip(s,6.34,"ЧОГО НЕ ВИСТАЧАЄ ДЛЯ ОСТАТОЧНОГО РІШЕННЯ",
 "Специфікації шести SKU · розшифровка собівартості за статтями · податковий режим ПДВ · умови лістингу хоча б однієї мережі · статус і ціни Choi’s у «Таврії В» · історія цін конкурентів.",NAVY,0.66)
slide_tail_sources(s,"Розрахунок за формулою H66 вкладок «Цена 120 грн», зміна одного параметра за раз; ціна закупівлі — на повну машину (33 палети), якщо не зазначено інше.",7.06)

# ============ ПІДСУМОК ============
s=slide(prs,"ПІДСУМОК","Що просимо в SIAS і що робити далі",num(),
 "Одне рішення: чи будуємо полицю 120 грн на новій закупівлі, чи приймаємо нижчу маржу. Чотири кроки до рішення.")
estat(s,0.62,1.76,"ЦІЛЬОВА ПОЛИЦЯ","120 грн","P75 пакетного рамену: дорожчих — 18 із 71;\nконкуренти в акції — 86–105 грн.",NAVY,3.88,1.52,26)
estat(s,4.72,1.76,"ЗАКУПІВЛЯ ДЛЯ МАРЖІ 30 %","€0,54 / €0,32","повна машина / 6 палет — замість\n€0,65–0,75 зараз.",TEAL,3.88,1.52,26)
estat(s,8.82,1.76,"ПОТРІБНЕ ЗНИЖЕННЯ","17–28 % / 51–57 %","машина / 6 палет. Альтернатива —\nполиця 120 грн з маржею 14–22 %.",RED,3.90,1.52,23)
def step(l,t,n,head,body,who,acc):
    _rect(s,l,t,5.92,1.36,LIGHT); _rect(s,l,t,0.045,1.36,acc)
    txt(s,l+0.30,t+0.12,0.70,0.40,n,28,True,acc)
    txt(s,l+1.00,t+0.16,4.70,0.26,head,12,True,NAVY,lh=1.15)
    txt(s,l+1.00,t+0.50,4.70,0.72,body,8.6,False,GREY,lh=1.36)
    txt(s,l+1.00,t+1.12,4.70,0.14,who,7.2,True,acc)
step(0.62,3.40,"1","Показати SIAS ринок і запросити ціну",
 "Слайди 32–33: конкуренти й акції, полиця 120 грн, ціна закупівлі €0,54 для машини. Попросити розшифровку СС за статтями.","ТЕРМІН: до переговорів · ВІДПОВІДАЄ: закупівлі",RED)
step(6.80,3.40,"2","Вирішити: машина чи 6 палет",
 "6 палет за поточних цін не дають полиці 120 грн за жодної маржі ≥ 15 %. Повна машина — єдиний реалістичний шлях.","ТЕРМІН: до переговорів · ВІДПОВІДАЄ: керівництво",RED)
step(0.62,4.88,"3","З’ясувати статус Choi’s у «Таврії В»",
 "Хто діючий імпортер, чи є ексклюзив, за якою ціною ввозили, історія цін. Без цього наша 120 грн може вступити в конфлікт.","ТЕРМІН: паралельно · ВІДПОВІДАЄ: категорійний менеджмент",DARKAMBER)
step(6.80,4.88,"4","Закрити невизначеності моделі",
 "Специфікації SKU (сегмент), режим ПДВ, лістинг і промо-бюджет мережі. Від них залежить, чи реальна маржа 30 %.","ТЕРМІН: до рішення · ВІДПОВІДАЄ: фінанси",TEAL)
verdict(s,6.34,"РІШЕННЯ","Вимагати закупівлю €0,54 за повну машину або свідомо йти на маржу 14–22 %: з 6 палет полиця 120 грн недосяжна.",RED,0.92)

out="/home/user/Chaplygin-Illya/SIAS_MARKET_RESEARCH_UA_V4.pptx"
prs.save(out); print("saved",out,"slides:",len(prs.slides._sldIdLst))
