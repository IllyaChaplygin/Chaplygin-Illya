import statistics as st, math, copy, collections
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

W, H = 13.3333, 7.5
M = 0.66
HDR = 1.16
NAVY = RGBColor(0x30, 0x3A, 0x5D); NAVY_L = RGBColor(0x46, 0x53, 0x82)
ORANGE = RGBColor(0xF9, 0xA5, 0x0B); AMBER = RGBColor(0xFF, 0xC9, 0x5C)
TEAL = RGBColor(0x1E, 0x9E, 0xA6); GREEN = RGBColor(0x37, 0xA1, 0x69)
PLUM = RGBColor(0x5E, 0x4B, 0x94); ROSE = RGBColor(0xC9, 0x4F, 0x7C)
GOLD = RGBColor(0xC4, 0x81, 0x0A); SLATE = RGBColor(0x5B, 0x6B, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); MIST = RGBColor(0xF1, 0xF4, 0xFA); MIST_D = RGBColor(0xE2, 0xE8, 0xF3)
INK = RGBColor(0x1C, 0x22, 0x35); GREY = RGBColor(0x71, 0x78, 0x90); LGREY = RGBColor(0xA9, 0xB0, 0xC4)
FONT = "Segoe UI"
V3 = '/tmp/claude-0/-home-user-Chaplygin-Illya/c5886f59-fb29-5de4-ac0f-e7da8a3536ab/scratchpad/v3/'
RTE = '/tmp/claude-0/-home-user-Chaplygin-Illya/c5886f59-fb29-5de4-ac0f-e7da8a3536ab/scratchpad/rte/'
LOGO = V3 + 'logo.png'

FMT_COL = {'pouch': GOLD, 'cup': TEAL, 'doypack': PLUM, 'box': ROSE}
SEG_COL = {1: GOLD, 2: TEAL, 3: ROSE, 4: PLUM}

prs = Presentation(V3 + 'user.pptx')
LAYOUT = prs.slide_layouts[0]


def rect(s, x, y, w, h, fill, line=None, lw=1.0, rounded=False, adj=0.10):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(max(w, 0.001)), Inches(max(h, 0.001)))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if rounded: sh.adjustments[0] = adj
    return sh


def dot(s, cx, cy, d, fill, line=None, lw=1.5):
    return rect(s, cx - d / 2, cy - d / 2, d, d, fill, line=line, lw=lw, rounded=True, adj=0.5)


def text(s, x, y, w, h, body, size=11, bold=False, italic=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=None):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    items = body if isinstance(body, list) else [body]
    for i, item in enumerate(items):
        txt, over = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = over.get("align", align)
        if line: p.line_spacing = line
        for chunk in (txt if isinstance(txt, list) else [txt]):
            ctxt, cover = chunk if isinstance(chunk, tuple) else (chunk, {})
            r = p.add_run(); r.text = ctxt; f = r.font
            f.name = FONT
            f.size = Pt(cover.get("size", over.get("size", size)))
            f.bold = cover.get("bold", over.get("bold", bold))
            f.italic = cover.get("italic", over.get("italic", italic))
            f.color.rgb = cover.get("color", over.get("color", color))
    return box


def pic(s, path, x, y, w, h):
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih); pw, ph = iw * sc, ih * sc
    return s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def new_slide():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    rect(s, 0, 0, W, H, WHITE)
    return s


def header(s, eyebrow, title, dek=None):
    rect(s, 0, 0, W, HDR, NAVY); rect(s, 0, HDR, W, 0.055, ORANGE)
    s.shapes.add_picture(LOGO, Inches(M), Inches(0.26), Inches(1.52), Inches(0.66))
    text(s, 2.52, 0.34, 9.3, 0.24, eyebrow, size=10, bold=True, color=AMBER)
    text(s, 2.52, 0.60, 9.3, 0.46, title, size=19, bold=True, color=WHITE)
    text(s, W - M - 0.70, 0.44, 0.70, 0.36, "00", size=17, bold=True, color=NAVY_L, align=PP_ALIGN.RIGHT)
    if dek:
        text(s, M, HDR + 0.20, 12.0, 0.34, dek, size=12.5, color=GREY, line=1.15)


def foot(s, src):
    text(s, M, 7.08, 12.0, 0.30, src, size=8.5, color=GREY, line=1.12)


def num(v, d=0):
    return f"{v:,.{d}f}".replace(",", " ").replace(".", ",")


def pl(n, one, few, many):
    n = int(n)
    if n % 10 == 1 and n % 100 != 11: return f"{n} {one}"
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14: return f"{n} {few}"
    return f"{n} {many}"


def rngs(lo, hi):
    a, b = num(lo), num(hi)
    return a if a == b else f"{a}–{b}"


def insight(s, x, y, w, h, title, body, c=ORANGE, size=11):
    rect(s, x, y, w, h, NAVY, rounded=True, adj=0.06); rect(s, x, y, 0.09, h, c, rounded=True, adj=0.5)
    if title: text(s, x + 0.34, y + 0.16, w - 0.6, 0.3, title, size=13, bold=True, color=AMBER)
    text(s, x + 0.34, y + (0.54 if title else 0.2), w - 0.6, h - (0.66 if title else 0.36), body,
         size=size, color=RGBColor(0xD5, 0xDB, 0xEA), line=1.25)


def chip(s, x, y, label, bg, fg=WHITE, size=9, w=None, h=0.26):
    w = w or (0.26 + 0.075 * len(label))
    rect(s, x, y, w, h, bg, rounded=True, adj=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=fg, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def sku_cell(s, x, y, w, h, img, name, grams, lo, hi, seller, c, brand=None, tag=None):
    rect(s, x, y, w, h, WHITE, line=MIST_D, lw=1.0, rounded=True, adj=0.06)
    ph = h - 1.46
    rect(s, x + 0.07, y + 0.07, w - 0.14, ph, MIST, rounded=True, adj=0.07)
    if img: pic(s, img, x + 0.12, y + 0.11, w - 0.24, ph - 0.08)
    yy = y + ph + 0.14
    text(s, x + 0.11, yy, w - 0.22, 0.40, name, size=9.5, bold=True, color=NAVY, line=1.0)
    gt = f"{num(grams, 0 if float(grams).is_integer() else 1)} г" if grams else "маса н/д"
    if brand:
        text(s, x + 0.11, yy + 0.42, w - 0.22, 0.2, [([(brand.upper() + " · ", {"bold": True, "color": c}), gt], {})],
             size=8 if len(brand) < 15 else 6.8, color=GREY)
    else:
        text(s, x + 0.11, yy + 0.42, w - 0.22, 0.2, gt, size=8, color=GREY)
    price = (num(lo) if lo == hi else f"{num(lo)}–{num(hi)}") + " грн"
    text(s, x + 0.11, yy + 0.62, w - 0.22, 0.3, price, size=12.5, bold=True, color=NAVY)
    text(s, x + 0.11, yy + 0.94, w - 0.22, 0.34, seller, size=7.5, color=GREY, line=1.05)


# ── службові дії над колодою ────────────────────────────────────────────────
def slide_ids():
    return list(prs.slides._sldIdLst)


def drop_slide(sldId):
    prs.part.drop_rel(sldId.rId)
    prs.slides._sldIdLst.remove(sldId)


def reorder(order_ids):
    lst = prs.slides._sldIdLst
    for el in list(lst): lst.remove(el)
    for el in order_ids: lst.append(el)


def renumber():
    n = 0
    for s in prs.slides:
        n += 1
        for sh in s.shapes:
            if sh.has_text_frame and abs(sh.left - 10945440) < 400000 and abs(sh.top - 402480) < 200000 \
                    and sh.text_frame.text.strip().isdigit():
                r = sh.text_frame.paragraphs[0].runs[0]; r.text = f"{n:02d}"
                for extra in sh.text_frame.paragraphs[0].runs[1:]: extra.text = ""
