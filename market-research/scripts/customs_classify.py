"""Розбір митної бази (codexmb, вивантаження 28.09.2026) по коду 1904901000:
відділяє справжній готовий рис для розігріву від заморожених сирих сумішей,
рисових чіпсів, долми й різото "довести до готовності", які йдуть під тим
самим 10-значним кодом УКТ ЗЕД. Джерело: RICE_PRICING... не цей файл, а
окремий файл митної бази, наданий користувачем 05.10.2026."""
import openpyxl
import collections

wb = openpyxl.load_workbook("customs.xlsx", data_only=True)
ws = wb["Сирі дані"]
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
idx = {h: i for i, h in enumerate(hdr)}
rows = [r for r in ws.iter_rows(min_row=2, values_only=True)
        if r[idx["hs_code_normalized"]] == "1904901000"]


def tier(desc, man):
    d = (desc or "").upper()
    m = (man or "").upper()
    if "ЗАМОРОЖ" in d:
        return "X_frozen"
    if "ЧІПС" in d:
        return "X_chips"
    if "ПОСИПК" in d:
        return "X_topping"
    if "ДОЛМ" in d:
        return "X_dolma_canned"
    if ("НЕПРИГОТОВАН" in d or "ДЛЯ ПРИГОТУВАННЯ РІЗОТТО" in d or "РИЗОТО CORDERO" in d
            or "CASA RINALDI" in d or "RISO GALLO" in d or "PRINCIPATO" in d):
        return "X_cook_it_yourself_risotto"
    if "RICE PORRIDGE" in d or "GYMBEAM" in m:
        return "X_sport_porridge"
    if "YOPOKKI" in d or "ТТОКПОКК" in d:
        return "X_tteokbokki_snack"
    if "РИС 10 ХВ" in d or "АШАН РИС" in d or "РИС ТАЙСЬКИЙ" in d or "РИС КРУГЛИЙ" in d:
        return "X_raw_rice"
    if "MARS AUSTRIA" in m or "UNCLE BEN" in d or "BEN'S" in d:
        return "A_bens"
    if "CJCHEIL" in m:
        return "A_cj_hetbahn"
    if "CLEARSPRING" in m:
        return "A_clearspring"
    if "INSTANT RICE WITH" in d or ("174 G" in d and "РИС ШВИДКОГО" in d) or \
            ("144" in d and "РИС ШВИДКОГО" in d):
        return "B_henan_boilwater"
    if "RTE BIRYANI" in d or "HALDIRAM" in m:
        return "B_haldiram_biryani"
    return "UNCLASSIFIED"


cnt, wt, val = collections.Counter(), collections.Counter(), collections.Counter()
for r in rows:
    t = tier(r[idx["product_description"]], r[idx["manufacturer_normalized"]])
    cnt[t] += 1
    wt[t] += r[idx["net_weight_kg"]] or 0
    val[t] += r[idx["invoice_value_usd"]] or 0

for k in sorted(wt, key=lambda k: -wt[k]):
    print(f"{k:32s} n={cnt[k]:4d} wt={wt[k]:10.1f}kg  usd={val[k]:10.1f}")
