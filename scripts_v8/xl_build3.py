# -*- coding: utf-8 -*-
import openpyxl,sys
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.workbook.properties import CalcProperties
SRC="/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/5927c2e7-RAMEN_FINMODEL_4_TABS_PRESENTATION.xlsx"
OUT=sys.argv[1]
NAVY='223A5E'; INK='172B46'; GREYT='5E6D7E'; KEY='E7EDF5'; BAND2='F2F5F9'; INP='FFF4D2'; BLUE='0000FF'; GREEN='008000'; AFTER='FCE3A6'
def font(sz=11,b=False,c=INK,i=False): return Font(name='Arial',size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill('solid',fgColor=c)
thin=Side(style='thin',color='D5DCE8'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
UAH='#,##0.00'; EUR='"€"0.000'; PCT='0.0%'; INT='#,##0'
wb=openpyxl.load_workbook(SRC)
CFG=[dict(name='Цена 120 грн — 6 паллет',src='Финмодель 6 паллет',head='6 паллет',freight=2600,pallets=6),
     dict(name='Цена 120 грн — 33 паллет',src='Финмодель 33 паллет',head='33 паллеты (Full Truck)',freight=3700,pallets=33)]
def build(cfg):
    idx=wb.sheetnames.index(cfg['name']); del wb[cfg['name']]
    ws=wb.create_sheet(cfg['name'],idx); S="'"+cfg['src']+"'"
    ws.sheet_properties.tabColor=NAVY; ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=2; ws.column_dimensions['B'].width=34
    for col in "CDEFGHIJKLMNOP": ws.column_dimensions[col].width=11
    ws.column_dimensions['Q'].width=2
    CT=Alignment(horizontal='center',vertical='center',wrap_text=True); LF=Alignment(horizontal='left',vertical='center',wrap_text=True)
    def put(a,v=None,f=None,fl=None,fmt=None,al=None,border=False,merge=None):
        c=ws[a]; c.value=v; c.font=f or font(); c.alignment=al or CT
        if fl: c.fill=fill(fl)
        if fmt: c.number_format=fmt
        if border: c.border=BOX
        if merge:
            ws.merge_cells(merge)
            for row in ws[merge]:
                for cc in row:
                    if fl: cc.fill=fill(fl)
                    if border: cc.border=BOX
        return c
    ws.row_dimensions[1].height=8
    put("B2",f"SIAS · было → стало · от себестоимости в Украине до полки · {cfg['head']}",font(16,True),al=LF,merge="B2:P2"); ws.row_dimensions[2].height=28
    # --- параметры (ниже), адреса:
    # строка 15 подписи / 16 значения: C:D полка, E:F маржа, G:H бонус, I:J множитель, K:L курс, M:N фин., O:P ПДВ
    # строка 18/19: C:D пошлина, E:F перевозка, G:H брокер, I:J прочие, K:L доля фрахта, M:N паллет, O:P пачек на паллете
    SH,MG,BO,MK,RT,FN,VT="C16","E16","G16","I16","K16","M16","O16"
    DT,FR,BR,OT,FS,PL,PP="C19","E19","G19","I19","K19","M19","O19"
    BATCH=f"({FR}+{BR}+{OT}+{FR}*{FS}*({DT}+{VT}*(1+{DT})))"; PACKS=f"({PL}*{PP})"
    ASK="C22"
    put("B3",f'="Чтобы полка стала "&ROUND({SH},0)&" ₴, закупка SIAS должна упасть с €"&MIN(C7:C8)&"–€"&MAX(C7:C8)&" до €"&ROUND({ASK},2)&" за пачку (−"&ROUND((1-{ASK}/MIN(C7:C8))*100,0)&"…"&ROUND((1-{ASK}/MAX(C7:C8))*100,0)&"%). Маржа остаётся "&ROUND({MG}*100,0)&"%."',font(11,True,NAVY),KEY,al=LF,merge="B3:P3"); ws.row_dimensions[3].height=28
    # --- главный блок
    groups=["Закупка SIAS","Себестоимость в Украине (СС)","Цена партнёру (магазину)","Бонус сети","Наша прибыль","Наценка магазина","ПОЛКА с НДС"]
    put("B5","",font(9,True,'FFFFFF'),'3E5A85'); put("B6","Вариант",font(9,True,'FFFFFF'),'3E5A85')
    for i,g in enumerate(groups):
        c1=openpyxl.utils.get_column_letter(3+2*i); c2=openpyxl.utils.get_column_letter(4+2*i)
        put(f"{c1}5",g,font(9,True,'FFFFFF'),'3E5A85',merge=f"{c1}5:{c2}5")
        put(f"{c1}6","€",font(9,True,'FFFFFF'),'5A74A0'); put(f"{c2}6","₴",font(9,True,'FFFFFF'),'5A74A0')
    ws.row_dimensions[5].height=32; ws.row_dimensions[6].height=18
    # порядок колонок: закупка € C, ₴ D? сделаем «€ слева, ₴ справа»
    # индексы: закупка C(€) D(₴); СС E F; партнёр G H; бонус I J; прибыль K L; наценка M N; полка O P
    rows=[(7,"ДО · Vegetable, Spicy, Beef, Chicken",0.65,f"={S}!E13",False),
          (8,"ДО · Gouda Cheese, Carbonara Spicy",0.75,f"={S}!E11",False),
          (9,"ПОСЛЕ · запрос SIAS (все 6 SKU)",None,None,True)]
    for r,lab,buy,cc_link,after in rows:
        fl=AFTER if after else (BAND2 if r==7 else None)
        put(f"B{r}",lab,font(10,True,NAVY if after else INK),fl,al=LF,border=True)
        # закупка
        if after: put(f"C{r}",f"={ASK}",font(11,True),fl,EUR,CT,border=True)
        else: put(f"C{r}",buy,font(11,True,BLUE),INP,EUR,CT,border=True)
        put(f"D{r}",f"=C{r}*{RT}",font(11),fl,UAH,CT,border=True)
        # себестоимость
        if after:
            ccuah=f"=(C{r}*{RT}+({FR}+{BR}+{OT})/{PACKS}*{RT}+(C{r}+{FR}*{FS}/{PACKS})*{DT}*{RT}+(C{r}+{FR}*{FS}/{PACKS})*(1+{DT})*{VT}*{RT})*(1+{FN})"
            put(f"F{r}",ccuah,font(11,True),fl,UAH,CT,border=True); put(f"E{r}",f"=F{r}/{RT}",font(11),fl,EUR,CT,border=True)
        else:
            put(f"E{r}",cc_link,font(11,False,GREEN),fl or BAND2,EUR,CT,border=True); put(f"F{r}",f"=E{r}*{RT}",font(11,True),fl,UAH,CT,border=True)
        # остальное — из СС одинаково для всех строк
        put(f"H{r}",f"=F{r}/(1-{BO}-{MG})",font(11,True),fl,UAH,CT,border=True)
        put(f"J{r}",f"=H{r}*{BO}",font(11),fl,UAH,CT,border=True)
        put(f"L{r}",f"=H{r}-F{r}-J{r}",font(11),fl,UAH,CT,border=True)
        put(f"N{r}",f"=H{r}*({MK}-1)",font(11),fl,UAH,CT,border=True)
        put(f"P{r}",f"=H{r}*{MK}",font(12,True),fl or KEY,UAH,CT,border=True)
        for ce,cu in (("G","H"),("I","J"),("K","L"),("M","N"),("O","P")):
            put(f"{ce}{r}",f"={cu}{r}/{RT}",font(10,False,GREYT),fl,EUR,CT,border=True)
        ws.row_dimensions[r].height=30
    # разница
    put("B10","ИЗМЕНЕНИЕ: ПОСЛЕ − ДО (4 SKU)",font(10,True,NAVY),KEY,al=LF,border=True)
    for ci in range(3,17):
        col=openpyxl.utils.get_column_letter(ci)
        put(f"{col}10",f"={col}9-{col}7",font(10,True,'B53A3A'),KEY,('+"€"0.000;-"€"0.000' if ci%2==1 else '+#,##0.00;-#,##0.00'),CT,border=True)
    ws.row_dimensions[10].height=24
    put("B11","Читайте по строке слева направо: закупка → себестоимость в Украине → цена партнёру = СС ÷ (1 − бонус − маржа) → бонус, наша прибыль, наценка → полка. Каждая цифра дана в € и ₴.",font(9,False,GREYT,True),al=LF,merge="B11:P11"); ws.row_dimensions[11].height=26
    # --- параметры
    put("B13","ИСХОДНЫЕ ДАННЫЕ — жёлтое можно менять, зелёное из вкладки «"+cfg['src']+"»",font(9,True,NAVY),al=LF,merge="B13:P13")
    for c in "BCDEFGHIJKLMNOP": ws[f"{c}13"].border=Border(bottom=Side(style='medium',color=NAVY))
    P1=[("Полка, ₴",120,UAH,'inp'),("Наша маржа",0.30,PCT,'inp'),("Бонус сети",f"={S}!D4",PCT,'lnk'),("Множитель полки",f"={S}!D6",'0.00"×"','lnk'),("Курс ₴/€",f"={S}!D3",'0.00','lnk'),("Финансирование",0.025,PCT,'inp'),("Импортный ПДВ",0.20,PCT,'inp')]
    P2=[("Пошлина",0,PCT,'inp'),("Перевозка партии, €",cfg['freight'],INT,'inp'),("Брокер, €",288,INT,'inp'),("Прочие, €",42,INT,'inp'),("Доля фрахта в таможне",0.7,'0.00','inp'),("Паллет в партии",cfg['pallets'],INT,'inp'),("Пачек на паллете",1600,INT,'inp')]
    for rl,rv,PP_ in ((15,16,P1),(18,19,P2)):
        for i,(lab,val,nf,kind) in enumerate(PP_):
            c1=openpyxl.utils.get_column_letter(3+2*i); c2=openpyxl.utils.get_column_letter(4+2*i)
            put(f"{c1}{rl}",lab,font(8,True,GREYT),merge=f"{c1}{rl}:{c2}{rl}")
            put(f"{c1}{rv}",val,font(11,True,BLUE if kind=='inp' else GREEN),INP if kind=='inp' else BAND2,nf,CT,border=True,merge=f"{c1}{rv}:{c2}{rv}")
        ws.row_dimensions[rl].height=22; ws.row_dimensions[rv].height=22
    put("B16","Цель расчёта",font(9,True,GREYT),al=LF); put("B19","Расходы партии (Self-Cost_Sias)",font(9,True,GREYT),al=LF)
    # --- запрос
    put("B21","КАК ПОЛУЧЕН ЗАПРОС: из полки 120 ₴ обратно к закупке",font(9,True,NAVY),al=LF,merge="B21:P21")
    for c in "BCDEFGHIJKLMNOP": ws[f"{c}21"].border=Border(bottom=Side(style='medium',color=NAVY))
    put("B22","Запрос SIAS, €/уп. (вниз до евроцента)",font(10,True),KEY,al=LF,border=True)
    put("C22",f"=ROUNDDOWN((({SH}/{MK}*(1-{BO}-{MG}))/{RT}/(1+{FN})-{BATCH}/{PACKS})/((1+{DT})*(1+{VT})),2)",font(14,True),AFTER,'"€"0.00',CT,border=True,merge="C22:D22")
    put("E22","Формула: (полка ÷ множитель × (1 − бонус − маржа) ÷ курс ÷ (1 + фин.) − расходы партии на пачку) ÷ (1 + ПДВ)(1 + пошлина), округлено вниз",font(9,False,GREYT),al=LF,merge="E22:P22")
    put("B23","Предельная закупка до округления, €/уп.",font(10),None,al=LF,border=True)
    put("C23",f"=(({SH}/{MK}*(1-{BO}-{MG}))/{RT}/(1+{FN})-{BATCH}/{PACKS})/((1+{DT})*(1+{VT}))",font(11),None,'"€"0.0000',CT,border=True,merge="C23:D23")
    put("B24","Допущения: бонус и множитель — по модели (не подтверждены сетями); импортный ПДВ входит в себестоимость, как в Self-Cost; без листинг-фи, промо и отсрочки; согласие SIAS не получено — цена для переговоров.",font(8,False,GREYT,True),al=LF,merge="B24:P24"); ws.row_dimensions[24].height=28
    ws.row_dimensions[22].height=30; ws.row_dimensions[23].height=30
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
for c in CFG: build(c)
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print("saved",OUT)
