import openpyxl,math,zipfile,sys
from lxml import etree
F=sys.argv[1]
nv=openpyxl.load_workbook(F,data_only=True); nf=openpyxl.load_workbook(F)
O=openpyxl.load_workbook('/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/5927c2e7-RAMEN_FINMODEL_4_TABS_PRESENTATION.xlsx',data_only=True)
SC=openpyxl.load_workbook('/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/87c2996d-Self-Cost_Sias.xlsx',data_only=True).worksheets[0]
fail=0
def chk(name,a,b,tol=1e-9):
    global fail
    ok=abs(a-b)<=tol*max(1,abs(b)); fail+=0 if ok else 1
    if not ok: print('FAIL',name,a,b)
    return ok
# 1. XML well-formed + ошибки
z=zipfile.ZipFile(F); assert z.testzip() is None
for n in z.namelist():
    if n.endswith('.xml') or n.endswith('.rels'): etree.fromstring(z.read(n))
errs=0
for ws in nf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='):
                v=nv[ws.title][c.coordinate].value
                if isinstance(v,str) and v.startswith('#'): errs+=1; print('ERR',ws.title,c.coordinate,v)
print('formula errors:',errs)
# 2. независимый пересчёт с нуля
RATE=51.2;BON=.25;MARG=.30;MARK=1.4;FIN=.025;VAT=.2
for sheet,pal,freight,rowsSC,srcname in (('Цена 120 грн — 33 паллет',33,3700,(32,30),'Финмодель 33 паллет'),('Цена 120 грн — 6 паллет',6,2600,(8,6),'Финмодель 6 паллет')):
    w=nv[sheet]; packs=pal*1600; batch=freight+288+42+freight*0.7*VAT
    cap=120/MARK*(1-BON-MARG); mx=(cap/RATE/(1+FIN)-batch/packs)/(1+VAT); ask=math.floor(mx*100+1e-9)/100
    prod=ask*RATE; logi=(freight+288+42)/packs*RATE; vat=(ask+freight*0.7/packs)*VAT*RATE; fin=(prod+logi+vat)*FIN; cc=prod+logi+vat+fin
    partner=cc/(1-BON-MARG); bonus=partner*BON; profit=partner-cc-bonus; markup=partner*(MARK-1); shelf=partner*MARK
    chk(sheet+' ask',w['C21'].value,ask); chk(sheet+' max',w['C22'].value,mx); chk(sheet+' CC',w['D9'].value,cc)
    for col,val in zip("DFHJLN",(cc,partner,bonus,profit,markup,shelf)): chk(f'{sheet} ПОСЛЕ ₴ {col}',w[f'{col}9'].value,val)
    for col,val in zip("CEGIKM",(cc,partner,bonus,profit,markup,shelf)): chk(f'{sheet} ПОСЛЕ € {col}',w[f'{col}9'].value,val/RATE)
    for r,(sk,price) in zip((7,8),((rowsSC[0],0.65),(rowsSC[1],0.75))):
        goods=SC[f'J{sk}'].value; assert abs(goods-price)<1e-9
        cc0=SC[f'AJ{sk}'].value*RATE; p0=cc0/(1-BON-MARG)
        for col,val in zip("DFHJLN",(cc0,p0,p0*BON,p0-cc0-p0*BON,p0*(MARK-1),p0*MARK)): chk(f'{sheet} ДО{r} ₴ {col}',w[f'{col}{r}'].value,val)
    for ci in range(3,15):
        col=openpyxl.utils.get_column_letter(ci); chk(f'{sheet} Δ {col}',w[f'{col}10'].value,w[f'{col}9'].value-w[f'{col}7'].value)
    # цель: полка после ≤ 120 и маржа 30 %
    assert w['N9'].value<=120 and abs((w['J9'].value/w['F9'].value)-0.30)<1e-9
    # сверка с исходной моделью
    o=O[sheet]
    for a,b in (('D9','H73'),('F9','H74'),('H9','H75'),('J9','H76'),('L9','H77'),('N9','H78'),('C21','A6'),('C22','H66')): chk(f'{sheet} vs model {a}/{b}',w[a].value,o[b].value)
    print(sheet,'ask',ask,'СС после',round(cc,4),'полка',round(shelf,4),'| headline:',w['B3'].value)
    # скрытые строки и видимые настройки
    hid=[r for r,d in nf[sheet].row_dimensions.items() if d.hidden]; print('  hidden rows',hid)
# 3. вкладки 1–2 не изменены
bad=sum(1 for n in O.sheetnames[:2] for row in O[n].iter_rows() for c in row if isinstance(c.value,(int,float)) and abs(c.value-nv[n][c.coordinate].value)>1e-9)
print('tabs1-2 numeric diffs vs original:',bad); fail+=bad
print('TOTAL FAILS',fail)
