# -*- coding: utf-8 -*-
"""V7: правки поверх файлу користувача (його видалення збережено). Нові слайди додаються в кінець, потім переставляються."""
import sys,os,re,math,json,statistics as st,collections as co
sys.path.insert(0,"/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad/v5")
import kit2
from kit2 import *
from pptx import Presentation
from pptx.util import Inches,Emu
from deckkit import _rect
UP="/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/b2d4af9e-SIAS_MARKET_RESEARCH_UA_V6.pptx"
prs=Presentation(UP)
N_USER=len(prs.slides._sldIdLst)
USLIDE=[sl.slide_id for sl in prs.slides]
CREAM=RGBColor(0xFF,0xF6,0xE3); SHADOW=RGBColor(0xDF,0x91,0x00)
GREEN=RGBColor(0x4E,0x8F,0x6B); BLUE2=RGBColor(0x3E,0x59,0xA8); GREYB=RGBColor(0x9A,0xA1,0xB8)
BC={"Nongshim":NAVY,"Samyang":RED,"Paldo":PURPLE,"Ottogi":TEAL,"Bibigo":BLUE2,"O'Food":DARKAMBER,"Otoki":GREYB,"fire bull":GREYB,"Choi's":AMBER}
def rr(s,l,t,w,h,fill,rad=0.08):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    sh.line.fill.background(); sh.shadow.inherit=False; sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.adjustments[0]=rad; return sh
def oval(s,l,t,d,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(l),Inches(t),Inches(d),Inches(d))
    sh.line.fill.background(); sh.shadow.inherit=False; sh.fill.solid(); sh.fill.fore_color.rgb=fill; return sh
def photo(s,path,l,t,w,h,pad=0.06):
    from PIL import Image
    im=Image.open(path); ar=im.width/im.height; bw,bh=w-2*pad,h-2*pad
    if bw/bh>ar: ph=bh; pw=ph*ar
    else: pw=bw; ph=pw/ar
    s.shapes.add_picture(path,Inches(l+(w-pw)/2),Inches(t+(h-ph)/2),Inches(pw),Inches(ph))
def newslide(kicker,title,sub=None): return slide(prs,kicker,title,"00",sub)
def legend(s,x,y,items,gap=0.25,fs=8.5):
    for lab,col in items:
        _rect(s,x,y+0.02,0.16,0.14,col); w=0.075*len(lab)+0.1
        txt(s,x+0.22,y,w,0.18,lab,fs,False,NAVY); x+=0.22+w+gap
    return x
def dig(x,n=1): return f"{x:.{n}f}".replace('.',',')

# ============ ОБКЛАДИНКА ============
s=prs.slides.add_slide(prs.slide_layouts[6])
_rect(s,0,0,13.333,7.5,CREAM)
_rect(s,7.25,0,6.083,7.5,AMBER)
oval(s,7.5,0.7,5.6,GOLD); oval(s,6.6,4.7,2.5,AMBER); oval(s,10.4,4.6,2.8,GOLD)
_rect(s,0,7.22,13.333,0.28,NAVY)
rr(s,0.80,0.62,2.55,1.12,NAVY,0.12)
s.shapes.add_picture(LOGO,Inches(0.98),Inches(0.72),Inches(2.19),Inches(0.95))
txt(s,0.85,2.55,6.2,0.25,"ДОСЛІДЖЕННЯ РИНКУ ТА ПОЗИЦІОНУВАННЯ · ЖОВТЕНЬ 2026",11,True,DARKAMBER)
txt(s,0.82,2.95,6.4,2.0,"Корейський\nрамен\nв Україні",58,True,NAVY,lh=1.0)
_rect(s,0.88,6.05,1.6,0.09,AMBER)
cover_ph=[("08801073110502",7.55,0.55,2.55,2.45),("08801043150620",7.55,3.2,2.55,2.45),
          ("00648436310685",10.5,1.45,2.55,2.45),("03760344350847",10.5,4.15,2.55,2.45)]
for e,l,t,w,h in cover_ph:
    rr(s,l+0.06,t+0.07,w,h,SHADOW,0.06); rr(s,l,t,w,h,WHITE,0.06)
    photo(s,f"{SC}/img/{e}.jpg",l,t,w,h,0.10)
NEW={'cover':s}

# ============ ПОХОДЖЕННЯ: УСІ НАПРЯМКИ ============
TOT=838.480003
CT=[("В’єтнам",166.687505),("Корея",118.497356),("Румунія",86.919064),("Велика Британія",58.564844),
    ("Латвія",57.30452),("Чехія",50.800755),("Польща",48.761462)]
OTH=TOT-sum(v for _,v in CT)
CT=CT+[("Інші країни разом",OTH)]
CPAL=[NAVY,BLUE2,TEAL,PURPLE,DARKAMBER,RED,GREEN,GREYB]
s=newslide("МИТНА БАЗА · ПОХОДЖЕННЯ","Звідки їде локшина: усі напрямки ввозу","Січень–травень 2026 · вага нетто, тонн · частка в загальному ввозі 838,5 т")
txt(s,0.62,1.78,6,0.16,"СТРУКТУРА ВВОЗУ, %",8.4,True,NAVY)
x=0.62; W=12.10
for (nm,v),col in zip(CT,CPAL):
    w=W*v/TOT; _rect(s,x,2.02,w-0.02,0.50,col)
    if w>0.6: txt(s,x,2.17,w-0.02,0.2,f"{v/TOT*100:.1f}".replace('.',',')+" %",9.5,True,WHITE,C)
    x+=w
txt(s,0.62,2.72,6,0.16,"ПО КРАЇНАХ ПОХОДЖЕННЯ",8.4,True,NAVY)
y=3.02; RH=0.44
for (nm,v),col in zip(CT,CPAL):
    _rect(s,0.62,y+0.10,0.14,0.14,col)
    txt(s,0.88,y+0.06,2.5,0.24,nm,10.5,True,NAVY)
    bw=6.4*v/300
    _rect(s,3.5,y+0.04,max(bw,0.03),0.27,col)
    txt(s,3.5+bw+0.12,y+0.06,1.3,0.24,dig(v)+" т",10.5,True,NAVY)
    txt(s,11.4,y+0.06,1.3,0.24,dig(v/TOT*100)+" %",10.5,True,GREY,R)
    y+=RH
_rect(s,0.62,y+0.04,12.10,0.012,NAVY)
txt(s,0.88,y+0.12,2.5,0.24,"Усього",10.5,True,DARKAMBER); txt(s,3.5,y+0.12,2.0,0.24,dig(TOT)+" т",10.5,True,DARKAMBER); txt(s,11.4,y+0.12,1.3,0.24,"100 %",10.5,True,DARKAMBER,R)
txt(s,0.62,7.04,12.1,0.2,"Митна база користувача, код УКТ ЗЕД 1902 30 10 00. «Інші країни» — одним рядком: розбивки за країнами в базі немає.",7.4,False,PALE)
NEW['origin']=s

# ============ СЕГМЕНТИ: ОГЛЯДИ І КАТАЛОГИ (великі фото) ============
NOW_LOW=M.forward_current("33")[2]['shelf']
def n_above(arr,x): return sum(1 for p in arr if p>x)
def bigcard(sl,l,t,w,h,c,acc,sq=1.55):
    _rect(sl,l,t,w,h,LIGHT); _rect(sl,l,t,w,0.04,acc)
    _rect(sl,l+0.08,t+0.12,w-0.16,sq,WHITE)
    if c.get('img') and os.path.exists(c['img']): photo(sl,c['img'],l+0.08,t+0.12,w-0.16,sq,0.04)
    else: txt(sl,l+0.08,t+0.12+sq/2-0.08,w-0.16,0.16,"фото немає",8,False,PALE,C)
    y=t+0.12+sq+0.08
    txt(sl,l+0.10,y,w-0.2,0.34,c['name'],8.4,True,INK,lh=1.1)
    txt(sl,l+0.10,y+0.35,w-0.2,0.24,c['price'].replace(' грн',' ₴'),13,True,INK)
    sel=c['sellers']; sel=sel if len(sel)<=24 else sel[:23]+"…"
    txt(sl,l+0.10,y+0.62,w-0.2,0.2,c['w'].replace(' Г',' г')+" · "+sel,7.2,False,GREY)
def sample_k(k,n=6):
    arr=sorted([c for c in SEG.get(k,[]) if c.get('pmed')],key=lambda c:c['pmed']); out=[]
    for i in range(n):
        j=round(i*(len(arr)-1)/(n-1)) if len(arr)>1 else 0
        if arr[j] not in out: out.append(arr[j])
    return out
def band_text(k):
    v=SS[k]; return f"СЕГМЕНТ {k} · {SNAME[k]}   ·   {SPREP[k]}   ·   вага {v['wlo']}–{v['whi']} г"
CX=[0.62+i*(1.92+0.116) for i in range(6)]
NEWSEG=[]
def overview(k,title,sub=None):
    v=SS[k]; col=SCOL[k]
    sl=slide(prs,f"СЕГМЕНТ {k} · {SNAME[k]}",title,"00",sub)
    _pp=[r['pmed'] for r in SD_ROWS if r['seg']==int(k) and r.get('pmed')]
    estat(sl,0.62,1.74,"ПОЗИЦІЙ У СЕГМЕНТІ",f"{v['n']} SKU",f"{v['brands']} брендів у {v['chains']} мережах.\n{round(v['n']*100/F['n_sku'])} % усієї категорії.",col,3.10,1.18,22)
    estat(sl,3.88,1.74,"МЕДІАНА ПОЛИЦІ",f"{v['med']} ₴",f"P25–P75: {v['p25']}–{v['p75']} ₴.\nДіапазон: {v['lo']}–{v['hi']} ₴.",col,3.10,1.18,22)
    estat(sl,7.14,1.74,"ДОРОЖЧІ ЗА ЦІЛЬ 120 ₴",f"{n_above(_pp,120)} з {len(_pp)}",f"позицій сегмента.\nДорожчі за 140,55 ₴ (зараз): {n_above(_pp,NOW_LOW)}.",col,3.10,1.18,22)
    estat(sl,10.40,1.74,"ВАГА ПАЧКИ",f"{v['wlo']}–{v['whi']} г","межі сегмента\nза зрізом 29.09.2026.",col,2.32,1.18,19)
    txt(sl,0.62,3.04,5.6,0.14,"ХТО ТРИМАЄ СЕГМЕНТ · SKU НА БРЕНД",8,True,NAVY)
    mx=max(n for _,n in v['blist'])
    minibar(sl,0.62,3.28,5.60,[(b,n,f"{n}") for b,n in v['blist'][:6]],mx,col,lab_w=1.30,val_w=0.50,rowh=0.195)
    txt(sl,6.95,3.04,5.8,0.14,"ПОЗИЦІЙ ЗА ЦІНОВОЮ СМУГОЮ · ₴ ЗА ПАЧКУ",8,True,NAVY)
    cnt=[sum(1 for p in _pp if lo<=p<hi) for lo,hi,_ in PB]; cm=max(cnt)
    for i,(c_,(lo,hi,lab)) in enumerate(zip(cnt,PB)):
        x=6.95+i*0.83; hh=0.78*c_/cm if c_ else 0.02
        _rect(sl,x,4.38-hh,0.62,hh,col if c_ else HAIR)
        txt(sl,x-0.1,4.38-hh-0.2,0.82,0.16,str(c_) if c_ else "—",9.5,True,NAVY,C)
        txt(sl,x-0.1,4.42,0.82,0.16,lab,7.6,False,GREY,C)
    y=section_head(sl,4.62,"ПРЕДСТАВНИКИ СЕГМЕНТА · ВІД НАЙДЕШЕВШОЇ ДО НАЙДОРОЖЧОЇ ПОЗИЦІЇ","")
    for i,c in enumerate(sample_k(k,6)): bigcard(sl,CX[i],4.92,1.92,2.46,c,col,sq=1.30)
    NEWSEG.append(sl); return sl
def catalog2(k,items,i,m):
    v=SS[k]; col=SCOL[k]; pm=[c['pmed'] for c in items]; lo,hi=min(pm),max(pm)
    sl=slide(prs,f"СЕГМЕНТ {k} · {SNAME[k]}   ·   {v['n']} ПОЗИЦІЙ",f"{STITLE[k]}: каталог {i} з {m} — {lo:.0f}–{hi:.0f} грн","00")
    band(sl,1.50,band_text(k),col)
    for j,c in enumerate(items[:12]):
        bigcard(sl,CX[j%6],1.82+(j//6)*2.82,1.92,2.70,c,col,sq=1.55)
    NEWSEG.append(sl); return sl
def catalog_series2(k):
    items=SEG[k]; m=math.ceil(len(items)/12); size=math.ceil(len(items)/m)
    for i in range(m): catalog2(k,items[i*size:(i+1)*size],i+1,m)
overview("1","Пакет з бульйоном: найбільший сегмент","Основна частина корейської полиці: локшина й порошковий бульйон, вариться в каструлі.")
catalog_series2("1")
SEG1_END=len(NEWSEG)
overview("2","Пакет із соусом: 30 позицій, медіана 106 грн")
catalog_series2("2")
SEG2_END=len(NEWSEG)
overview("3","Стакан: заливається окропом, без плити")
catalog_series2("3")
NEW['seg_slides']=list(NEWSEG)
NEW['seg_split']=(SEG1_END,SEG2_END)

# ============ ЗВЕДЕННЯ БРЕНДИ × МЕРЕЖІ ============
SV_CH=[("auchan","Ашан"),("grono","Grono"),("onde","Onde"),("metro","METRO"),("zaraz","«Зараз»"),("kharkiv","«Клас»"),
       ("vostorg","«Восторг»"),("cosmos","«Космос»"),("tavriav","«Таврія»"),("novus","NOVUS"),("chudomarket","«ЧудоМ.»"),
       ("torba","«Торба»"),("silpo","«Сільпо»")]
SV_BR=["Nongshim","Samyang","Paldo","Ottogi","O'Food","Bibigo","Otoki","fire bull","Choi's"]
SILPO_P={"Samyang":[89.99,89.99,199.0],"Nongshim":[139.0]}
SILPO_BR={b:(len(v),st.median(v)) for b,v in SILPO_P.items()}
_EAN={r['ean']:r for r in SD_ROWS}
def _cell(b,ch):
    if ch=="silpo": return SILPO_BR.get(b,(0,None))
    v=[z for z in ZREC if z['chain']==ch and _EAN[z['ean']]['tm']==b]
    return (len({z['ean'] for z in v}),st.median([z['price'] for z in v]) if v else None)
SV=[[_cell(b,ch) for ch,_ in SV_CH] for b in SV_BR]
def _tot(b):
    rb=[r for r in SD_ROWS if r['tm']==b]
    chs={c for r in rb for c in r['chains']}|({"silpo"} if b in SILPO_BR else set())
    n=len(rb)+(SILPO_BR[b][0] if b in SILPO_BR else 0)
    ws=sorted(r['w'] for r in rb); pm=[r['pmed'] for r in rb if r.get('pmed')]
    wt=(f"{ws[0]:g} г" if ws[0]==ws[-1] else f"{ws[0]:g}–{ws[-1]:g} г").replace('.',',')
    return n,len(chs),(st.median(pm) if pm else None),wt
s=newslide("ЗВЕДЕННЯ · БРЕНДИ × МЕРЕЖІ","Усі бренди в усіх мережах")
# легенда: що означає клітинка
_rect(s,0.62,1.38,0.62,0.44,SEQT[4]); txt(s,0.62,1.42,0.62,0.2,"17",11,True,WHITE,C); txt(s,0.62,1.62,0.62,0.16,"89 ₴",8,False,WHITE,C)
txt(s,1.34,1.38,5.6,0.2,"верхнє число — кількість позицій (SKU) бренду в мережі",9,True,NAVY)
txt(s,1.34,1.60,5.6,0.2,"нижнє число — медіанна ціна пачки в ₴  ·  чим темніше, тим більше позицій",9,False,GREY)
L0=0.62; RL=1.30; CW=0.60; T0=1.98; HH=0.42; RH=0.405
x_tail=L0+RL+CW*13+0.10; TWS=[0.62,0.70,0.82,0.98]
for j,(ch,nm) in enumerate(SV_CH):
    _rect(s,L0+RL+j*CW+0.02,T0,CW-0.04,HH-0.04,NAVY); txt(s,L0+RL+j*CW+0.02,T0+0.12,CW-0.04,0.18,nm,6.8,True,WHITE,C)
xt=x_tail
for hd,tw in zip(("SKU","МЕРЕЖ","МЕДІАНА, ₴","ВАГА ПАЧКИ"),TWS):
    _rect(s,xt,T0,tw-0.03,HH-0.04,DARKAMBER); txt(s,xt,T0+0.12,tw-0.03,0.18,hd,6.6,True,WHITE,C); xt+=tw
mx=max(n for row in SV for n,_ in row); y=T0+HH
for i,b in enumerate(SV_BR):
    ours=(b=="Choi's")
    txt(s,L0,y+RH/2-0.09,RL-0.10,0.18,("Choi’s (SIAS)" if ours else b),8.8,True,(DARKAMBER if ours else NAVY),R)
    for j,(n,p) in enumerate(SV[i]):
        xx=L0+RL+j*CW; step=0 if n==0 else 1+int(round((len(SEQT)-2)*(n/mx)))
        _rect(s,xx+0.02,y+0.02,CW-0.04,RH-0.04,SEQT[step])
        if n:
            col=WHITE if step>=4 else NAVY
            txt(s,xx+0.02,y+0.035,CW-0.04,0.18,str(n),10.5,True,col,C); txt(s,xx+0.02,y+0.23,CW-0.04,0.14,f"{p:.0f} ₴",7.6,False,col,C)
        else: txt(s,xx+0.02,y+0.12,CW-0.04,0.18,"—",9,False,PALE,C)
    n,nc,pm,wt=_tot(b); xt=x_tail
    for val,tw in zip((str(n),str(nc),f"{pm:.0f} ₴" if pm else "—",wt),TWS):
        _rect(s,xt,y+0.02,tw-0.03,RH-0.04,LIGHT); txt(s,xt,y+0.12,tw-0.03,0.18,val,8.8,True,NAVY,C); xt+=tw
    y+=RH
_rect(s,L0+RL,y+0.04,CW*13+sum(TWS)+0.10,0.012,NAVY)
txt(s,L0,y+0.14,RL-0.10,0.18,"УСЬОГО",8.8,True,DARKAMBER,R)
for j,(ch,_) in enumerate(SV_CH):
    n=sum(SV[i][j][0] for i in range(len(SV_BR)))
    pr=[z['price'] for z in ZREC if z['chain']==ch] if ch!="silpo" else SILPO_P["Samyang"]+SILPO_P["Nongshim"]
    txt(s,L0+RL+j*CW,y+0.10,CW,0.18,str(n),10.5,True,DARKAMBER,C); txt(s,L0+RL+j*CW,y+0.30,CW,0.14,f"{st.median(pr):.0f} ₴",7.6,False,GREY,C)
tn=len(SD_ROWS)+4; xt=x_tail
for val,tw in zip((str(tn),"13",f"{st.median([r['pmed'] for r in SD_ROWS]):.0f} ₴",f"{min(r['w'] for r in SD_ROWS):g}–{max(r['w'] for r in SD_ROWS):g} г".replace('.',',')),TWS):
    txt(s,xt,y+0.20,tw-0.03,0.18,val,8.8,True,DARKAMBER,C); xt+=tw
slide_tail_sources(s,"Каталоги мереж (zakaz.ua) і «Сільпо», 29.09.2026. Клітинка = унікальні EAN × медіана цін по магазинах мережі. «Сільпо» — 4 позиції без EAN (3 Samyang, 1 Nongshim); 93 = 89 позицій з EAN + 4 «Сільпо». Вага — мін.–макс. пачки бренду.",7.04)
NEW['svod']=s

# ============ ФОРМАТ × ВАГА: ГРАФІК ============
B=F['bands']; MX=F['matrix']
s=newslide("ФОРМАТ × ВАГА","Де насправді щільна полиця","Кількість позицій і медіанна ціна пачки за ваговими смугами")
legend(s,0.62,1.74,[("Пакет",NAVY),("Стакан",PURPLE)])
txt(s,0.62,2.05,6.8,0.16,"ПОЗИЦІЙ У ВАГОВІЙ СМУЗІ",8.4,True,NAVY)
txt(s,8.35,2.05,3.4,0.16,"МЕДІАННА ЦІНА ПАЧКИ, ₴",8.4,True,NAVY)
T=2.40; RH=0.62
_rect(s,0.62,T+RH*2-0.02,12.10,RH*2,AMBLIGHT)
for i,b in enumerate(B):
    y=T+i*RH; pk,cp=MX[0][i],MX[1][i]; unit=4.3/38
    txt(s,0.72,y+0.19,1.3,0.24,b['lab'],11,True,NAVY)
    x=2.05
    for v,col in ((pk,NAVY),(cp,PURPLE)):
        if v:
            w=unit*v; _rect(s,x,y+0.12,w,0.38,col)
            if w>0.28: txt(s,x,y+0.20,w,0.2,str(v),10,True,WHITE,C)
            x+=w
    txt(s,x+0.10,y+0.20,1.4,0.22,f"{pk+cp}  ·  {b['share']} %",10,True,NAVY)
    pw=3.05*b['pmed']/150; _rect(s,8.35,y+0.12,pw,0.38,DARKAMBER)
    txt(s,8.35,y+0.20,pw-0.12,0.22,f"{b['pmed']} ₴",10.5,True,WHITE,R)
xl=8.35+3.05*120/150
_rect(s,xl-0.006,T-0.02,0.014,RH*5+0.04,RED); txt(s,xl-0.45,T+RH*5+0.04,0.9,0.16,"120 грн",8,True,RED,C)
txt(s,11.55,T+RH*2+0.45,1.15,0.4,"наші SKU:\n112,5–131 г",8.5,True,DARKAMBER,C,lh=1.1)
ecall(s,0.62,6.14,5.92,0.86,"ВАГА НАШИХ SKU — РИНКОВА",
 f"112,5–131 г потрапляє у дві найщільніші смуги: {B[2]['share']+B[3]['share']} % усіх позицій ринку.",TEAL,9.6,8.6)
ecall(s,6.80,6.14,5.92,0.86,"ЦІНА НАШИХ SKU ПРОТИ ЦИХ СМУГ",
 f"При полиці 120 грн наші SKU на {(120/B[2]['pmed']-1)*100:.0f} % дорожчі за медіану смуги 110–125 г ({B[2]['pmed']} грн), при поточних 140,55 грн — на {(140.55/B[2]['pmed']-1)*100:.0f} %.",RED,9.6,8.6)
NEW['fmt']=s

# ============ ДАНІ КАНАЛІВ ============
CHP=co.defaultdict(dict)
_tmp=co.defaultdict(list)
for z in ZREC: _tmp[(z['chain'],z['ean'])].append(z['price'])
for (c,e),v in _tmp.items(): CHP[c][e]=st.median(v)
CHP['silpo']={f"sp{i}":p for i,p in enumerate([89.99,89.99,199.0,139.0])}
BRX=[("Samyang",r"samyang|бульдак|buldak"),("Nongshim",r"nong\s?shim|шин рамьон|neoguri|chapagetti|toomba"),("Paldo",r"paldo"),("Ottogi",r"ottogi"),
     ("Bibigo",r"bibigo"),("Otoki",r"otoki"),("O'Food",r"o[`'’]?food")]
def brand_of(t):
    for b,rx in BRX:
        if re.search(rx,t,re.I): return b
    return "інші"
_sh=json.load(open(f"{SC}/shops_raw.json",encoding='utf-8'))
_sp=json.load(open(f"{SC}/sp_named.json",encoding='utf-8'))
SHOPN={"Суші Повар":[(r['title'],r['price']) for r in _sp if r['price']<=250],
       "Asia Foods":[(r['title'],r['price']) for r in _sh if r['shop']=="AsiaFoods" and r['price']<=250],
       "Тайякі Март":[(r['title'],r['price']) for r in _sh if r['shop']=="TaiyakiMart" and r['price']<=250]}
PROMN=[(t,p) for t,p in json.load(open(f"{SC}/prom_named.json",encoding='utf-8')) if p<=250]
SHOPN["Prom.ua"]=PROMN
def brandmix(items,n=4):
    c=co.Counter(brand_of(t) for t,_ in items); return " · ".join(f"{b} {k}" for b,k in c.most_common(n))
CHAIN_BR={}
for c in list(CHP):
    d=co.Counter(_EAN[e]['tm'] for e in CHP[c] if e in _EAN)
    if c=='silpo': d=co.Counter({"Samyang":3,"Nongshim":1})
    CHAIN_BR[c]=d
NAT_={'auchan','novus','metro','zaraz','silpo'}
def q4(a):
    q=st.quantiles(sorted(a),n=4); return q[0],q[1],q[2]
NOW33=(M.forward_current("33")[2]['shelf'],M.forward_current("33")[0]['shelf'])
NOW6=(M.forward_current("6")[2]['shelf'],M.forward_current("6")[0]['shelf'])
BAND33=RGBColor(0xFD,0xF0,0xD2); BAND6=RGBColor(0xF6,0xDD,0xDD)

# ============ СЛАЙД: МЕРЕЖІ (жива візуалізація) ============
s=newslide("РИТЕЙЛ-АУДИТ · МЕРЕЖІ","Що з корейського рамену стоїть у мережах","Суцільна перевірка онлайн-каталогів 18 мереж, зріз 29.09.2026. SKU — унікальні позиції корейського рамену в каталозі мережі.")
x=0.62
for b in ("Samyang","Nongshim","Paldo","Ottogi","O'Food","Bibigo","Otoki","Choi's"):
    _rect(s,x,1.66,0.16,0.14,BC[b]); txt(s,x+0.22,1.64,1.0,0.18,b,8.4,False,NAVY); x+=0.32+0.082*len(b)+0.22
txt(s,8.75,1.64,4.0,0.18,"ЦІНА ПАЧКИ, ₴: мін.–макс. (лінія) · медіана (мітка)",8,True,GREY,R)
rowsC=[]
for r in F['chain_rows']: rowsC.append(dict(ch=r['chain'],name=NM.get(r['chain'],r['chain']),n=r['n'],med=r['med'],lo=r['lo'],hi=r['hi']))
rowsC.append(dict(ch='silpo',name="«Сільпо»",n=4,med=114,lo=90,hi=199))
rowsC.sort(key=lambda r:-r['n'])
T=2.02; RH=0.272; BX=3.05; UNIT=3.3/60
RX=8.45; RW=3.3; X0,X1=40,220
RXf=lambda v: RX+RW*(min(max(v,X0),X1)-X0)/(X1-X0)
for i,r in enumerate(rowsC):
    y=T+i*RH
    if i%2==0: _rect(s,0.62,y,12.10,RH-0.02,LIGHT)
    nat=r['ch'] in NAT_
    txt(s,0.70,y+0.045,1.55,0.2,r['name'],9.2,True,NAVY)
    rr(s,2.15,y+0.05,0.78,0.17,NAVY if nat else TEAL,0.3); txt(s,2.15,y+0.062,0.78,0.14,"НАЦ." if nat else "РЕГІОН.",6.6,True,WHITE,C)
    xx=BX
    for b,k in sorted(CHAIN_BR[r['ch']].items(),key=lambda kv:-kv[1]):
        w=UNIT*k; _rect(s,xx,y+0.04,w-0.01,RH-0.09,BC.get(b,GREYB)); 
        if w>0.2: txt(s,xx,y+0.052,w-0.01,0.16,str(k),7.6,True,WHITE,C)
        xx+=w
    txt(s,xx+0.05,y+0.045,0.5,0.2,str(r['n']),9,True,NAVY)
    _rect(s,RXf(r['lo']),y+RH/2-0.02,max(RXf(r['hi'])-RXf(r['lo']),0.03),0.04,GREYB)
    _rect(s,RXf(r['med'])-0.03,y+0.045,0.06,RH-0.1,DARKAMBER)
    txt(s,RXf(r['hi'])+0.06,y+0.045,0.9,0.2,f"{r['lo']}–{r['hi']} ₴",7.6,False,GREY)
    txt(s,11.95,y+0.045,0.77,0.2,f"{r['med']} ₴",9,True,NAVY,R)
y=T+len(rowsC)*RH
_rect(s,0.62,y,12.10,RH-0.02,LIGHT)
txt(s,0.70,y+0.045,8.5,0.2,"МегаМаркет · Ultramarket · ЕКО · Епіцентр · «Ідеал» — 0 позицій у зрізі",8.8,True,GREY)
_rect(s,RXf(120)-0.008,T-0.03,0.016,RH*len(rowsC)+0.05,RED); txt(s,RXf(120)-0.5,1.88,1.0,0.14,"120 ₴",7.6,True,RED,C)
txt(s,11.55,1.88,1.17,0.14,"МЕДІАНА",7.4,True,GREY,R)
ecall(s,0.62,6.52,3.88,0.88,"ДІРКА №1 · «СІЛЬПО»","Найбільша мережа країни тримає 4 позиції. Полиця відкрита, але вхід там 90–199 ₴.",AMBER,9.2,8.4)
ecall(s,4.72,6.52,3.88,0.88,"ДІРКА №2 · П’ЯТЬ МЕРЕЖ З НУЛЕМ","МегаМаркет, Ultramarket, ЕКО, Епіцентр, «Ідеал» — жодної позиції у зрізі.",AMBER,9.2,8.4)
_cs={z['ean'] for z in ZREC if z['chain']=='cosmos'}; _ts={z['ean'] for z in ZREC if z['chain']=='tavriav'}
ecall(s,8.82,6.52,3.90,0.88,"ОДИН ПОСТАЧАЛЬНИК НА ДВІ МЕРЕЖІ",f"«Космос» і «Таврія В» — по {len(_cs)} позицій; {len(_cs&_ts)} із них збігаються за EAN.",NAVY,9.2,8.4)
NEW['chains']=s

# ============ СЛАЙД: КАНАЛИ (розширений) ============
s=newslide("КАНАЛИ ПРОДАЖУ","Де продається категорія і кому","Ліворуч — усі перевірені точки: де категорія є і скільки там позицій. Праворуч — хто туди приходить і чи це наш покупець.")
txt(s,0.62,1.62,7,0.16,"ПЕРЕВІРЕНІ ТОЧКИ ПРОДАЖУ · ПОЗИЦІЙ У КОЖНІЙ",8.4,True,NAVY)
none_ch=["МегаМаркет","Ultramarket","ЕКО","Епіцентр","«Ідеал»"]
groups=[("ПРОДУКТОВІ МЕРЕЖІ · 18 ПЕРЕВІРЕНО · 13 МАЮТЬ РАМЕН",NAVY,[(r['name'],r['n'],(NAVY if r['ch'] in NAT_ else BLUE2)) for r in rowsC]+[(n,0,GREYB) for n in none_ch]),
        ("АЗІЙСЬКІ ФУДШОПИ · 5 ПЕРЕВІРЕНО · 3 З ЦІНАМИ",TEAL,[(k,len(SHOPN[k]),TEAL) for k in ("Суші Повар","Тайякі Март","Asia Foods")]+[("OMG! Asia",None,GREYB),("Panda Asian",None,GREYB)]),
        ("МАРКЕТПЛЕЙСИ · 4 ПЕРЕВІРЕНО · 1 З ЦІНАМИ",DARKAMBER,[("Prom.ua",len(PROMN),DARKAMBER),("Rozetka",None,GREYB),("MAUDAU",None,GREYB),("Bigl",None,GREYB)])]
y=1.84; RH2=0.154; LX0=0.62; BXc=2.35; UN=3.2/62
for head,hc,rows_ in groups:
    _rect(s,LX0,y,6.95,0.2,hc); txt(s,LX0+0.08,y+0.03,6.8,0.14,head,7.4,True,WHITE); y+=0.23
    for nm,n,col in rows_:
        txt(s,LX0+0.05,y+0.01,1.65,0.15,nm,7.8,False,NAVY,R if False else L)
        if n is None: txt(s,BXc,y+0.01,3.0,0.15,"цінових даних немає",7.4,False,PALE)
        elif n==0: txt(s,BXc,y+0.01,3.0,0.15,"категорії немає",7.4,False,PALE)
        else:
            _rect(s,BXc,y+0.025,UN*min(n,77),RH2-0.05,col); txt(s,BXc+UN*min(n,77)+0.06,y+0.0,0.8,0.15,str(n),7.8,True,NAVY)
        y+=RH2
    y+=0.05
cards_=[("Продуктовий роздріб","ЦІЛЬОВИЙ КАНАЛ",RED,"МЕРЕЖІ НАЦІОНАЛЬНОГО ПОКРИТТЯ · 18 ПЕРЕВІРЕНО","13 мереж мають корейський рамен, 5 — жодної позиції. «Сільпо» — 4 позиції на всю мережу.","Масовий покупець · щоденне харчування"),
        ("Онлайн-маркетплейси","ЦІЛЬОВИЙ КАНАЛ",DARKAMBER,"УНІВЕРСАЛЬНІ МАЙДАНЧИКИ · 4 МАЙДАНЧИКИ",f"Prom · Rozetka · MAUDAU · Bigl. Ціни є лише з Prom: {len(PROMN)} пропозицій, медіана {st.median([p for _,p in PROMN]):.0f} ₴, мінімум {min(p for _,p in PROMN):.0f} ₴.","Масовий покупець · планована закупівля"),
        ("Азійські фудшопи","ВХІД",TEAL,"СПЕЦІАЛІЗОВАНИЙ ФУДРІТЕЙЛ · 5 МАГАЗИНІВ","OMG! Asia · Тайякі Март · Суші Повар · Asia Foods · Panda Asian. Ціни є з трьох: "+f"{min(p for k in ('Суші Повар','Тайякі Март','Asia Foods') for _,p in SHOPN[k]):.0f}–{max(p for k in ('Суші Повар','Тайякі Март','Asia Foods') for _,p in SHOPN[k]):.0f} ₴.","Покупець азійської кухні · знає категорію"),
        ("Outdoor- і мілітарі-рітейл","НЕ НАШ КАНАЛ",PALE,"ТУРИСТИЧНІ ТА ВІЙСЬКОВІ МАГАЗИНИ · 0 ПОЗИЦІЙ","Рамену там немає: він потребує окропу й каструлі.","Турист і військовий · автономне харчування")]
for k,(head,tag,tc,sub_,body,who) in enumerate(cards_):
    t2=1.62+k*1.32
    _rect(s,7.85,t2,4.87,1.24,LIGHT); _rect(s,7.85,t2,0.045,1.24,tc)
    txt(s,8.05,t2+0.09,3.0,0.2,head,10.5,True,NAVY)
    rr(s,10.95,t2+0.09,1.65,0.21,tc,0.3); txt(s,10.95,t2+0.12,1.65,0.16,tag,6.4,True,WHITE,C)
    txt(s,8.05,t2+0.36,4.55,0.13,sub_,6.6,True,tc)
    txt(s,8.05,t2+0.55,4.55,0.5,body,8,False,NAVY,lh=1.25)
    txt(s,8.05,t2+1.02,4.55,0.14,"ПОКУПЕЦЬ:  "+who,6.8,True,GREY)
slide_tail_sources(s,"Мережі — каталоги zakaz.ua і «Сільпо» 29.09.2026. Фудшопи — сторінки категорій (Суші Повар, Asia Foods, Тайякі Март), Prom.ua — одна сторінка видачі, жовтень 2026; без наборів дорожче 250 ₴. Rozetka, MAUDAU, Bigl, OMG! Asia, Panda Asian цінових даних не дали.",7.06)
NEW['channels']=s

# ============ КАНАЛИ: ДІАПАЗОН ЦІН (2 варіанти «зараз») ============
CHANNELS=[("Національні мережі","5 мереж · Ашан, METRO, NOVUS…",[p for c in NAT_ for p in CHP[c].values()],NAVY),
          ("Регіональні мережі","8 мереж · Grono, Onde, «Космос»…",[p for c in CHP if c not in NAT_ for p in CHP[c].values()],BLUE2),
          ("Суші Повар","фудшоп",[p for _,p in SHOPN["Суші Повар"]],TEAL),
          ("Asia Foods","фудшоп",[p for _,p in SHOPN["Asia Foods"]],TEAL),
          ("Тайякі Март","фудшоп",[p for _,p in SHOPN["Тайякі Март"]],TEAL),
          ("Prom.ua","маркетплейс",[p for _,p in PROMN],DARKAMBER)]
s=newslide("КАНАЛИ · ЦІНИ","Ціна пачки по каналах і наша поточна полиця","Мін.–макс. (лінія), 25–75 % пропозицій (блок), медіана (риска) · ціна пачки, ₴ · смуги — наша поточна полиця у двох варіантах закупівлі")
X0,X1=40,220; PX=3.35; PW=6.7
X=lambda v: PX+PW*(min(max(v,X0),X1)-X0)/(X1-X0)
T=2.50; RH=0.62; H=RH*len(CHANNELS)
for i in range(0,len(CHANNELS),2): _rect(s,0.62,T+i*RH,12.10,RH-0.02,LIGHT)
_rect(s,X(NOW33[0]),T-0.05,X(NOW33[1])-X(NOW33[0]),H+0.05,BAND33)
_rect(s,X(NOW6[0]),T-0.05,X(NOW6[1])-X(NOW6[0]),H+0.05,BAND6)
for tv in range(40,221,20):
    _rect(s,X(tv),T,0.006,H,HAIR); txt(s,X(tv)-0.3,T+H+0.05,0.6,0.14,str(tv),7.5,False,GREY,C)
txt(s,10.55,2.0,1.0,0.3,"ПРОПОЗИЦІЙ",7.4,True,GREY,R); txt(s,11.65,2.0,1.07,0.3,"ДО 120 ₴",7.4,True,GREY,R)
for i,(nm,sub,pr,col) in enumerate(CHANNELS):
    y=T+i*RH; lo,hi=min(pr),max(pr); a,m,b_=q4(pr)
    txt(s,0.74,y+0.09,2.6,0.2,nm,10,True,NAVY); txt(s,0.74,y+0.31,2.6,0.2,sub,7.4,False,GREY)
    _rect(s,X(lo),y+RH/2-0.025,max(X(hi)-X(lo),0.03),0.05,tint(col,0.4))
    _rect(s,X(a),y+RH/2-0.12,max(X(b_)-X(a),0.03),0.24,tint(col,0.15))
    _rect(s,X(m)-0.02,y+RH/2-0.17,0.04,0.34,col)
    txt(s,X(m)-0.4,y+0.01,0.8,0.16,f"{m:.0f} ₴",8.6,True,NAVY,C)
    txt(s,X(hi)+0.08,y+RH/2-0.09,0.6,0.18,f"{hi:.0f}",7.6,False,GREY)
    txt(s,10.55,y+0.20,1.0,0.22,str(len(pr)),9.6,False,GREY,R)
    txt(s,11.65,y+0.20,1.07,0.22,f"{sum(1 for p in pr if p<=120)/len(pr)*100:.0f} %",10.5,True,NAVY,R)
_rect(s,X(120)-0.008,T-0.05,0.018,H+0.08,RED); txt(s,X(120)-1.0,T-0.27,0.95,0.16,"ціль 120 ₴",7.8,True,RED,R)
txt(s,X(NOW33[0])+0.02,T-0.27,2.2,0.16,"зараз · 33 палети",7.8,True,DARKAMBER,L); txt(s,X(NOW33[0])+0.02,T-0.12,2.2,0.14,f"{NOW33[0]:.0f}–{NOW33[1]:.0f} ₴",7.6,False,DARKAMBER,L)
txt(s,X(NOW6[0])+0.02,T-0.27,2.2,0.16,"зараз · 6 палет",7.8,True,RED,L); txt(s,X(NOW6[0])+0.02,T-0.12,2.2,0.14,f"{NOW6[0]:.0f}–{NOW6[1]:.0f} ₴",7.6,False,RED,L)
slide_tail_sources(s,"Мережі: каталоги 29.09.2026, медіана ціни позиції в мережі. Фудшопи й Prom.ua: сторінки категорій, жовтень 2026, без наборів дорожче 250 ₴. Смуги: Vegetable/Spicy/Beef/Chicken (нижня межа) і Gouda/Carbonara (верхня межа) за фінмоделлю, маржа 30 %.",7.04)
NEW['chA']=s

# ============ КАНАЛИ: ЦІНОВІ РІВНІ + що за позиції ============
ROWS=[]
for c,d in CHP.items(): ROWS.append((NM.get(c,c),list(d.values())," · ".join(f"{b} {k}" for b,k in CHAIN_BR[c].most_common(4))))
for k in ("Суші Повар","Asia Foods","Тайякі Март","Prom.ua"): ROWS.append((k,[p for _,p in SHOPN[k]],brandmix(SHOPN[k],4)))
TIERS=[("до 100 ₴",lambda p:p<=100,TEAL),("100–120",lambda p:100<p<=120,tint(TEAL,0.55)),("120–140",lambda p:120<p<=140,tint(AMBER,0.45)),("понад 140 ₴",lambda p:p>140,RED)]
def shares(pr): return [sum(1 for p in pr if f(p))/len(pr) for _,f,_ in TIERS]
ROWS=[(n,p,shares(p),bm) for n,p,bm in ROWS if p]
ROWS.sort(key=lambda r:-(r[2][0]+r[2][1]))
s=newslide("КАНАЛИ · ЦІНОВІ РІВНІ","Частка пропозицій за ціною пачки в кожному каналі","Мережі — 89 SKU корейського рамену (пакети й стакани); фудшопи й Prom.ua — рамен зі сторінок категорій, без наборів. Числа в смугах — % пропозицій")
legend(s,0.62,1.64,[(l,c) for l,_,c in TIERS],gap=0.35,fs=9)
txt(s,8.35,1.66,0.8,0.16,"ДО 120 ₴",7.6,True,GREY,R); txt(s,9.2,1.66,0.5,0.16,"ПОЗ.",7.6,True,GREY,R); txt(s,9.9,1.66,2.8,0.16,"ЯКІ БРЕНДИ (ПОЗИЦІЙ)",7.6,True,GREY)
T=1.98; RH=0.283
for i,(nm,pr,sh,bm) in enumerate(ROWS):
    y=T+i*RH
    txt(s,0.62,y+0.04,1.9,0.2,nm,9.2,True,NAVY,R)
    x=2.65
    for v,(lab,_,col) in zip(sh,TIERS):
        w=5.0*v
        if w>0.005:
            _rect(s,x,y+0.03,w,RH-0.06,col)
            if w>0.36: txt(s,x,y+0.055,w,0.18,f"{v*100:.0f}",8.2,True,(WHITE if col in (TEAL,RED) else NAVY),C)
        x+=w
    txt(s,7.65,y+0.04,1.5,0.2,f"{(sh[0]+sh[1])*100:.0f} %",10,True,NAVY,R)
    txt(s,9.2,y+0.05,0.5,0.2,str(len(pr)),8.4,False,GREY,R)
    txt(s,9.9,y+0.05,2.85,0.2,bm,7.6,False,NAVY)
slide_tail_sources(s,"Мережі — унікальні позиції (EAN) за медіанною ціною в мережі, 29.09.2026; «Сільпо» — 4 позиції. Фудшопи й Prom.ua — сторінки категорій, жовтень 2026 (бренд визначено за назвою; «інші» — не корейські або без бренду).",7.04)
NEW['chB']=s

# ============ ЦІНА НА ПОЛИЦІ ЗАРАЗ · ПО ВСІХ ПОЗИЦІЯХ ============
F33=M.forward_current("33"); F6=M.forward_current("6")
C_CC=RGBColor(0x46,0x53,0x82); C_BON=DARKAMBER; C_MAR=AMBER; C_MK=RGBColor(0xA9,0xB0,0xC8)
BONUS_,MARGIN_,MARKUP_=M.BONUS,M.MARGIN,M.MARKUP
ORDER=[2,3,4,5,0,1]            # Vegetable, Spicy, Beef, Chicken, Gouda, Carbonara
def shelf_slide(sc,fc,head):
    s=newslide(f"ФІНМОДЕЛЬ · ПОТОЧНА ПОЛИЦЯ · {head}",f"Ціна на полиці зараз · {head.lower()}")
    legend(s,0.62,1.42,[("Собівартість (СС)",C_CC),("Бонус мережі 25 %",C_BON),("Наша маржа 30 %",C_MAR),("Націнка магазину 40 %",C_MK)],gap=0.4,fs=9)
    BX=3.55; PWd=7.2; XM=230.0
    bx=lambda v: BX+PWd*v/XM
    T=1.92; RH=0.70
    for k,idx in enumerate(ORDER):
        nm,wt,_,cur=M.SKUS[idx]; f=fc[idx]; p=f['partner']; y=T+k*RH
        if k%2==0: _rect(s,0.62,y-0.03,12.10,RH-0.02,LIGHT)
        txt(s,0.72,y+0.06,2.8,0.2,nm,10,True,NAVY); txt(s,0.72,y+0.29,2.8,0.2,f"{wt} · закупівля €{cur:.2f} · СС {dig(f['cc'],2)} ₴".replace('.',','),8,False,GREY)
        x=BX
        for v,c_,tc in ((f['cc'],C_CC,WHITE),(p*BONUS_,C_BON,WHITE),(p*MARGIN_,C_MAR,NAVY),(p*(MARKUP_-1),C_MK,NAVY)):
            w=PWd*v/XM; _rect(s,x,y+0.06,w-0.015,0.46,c_); txt(s,x,y+0.18,w-0.015,0.22,f"{v:.0f}",10.5,True,tc,C); x+=w
        txt(s,bx(f['shelf'])+0.12,y+0.04,1.7,0.28,dig(f['shelf'],2)+" ₴",16,True,RED)
        txt(s,bx(f['shelf'])+0.12,y+0.34,2.0,0.18,("+"+dig(f['shelf']-120,2)).replace('+','+')+f" ₴ · +{(f['shelf']/120-1)*100:.0f} % до цілі",8.4,False,GREY)
    H=RH*6
    _rect(s,bx(120)-0.008,T-0.08,0.018,H+0.12,RED); txt(s,bx(120)-0.7,T+H+0.06,1.4,0.16,"ціль 120 ₴",8.4,True,RED,C)
    for tv in (0,50,100,150,200): txt(s,bx(tv)-0.3,T+H+0.06,0.6,0.14,str(tv),7.5,False,GREY,C) if abs(bx(tv)-bx(120))>0.5 else None
    _rect(s,0.62,6.52,12.10,0.50,LIGHT); _rect(s,0.62,6.52,0.045,0.50,NAVY)
    txt(s,0.9,6.58,11.6,0.2,"Ціна партнеру = СС ÷ 0,45  (1 − бонус 25 % − маржа 30 %)   ·   полиця = ціна партнеру × 1,40   ·   числа в смугах — ₴ за пачку",9,True,NAVY)
    txt(s,0.9,6.80,11.6,0.18,"СС — із Self-Cost (товар, логістика, брокер, ПДВ, фінансування); курс 51,20 ₴/€. Показано маржу 30 %.",8,False,GREY)
    slide_tail_sources(s,"Фінмодель SIAS: вкладки «Финмодель 6 паллет» / «33 паллет», колонки маржі 30 %; СС — Self-Cost_Sias.xlsx.",7.1)
    return s
NEW['shelf33']=shelf_slide("33",F33,"33 ПАЛЕТИ")
NEW['shelf6']=shelf_slide("6",F6,"6 ПАЛЕТ")

# ============ ПОПУЛЯРНІ ПАКЕТИ: ПОЗИЦІЇ ============
def _clean_name(r):
    t=r['title']; tm=r['tm']
    pat=re.compile(r"nong\s?shim" if tm=="Nongshim" else re.escape(tm).replace("'","['’]?"),re.I)
    m=pat.search(t); rest=t[m.end():] if m else t
    rest=re.sub(r"\s*\d+[.,]?\d*\s*(г|гр|g)\b\s*$","",rest.strip(),flags=re.I)
    rest=re.sub(r"швидкого приготування|локшина|рамьон|рамен","",rest,flags=re.I)
    rest=re.sub(r"\s+"," ",rest).strip(" ,.")
    return rest or "рамен"
def _short(x,n): return x if len(x)<=n else x[:n-1].rstrip()+"…"
PK=[r for r in SD_ROWS if r['seg'] in (1,2) and r.get('pmed')]
OURS={120:[("SIAS · Vegetable/Spicy/Beef/Chicken 113 г",NOW33[0],NOW6[0]),("SIAS · Gouda Cheese 123 г",NOW33[1],NOW6[1])],
      130:[("SIAS · Carbonara Spicy 131 г",NOW33[1],NOW6[1])],140:[]}
def ladder_panel(l,t,w,wt,rows,ours,xmax=210,rowh=0.185,oh=0.34):
    items=[dict(kind='m',p=r['pmed'],brand=r['tm'],name=_short(_clean_name(r),27)) for r in rows]
    items+=[dict(kind='o',p=120.0,name=n,n33=a,n6=b) for n,a,b in ours]
    items.sort(key=lambda d:(d['p'],d['kind']=='m'))
    pm=[r['pmed'] for r in rows]
    txt(s,l,t,w,0.2,f"ПАКЕТ {wt} Г · {len(rows)} ПОЗИЦІЙ",9.2,True,NAVY)
    txt(s,l,t+0.2,w,0.16,f"{min(pm):.0f}–{max(pm):.0f} ₴ · до 120 ₴ включно: {sum(1 for p in pm if p<=120)}",7.8,False,GREY)
    lab_w=2.8; bl=l+lab_w+0.05; bw=w-lab_w-0.05-0.62
    X=lambda v: bl+bw*v/xmax
    y=t+0.46; H=sum(oh if d['kind']=='o' else rowh for d in items)
    _rect(s,bl,y-0.02,0.006,H+0.04,HAIR); _rect(s,X(120)-0.007,y-0.05,0.015,H+0.08,RED)
    for d in items:
        if d['kind']=='m':
            txt(s,l,y+0.015,lab_w,0.16,f"{d['brand']} · {d['name']}",7.9,False,NAVY)
            _rect(s,bl,y+0.03,X(d['p'])-bl,rowh-0.07,BC.get(d['brand'],GREYB))
            txt(s,X(d['p'])+0.06,y+0.01,0.6,0.16,f"{d['p']:.0f} ₴",8,True,NAVY); y+=rowh
        else:
            _rect(s,l,y-0.005,w,oh,AMBLIGHT)
            txt(s,l+0.04,y+0.01,lab_w,0.16,d['name'],7.9,True,DARKAMBER)
            _rect(s,bl,y+0.03,X(120)-bl,0.12,AMBER); _rect(s,X(120),y+0.03,X(d['n33'])-X(120),0.12,tint(AMBER,0.55))
            _rect(s,X(d['n33']),y+0.03,X(d['n6'])-X(d['n33']),0.12,tint(RED,0.6))
            txt(s,l+0.04,y+0.175,w-0.1,0.14,f"ціль 120 ₴  ·  зараз: 33 палети {dig(d['n33'],0)} ₴  ·  6 палет {dig(d['n6'],0)} ₴",7.2,False,GREY)
            y+=oh
    for tv in (0,60,120,180): txt(s,X(tv)-0.25,y+0.04,0.5,0.14,str(tv),7,False,GREY,C)
    return y+0.2
s=newslide("КОНКУРЕНТИ · ПОПУЛЯРНІ ПАКЕТИ","Популярні пакети: кожна позиція та її ціна")
x=0.62
for b in ("Nongshim","Samyang","Paldo","Ottogi","Bibigo","O'Food","Otoki"):
    _rect(s,x,1.44,0.16,0.14,BC[b]); txt(s,x+0.22,1.42,1.0,0.18,b,8.8,False,NAVY); x+=0.32+0.085*len(b)+0.25
_rect(s,x,1.44,0.16,0.14,AMBER); txt(s,x+0.22,1.42,0.5,0.18,"SIAS:",8.8,True,DARKAMBER)
for lab,col,dx in (("ціль 120",AMBER,0.62),("зараз 33 палети",tint(AMBER,0.55),1.0),("зараз 6 палет",tint(RED,0.6),1.1)):
    pass
txt(s,x+0.7,1.42,3.4,0.18,"смуга: ціль 120 → зараз 33 палети → 6 палет",8.4,False,DARKAMBER)
byw=lambda w: [r for r in PK if round(r['w'])==w]
ladder_panel(0.62,1.80,5.95,120,byw(120),OURS[120],rowh=0.175,oh=0.32)
yy=ladder_panel(6.82,1.80,5.90,130,byw(130),OURS[130])
ladder_panel(6.82,yy+0.02,5.90,140,byw(140),OURS[140])
slide_tail_sources(s,"Каталоги мереж 29.09.2026: медіанна ціна пачки по мережах, де позиція є. Червона лінія — 120 ₴. SIAS: поточна полиця за фінмоделлю, маржа 30 %.",7.12)
NEW['pop']=s

# ============ КОНКУРЕНТИ НИЖЧЕ 120 ============
KEY=[('08801073113428','Samyang Buldak 2× spicy','140 г'),('08801073110502','Samyang Buldak Hot Chicken','140 г'),
     ('08801073113381','Samyang Buldak тушкована курка','145 г'),('08801073116474','Samyang Buldak з сиром','130 г'),
     ('08801043150620','Nongshim Рамен зі спеціями','120 г'),('08801043157742','Nongshim Kimchi Ramyun','120 г'),
     ('08801043157728','Nongshim Chapaghetti','140 г'),('08801043018470','Nongshim Toomba Spicy & Creamy','137 г'),
     ('00648436310685','Paldo Volcano Carbonara','130 г')]
_pk_e={r['ean'] for r in SD_ROWS if r['seg'] in (1,2)}
krow=[]
for e,lab,wt in KEY:
    d={c:CHP[c][e] for c in CHP if e in CHP[c]}
    lo=min(d,key=d.get); hi=max(d,key=d.get)
    krow.append(dict(label=lab,wt=wt,lo=d[lo],lochain=NM.get(lo,lo),hi=d[hi],hichain=NM.get(hi,hi),n=len(d)))
krow.sort(key=lambda r:r['lo'])
pr={}
for z in ZREC:
    if z['ean'] in _pk_e and z.get('old_price') and z['price']<z['old_price']-0.005:
        dsc=1-z['price']/z['old_price']
        if z['ean'] not in pr or dsc>pr[z['ean']][0]: pr[z['ean']]=(dsc,z['price'],z['old_price'],z['chain'])
promo=sorted(pr.items(),key=lambda kv:-kv[1][0])[:4]
promo_rows=[dict(label="Samyang Buldak курка та сир · 140 г",chain="«Сільпо»",new=89.99,old=174.0),
            dict(label="Samyang Buldak карбонара · 130 г",chain="«Сільпо»",new=89.99,old=174.0)]
for e,(dsc,new,old,ch) in promo:
    r=_EAN[e]; promo_rows.append(dict(label=f"{r['tm']} · {_short(_clean_name(r),24)} · {r['w']:g} г".replace('.',','),chain=NM.get(ch,ch),new=new,old=old))
s=newslide("КОНКУРЕНТИ · НИЖЧЕ 120 ₴","Ключові конкуренти продаються нижче 120 ₴")
txt(s,0.62,1.52,7,0.18,"ДІАПАЗОН ЦІНИ ПАЧКИ ПО МЕРЕЖАХ, ₴",8.8,True,NAVY)
legend(s,0.62,1.78,[("дешевше за ціль 120 ₴",TEAL),("дорожче за ціль",GREYB)],gap=0.3,fs=8.5)
LX=3.05; LW=4.6; X0,X1=60,220
X=lambda v: LX+LW*(v-X0)/(X1-X0)
T=2.42; RH=0.45; H=RH*len(krow)
_rect(s,X(NOW33[0]),T-0.05,X(NOW33[1])-X(NOW33[0]),H+0.05,BAND33); _rect(s,X(NOW6[0]),T-0.05,X(NOW6[1])-X(NOW6[0]),H+0.05,BAND6)
for tv in range(60,221,20):
    _rect(s,X(tv),T,0.006,H,HAIR); txt(s,X(tv)-0.3,T+H+0.05,0.6,0.14,str(tv),7.5,False,GREY,C)
for i,r in enumerate(krow):
    y=T+i*RH
    txt(s,0.62,y+0.04,2.4,0.2,r['label'],8.8,True,NAVY); txt(s,0.62,y+0.24,2.4,0.16,f"{r['wt']} · {r['n']} мереж",7.4,False,GREY)
    if r['lo']<120: _rect(s,X(r['lo']),y+0.12,X(min(120,r['hi']))-X(r['lo']),0.2,TEAL)
    if r['hi']>120: _rect(s,X(max(120,r['lo'])),y+0.12,X(r['hi'])-X(max(120,r['lo'])),0.2,GREYB)
    txt(s,X(r['lo'])-1.0,y+0.01,0.95,0.16,f"{r['lo']:.0f} ₴",8.6,True,(TEAL if r['lo']<120 else NAVY),R)
    txt(s,X(r['lo'])-1.3,y+0.2,1.25,0.15,r['lochain'],7,False,GREY,R)
    txt(s,X(r['hi'])+0.06,y+0.01,0.8,0.16,f"{r['hi']:.0f} ₴",8.6,True,NAVY)
    txt(s,X(r['hi'])+0.06,y+0.2,1.1,0.15,r['hichain'],7,False,GREY)
_rect(s,X(120)-0.008,T-0.05,0.018,H+0.08,RED); txt(s,X(120)-1.0,T-0.27,0.95,0.16,"ціль 120 ₴",7.8,True,RED,R)
txt(s,X(NOW33[0])+0.02,T-0.27,1.6,0.16,"зараз · 33 палети",7.6,True,DARKAMBER,L); txt(s,X(NOW6[0])+0.02,T-0.27,1.6,0.16,"зараз · 6 палет",7.6,True,RED,L)
RX=8.75
txt(s,RX,1.52,4.0,0.18,"АКЦІЇ: ЗВИЧАЙНА ЦІНА → АКЦІЙНА, ₴",8.8,True,NAVY)
legend(s,RX,1.78,[("звичайна",GREYB),("акційна",AMBER)],gap=0.3,fs=8.5)
PH_=3.4; PXM=190; y=2.28
for r in promo_rows:
    txt(s,RX,y,4.0,0.18,r['label'],8.2,True,NAVY)
    _rect(s,RX,y+0.22,PH_*r['old']/PXM,0.15,GREYB); txt(s,RX+PH_*r['old']/PXM+0.06,y+0.2,0.6,0.16,f"{r['old']:.0f} ₴",8,False,GREY)
    _rect(s,RX,y+0.40,PH_*r['new']/PXM,0.15,AMBER); txt(s,RX+PH_*r['new']/PXM+0.06,y+0.38,2.0,0.16,f"{r['new']:.0f} ₴  (−{(1-r['new']/r['old'])*100:.0f} %) · {r['chain']}",8,True,DARKAMBER)
    y+=0.70
slide_tail_sources(s,"Каталоги мереж 29.09.2026: ціна позиції в мережі (медіана по магазинах); «стара ціна» віддається не всіма мережами, тому акцій насправді більше. Смуги — наша поточна полиця: Vegetable/Spicy/Beef/Chicken (нижня межа) і Gouda/Carbonara (верхня).",7.04)
NEW['comp']=s

# ============ З ЧОГО СКЛАДАЄТЬСЯ ЦІНА НА ПОЛИЦІ ============
import openpyxl
_sc=openpyxl.load_workbook("/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/87c2996d-Self-Cost_Sias.xlsx",data_only=True).worksheets[0]
def scost(row):
    price=_sc[f"J{row}"].value; vat=_sc[f"U{row}"].value/_sc[f"K{row}"].value; fin=_sc[f"AH{row}"].value; tot=_sc[f"AJ{row}"].value
    return dict(prod=price*M.RATE,vat=vat*M.RATE,fin=fin*M.RATE,logi=(tot-price-vat-fin)*M.RATE,cc=tot*M.RATE)
def parts(d):
    p=d['cc']/(1-M.BONUS-M.MARGIN)
    return [d['prod'],d['cc']-d['prod'],p*M.BONUS,p*M.MARGIN,p*(M.MARKUP-1)],p*M.MARKUP
def parts_t(sc,a):
    b=M.breakdown(sc,a); return [b['prod'],b['cc']-b['prod'],b['bonus'],b['profit'],b['markup']],b['shelf']
A33=M.ask_price("33"); A6=M.ask_price("6")
ST=[("Ціль · 33 палети","закупівля €%s"%f"{A33:.2f}".replace('.',','),)+parts_t("33",A33)+(TEAL,),
    ("Ціль · 6 палет","закупівля €%s"%f"{A6:.2f}".replace('.',','),)+parts_t("6",A6)+(TEAL,)]
for sc,nm_,rows_ in (("33","33 палети",((32,"Vegetable/Spicy/Beef/Chicken · €0,65"),(30,"Gouda/Carbonara · €0,75"))),("6","6 палет",((8,"Vegetable/Spicy/Beef/Chicken · €0,65"),(6,"Gouda/Carbonara · €0,75")))):
    for row,sub_ in rows_:
        ST.append((f"Зараз · {nm_}",sub_)+parts(scost(row))+(RED,))
CS=[TEAL,C_CC,C_BON,C_MAR,C_MK]
s=newslide("ФІНМОДЕЛЬ · СТРУКТУРА ЦІНИ","З чого складається ціна на полиці: куди йде кожна гривня")
legend(s,0.62,1.42,[("Товар SIAS",CS[0]),("Логістика, брокер, ПДВ, фінансування",CS[1]),("Бонус мережі 25 %",CS[2]),("Наша маржа 30 %",CS[3]),("Націнка магазину 40 %",CS[4])],gap=0.3,fs=8.6)
BX=3.55; PWd=7.0; XM=230.0; bx=lambda v: BX+PWd*v/XM
T=1.98; RH=0.64
for k,(h1,h2,vals,shelf,col) in enumerate(ST):
    y=T+k*RH+(0.22 if k>=2 else 0)
    if k==0: txt(s,0.62,T-0.2,5,0.16,"ЦІЛЬ: ПОЛИЦЯ 120 ₴ (ЗАПИТ ДО SIAS)",8.2,True,TEAL)
    if k==2: txt(s,0.62,y-0.2,5,0.16,"ЗАРАЗ: ПОТОЧНА ЗАКУПІВЛЯ",8.2,True,RED)
    txt(s,0.72,y+0.04,2.8,0.2,h1,10,True,NAVY); txt(s,0.72,y+0.27,2.8,0.2,h2,8,False,GREY)
    x=BX
    for v,c_ in zip(vals,CS):
        w=PWd*v/XM; _rect(s,x,y+0.04,w-0.015,0.44,c_)
        if w>0.3: txt(s,x,y+0.15,w-0.015,0.22,f"{v:.0f}",10,True,(WHITE if c_ in (CS[0],CS[1],CS[2]) else NAVY),C)
        x+=w
    txt(s,bx(shelf)+0.12,y+0.03,1.5,0.28,dig(shelf,2)+" ₴",15,True,(TEAL if k<2 else RED))
    txt(s,bx(shelf)+0.12,y+0.33,2.3,0.16,f"товар SIAS = {vals[0]/shelf*100:.0f} % полиці",8,False,GREY)
_rect(s,bx(120)-0.008,T-0.04,0.018,RH*6+0.28,RED); txt(s,bx(120)-0.7,T+RH*6+0.28,1.4,0.16,"120 ₴",8.4,True,RED,C)
_rect(s,0.62,6.56,12.10,0.46,LIGHT); _rect(s,0.62,6.56,0.045,0.46,NAVY)
txt(s,0.9,6.62,11.6,0.2,"Ціна партнеру (перші чотири блоки) = полиця ÷ 1,40.  Бонус 25 % і маржа 30 % — частки ціни партнеру.  Націнка магазину 40 % — понад ціну партнеру.",8.6,True,NAVY)
txt(s,0.9,6.82,11.6,0.16,"Усі числа — ₴ за пачку, з ПДВ. «Зараз» — СС із Self-Cost; «ціль» — за запитом €0,54 / €0,32 (вкладки «Цена 120 грн»).",7.8,False,GREY)
slide_tail_sources(s,"Фінмодель SIAS і Self-Cost_Sias.xlsx; курс 51,20 ₴/€. Логістика в «зараз» — залишок СС після товару, ПДВ і фінансування.",7.12)
NEW['struct']=s

# ============ ПОТРІБНА ЗАКУПІВЛЯ · ПО ВСІХ ПОЗИЦІЯХ ============
_x120=120/M.MARKUP; _ccm=_x120*(1-M.BONUS-M.MARGIN); _e1=_ccm/M.RATE
def need_slide(sc,ask,head,col):
    s=newslide(f"ФІНМОДЕЛЬ · ПОТРІБНА ЗАКУПІВЛЯ · {head}",f"Яка закупівля потрібна для полиці 120 ₴ · {head.lower()}")
    n=M.units(sc); b=M.batch_costs(sc)/n; v1=_e1/(1+M.FIN); v3=(v1-b)/(1+M.VAT)
    chips=[("120 ₴","цільова полиця"),(f"{dig(_x120,2)} ₴","ціна партнеру · 120 ÷ 1,40"),(f"{dig(_ccm,2)} ₴","макс. СС · × 45 %"),(f"€{dig(_e1,3)}","макс. СС в € · ÷ 51,20"),(f"€{dig(v3,3)}",f"мінус фінансування 2,5 %,\nвитрати партії €{dig(b,3)}, ПДВ 20 %"),(f"€{ask:.2f}".replace('.',','),"запит до SIAS · вниз до цента")]
    x=0.62; cw=1.82
    for i,(big,small) in enumerate(chips):
        last=i==len(chips)-1
        rr(s,x,1.45,cw,0.82,col if last else LIGHT,0.1)
        txt(s,x+0.12,1.52,cw-0.2,0.28,big,15,True,WHITE if last else NAVY)
        txt(s,x+0.12,1.86,cw-0.2,0.38,small,7,False,WHITE if last else GREY,lh=1.1)
        if i<len(chips)-1: txt(s,x+cw,1.68,0.2,0.2,"→",11,True,PALE,C)
        x+=cw+0.236
    txt(s,0.62,2.45,5.4,0.16,"ПО ВСІХ ПОЗИЦІЯХ · ЗАКУПІВЛЯ SIAS, €/ПАЧКУ",8.2,True,NAVY)
    legend(s,6.2,2.43,[("зараз",GREYB),("потрібно",col)],gap=0.3,fs=8.6)
    txt(s,9.2,2.45,3.5,0.16,"ПАРТІЯ, €: ЗАРАЗ → ПОТРІБНО",8.2,True,NAVY,R)
    T=2.75; RH=0.60; BX=3.3; BW=3.3; XM=0.82; pal=M.SC[sc]["pallets"]
    tot_now=0; tot_need=0
    for k,idx in enumerate(ORDER):
        nm,wt,_,cur=M.SKUS[idx]; u=pal[idx]*M.PACK_PALLET; y=T+k*RH
        if k%2==0: _rect(s,0.62,y-0.02,12.10,RH-0.02,LIGHT)
        txt(s,0.72,y+0.05,2.5,0.2,nm,10,True,NAVY); txt(s,0.72,y+0.28,2.5,0.18,f"{wt} · {pal[idx]} пал. · {th(u)} пачок",7.8,False,GREY)
        _rect(s,BX,y+0.06,BW*cur/XM,0.18,GREYB); txt(s,BX+BW*cur/XM+0.06,y+0.05,0.7,0.18,f"€{cur:.2f}".replace('.',','),9,True,GREY)
        _rect(s,BX,y+0.30,BW*ask/XM,0.18,col); txt(s,BX+BW*ask/XM+0.06,y+0.29,0.7,0.18,f"€{ask:.2f}".replace('.',','),9,True,NAVY)
        rr(s,7.45,y+0.12,1.1,0.32,col,0.25); txt(s,7.45,y+0.18,1.1,0.2,f"−{(1-ask/cur)*100:.0f} %",11,True,WHITE,C)
        now=u*cur; need=u*ask; tot_now+=now; tot_need+=need
        txt(s,8.8,y+0.16,3.9,0.24,f"€{th(now)}  →  €{th(need)}",10,True,NAVY,R)
    y=T+6*RH
    _rect(s,0.62,y+0.02,12.10,0.012,NAVY)
    txt(s,0.72,y+0.12,4.0,0.22,f"УСЬОГО · {sum(pal)} палет · {th(sum(pal)*M.PACK_PALLET)} пачок",10,True,DARKAMBER)
    txt(s,6.0,y+0.12,6.7,0.22,f"партія: €{th(tot_now)}  →  €{th(tot_need)}   (−€{th(tot_now-tot_need)} · −{(1-tot_need/tot_now)*100:.0f} %)",10.5,True,DARKAMBER,R)
    slide_tail_sources(s,"Фінмодель SIAS, вкладка «Цена 120 грн — %s»: бонус мережі 25 %%, маржа 30 %%, націнка 1,40, курс 51,20, ПДВ 20 %%, мито 0 %%. Запит однаковий для всіх шести SKU; згоди SIAS немає."%("6 паллет" if sc=="6" else "33 паллет"),7.12)
    return s
NEW['need33']=need_slide("33",A33,"33 ПАЛЕТИ",TEAL)
NEW['need6']=need_slide("6",A6,"6 ПАЛЕТ",RED)

# ============ ЗБІРКА ============
lst=prs.slides._sldIdLst
eid=lambda sid: next(e for e in lst if e.id==sid)
U=lambda k: USLIDE[k-1]
sl=lambda key: NEW[key].slide_id
SP=[x.slide_id for x in NEW['seg_slides']]; a1,a2=NEW['seg_split']
order=[sl('cover'),U(2),sl('origin'),U(4)]+SP[:a1]+[U(9)]+SP[a1:a2]+[U(14)]+SP[a2:]+\
      [sl('svod'),sl('fmt'),sl('chains'),sl('channels'),sl('chA'),sl('chB'),U(23),sl('shelf33'),sl('shelf6'),sl('pop'),sl('comp'),sl('struct'),sl('need33'),sl('need6'),U(27)]
keep=set(order)
for e in list(lst):
    if e.id not in keep: prs.part.drop_rel(e.rId); lst.remove(e)
els={e.id:e for e in lst}
for e in list(lst): lst.remove(e)
for sid in order: lst.append(els[sid])
for i,s_ in enumerate(prs.slides,1):
    for sh in s_.shapes:
        if sh.has_text_frame and re.fullmatch(r"\d\d",sh.text_frame.text.strip()) and abs(Emu(sh.left).inches-11.55)<0.05:
            sh.text_frame.paragraphs[0].runs[0].text=f"{i:02d}"
OUT=os.environ.get("OUT","/home/user/Chaplygin-Illya/SIAS_MARKET_RESEARCH_UA_V8.pptx")
prs.save(OUT); print("saved",OUT,len(prs.slides._sldIdLst))
