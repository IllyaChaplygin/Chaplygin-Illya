# -*- coding: utf-8 -*-
import openpyxl,sys
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.workbook.properties import CalcProperties
SRC="/root/.claude/uploads/8bc0034a-4b53-5022-9bf5-e5697c987387/5927c2e7-RAMEN_FINMODEL_4_TABS_PRESENTATION.xlsx"
OUT=sys.argv[1]
NAVY='223A5E'; INK='172B46'; GREYT='5E6D7E'; KEY='E7EDF5'; BAND2='F2F5F9'; INP='FFF4D2'; BLUE='0000FF'; GREEN='008000'
def font(sz=11,b=False,c=INK,i=False): return Font(name='Arial',size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill('solid',fgColor=c)
thin=Side(style='thin',color='D5DCE8'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
UAH='#,##0.00" ₴"'; E4='"€"0.0000'; E2='"€"0.00'; PCT='0.0%'; INT='#,##0'
wb=openpyxl.load_workbook(SRC)
CFG=[dict(name='Цена 120 грн — 6 паллет',src='Финмодель 6 паллет',head='6 паллет',freight=2600,pallets=6),
     dict(name='Цена 120 грн — 33 паллет',src='Финмодель 33 паллет',head='33 паллеты (Full Truck)',freight=3700,pallets=33)]
NOW=[0.75,0.75,0.65,0.65,0.65,0.65]
COLS="BDFHJLN"; ARR="CEGIKM"
def build(cfg):
    idx=wb.sheetnames.index(cfg['name']); del wb[cfg['name']]
    ws=wb.create_sheet(cfg['name'],idx); S="'"+cfg['src']+"'"
    ws.sheet_properties.tabColor=NAVY; ws.sheet_view.showGridLines=False
    for col in "ABCDEFGHIJKLMNO": ws.column_dimensions[col].width=2 if col in "AO" else (3.5 if col in ARR else (22 if col=="B" else 17))
    CT=Alignment(horizontal='center',vertical='center',wrap_text=True); WR=Alignment(vertical='center',wrap_text=True); LF=Alignment(horizontal='left',vertical='center',wrap_text=True)
    def put(a,v=None,f=None,fl=None,fmt=None,al=None,border=False,merge=None):
        c=ws[a]; c.value=v; c.font=f or font(); c.alignment=al or CT
        if fl: c.fill=fill(fl)
        if fmt: c.number_format=fmt
        if border: c.border=BOX
        if merge:
            ws.merge_cells(merge)
            if fl:
                for row in ws[merge]:
                    for cc in row: cc.fill=fill(fl)
        return c
    def label(r,t):
        put(f"B{r}",t,font(9,True,NAVY),al=LF,merge=f"B{r}:N{r}")
        for col in COLS+ARR: ws[f"{col}{r}"].border=Border(bottom=Side(style='medium',color=NAVY))
        ws.row_dimensions[r].height=20
    ws.row_dimensions[1].height=8
    put("B2",f"SIAS · закупка для полки 120 грн · {cfg['head']}",font(16,True),al=LF,merge="B2:N2"); ws.row_dimensions[2].height=28
    BATCH="($F$28+$H$28+$J$28+$F$28*$L$28*($D$28+$B$28*(1+$D$28)))"; PACKS="($B$30*$D$30)"
    put("B3",'="Запрос €"&ROUND(L7,2)&" за пачку → полка "&ROUND(N12,2)&" ₴ при марже "&ROUND(J12/(F12/(1-F26-D26))*100,0)&"%. Сейчас у SIAS €"&MIN(D17:D22)&"–€"&MAX(D17:D22)&" → снижение на "&ROUND((1-L7/MIN(D17:D22))*100,0)&"…"&ROUND((1-L7/MAX(D17:D22))*100,0)&"%."',font(11,True,NAVY),KEY,al=LF,merge="B3:N3"); ws.row_dimensions[3].height=26
    # --- линия расчёта
    label(5,"РАСЧЁТ ОДНОЙ ЛИНИЕЙ — от полки к закупочной цене")
    chain=[("Полка (цель)","=B26",UAH,'="цель, с НДС"'),
           ("Цена партнёру","=B7/H26",UAH,'="÷ множитель полки "&H26'),
           ("Макс. себестоимость","=D7*(1-F26-D26)",UAH,'="× "&ROUND((1-F26-D26)*100,0)&"% (100 − бонус − маржа)"'),
           ("Она же в €, без фин.","=F7/J26/(1+L26)",E4,'="÷ курс "&J26&" и ÷ (1+"&ROUND(L26*100,1)&"% фин.)"'),
           ("Предельная закупка",f"=(H7-{BATCH}/{PACKS})/((1+D28)*(1+B28))",E4,f'="− расходы партии "&ROUND({BATCH}/{PACKS},3)&" €/пачку; ÷ "&ROUND((1+D28)*(1+B28),2)&" (ПДВ)"'),
           ("ЗАПРОС SIAS","=ROUNDDOWN(J7,2)",E2,'="вниз до евроцента"')]
    for i,(lab,fm,nf,how) in enumerate(chain):
        col=COLS[i]; last=i==5
        put(f"{col}6",lab,font(9,True,GREYT),al=CT)
        put(f"{col}7",fm,font(17 if last else 14,True),KEY if not last else 'FCE3A6',nf,CT,border=True)
        put(f"{col}8",how,font(8,False,GREYT),al=CT)
        if i<5: put(f"{ARR[i]}7","→",font(14,True,NAVY),al=CT)
    ws.row_dimensions[6].height=28; ws.row_dimensions[7].height=40; ws.row_dimensions[8].height=34
    # --- структура полки
    label(10,"ИЗ ЧЕГО СОСТОИТ ПОЛКА ПРИ ЗАПРОСЕ — ₴ за пачку и доля полки")
    CC=f"(L7*J26+($F$28+$H$28+$J$28)/{PACKS}*J26+(L7+$F$28*$L$28/{PACKS})*$D$28*J26+(L7+$F$28*$L$28/{PACKS})*(1+$D$28)*$B$28*J26)*(1+L26)"
    st=[("Товар SIAS","=L7*J26",'3E8E96',None),("Логистика, ПДВ, фін.","=F12-B12",'46538A',"+"),("ИТОГО: себестоимость (СС)","="+CC,'223A5E',"="),
        ("Бонус сети","=F12/(1-F26-D26)*F26",'C4820A',"+"),("Наша прибыль","=F12/(1-F26-D26)-F12-H12",'F9A50B',"+"),
        ("Наценка магазина","=F12/(1-F26-D26)*(H26-1)",'8890B0',"+"),("ИТОГО: ПОЛКА","=F12/(1-F26-D26)*H26",'223A5E',"=")]
    # порядок значений: B,D,F,H,J,L,N
    for i,(lab,fm,colr,sym) in enumerate(st):
        col=COLS[i]; tot=lab.startswith("ИТОГО")
        put(f"{col}11",lab,font(9,True,GREYT),al=CT)
        put(f"{col}12",fm,font(13 if not tot else 14,True,'FFFFFF' if False else INK),KEY if tot else None,UAH,CT,border=True)
        put(f"{col}13",f"={col}12/$N$12",font(10,tot,colr),None,PCT,CT)
        if i>0: put(f"{ARR[i-1]}12",sym,font(13,True,NAVY),al=CT)
    ws.row_dimensions[11].height=28; ws.row_dimensions[12].height=32; ws.row_dimensions[13].height=18
    # --- SKU
    label(15,"ПО КАЖДОМУ SKU — цена закупки сейчас и в запросе")
    for col,t in zip("BDFHJL",["Продукт SIAS","Цена сейчас, €/уп.","Запрос, €/уп.","Снижение","Полка сейчас (маржа 30 %), ₴","Отрыв от цели 120, ₴"]):
        put(f"{col}16",t,font(9,True,'FFFFFF'),'3E5A85',al=CT)
    ws.row_dimensions[16].height=32
    for i in range(6):
        r=17+i; b=BAND2 if i%2==0 else None
        put(f"B{r}",f"={S}!C{11+i}",font(10,False,GREEN),b,al=LF)
        put(f"D{r}",NOW[i],font(11,True,BLUE),INP,E2,CT,border=True)
        put(f"F{r}","=$L$7",font(11),b,E2,CT); put(f"H{r}",f"=1-F{r}/D{r}",font(11,True),b,'0%',CT)
        put(f"J{r}",f"={S}!T{11+i}",font(11,False,GREEN),b,UAH,CT); put(f"L{r}",f"=J{r}-$B$26",font(11),b,'+#,##0.00" ₴";-#,##0.00" ₴"',CT)
        ws.row_dimensions[r].height=22
    # --- параметры
    label(24,"ИСХОДНЫЕ ДАННЫЕ — жёлтое можно менять, зелёное — из вкладки «"+cfg['src']+"», остальное — формулы")
    P1=[("Полка, ₴",120,UAH,'inp'),("Наша маржа",0.30,PCT,'inp'),("Бонус сети",f"={S}!D4",PCT,'lnk'),("Множитель полки",f"={S}!D6",'0.00"×"','lnk'),("Курс, ₴/€",f"={S}!D3",'0.00','lnk'),("Финансирование",0.025,PCT,'inp')]
    P2=[("Импортный ПДВ",0.20,PCT,'inp'),("Пошлина",0,PCT,'inp'),("Перевозка партии, €",cfg['freight'],'#,##0','inp'),("Брокер, €",288,'#,##0','inp'),("Прочие расходы, €",42,'#,##0','inp'),("Доля фрахта в таможне",0.7,'0.00','inp')]
    P3=[("Паллет в партии",cfg['pallets'],INT,'inp'),("Пачек на паллете",1600,INT,'inp')]
    for rl,rv,PP in ((25,26,P1),(27,28,P2),(29,30,P3)):
        for i,(lab,val,nf,kind) in enumerate(PP):
            col=COLS[i]
            put(f"{col}{rl}",lab,font(8,True,GREYT),al=CT)
            put(f"{col}{rv}",val,font(11,True,BLUE if kind=='inp' else GREEN),INP if kind=='inp' else BAND2,nf,CT,border=True)
        ws.row_dimensions[rl].height=18; ws.row_dimensions[rv].height=22
    put("F30","Источник чисел: Self-Cost_Sias (перевозка, брокер, прочие, ПДВ, фінансування, пачек на паллете).",font(8,False,GREYT,True),al=LF,merge="F30:N30")
    # --- мини-чувствительность
    label(32,"ЕСЛИ НАША МАРЖА ДРУГАЯ — запрос при той же полке")
    for i,m in enumerate([0.35,0.25,0.20,0.15]):
        col=COLS[i]
        put(f"{col}33",m,font(11,True,BLUE),INP,'"маржа "0%',CT,border=True)
        put(f"{col}34",f"=ROUNDDOWN((($B$26/$H$26*(1-$F$26-{col}33))/$J$26/(1+$L$26)-{BATCH}/{PACKS})/((1+$D$28)*(1+$B$28)),2)",font(13,True),KEY,E2,CT,border=True)
    put("J33","запрос, €/уп.: чем ниже наша маржа, тем выше допустимая закупка",font(9,False,GREYT,True),al=LF,merge="J33:N34")
    ws.row_dimensions[33].height=22; ws.row_dimensions[34].height=28
    put("B36","Допущения: бонус и множитель — по модели (не подтверждены сетями); импортный ПДВ входит в себестоимость, как в Self-Cost (если возмещается — допустимая закупка выше); без листинг-фи, промо и отсрочки; согласие SIAS не получено — цена для переговоров.",font(8,False,GREYT,True),al=LF,merge="B36:N36"); ws.row_dimensions[36].height=30
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
for c in CFG: build(c)
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print("saved",OUT)
