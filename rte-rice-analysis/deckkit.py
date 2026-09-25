# -*- coding: utf-8 -*-
"""Slide primitives matching RTE_Rice_Market_Research..pptx design language."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import data as d

def C(h): return RGBColor.from_string(h.lstrip("#").upper())
FONT="Segoe UI"
W, H = 13.333, 7.5

def new_deck():
    p=Presentation(); p.slide_width=Emu(12192000); p.slide_height=Emu(6858000); return p

def slide(prs, bg=d.WHITE):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb=C(bg); r.line.fill.background(); r.shadow.inherit=False
    return s

def rect(s,x,y,w,h,fill=None,line=None,lw=0.75,radius=None):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                           Inches(x),Inches(y),Inches(w),Inches(h))
    if radius is not None:
        try: shp.adjustments[0]=radius
        except Exception: pass
    if fill: shp.fill.solid(); shp.fill.fore_color.rgb=C(fill)
    else: shp.fill.background()
    if line: shp.line.color.rgb=C(line); shp.line.width=Pt(lw)
    else: shp.line.fill.background()
    shp.shadow.inherit=False
    return shp

def text(s,x,y,w,h,runs,size=10,color=d.INK,bold=False,align="l",
         anchor="t",spacing=1.15,space_after=0,caps=False):
    """runs: str or list of (txt,{opts}) ; returns textbox"""
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor={"t":MSO_ANCHOR.TOP,"m":MSO_ANCHOR.MIDDLE,"b":MSO_ANCHOR.BOTTOM}[anchor]
    paras = runs if isinstance(runs,list) else [runs]
    for i,item in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment={"l":PP_ALIGN.LEFT,"c":PP_ALIGN.CENTER,"r":PP_ALIGN.RIGHT,
                     "j":PP_ALIGN.JUSTIFY}[align]
        p.line_spacing=spacing
        if space_after: p.space_after=Pt(space_after)
        segs = item if isinstance(item,list) else [(item,{})]
        for txt,o in segs:
            r=p.add_run(); r.text=txt.upper() if o.get("caps",caps) else txt
            f=r.font; f.name=FONT; f.size=Pt(o.get("size",size))
            f.bold=o.get("bold",bold); f.color.rgb=C(o.get("color",color))
            if o.get("italic"): f.italic=True
    return tb

def header(s, kicker, title, num, sub=None):
    rect(s,0,0,W,0.055,fill=d.C_BRAND)
    text(s,0.62,0.42,9.5,0.3,kicker,size=8,color=d.C_BRAND,bold=True,caps=True)
    text(s,0.62,0.70,10.6,0.55,title,size=21,color=d.INK,bold=True)
    text(s,12.1,0.40,0.7,0.4,num,size=17,color=d.BG2,bold=True,align="r")
    if sub: text(s,0.62,1.30,11.3,0.3,sub,size=9,color=d.MUTED)

def footnote(s, txt, y=6.95):
    text(s,0.62,y,12.1,min(0.35,H-y-0.03),txt,size=6.5,color=d.MUTED2)

def stat(s,x,y,w,label,value,note,accent=d.INK,vsize=22,bg=d.BG):
    h=1.26
    rect(s,x,y,w,h,fill=bg,radius=0.10)
    text(s,x+0.22,y+0.17,w-0.4,0.2,label,size=6.5,color=d.MUTED,bold=True,caps=True)
    text(s,x+0.22,y+0.40,w-0.4,0.45,value,size=vsize,color=accent,bold=True)
    text(s,x+0.22,y+0.92,w-0.4,0.3,note,size=6.8,color=d.MUTED,spacing=1.2)

def table(s,x,y,w,cols,rows,colw,head_size=6.8,row_size=7.6,rh=0.255,
          zebra=True,accent_col=None,bold_cols=()):
    """cols: list of header strings; rows: list of list of (txt,color,bold) or str"""
    cx=x
    text(s,x,y,w,0.2,"",size=6)
    for c,cw in zip(cols,colw):
        text(s,cx,y,cw,0.2,c,size=head_size,color=d.MUTED,bold=True,caps=True)
        cx+=cw
    ln=rect(s,x,y+0.26,w,0.012,fill=d.LINE)
    yy=y+0.34
    for i,r in enumerate(rows):
        if zebra and i%2==1: rect(s,x-0.08,yy-0.045,w+0.16,rh,fill=d.BG,radius=0.18)
        cx=x
        for j,(cell,cw) in enumerate(zip(r,colw)):
            if isinstance(cell,tuple): txt,col,bd = (cell+(False,))[:3]
            else: txt,col,bd = cell,d.INK2,(j in bold_cols)
            text(s,cx,yy,cw,rh,str(txt),size=row_size,color=col,bold=bd,anchor="t")
            cx+=cw
        yy+=rh
    return yy

def pill(s,x,y,txt,fill,color=d.WHITE,size=6.5,pad=0.12):
    w=max(0.42,len(txt)*0.055+pad*2)
    rect(s,x,y,w,0.20,fill=fill,radius=0.5)
    text(s,x,y+0.012,w,0.18,txt,size=size,color=color,bold=True,align="c",caps=True)
    return w

def img(s,path,x,y,w=None,h=None):
    k={}
    if w: k["width"]=Inches(w)
    if h: k["height"]=Inches(h)
    return s.shapes.add_picture(path,Inches(x),Inches(y),**k)

def callout(s,x,y,w,h,title,body,accent=d.C_BRAND,bg=d.BG):
    rect(s,x,y,w,h,fill=bg,radius=0.06)
    rect(s,x,y,0.045,h,fill=accent)
    text(s,x+0.26,y+0.17,w-0.5,0.25,title,size=8,color=accent,bold=True,caps=True)
    text(s,x+0.26,y+0.46,w-0.5,h-0.6,body,size=8,color=d.INK2,spacing=1.35)
