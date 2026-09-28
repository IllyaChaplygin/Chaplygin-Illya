# -*- coding: utf-8 -*-
"""Copy the supplier deck's slides into the main presentation, images and all."""
import copy, io
from pptx.oxml.ns import qn

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_ATTRS = [f"{{{R_NS}}}embed", f"{{{R_NS}}}link", f"{{{R_NS}}}id"]
SKIP_RELTYPES = ("slideLayout", "notesSlide", "slideMaster")


def _blank_layout(prs):
    """The layout with the fewest placeholders — closest to a blank canvas."""
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


def _rebind(el, src_slide, dest_slide, rid_map):
    """Repoint every relationship reference inside a copied element."""
    for node in el.iter():
        for attr in REL_ATTRS:
            old = node.get(attr)
            if not old:
                continue
            if old not in rid_map:
                rid_map[old] = _clone_rel(src_slide, dest_slide, old)
            new = rid_map[old]
            if new:
                node.set(attr, new)
            else:
                del node.attrib[attr]


def copy_slide(src_slide, dest_prs):
    """Append a copy of src_slide to dest_prs, rebinding its image relationships."""
    dest = dest_prs.slides.add_slide(_blank_layout(dest_prs))
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for ph in list(dest.placeholders):
        ph._element.getparent().remove(ph._element)

    rid_map = {}

    # the slide background can itself be a picture fill — it needs rebinding too
    src_bg = src_slide._element.find(qn("p:cSld")).find(qn("p:bg"))
    if src_bg is not None:
        new_bg = copy.deepcopy(src_bg)
        _rebind(new_bg, src_slide, dest, rid_map)
        dest._element.find(qn("p:cSld")).insert(0, new_bg)

    src_tree = src_slide.shapes._spTree
    dest_tree = dest.shapes._spTree
    for el in src_tree:
        tag = el.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        new_el = copy.deepcopy(el)
        _rebind(new_el, src_slide, dest, rid_map)
        dest_tree.append(new_el)
    return dest


def _clone_rel(src_slide, dest_slide, rId):
    """Copy one relationship target into the destination, return its new rId."""
    try:
        rel = src_slide.part.rels[rId]
    except KeyError:
        return None
    kind = rel.reltype.rsplit("/", 1)[-1]
    if kind in SKIP_RELTYPES:
        return None
    if rel.is_external:
        return dest_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
    if kind == "image":
        blob = rel.target_part.blob
        image_part, new_rid = dest_slide.part.get_or_add_image_part(io.BytesIO(blob))
        return new_rid
    # anything else (charts, embedded objects) — carry the part over as-is
    return dest_slide.part.rels._add_relationship(
        rel.reltype, rel.target_part, dest_slide.part.rels._next_rId)


def merge(dest_prs, src_prs, indices):
    """Copy the given source slide indices; returns the new slide objects."""
    return [copy_slide(src_prs.slides[i], dest_prs) for i in indices]


def normalize_partnames(prs):
    """Renumber slide parts sequentially so no two share a package partname.

    python-pptx picks the partname for a new slide from a counter that can
    collide with parts already in a deck assembled by reordering. Two passes,
    via temporary names, so no transient collision either.
    """
    from pptx.opc.packuri import PackURI
    parts = [prs.slides[i].part for i in range(len(prs.slides._sldIdLst))]
    for n, part in enumerate(parts, 1):
        part.partname = PackURI(f"/ppt/slides/_tmp{n}.xml")
    for n, part in enumerate(parts, 1):
        part.partname = PackURI(f"/ppt/slides/slide{n}.xml")
    return len(parts)


def drop_orphan_slides(prs):
    """Remove slide parts no longer referenced by the slide id list."""
    keep = {prs.slides[i].part for i in range(len(prs.slides._sldIdLst))}
    pres_part = prs.part
    dropped = 0
    for rId, rel in list(pres_part.rels.items()):
        if rel.reltype.endswith("/slide") and rel.target_part not in keep:
            pres_part.rels.pop(rId)
            dropped += 1
    return dropped
