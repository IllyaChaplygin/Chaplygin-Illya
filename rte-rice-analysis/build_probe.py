# -*- coding: utf-8 -*-
"""Slides 18-31 on their own, to tell a bad slide from a cumulative ceiling.

The full deck renders to slide 17 and then shows white. Either those later
slides carry something PowerPoint rejects, or the deck as a whole crosses a
limit. Building exactly those slides as their own file separates the two:
if they render here, the slides are sound and the volume is the problem.
"""
import re
from pptx.util import Emu
import deckkit as K
import slides as S
from build_final import PLAN          # same builders, same order

TAIL = PLAN[17:]                      # 18-й слайд і далі

prs = K.new_deck()
for kind, arg in TAIL:
    S.card_page(prs, arg) if kind == "card" else getattr(S, kind)(prs)

for pos, sl in enumerate(prs.slides, 18):     # keep the numbering of the full deck
    for sh in sl.shapes:
        if (sh.has_text_frame
                and re.fullmatch(r"\d\d?", sh.text_frame.text.strip())
                and Emu(sh.left).inches > 11.5 and Emu(sh.top).inches < 0.9):
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = f"{pos:02}"
                for r in p.runs[1:]:
                    r.text = ""

prs.save("TEST_slides_18-31.pptx")
print("TEST_slides_18-31.pptx:", len(prs.slides._sldIdLst), "слайдів")
