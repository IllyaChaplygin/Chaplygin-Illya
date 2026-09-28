# -*- coding: utf-8 -*-
"""The complete deck, assembled in a fresh package.

Every build that failed was Presentation(their_file) saved back out. The one
that opens was built from a fresh presentation. So build the whole thing
fresh and copy their 19 slides into it, shapes and images and all, instead
of re-serialising their package.

Their slides carry explicit colours and fonts (1674 srgbClr, Segoe UI
everywhere); the schemeClr references sit in the <p:style> defaults that the
explicit fills override, so moving them to another theme does not change how
they look.
"""
from pptx import Presentation
import re
from pptx.util import Emu
import deckkit as K
import slides as S
import merge_suppliers as MS

ORIG = "/root/.claude/uploads/df3fafbd-cd5a-5689-8909-9c95fa4cbf11/c731572d-RTE_Rice_Market_Research..pptx"
OUT = "RTE_Rice_Market_Research_FINAL.pptx"

src = Presentation(ORIG)
out = K.new_deck()

# their slide index -> where it goes, and which of mine follow it
PLAN = [
    (0, []), (1, []), (2, []), (3, ["sl_04", "sl_05"]),
    (4, ["sl_06"]), (5, ["sl_15", "sl_16"]), (6, []),
    (7, []), (8, []), (9, []), (10, []), (11, []), (12, []),
    (13, []), (14, []), (15, []),
    (16, ["sl_17", "sl_18", "sl_03"]),
    (17, []),
    (18, ["sl_02", "sl_07", "sl_09", "sl_10", "sl_11", "sl_12", "sl_13", "sl_14"]),
]

mine = set()
for oi, after in PLAN:
    MS.copy_slide(src.slides[oi], out)
    for fn in after:
        sl = getattr(S, fn)(out)
        mine.add(id(sl._element))

# page numbers on my slides; theirs keep whatever they had
for pos, sl in enumerate(out.slides, 1):
    if id(sl._element) not in mine:
        continue
    for sh in sl.shapes:
        if (sh.has_text_frame
                and re.fullmatch(r"\d\d", sh.text_frame.text.strip())
                and Emu(sh.left).inches > 11.5 and Emu(sh.top).inches < 0.9):
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = f"{pos:02}"
                for r in p.runs[1:]:
                    r.text = ""

out.save(OUT)
print(f"{OUT}: {len(out.slides._sldIdLst)} слайдів "
      f"({len(PLAN)} ваших + {len(mine)} нових)")
