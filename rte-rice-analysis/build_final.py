# -*- coding: utf-8 -*-
"""The complete deck, every slide generated — nothing copied from the original.

Copying their slides into a package is what PowerPoint rejects; a deck built
from scratch opens. So their slides are rebuilt rather than copied: the
overview slides from the research data, the ten brand-card pages from the
card text and photos pulled out of the original with extract_cards.py.
"""
import re
from pptx.util import Emu
import deckkit as K
import slides as S

OUT = "RTE_Rice_Market_Research_FINAL.pptx"

# (builder, argument) in final order
PLAN = [
    ("sl_t1", None),      # титул
    ("sl_t2", None),      # головне — шість цифр
    ("sl_t3", None),      # обсяг ринку
    ("sl_t4", None),      # карта ринку — чотири технології
    ("sl_04", None),      # матриця брендів 1
    ("sl_05", None),      # матриця брендів 2
    ("sl_06", None),      # аналітика форматів
    ("sl_15", None),      # формат × грамаж
    ("sl_16", None),      # найпопулярніші поєднання
    ("sl_t5", None),      # цінові сходи 23 брендів
    ("sl_19", None),      # карта сегментів
    ("card", 7), ("card", 8), ("card", 9), ("card", 10),
    ("card", 11), ("card", 12), ("card", 13), ("card", 14),
    ("card", 15), ("card", 16),
    ("sl_17", None),      # ритейл-аудит
    ("sl_18", None),      # азійські бренди в мережах
    ("sl_03", None),      # канали продажу
    ("sl_02", None),      # шість висновків
    ("sl_07", None),      # ранжир за упаковку
    ("sl_09", None),      # фінмодель проти ринку
    ("sl_11", None),      # рекомендація формату
    ("sl_12", None),      # рекомендація ціни
    ("sl_13", None),      # постачальник
    ("sl_14", None),      # план дій
]

prs = K.new_deck()
for kind, arg in PLAN:
    if kind == "card":
        S.card_page(prs, arg)
    else:
        getattr(S, kind)(prs)

# continuous page numbers, title page excluded
for pos, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if (sh.has_text_frame
                and re.fullmatch(r"\d\d?", sh.text_frame.text.strip())
                and Emu(sh.left).inches > 11.5 and Emu(sh.top).inches < 0.9):
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = f"{pos:02}"
                for r in p.runs[1:]:
                    r.text = ""

prs.save(OUT)
print(f"{OUT}: {len(prs.slides._sldIdLst)} слайдів, усі згенеровані")
