# -*- coding: utf-8 -*-
"""Merge the part-2 analysis slides into the original market research deck."""
import re, copy
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
import data as d, deckkit as K, slides as S

ORIG="/root/.claude/uploads/df3fafbd-cd5a-5689-8909-9c95fa4cbf11/c731572d-RTE_Rice_Market_Research..pptx"
prs=Presentation(ORIG)
n_orig=len(prs.slides._sldIdLst)
assert n_orig==19, n_orig

def badge_shape(sl):
    """The page-number textbox: top-right corner, two digits."""
    for sh in sl.shapes:
        if sh.has_text_frame and re.fullmatch(r"\d\d", sh.text_frame.text.strip()):
            if Emu(sh.left).inches>11.5 and Emu(sh.top).inches<0.9:
                return sh
    return None

def set_badge(sl, num):
    sh=badge_shape(sl)
    if sh is None: return False
    p=sh.text_frame.paragraphs[0]
    if p.runs: p.runs[0].text=f"{num:02d}"
    for r in p.runs[1:]: r.text=""
    return True

# ── section divider ────────────────────────────────────────────────
def divider(prs, kicker, title, sub, items):
    s=K.slide(prs, d.INK)
    K.rect(s,0,0,K.W,0.09,fill=d.C_BRAND)
    K.text(s,0.9,2.35,9,0.3,kicker,size=9,color=d.C_BRAND,bold=True)
    K.text(s,0.9,2.78,10.6,1.35,title,size=36,color=d.WHITE,bold=True,spacing=1.06)
    K.text(s,0.9,4.42,9.6,0.7,sub,size=11.5,color="#B9C2DA",spacing=1.5)
    for i,(k,v) in enumerate(items):
        x=0.9+i*2.95
        K.text(s,x,5.52,2.7,0.2,k,size=7,color="#717890",bold=True,caps=True)
        K.text(s,x,5.74,2.7,0.42,v,size=19,color=d.C_BRAND,bold=True)
    return s

# ── build the new slides onto the original deck ────────────────────
built={}
for fn in ("sl_02","sl_03","sl_04","sl_05","sl_06","sl_07","sl_08",
           "sl_09","sl_10","sl_11","sl_12","sl_13","sl_14"):
    built[fn]=getattr(S,fn)(prs)
DIV=divider(prs,"ЧАСТИНА 2 · НАША ПОЗИЦІЯ",
  "Де ми стоїмо\nі що заводити",
  "Зіставлення фінансової моделі з полицею, рекомендація формату й ціни, цільовий FOB.",
  [("НАШИХ SKU","32"),("ТОВАРНИХ ГРУП","7"),("ПОСТАЧАЛЬНИКІВ","3"),("СЦЕНАРІЇВ ЦІНИ","3")])

ids=list(prs.slides._sldIdLst)                 # 0..18 original, 19.. new
orig=ids[:n_orig]
def nid(fn): return ids[n_orig+list(built).index(fn)] if fn in built else None
new={fn:ids[n_orig+i] for i,fn in enumerate(built)}
new["DIV"]=ids[-1]

# ── target narrative order ─────────────────────────────────────────
order=[
 ("o",0,None),      # 01 титул
 ("o",1,2),         # 02 головне — шість цифр
 ("o",2,3),         # 03 обсяг ринку
 ("o",3,4),         # 04 карта ринку
 ("n","sl_04",5),   # 05 матриця брендів 1/2
 ("n","sl_05",6),   # 06 матриця брендів 2/2
 ("o",4,7),         # 07 формати упаковки (оригінал)
 ("n","sl_06",8),   # 08 формати — аналітика і рекомендація
 ("o",5,9),         # 09 грамаж
 ("o",6,10),        # 10 ціна за упаковку — сходи брендів
 ("o",7,11),("o",8,12),("o",9,13),("o",10,14),("o",11,15),
 ("o",12,16),("o",13,17),("o",14,18),("o",15,19),   # 11–19 картки брендів
 ("o",16,20),       # 20 що стоїть у мережі замість нас
 ("o",17,21),       # 21 де представлено (оригінал)
 ("n","sl_03",22),  # 22 канали — розширений перелік
 ("o",18,23),       # 23 висновки дослідження
 ("n","DIV",None),  # 24 роздільник частини 2
 ("n","sl_02",25),  # 25 сім висновків
 ("n","sl_07",26),  # 26 ранжир за упаковку
 ("n","sl_08",27),  # 27 ціна за 100 г
 ("n","sl_09",28),  # 28 фінмодель проти ринку
 ("n","sl_10",29),  # 29 аудит моделі
 ("n","sl_11",30),  # 30 рекомендація формату
 ("n","sl_12",31),  # 31 рекомендація ціни
 ("n","sl_13",32),  # 32 постачальник
 ("n","sl_14",33),  # 33 план дій
]
assert len(order)==33, len(order)

sldIdLst=prs.slides._sldIdLst
seq=[]
for kind,ref,badge in order:
    seq.append(orig[ref] if kind=="o" else new[ref])
for el in list(sldIdLst): sldIdLst.remove(el)
for el in seq: sldIdLst.append(el)

# ── renumber every badge to its new position ───────────────────────
missing=[]
for i,(sl,(kind,ref,badge)) in enumerate(zip(prs.slides,order),1):
    if badge is None: continue
    if not set_badge(sl,badge): missing.append(i)
if missing: print("badge not found on:",missing)

# ── refresh the title slide for the expanded scope ─────────────────
t=prs.slides[0]
for sh in t.shapes:
    if sh.has_text_frame and "ДОСЛІДЖЕННЯ РИНКУ" in sh.text_frame.text:
        p=sh.text_frame.paragraphs[0]
        if p.runs:
            p.runs[0].text="ДОСЛІДЖЕННЯ РИНКУ ТА ПОЗИЦІОНУВАННЯ · ВЕРЕСЕНЬ 2026"
            for r in p.runs[1:]: r.text=""

prs.save("RTE_Rice_Market_Research_FULL.pptx")
print("saved:",len(prs.slides._sldIdLst),"slides")
