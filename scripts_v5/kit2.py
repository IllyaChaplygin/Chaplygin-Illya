# -*- coding: utf-8 -*-
"""Спільний модуль колоди «Корейський рамен в Україні» за моделлю рисової колоди."""
import sys,json,os,re,math,statistics as st,collections as co
SC="/tmp/claude-0/-home-user-Chaplygin-Illya/8bc0034a-4b53-5022-9bf5-e5697c987387/scratchpad"
sys.path.insert(0,SC)
from deckkit import *
from deckkit import _rect,_S
from pptx.enum.text import PP_ALIGN as A
import calc120 as M
L=A.LEFT; R=A.RIGHT; C=A.CENTER

# ---------------- дані ----------------
F=json.load(open(f"{SC}/facts.json",encoding='utf-8'))
SEG=json.load(open(f"{SC}/cards_seg.json",encoding='utf-8'))
CK=json.load(open(f"{SC}/cards_kor.json",encoding='utf-8'))
CM=json.load(open(f"{SC}/cards_mass.json",encoding='utf-8'))
SS=json.load(open(f"{SC}/segstats.json",encoding='utf-8'))
_SJ=json.load(open(f"{SC}/segmented.json",encoding='utf-8')); SD_ROWS=_SJ['rows']
BT={r['tm']:r for r in json.load(open(f"{SC}/brandtiers.json",encoding='utf-8'))}
if "без тм" in CK: CK["Choi's"]=CK.pop("без тм")

# ---- v5: без рисових галушок (сегмент 4) — прибрано з усіх розрахунків ----
SD_ROWS=[r for r in SD_ROWS if r['seg']!=4]
SEG.pop('4',None); SS.pop('4',None)
_ZR=json.load(open(f"{SC}/zakaz_raw.json",encoding='utf-8'))
_EANS={r['ean']:r for r in SD_ROWS}
ZREC=[z for z in _ZR if z['ean'] in _EANS]                       # записи «магазин × SKU» для 89 позицій
F['n_sku']=len(SD_ROWS); F['n_brands']=len({r['tm'] for r in SD_ROWS})
F['n_chains']=len({c for r in SD_ROWS for c in r['chains']}|{'silpo'})
def _med(a): return st.median(a) if a else 0
_bands=[(40,90,"40–90 г"),(90,110,"90–110 г"),(110,125,"110–125 г"),(125,145,"125–145 г"),(145,900,"145 г +")]
F['bands']=[]
for a_,b_,lab in _bands:
    v=[r for r in SD_ROWS if a_<=r['w']<b_]
    F['bands'].append({"lab":lab,"n":len(v),"share":round(len(v)/len(SD_ROWS)*100),
                       "pmed":round(_med([r['pmed'] for r in v])),"g":round(_med([r['per100'] for r in v]))})
F['matrix']=[[sum(1 for r in SD_ROWS if r['seg'] in fm and a_<=r['w']<b_) for a_,b_,_ in _bands] for fm in ((1,2),(3,))]
F['chain_rows']=[]
for ch_,v in sorted(co.defaultdict(list,{c:[z for z in ZREC if z['chain']==c] for c in {z['chain'] for z in ZREC}}).items(),key=lambda x:-len({z['ean'] for z in x[1]})):
    pr=[z['price'] for z in v]
    F['chain_rows'].append({"chain":ch_,"n":len({z['ean'] for z in v}),"med":round(_med(pr)),"lo":round(min(pr)),"hi":round(max(pr))})
TOP2=[b for b in ("Samyang","Nongshim")]
N_TOP2=sum(1 for r in SD_ROWS if r['tm'] in TOP2)
def g(b): return CK.get(b,[])
def tl(tm,label=None):
    r=BT[tm]; p=f"{SC}/img/{r['best']['ean']}.jpg"
    return {"img":p if os.path.exists(p) else None,"name":label or tm,
            "price":f"{r['pmed']} грн","sub":f"{r['g']} грн/100 г · {r['nch']} мереж · {r['sku']} SKU"}
def img_for(ean):
    p=f"{SC}/img/{ean}.jpg"; return p if os.path.exists(p) else None

# ---------------- форматування ----------------
def d1(x): return f"{x:.1f}".replace('.',',')
def d2(x): return f"{x:.2f}".replace('.',',')
def d0(x): return f"{x:.0f}"
def th(x): return f"{x:,.0f}".replace(',',' ')
def th1(x): return f"{x:,.1f}".replace(',',' ').replace('.',',')
def eur(x,n=2): return "€"+f"{x:.{n}f}".replace('.',',')
def pct(x,n=1): return (f"{x:.{n}f}".replace('.',',').replace('-','−'))+" %"
def sgn(x,n=1): return (f"{x:+.{n}f}".replace('.',',').replace('-','−'))+" %"
NM={'auchan':'Ашан','novus':'NOVUS','metro':'METRO','cosmos':'«Космос»','vostorg':'«Восторг»',
    'tavriav':'«Таврія В»','torba':'«Торба»','zaraz':'«Зараз»','grono':'Grono','onde':'Onde',
    'chudomarket':'«ЧудоМаркет»','kharkiv':'«Клас»','silpo':'«Сільпо»'}
NATIONAL={'auchan','novus','metro','zaraz','silpo'}

# сегменти
SNAME={'1':'РАМЕН У ПАКЕТІ · З БУЛЬЙОНОМ','2':'РАМЕН У ПАКЕТІ · ІЗ СОУСОМ','3':'РАМЕН У СТАКАНІ','4':'РИСОВІ ГАЛУШКИ ТА ЛОКШИНА'}
SSHORT={'1':'Пакет · з бульйоном','2':'Пакет · із соусом','3':'Стакан','4':'Рисові'}
SCOL={'1':NAVY,'2':DARKAMBER,'3':PURPLE,'4':TEAL}
STITLE={'1':'Пакет з бульйоном','2':'Пакет із соусом','3':'Стакан','4':'Рисові галушки'}
SPREP={'1':'Варіння 4–5 хв · подається з бульйоном','2':'Варіння 4–5 хв · воду злити, змішати із соусом',
       '3':'Залити окропом 3–4 хв · без плити','4':'Рисова основа · окремий привід споживання'}
PB=[(0,80,"до 80"),(80,95,"80–95"),(95,110,"95–110"),(110,130,"110–130"),(130,150,"130–150"),(150,180,"150–180"),(180,9e9,"180 +")]
def band_of(p):
    for i,(lo,hi,_) in enumerate(PB):
        if lo<=p<hi: return i
    return len(PB)-1

PACK=[r for r in SD_ROWS if r['seg'] in (1,2) and r.get('pmed')]
PACK_P=sorted(r['pmed'] for r in PACK)
PACK_MED=st.median(PACK_P)
def quart(arr):
    """Квартили методом statistics.quantiles (exclusive) — як у статистиці сегментів."""
    return st.quantiles(sorted(arr),n=4)
PACK_Q=quart(PACK_P)

# ---------------- базові примітиви колоди ----------------
def estat(s,l,t,label,big,note,acc=DARKAMBER,w=3.88,h=1.46,bigsz=26):
    _rect(s,l,t,w,h,LIGHT); _rect(s,l,t,0.045,h,acc)
    txt(s,l+0.28,t+0.17,w-0.52,0.14,label,7.6,True,GREY)
    txt(s,l+0.28,t+0.40,w-0.52,0.46,big,bigsz,True,acc)
    ny=(t+h-0.52) if h>=1.40 else (t+0.40+bigsz*1.25/72+0.06)
    txt(s,l+0.28,ny,w-0.52,0.40,note,8.4,False,GREY,lh=1.3)

def ecall(s,l,t,w,h,head,body,acc=AMBER,hs=10,bs=8.8):
    _rect(s,l,t,w,h,LIGHT); _rect(s,l,t,0.045,h,acc)
    txt(s,l+0.30,t+0.18,w-0.60,0.18,head,hs,True,NAVY)
    txt(s,l+0.30,t+0.48,w-0.60,h-0.66,body,bs,False,GREY,lh=1.42)

def estrip(s,t,head,body,acc=NAVY,h=0.60):
    _rect(s,0.62,t,12.10,h,LIGHT); _rect(s,0.62,t,0.045,h,acc)
    txt(s,0.92,t+0.10,11.50,0.16,head,9.2,True,NAVY)
    txt(s,0.92,t+0.32,11.50,0.20,body,8.6,False,GREY,lh=1.3)

def verdict(s,t,head,body,acc=RED,h=0.92):
    _rect(s,0.62,t,12.10,h,NAVY); _rect(s,0.62,t,0.06,h,acc)
    txt(s,0.96,t+0.14,11.50,0.22,head,11,True,GOLD)
    txt(s,0.96,t+0.42,11.50,0.40,body,13.5,True,WHITE,lh=1.22)

def cardrow(s,t,items,acc=BLUE,n=7):
    for i,c in enumerate(items[:n]):
        card(s,CARD_X[i],t,c.get("img"),c["name"],c["w"],c["price"],c["sellers"],acc)

def minibar(s,l,t,w,rows,xmax,acc=BLUE,lab_w=1.05,val_w=0.62,rowh=0.205):
    bw=w-lab_w-val_w
    for lab,v,extra in rows:
        txt(s,l,t+0.02,lab_w-0.08,0.14,lab,7.8,False,NAVY)
        _rect(s,l+lab_w,t+0.045,max(bw*v/xmax,0.02),0.105,acc)
        txt(s,l+w-val_w,t+0.02,val_w,0.14,extra,7.8,True,GREY,PP_ALIGN.RIGHT)
        t+=rowh
    return t

def tint(c,k=0.55):
    r,g_,b=c[0],c[1],c[2]
    return RGBColor(int(r+(255-r)*k),int(g_+(255-g_)*k),int(b+(255-b)*k))
AMBLIGHT=RGBColor(0xFD,0xF3,0xDF)

def rangebars(s,l,t,w,rows,xmax,ticks,rowh=0.285,lab_w=2.60,w_col=1.05,n_col=0.45,val_w=1.15,band=None,bandlab=None,
              head=None,headcol=NAVY):
    """Діапазонні смуги за брендами: тонка лінія min–max, блок — медіана. rows=[dict(label,wt,n,lo,med,hi,col,bold)]."""
    px=l+lab_w+w_col+n_col+0.12; pw=w-(lab_w+w_col+n_col+0.12)-val_w
    def X(v): return px+pw*min(max(v,0),xmax)/xmax
    H=len(rows)*rowh
    if head:
        _rect(s,l,t,w,0.30,headcol)
        txt(s,l+0.14,t+0.07,lab_w,0.18,head,8.2,True,WHITE)
        txt(s,l+lab_w,t+0.08,w_col,0.14,"ВАГА",6.6,True,WHITE)
        txt(s,l+lab_w+w_col,t+0.08,n_col,0.14,"SKU",6.6,True,WHITE,PP_ALIGN.RIGHT)
        t+=0.36
    if band:
        _rect(s,X(band[0]),t-0.02,X(band[1])-X(band[0]),H+0.04,AMBLIGHT)
    for tv in ticks:
        _rect(s,X(tv),t-0.02,0.006,H+0.04,HAIR)
    y=t
    for rw in rows:
        b=rw.get('bold',False); col=rw['col']
        if b: _rect(s,l,y,w-0.02,rowh-0.02,AMBLIGHT)
        txt(s,l+0.10,y+rowh/2-0.095,lab_w-0.12,0.19,rw['label'],8.4,True,DARKAMBER if b else NAVY)
        txt(s,l+lab_w,y+rowh/2-0.085,w_col,0.17,rw.get('wt',''),7.6,False,GREY)
        txt(s,l+lab_w+w_col,y+rowh/2-0.085,n_col,0.17,str(rw.get('n','')),8,True,GREY,PP_ALIGN.RIGHT)
        lo,med,hi=rw['lo'],rw['med'],rw['hi']
        if hi>lo: _rect(s,X(lo),y+rowh/2-0.03,max(X(hi)-X(lo),0.02),0.06,tint(col,0.45) if not b else tint(col,0.25))
        _rect(s,X(med)-0.07,y+rowh/2-0.095,0.14,0.19,col)
        txt(s,X(hi)+0.10,y+rowh/2-0.09,val_w+0.3,0.18,rw.get('txt') or (f"{lo:.0f}–{hi:.0f}" if hi>lo+0.5 else f"{med:.0f}"),8.2,True,DARKAMBER if b else NAVY)
        y+=rowh
    for tv in ticks:
        txt(s,X(tv)-0.35,y+0.02,0.70,0.14,f"{tv}",7.2,False,GREY,PP_ALIGN.CENTER)
    if bandlab:
        txt(s,X(band[0])-1.6,y+0.18,X(band[1])-X(band[0])+3.2,0.14,bandlab,7.2,True,DARKAMBER,PP_ALIGN.CENTER)
    return y+0.34

def brand_rows(rows_,col,cap=None):
    """Групування SKU за брендом → рядки для rangebars (за зростанням медіани)."""
    d=co.defaultdict(list)
    for r in rows_:
        if r.get('pmed'): d[r['tm']].append(r)
    out=[]
    for tm,v in d.items():
        pm=[x['pmed'] for x in v]; ws=sorted(set(round(x['w']) for x in v))
        wt=(f"{ws[0]} г" if len(ws)==1 else f"{ws[0]}–{ws[-1]} г")
        out.append(dict(label=tm,wt=wt,n=len(v),lo=min(pm),med=st.median(pm),hi=max(pm),col=col))
    out.sort(key=lambda x:x['med'])
    return out[:cap] if cap else out

def slide_tail_sources(s,text,y=7.12):
    txt(s,0.62,y,12.10,0.30,text,7.2,False,PALE,lh=1.25)

prs=new_deck(); N=[0]
def num(): N[0]+=1; return f"{N[0]:02d}"
TOP2CH={t:len({c for r in SD_ROWS if r['tm']==t for c in r['chains']}|({'silpo'} if t in('Samyang','Nongshim') else set())) for t in TOP2}
